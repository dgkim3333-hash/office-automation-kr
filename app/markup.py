"""마크업표.xlsx: 파트넘버 → 마크업%(· 반올림자리). 본사 PO 시트의 본사 단가 = 고객 단가 × (1 + 마크업% ÷ 100).

작업 공용 파일(고객마스터 옆)이고 친구가 직접 고친다. 돈이 걸린 값이므로 애매하면 확인필요로 보낸다:
마크업표에 없는 파트를 0% 로 넘기지 않고, 반올림은 그 파트에 반올림자리가 있을 때만 한다(설계 13-2).
"""
import hashlib
import io
from dataclasses import dataclass
from decimal import ROUND_HALF_UP, Decimal, Inexact, InvalidOperation, localcontext
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from app.excel_extract import excel_number
from app.rules import MarkupColumns
from app.verify import Check

MARKUP_FILE = "마크업표.xlsx"  # 작업 폴더 맨 위(고객마스터 옆). 작업들이 같이 쓴다
PART_COLUMN = "파트넘버"
RATE_COLUMN = "마크업%"
PLACES_COLUMN = "반올림자리"
MAX_RATE = Decimal(10)
MAX_PLACES = 10
RATE = "마크업.율"
PRICE = "마크업.단가"
AMOUNT = "마크업.금액"
BLANK_PART = "(빈칸)"

# 가짜 예시(samples 의 가짜 품번). 실제로 쓸 때는 여러분 파트로 바꾼다. [파트넘버, 마크업%, 반올림자리]
EXAMPLE_ROWS = [
    ["VTS-CDU-300", 3, None],
    ["VTS-RDHX-42U", 3, None],
    ["VTS-MFD-12", 2, None],
    ["VTS-QD-DN25", 3, None],
    ["VTS-SENS-LEAK", 0, None],
    ["VTS-CHL-500RT", 2, None],
    ["VTS-PMP-CDU-K", 3, None],
    ["VTS-FLT-50", 2, 2],
]


@dataclass
class MarkupRow:
    part: str               # 마크업표에 적힌 파트넘버(앞뒤 공백 뺌)
    rate: Decimal           # % 숫자. 3 = 3%
    places: int | None = None  # 본사 단가 반올림 자리(사사오입). None = 반올림하지 않음


class MarkupError(Exception):
    """마크업표.xlsx 를 쓸 수 없을 때. 메시지는 사용자에게 그대로 보여 준다."""


def part_key(value) -> str:
    """파트넘버 비교 열쇠: 앞뒤 공백을 빼고 대소문자를 무시한다. 엑셀 숫자 칸 12345.0 은 12345."""
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip().casefold()


def _label(value) -> str:
    if not part_key(value):
        return BLANK_PART
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


# --- 마크업표 읽기 ---

def load_markup(path) -> dict[str, MarkupRow]:
    """마크업표.xlsx 를 읽는다. 열쇠는 part_key(파트넘버). 잘못된 줄이 하나라도 있으면 MarkupError."""
    path = Path(path)
    try:
        data = path.read_bytes()
    except FileNotFoundError:
        raise MarkupError(f"{MARKUP_FILE} 가 없습니다: {path}") from None
    except OSError as e:
        raise MarkupError(f"{MARKUP_FILE} 를 읽지 못했습니다: {e}") from None
    return markup_from_bytes(data)


def markup_from_bytes(data: bytes) -> dict[str, MarkupRow]:
    """한 번 읽은 파일 바이트에서 마크업표를 만든다(실행 중 파일이 바뀌어도 지문과 같은 내용을 쓰게)."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception as e:
        raise MarkupError(f"{MARKUP_FILE} 를 열지 못했습니다(엑셀 통합 문서 .xlsx 인지 확인): {e}") from None
    try:
        rows = list(wb.worksheets[0].iter_rows())
    finally:
        wb.close()
    header = [str(c.value).strip() if c.value is not None else "" for c in (rows[0] if rows else [])]
    if PART_COLUMN not in header or RATE_COLUMN not in header:
        raise MarkupError(f"{MARKUP_FILE} 첫 줄(머리글)에 '{PART_COLUMN}' 와 '{RATE_COLUMN}' 열이 있어야 합니다.")
    part_at, rate_at = header.index(PART_COLUMN), header.index(RATE_COLUMN)
    places_at = header.index(PLACES_COLUMN) if PLACES_COLUMN in header else None

    table: dict[str, MarkupRow] = {}
    first_line: dict[str, int] = {}
    errors: list[str] = []
    for cells in rows[1:]:
        line = cells[0].row if cells else 0

        def cell(i):
            return cells[i] if i is not None and i < len(cells) else None

        part, rate, places = cell(part_at), cell(rate_at), cell(places_at)
        values = [c.value for c in (part, rate, places) if c is not None]
        if all(_empty(v) for v in values):
            continue  # 빈 줄
        where = f"{MARKUP_FILE} {line}번째 줄"
        key = part_key(part.value if part is not None else None)
        if not key:
            errors.append(f"{where}: 파트넘버가 비어 있습니다")
            continue
        label = _label(part.value)
        try:
            row = MarkupRow(label, _rate(rate), _places(places))
        except ValueError as e:
            errors.append(f"{where} ({label}): {e}")
            continue
        if key in table:
            errors.append(f"{where}: 파트넘버 '{label}' 가 {first_line[key]}번째 줄('{table[key].part}')에도 "
                          "있습니다(대소문자 · 앞뒤 공백은 같은 파트로 봄. 한 파트는 한 줄만)")
            continue
        table[key] = row
        first_line[key] = line
    if errors:
        shown = "; ".join(errors[:5]) + (f" 외 {len(errors) - 5}건" if len(errors) > 5 else "")
        raise MarkupError(f"{MARKUP_FILE} 를 고쳐야 합니다: {shown}")
    return table


def _empty(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _plain(number: Decimal) -> Decimal:
    """뒤의 0 을 뺀다(3.00 → 3, 2.50 → 2.5). 300 이 3E+2 로 바뀌지 않게."""
    number = number.normalize()
    return number.quantize(Decimal(1)) if number.as_tuple().exponent > 0 else number


def _rate(cell) -> Decimal:
    value = cell.value if cell is not None else None
    if _empty(value):
        raise ValueError(f"{RATE_COLUMN} 가 비어 있습니다(0% 면 0 을 적으세요)")
    if isinstance(value, bool):
        raise ValueError(f"{RATE_COLUMN} '{value}' 가 숫자가 아닙니다")
    if isinstance(value, (int, float, Decimal)):
        number = excel_number(value)
        if number is not None and "%" in (cell.number_format or ""):
            number = number * 100  # 백분율 서식 칸: 엑셀은 3% 를 0.03 으로 보관한다
    else:
        text = str(value).strip()
        text = text[:-1].strip() if text.endswith("%") else text
        try:
            number = Decimal(text.replace(",", ""))
        except InvalidOperation:
            number = None
    if number is None or not number.is_finite():
        raise ValueError(f"{RATE_COLUMN} '{value}' 가 숫자가 아닙니다")
    number = _plain(number)
    if not (0 <= number <= MAX_RATE):
        raise ValueError(f"{RATE_COLUMN} {number} 는 0 이상 {MAX_RATE} 이하여야 합니다")
    return number


def _places(cell) -> int | None:
    value = cell.value if cell is not None else None
    if _empty(value):
        return None
    number = None
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        number = excel_number(value)
    elif isinstance(value, str):
        try:
            number = Decimal(value.strip())
        except InvalidOperation:
            number = None
    if number is None or not number.is_finite() or number != number.to_integral_value() \
            or not (0 <= number <= MAX_PLACES):
        raise ValueError(f"{PLACES_COLUMN} '{value}' 는 0 부터 {MAX_PLACES} 까지의 정수여야 합니다(비우면 반올림하지 않음)")
    return int(number)


# --- 계산 ---

def _exact(compute):
    """정확한 Decimal 계산. 자리가 모자라 조용히 반올림될 일이 생기면 오류로 멈춘다."""
    with localcontext() as ctx:
        ctx.prec = 60
        ctx.traps[Inexact] = True
        return compute()


def _round(value: Decimal, places: int | None) -> Decimal:
    return value if places is None else value.quantize(Decimal(1).scaleb(-places), rounding=ROUND_HALF_UP)


def apply_markup(price: Decimal, row: MarkupRow) -> Decimal:
    """본사 단가 = price × (1 + rate/100). 반올림자리가 있으면 그 자리에서 사사오입, 없으면 정확한 값 그대로."""
    return _round(_exact(lambda: price * (1 + row.rate / 100)), row.places)


def _number(value) -> Decimal | None:
    """품목 칸 값을 Decimal 로(엑셀 float 은 유효 15자리). 숫자가 아니면 None."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        return excel_number(value)
    try:
        number = Decimal(str(value).replace(",", "").replace("$", "").strip())
    except InvalidOperation:
        return None
    return number if number.is_finite() else None


def lookup_row(item, columns: MarkupColumns, table) -> MarkupRow | None:
    return (table or {}).get(part_key((item or {}).get(columns.part))) or None


def markup_value(source: str, item, columns: MarkupColumns, table):
    """회사 양식 열 값 마크업.율 · 마크업.단가 · 마크업.금액. 계산할 수 없으면 None(검산이 확인필요로 보낸다)."""
    row = lookup_row(item, columns, table)
    if row is None:
        return None
    if source == RATE:
        return row.rate
    price = _number((item or {}).get(columns.price))
    if price is None:
        return None
    unit = apply_markup(price, row)
    if source == PRICE:
        return unit
    quantity = _number((item or {}).get(columns.quantity))
    return None if quantity is None else _exact(lambda: quantity * unit)  # 금액은 반올림하지 않는다


def missing_parts(result, form, table) -> list[str]:
    """회사 양식이 마크업을 쓰는데 마크업표에 없는 파트넘버(나온 순서, 한 번씩). 빈 품번은 (빈칸)."""
    if form is None or not form.uses_markup:
        return []
    missing = []
    for item in result.table or []:
        if lookup_row(item, form.markup, table) is None:
            label = _label((item or {}).get(form.markup.part))
            if label not in missing:
                missing.append(label)
    return missing


def written_markup(form, rows) -> list[dict]:
    """회사 양식 행(company_rows)에서 마크업 열에 쓴 값만 줄마다 {값 출처: 값} 으로."""
    sources = [c.source for c in form.columns]
    return [{s: v for s, v in zip(sources, row) if s in (RATE, PRICE, AMOUNT)} for row in rows]


# --- 검산 ---

def _show(value) -> str:
    if value is None:
        return ""
    return format(value, "f") if isinstance(value, Decimal) else str(value)


def markup_checks(result, rows_written, table, columns: MarkupColumns | None = None) -> list[Check]:
    """줄마다 마크업% · 본사 단가(· 쓴 경우 본사 금액)를 따로 계산해 쓴 값과 대조한다.

    본사 단가 = 고객 단가 × (100 + 마크업%) ÷ 100 (반올림자리가 있으면 그 자리에서 사사오입),
    본사 금액 = 수량 × 본사 단가(반올림 없음). 마크업% 는 0 이상 10 이하.
    """
    columns = columns or MarkupColumns()
    checks = []
    for n, item in enumerate(result.table or [], 1):
        written = rows_written[n - 1] if n - 1 < len(rows_written) else {}
        label = _label((item or {}).get(columns.part))
        row = lookup_row(item, columns, table)
        if row is None:
            checks.append(Check(f"{n}줄 마크업% ({label})", "마크업표에 없음", "", False))
            continue
        rate = written.get(RATE, row.rate)
        checks.append(Check(f"{n}줄 마크업% ({label})", f"마크업표 {_show(row.rate)}%", _show(rate),
                            0 <= row.rate <= MAX_RATE and rate == row.rate))

        price = _number((item or {}).get(columns.price))
        if price is None:
            checks.append(Check(f"{n}줄 본사 단가", f"고객 단가({columns.price}) 없음",
                                _show(written.get(PRICE)), False))
            continue
        exact = _exact(lambda: price * (100 + row.rate) / 100)
        expected = _round(exact, row.places)
        text = f"{_show(price)} x (100 + {_show(row.rate)}) / 100 = {_show(exact)}"
        if row.places is not None:
            text += f", 소수 {row.places}자리 반올림 {_show(expected)}"
        unit = written[PRICE] if PRICE in written else apply_markup(price, row)
        checks.append(Check(f"{n}줄 본사 단가", text, _show(unit), unit is not None and unit == expected))

        if AMOUNT in written:
            quantity = _number((item or {}).get(columns.quantity))
            if quantity is None:
                checks.append(Check(f"{n}줄 본사 금액", f"수량({columns.quantity}) 없음",
                                    _show(written[AMOUNT]), False))
                continue
            amount = _exact(lambda: quantity * expected)
            checks.append(Check(f"{n}줄 본사 금액", f"{_show(quantity)} x {_show(expected)} = {_show(amount)}",
                                _show(written[AMOUNT]), written[AMOUNT] is not None and written[AMOUNT] == amount))
    return checks


def markup_stamp(result, form, table) -> str:
    """이 문서가 쓰는 마크업표 줄들의 지문. 승인대기 대조표에 남겨, 친구가 본 뒤 그 파트의 마크업이 바뀌면
    대조표를 새로 만든다(다른 파트를 더하거나 고친 것은 상관없음). 승인 수(양식 지문)에는 넣지 않는다."""
    used = {}
    for item in result.table or []:
        row = lookup_row(item, form.markup, table)
        key = part_key((item or {}).get(form.markup.part))
        used[key] = "-" if row is None else f"{_show(row.rate)}|{row.places}"
    text = "\n".join(f"{k}\t{v}" for k, v in sorted(used.items()))
    return hashlib.sha256(text.encode("utf-8")).hexdigest()[:16]


# --- 예시 ---

def write_example_markup(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "마크업표"
    ws.append([PART_COLUMN, RATE_COLUMN, PLACES_COLUMN])
    for row in EXAMPLE_ROWS:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for letter, width in zip("ABC", (22, 12, 12)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    wb.save(path)
