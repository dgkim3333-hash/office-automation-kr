"""누적장부.xlsx — 열 동기화, 행 추가, 중복 판정, 잠김 감지."""
import os
import sys
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ITEMS_SHEET = "품목"
LOG_SHEET = "처리로그"
DOC_HEAD = ["처리일시", "원본파일"]
DOC_TAIL = ["정리된경로"]
ITEM_HEAD = ["문서종류", "문서번호"]
LOG_COLUMNS = ["일시", "파일", "결과", "사유"]

DATETIME_FORMAT = "yyyy-mm-dd hh:mm:ss"
DATE_FORMAT = "yyyy-mm-dd"
AMOUNT_FORMAT = "#,##0.00"


class LedgerLocked(Exception):
    """장부가 엑셀에서 열려 있어 쓸 수 없을 때."""


def _is_locked(path: Path) -> bool:
    """다른 프로그램(엑셀)이 파일을 열고 있으면 True."""
    if not path.exists():
        return False
    if sys.platform == "win32":
        import ctypes
        from ctypes import wintypes

        create = ctypes.windll.kernel32.CreateFileW
        create.restype = wintypes.HANDLE
        create.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID,
                           wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
        generic_read_write, share_none, open_existing = 0xC0000000, 0, 3
        handle = create(str(path), generic_read_write, share_none, None, open_existing, 0, None)
        if handle == wintypes.HANDLE(-1).value:
            return True
        ctypes.windll.kernel32.CloseHandle(handle)
        return False
    try:
        with open(path, "r+b"):
            return False
    except PermissionError:
        return True


class Ledger:
    def __init__(self, path, rules, wb):
        self.path = Path(path)
        self.rules = rules
        self.wb = wb

    @classmethod
    def open(cls, path, rules):
        path = Path(path)
        if _is_locked(path):
            raise LedgerLocked(str(path))
        wb = openpyxl.load_workbook(path) if path.exists() else _new_workbook()
        ledger = cls(path, rules, wb)
        ledger._sync_columns()
        return ledger

    # --- 열 동기화 ---

    def _sync_columns(self):
        for doc_type in self.rules.doc_types:
            if doc_type.name not in self.wb.sheetnames:
                self.wb.create_sheet(doc_type.name, index=self.wb.sheetnames.index(ITEMS_SHEET))
            _ensure_columns(self.wb[doc_type.name], DOC_HEAD + list(doc_type.fields) + DOC_TAIL)
        item_columns = list(ITEM_HEAD)
        for doc_type in self.rules.doc_types:
            for col in doc_type.table.columns if doc_type.table else []:
                if col not in item_columns:
                    item_columns.append(col)
        _ensure_columns(self.wb[ITEMS_SHEET], item_columns)
        _ensure_columns(self.wb[LOG_SHEET], LOG_COLUMNS)

    # --- 조회·추가·삭제 ---

    def is_duplicate(self, doc_type_name, number) -> bool:
        return bool(self._doc_rows(doc_type_name, number))

    def add(self, result, source_name, organized_path, processed_at=None):
        doc_type = result.doc_type
        values = {"처리일시": processed_at or datetime.now().replace(microsecond=0),
                  "원본파일": source_name, "정리된경로": str(organized_path), **result.fields}
        _append(self.wb[doc_type.name], values)
        for row in result.table or []:
            _append(self.wb[ITEMS_SHEET], {"문서종류": doc_type.name, "문서번호": result.number, **row})

    def remove(self, doc_type_name, number):
        ws = self.wb[doc_type_name]
        for r in reversed(self._doc_rows(doc_type_name, number)):
            ws.delete_rows(r)
        items = self.wb[ITEMS_SHEET]
        header = _header(items)
        kind_col, num_col = header.index("문서종류") + 1, header.index("문서번호") + 1
        for r in range(_last_row(items), 1, -1):
            if (items.cell(r, kind_col).value == doc_type_name
                    and str(items.cell(r, num_col).value) == str(number)):
                items.delete_rows(r)

    def add_log(self, when, file_name, outcome, reason):
        _append(self.wb[LOG_SHEET], {"일시": when, "파일": file_name, "결과": outcome, "사유": reason or None})

    def _doc_rows(self, doc_type_name, number):
        if doc_type_name not in self.wb.sheetnames:
            return []
        ws = self.wb[doc_type_name]
        col = _header(ws).index(self.rules.get(doc_type_name).number_field) + 1
        return [r for r in range(2, _last_row(ws) + 1) if str(ws.cell(r, col).value) == str(number)]

    # --- 저장 ---

    def save(self):
        for i, ws in enumerate(self.wb.worksheets):
            _refresh_table(ws, f"Ledger_{i + 1}")
        tmp = self.path.with_name(f".{self.path.stem}.tmp.xlsx")
        try:
            self.wb.save(tmp)
            os.replace(tmp, self.path)
        except PermissionError:
            raise LedgerLocked(str(self.path)) from None
        finally:
            if tmp.exists():
                tmp.unlink()


def _new_workbook():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.create_sheet(ITEMS_SHEET)
    wb.create_sheet(LOG_SHEET)
    return wb


def _header(ws):
    return [c.value for c in ws[1] if c.value is not None]


def _ensure_columns(ws, columns):
    """없는 열만 오른쪽 끝에 더한다. 기존 열은 순서·이름을 바꾸지 않는다."""
    header = _header(ws)
    for col in columns:
        if col not in header:
            header.append(col)
            ws.cell(1, len(header), col)
            ws.column_dimensions[get_column_letter(len(header))].width = max(12, len(str(col)) * 2 + 4)
    ws.freeze_panes = "A2"


def _last_row(ws):
    """마지막 데이터 행 번호(없으면 1 = 헤더)."""
    for r in range(ws.max_row, 1, -1):
        if any(c.value is not None for c in ws[r]):
            return r
    return 1


def _append(ws, values):
    header = _header(ws)
    r = _last_row(ws) + 1
    for name, value in values.items():
        if name not in header or value is None:
            continue
        cell = ws.cell(r, header.index(name) + 1, value)
        if isinstance(value, datetime):
            cell.number_format = DATETIME_FORMAT
        elif isinstance(value, date):
            cell.number_format = DATE_FORMAT
        elif isinstance(value, Decimal):
            cell.number_format = AMOUNT_FORMAT


def _refresh_table(ws, name):
    """엑셀 표(헤더 고정·필터)를 현재 열·행 범위로 다시 만든다."""
    for old in list(ws.tables):
        del ws.tables[old]
    width = len(_header(ws))
    last = max(_last_row(ws), 2)  # 데이터가 없어도 빈 행 하나를 둔다(엑셀 표 규칙)
    table = Table(displayName=name, ref=f"A1:{get_column_letter(width)}{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
