"""문서별 결과 엑셀(출력파일). 엑셀 PO → 회사 SO 양식, PDF → 추출결과."""
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font

from app.excel_extract import excel_number
from app.ledger import AMOUNT_FORMAT, DATE_FORMAT, DATETIME_FORMAT, fit_columns
from app.organize import free_name, render_text

LINE_STEP = 10  # 줄번호: 10, 20, 30 …


def _number(value) -> Decimal:
    """엑셀 칸에서 온 float 은 엑셀처럼 유효 15자리로(0.30000000000000004 → 0.3)."""
    return excel_number(value) if isinstance(value, float) else Decimal(str(value))


def _value(source, result, item, master_row, index):
    if source == "줄번호":
        return (index + 1) * LINE_STEP
    if source == "빈칸":
        return None
    if source == "수량x단가":
        qty, price = (item or {}).get("수량"), (item or {}).get("단가")
        return None if qty is None or price is None else _number(qty) * _number(price)
    if source.startswith("고정:"):
        return source[len("고정:"):]
    group, _, name = source.partition(".")
    if group == "머리":
        return result.fields.get(name)
    if group == "품목":
        return (item or {}).get(name)
    if group == "고객마스터":
        return (master_row or {}).get(name)
    return None


def company_rows(result, master_row, form):
    """회사 양식의 머리글과 행(품목 한 줄 = 한 행)."""
    header = [c.name for c in form.columns]
    items = result.table or []
    rows = [[_value(c.source, result, item, master_row, i) for c in form.columns] for i, item in enumerate(items)]
    return header, rows


def _write_sheet(ws, header, rows):
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in row])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, datetime):
                cell.number_format = DATETIME_FORMAT
            elif isinstance(cell.value, date):
                cell.number_format = DATE_FORMAT
            elif isinstance(cell.value, Decimal):
                cell.number_format = AMOUNT_FORMAT
    ws.freeze_panes = "A2"
    fit_columns(ws)


def _save(wb, folder, name) -> Path:
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    path = free_name(folder / name)
    wb.save(path)
    return path


def company_workbook(result, master_row, form):
    """회사 양식 통합문서(저장 전). 저장 위치 · 이름은 부르는 쪽이 정한다."""
    header, rows = company_rows(result, master_row, form)
    wb = openpyxl.Workbook()
    wb.active.title = form.sheet
    _write_sheet(wb.active, header, rows)
    return wb


def company_form_name(result, form) -> str:
    return render_text(form.filename_tpl, result.fields)


def write_company_form(folder, result, master_row, form) -> Path:
    return _save(company_workbook(result, master_row, form), folder, company_form_name(result, form))


def extract_result_rows(result, source_name):
    """추출결과 시트의 머리글과 행(필드 · 값 · 비고). 저장과 다시 읽기 검사가 같은 것을 쓴다."""
    rows = [["문서종류", result.doc_type.name, None], ["원본파일", source_name, None]]
    rows += [[name, value, None if value is not None else "못 찾음"] for name, value in result.fields.items()]
    return ["필드", "값", "비고"], rows


def write_extract_result(folder, result, source_name) -> Path:
    """PDF 한 건의 추출 결과: 추출결과 시트(필드·값·비고) + 품목 시트."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "추출결과"
    header, rows = extract_result_rows(result, source_name)
    _write_sheet(ws, header, rows)
    if result.doc_type.table is not None:
        columns = result.doc_type.table.columns
        _write_sheet(wb.create_sheet("품목"), columns, [[r.get(c) for c in columns] for r in result.table or []])
    return _save(wb, folder, render_text(f"추출결과_{{{result.doc_type.number_field}}}.xlsx", result.fields))


CHECK_HEADER = ["항목", "원본 값", "결과 값", "일치"]


def _mark(ok) -> str:
    return "해당 없음" if ok is None else "일치" if ok else "불일치"


def write_check_file(folder, result, checks, source_name) -> Path:
    """검산 기록: 검산_{번호}.xlsx. 맨 위 줄에 원본파일, 그 아래 항목별 대조 결과. 회사 양식 파일에는 쓰지 않는다."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "검산"
    ws.append(["원본파일", ILLEGAL_CHARACTERS_RE.sub("", source_name)])
    ws.append(CHECK_HEADER)
    for cell in ws[2]:
        cell.font = Font(bold=True)
    for c in checks:
        ws.append([ILLEGAL_CHARACTERS_RE.sub("", v) for v in (c.item, c.source, c.output)] + [_mark(c.ok)])
    ws.freeze_panes = "A3"
    fit_columns(ws)
    return _save(wb, folder, render_text(f"검산_{{{result.doc_type.number_field}}}.xlsx", result.fields))


# --- 저장 뒤 다시 읽기 ---

def read_back(path, sheet, header, rows) -> list[str]:
    """저장된 파일을 다시 열어 머리글과 각 칸이 쓴 값과 같은지 본다. 다른 칸의 설명 목록(같으면 빈 목록).

    한 파일 = 한 문서인 시트 전체를 비교한다. 누적장부처럼 여러 문서가 든 파일은 쓰지 않는다.
    """
    wb = openpyxl.load_workbook(Path(path), data_only=True)
    try:
        if sheet not in wb.sheetnames:
            return [f"시트 '{sheet}' 가 없습니다"]
        saved = [list(r) for r in wb[sheet].iter_rows(values_only=True)]
    finally:
        wb.close()
    while saved and all(_blank(v) for v in saved[-1]):
        saved.pop()
    problems = []
    got_header = saved[0] if saved else []
    for col in range(max(len(header), _used(got_header))):
        want = header[col] if col < len(header) else None
        got = got_header[col] if col < len(got_header) else None
        if not _same_value(want, got):
            problems.append(f"머리글 {col + 1}열: 쓴 값 {_text(want)}, 읽은 값 {_text(got)}")
    names = [h if h is not None else f"{i + 1}열" for i, h in enumerate(header)]
    for index, row in enumerate(rows):
        line = index + 2
        got_row = saved[index + 1] if index + 1 < len(saved) else None
        if got_row is None and all(_blank(v) for v in row):
            continue  # 빈 줄은 파일 끝에서 사라질 수 있다
        if got_row is None:
            problems.append(f"{line}행: 쓴 줄이 파일에 없습니다")
            continue
        for col in range(max(len(row), _used(got_row))):
            want = row[col] if col < len(row) else None
            got = got_row[col] if col < len(got_row) else None
            if not _same_value(want, got):
                name = names[col] if col < len(names) else f"{col + 1}열"
                problems.append(f"{line}행 {name}: 쓴 값 {_text(want)}, 읽은 값 {_text(got)}")
    for index in range(len(rows) + 1, len(saved)):
        problems.append(f"{index + 1}행: 쓰지 않은 줄이 파일에 있습니다")
    return problems


def _blank(value) -> bool:
    return value is None or value == ""


def _used(row) -> int:
    """뒤쪽 빈 칸을 뺀 칸 수."""
    n = len(row)
    while n and _blank(row[n - 1]):
        n -= 1
    return n


def _same_value(want, got) -> bool:
    if _blank(want) or _blank(got):
        return _blank(want) and _blank(got)
    if isinstance(want, str):
        return isinstance(got, str) and ILLEGAL_CHARACTERS_RE.sub("", want) == got
    if isinstance(want, bool) or isinstance(got, bool):
        return want is got
    if isinstance(want, datetime):
        return isinstance(got, datetime) and want == got
    if isinstance(want, date):  # 날짜 칸은 엑셀에서 자정 시각의 날짜+시각으로 돌아온다. 시각이 있으면 다른 값
        return (got == datetime.combine(want, time.min)) if isinstance(got, datetime) else got == want
    if isinstance(want, (Decimal, int, float)) and isinstance(got, (int, float, Decimal)):
        try:
            number, back = Decimal(str(want)), Decimal(str(got))
            if number == back:
                return True
            # 15자리를 넘는 소수는 엑셀이 15자리로 저장하므로 그때만 float 로 비교. 정수는 정확히 같아야 한다
            return number != number.to_integral_value() and len(number.as_tuple().digits) > 15                 and float(number) == float(back)
        except (InvalidOperation, ValueError, OverflowError):
            return False
    return want == got


def _text(value) -> str:
    if _blank(value):
        return "(빈칸)"
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, datetime):
        return value.isoformat(" ")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)
