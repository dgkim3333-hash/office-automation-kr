"""작업 하나 실행: 입력 폴더의 PDF · 엑셀을 읽어 결과(회사 양식 · 누적장부)를 만든다.

입력 폴더의 원본은 읽기만 한다(옮기기 · 이름 바꾸기 · 지우기 없음). 어떤 파일을 처리했는지는
작업 폴더의 처리기록.csv 에 파일 내용 지문(sha256)으로 남긴다. 이 파일은 줄을 더하기만 하며,
한 지문의 지금 상태는 그 지문의 마지막 줄이다.
"""
import csv
import hashlib
import os
import shutil
import sys
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from app.excel_extract import read_excel
from app.extract import extract
from app.forms import form_fingerprint, form_key
from app.jobs import Job, _folder_title, list_jobs, save_job
from app.ledger import ITEMS_SHEET, Ledger, LedgerLocked, is_locked
from app.local_guard import LocalOnlyError, check_work_dir
from app.log import append_log
from app.master import MasterError, load_master, lookup
from app.organize import free_name
from app.outputs import (
    _same_value, _text, company_form_name, company_rows, company_workbook, read_back, write_check_file,
)
from app.rules import RuleError, load_rules
from app.verify import failed, verify

LEDGER_NAME = "누적장부_{제목}.xlsx"  # 출력 폴더에 둔다
RECORD_FILE = "처리기록.csv"  # 작업 폴더
RECORD_COLUMNS = ["지문", "파일", "문서종류", "번호", "상태", "규칙지문", "날짜"]
LOG_FILE = "처리로그.csv"  # 작업 폴더. 아무도 안 볼 때(자동 실행)를 위한 건별 기록
LOCK_FILE = "실행중.lock"
MASTER_FILE = "고객마스터.xlsx"  # 작업 폴더들의 위(work) 에 있는 공용 파일
WAITING = "승인대기"
WAITING_CHECKS = "대조표"
REVIEW = "확인필요"
CHECKS = "검산"

INPUT_SUFFIXES = (".pdf", ".xlsx")  # .xls 는 판별하지 못하므로 건너뛴다
UNMATCHED_SUFFIXES = (".pdf", ".xlsx", ".xls")
SETTLE_SECONDS = 60  # 이보다 최근에 바뀐 파일은 아직 저장 중일 수 있어 다음 실행으로 미룬다
UNMATCHED_SECONDS = 24 * 3600
COUNT_KEYS = ("성공", "중복", "확인필요", "승인대기", "보류")
NEWS_KEYS = ("성공", "중복", "확인필요", "승인대기")
LEDGER_MODES = ("누적장부", "둘다")
COMPANY_MODES = ("회사양식", "둘다")
MASTER_PREFIX = "고객마스터."

RETRY_NOTE = "규칙.yaml 을 고치면 다음 실행 때 다시 시도합니다. 규칙을 그대로 두어도 다음 날 한 번 더 시도합니다."
HELD_REASON = "다른 프로그램(PDF 뷰어 · 엑셀 등)에서 열려 있음. 닫으면 다음 실행 때 처리합니다"
SETTLING_REASON = "방금 바뀐 파일(1분 안). 저장이 끝난 뒤 다음 실행 때 처리합니다"
LEDGER_LOCKED_REASON = "누적장부가 엑셀에서 열려 있음. 엑셀을 닫으면 다음 실행 때 처리합니다"
INTERRUPTED = "처리기록에 성공 기록이 없음. 지난 실행이 중간에 끊겼거나 되돌리기가 끝나지 못했을 수 있음"
KEEP_HINT = "입력 폴더에서 원본을 치우세요(그대로 두면 날마다 다시 확인필요가 됩니다)."
REMAKE_HINT = "그러면 내일 실행 때(또는 규칙.yaml 을 고친 뒤) 새로 만듭니다(처리기록은 고칠 필요 없음)."
STAMP_FORM = "form_fingerprint"  # 승인대기 대조표 속성: 만들 때의 양식 지문
STAMP_FILE = "file_fingerprint"  # 승인대기 대조표 속성: 만들 때의 복사본 지문
REMADE_RULES = "규칙이 바뀌어 대조표를 새로 만들었습니다. 다시 보고 승인하세요"
REMADE_MISSING = "대조표가 없어 새로 만들었습니다. 보고 승인하세요"
REMADE_FILE = "승인대기의 파일이 대조표를 만든 뒤 바뀌어 대조표를 새로 만들었습니다. 다시 보고 승인하세요"
APPROVAL_RETRY_NOTE = ("고친 뒤(규칙.yaml · 고객마스터 등) 다시 승인하면 다시 확인합니다. "
                       f"승인하지 않을 문서면 {WAITING} 폴더의 그 복사본을 지우세요.")


class JobLocked(Exception):
    """같은 작업이 이미 실행 중일 때(자동 실행 중에 지금 실행을 누른 경우 등)."""


@dataclass
class RunReport:
    counts: dict[str, int] = field(default_factory=lambda: dict.fromkeys(COUNT_KEYS, 0))
    unmatched_old: int = 0
    message: str = ""
    output_files: list[Path] = field(default_factory=list)
    ledger_locked: bool = False  # 누적장부가 엑셀에서 열려 있어 보류한 파일이 있음
    stopped: bool = False        # 처리기록을 쓰지 못해 중간에 멈춤(message 맨 앞에 「중요:」)

    @property
    def has_news(self) -> bool:
        """알릴 것이 있나. 보통 보류만 있으면 알리지 않지만(다음 실행 때 처리됨), 누적장부가 열려 있어
        보류한 것은 사람이 엑셀을 닫아야 풀리므로 알린다."""
        return (any(self.counts.get(k, 0) > 0 for k in NEWS_KEYS) or self.unmatched_old > 0
                or self.ledger_locked or self.stopped)


# --- 같은 작업 두 번 실행 막기 ---

def _lock(f):
    f.seek(0)
    if sys.platform == "win32":
        import msvcrt

        msvcrt.locking(f.fileno(), msvcrt.LK_NBLCK, 1)
    else:
        import fcntl

        fcntl.flock(f.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)


def _unlock(f):
    f.seek(0)
    if sys.platform == "win32":
        import msvcrt

        msvcrt.locking(f.fileno(), msvcrt.LK_UNLCK, 1)
    else:
        import fcntl

        fcntl.flock(f.fileno(), fcntl.LOCK_UN)


@contextmanager
def job_lock(job: Job):
    """작업 폴더의 실행중.lock 을 잡는다. 이미 누가 잡고 있으면 JobLocked."""
    f = open(Path(job.folder) / LOCK_FILE, "a+b")
    try:
        try:
            _lock(f)
        except OSError:
            raise JobLocked(f"이 작업이 이미 실행 중입니다({job.name}). 끝난 뒤 다시 실행하세요.") from None
        try:
            yield
        finally:
            try:
                _unlock(f)
            except OSError:
                pass  # 파일을 닫으면 잠금도 풀린다
    finally:
        f.close()


# --- 지문 · 처리기록 ---

def file_fingerprint(path) -> str:
    """파일 내용의 sha256(16진수). 이름이 바뀌어도 같은 파일이면 같은 값."""
    digest = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def rules_fingerprint(job: Job) -> str:
    """규칙.yaml 바이트의 sha256 앞 16자. 규칙이 바뀌었는지 볼 때 쓴다."""
    return hashlib.sha256(Path(job.rules_path).read_bytes()).hexdigest()[:16]


def load_records(job: Job) -> list[dict]:
    path = Path(job.folder) / RECORD_FILE
    if not path.exists():
        return []
    with open(path, encoding="utf-8-sig", newline="") as f:
        return [dict(row) for row in csv.DictReader(f)]


def _append_record(job: Job, row: dict):
    path = Path(job.folder) / RECORD_FILE
    new = not path.exists() or path.stat().st_size == 0
    # 새 파일만 utf-8-sig(BOM 포함), 이어 쓸 때는 BOM 없이 utf-8
    with open(path, "a", encoding="utf-8-sig" if new else "utf-8", newline="") as f:
        writer = csv.writer(f)
        if new:
            writer.writerow(RECORD_COLUMNS)
        writer.writerow([row.get(c, "") for c in RECORD_COLUMNS])


def _current_state(records) -> dict[str, dict]:
    """지문 → 그 지문의 마지막 줄."""
    state = {}
    for row in records:
        state[row.get("지문", "")] = row
    return state


# --- 바깥에서 부르는 것 ---

def run_job(job: Job, work: Path, *, now=None, source_dir: Path | None = None, gate=None,
            on_success=None, after=None) -> RunReport:
    """작업 하나를 실행한다. source_dir 기본 = 입력 폴더(승인 실행은 작업 폴더의 승인대기\\ 를 넘긴다).

    gate(job, rules, result) 가 True 면 결과를 내지 않고 승인대기로 보낸다.
    on_success(job, result) 는 출력까지 끝난 건마다 부른다.
    after(job, report) 는 실행이 끝난 뒤 잠금을 쥔 채 부른다(승인대기 정리처럼 다른 실행과 겹치면 안 되는 일).
    """
    now = (now or datetime.now()).replace(microsecond=0)
    source = Path(source_dir) if source_dir is not None else Path(job.input_dir)
    with job_lock(job):
        report = _Run(job, Path(work), now, source, gate, on_success).run()
        if after is not None:
            try:
                after(job, report)
            except Exception as e:  # 정리가 실패해도 이번 실행 결과(보고)는 잃지 않는다
                report.message += f"\n경고: 실행 뒤 정리를 하지 못함 ({e})"
        return report


def check_file_info(path) -> tuple[str | None, dict]:
    """대조표의 (맨 위 원본파일 칸 = 승인대기 복사본 이름, 지문 속성). 읽지 못하면 (None, {})."""
    try:
        wb = openpyxl.load_workbook(Path(path), read_only=True)
    except Exception:
        return None, {}
    try:
        stamps = {p.name: str(p.value) for p in wb.custom_doc_props.props}
        row = next(wb.worksheets[0].iter_rows(min_row=1, max_row=1, values_only=True), ())
        name = str(row[1]) if len(row) >= 2 and row[0] == "원본파일" and row[1] else None
        return name, stamps
    except Exception:
        return None, {}
    finally:
        wb.close()


def waiting_checks(job: Job) -> dict[str, list[tuple[Path, dict]]]:
    """승인대기 대조표들: 복사본 이름(대소문자 무시) → [(대조표 경로, 지문 속성)]."""
    folder = Path(job.folder) / WAITING / WAITING_CHECKS
    index: dict[str, list[tuple[Path, dict]]] = {}
    if not folder.is_dir():
        return index
    for check in sorted(folder.glob("*.xlsx")):
        if check.name.startswith("~$"):
            continue
        name, stamps = check_file_info(check)
        if name is not None:
            index.setdefault(name.casefold(), []).append((check, stamps))
    return index


def unmatched_old(job: Job, work: Path, now: datetime) -> int:
    """입력 폴더에 하루 넘게 남아 있고, 같은 입력 폴더를 쓰는 어느 작업의 처리기록에도 없는 파일 수."""
    folder = Path(job.input_dir)
    if not folder.is_dir():
        return 0
    jobs = [j for j in list_jobs(work) if _same_dir(j.input_dir, folder)]
    if not any(_same_dir(j.folder, job.folder) for j in jobs):
        jobs.append(job)
    known = set()
    for j in jobs:
        try:
            known |= {r.get("지문") for r in load_records(j)}
        except (OSError, UnicodeDecodeError, csv.Error):
            pass
    count = 0
    for path in folder.iterdir():
        if not (path.is_file() and path.suffix.lower() in UNMATCHED_SUFFIXES and not path.name.startswith("~$")):
            continue
        try:
            if now.timestamp() - path.stat().st_mtime <= UNMATCHED_SECONDS:
                continue
            if file_fingerprint(path) not in known:
                count += 1
        except OSError:
            continue
    return count


def _documents(folder) -> list[Path]:
    """폴더의 처리 대상 파일(.pdf · .xlsx, 엑셀 임시 파일 ~$ 제외). 폴더가 없으면 빈 목록."""
    folder = Path(folder)
    if not folder.is_dir():
        return []
    return sorted(p for p in folder.iterdir()
                  if p.is_file() and p.suffix.lower() in INPUT_SUFFIXES and not p.name.startswith("~$"))


def _same_dir(a, b) -> bool:
    return os.path.normcase(os.path.abspath(a)) == os.path.normcase(os.path.abspath(b))


def ledger_path(job: Job) -> Path:
    return Path(job.output_dir) / LEDGER_NAME.format(제목=_folder_title(job.title))


# --- 저장한 누적장부 다시 읽기 ---

def _sheet_rows(ws):
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    return (rows[0] if rows else []), rows[1:]


def _row_problems(where, header, row, want: dict) -> list[str]:
    problems = []
    for name, value in want.items():
        if name not in header:
            continue  # 장부에 열이 없는 값은 쓰지 않았다
        got = row[header.index(name)] if header.index(name) < len(row) else None
        if isinstance(value, str):
            value = ILLEGAL_CHARACTERS_RE.sub("", value)
        if not _same_value(value, got):
            problems.append(f"{where} {name}: 쓴 값 {_text(value)}, 읽은 값 {_text(got)}")
    return problems


def _ledger_problems(path, result, doc_values: dict) -> list[str]:
    """저장한 누적장부에서 이번 문서의 줄(문서종류 시트 · 품목 시트)을 찾아 쓴 값과 같은지 본다."""
    doc = result.doc_type
    number = str(result.number)
    wb = openpyxl.load_workbook(Path(path), data_only=True)
    try:
        if doc.name not in wb.sheetnames or ITEMS_SHEET not in wb.sheetnames:
            return [f"누적장부에 '{doc.name}' 또는 '{ITEMS_SHEET}' 시트가 없습니다"]
        header, rows = _sheet_rows(wb[doc.name])
        if doc.number_field not in header:
            return [f"누적장부 {doc.name} 시트에 '{doc.number_field}' 열이 없습니다"]
        col = header.index(doc.number_field)
        mine = [r for r in rows if col < len(r) and r[col] is not None and str(r[col]) == number]
        if len(mine) != 1:
            return [f"누적장부 {doc.name} 시트에 번호 {number} 줄이 {len(mine)}개입니다"]
        problems = _row_problems(f"누적장부 {doc.name}", header, mine[0], doc_values)

        header, rows = _sheet_rows(wb[ITEMS_SHEET])
        kind, num = header.index("문서종류"), header.index("문서번호")
        mine = [r for r in rows if r[kind] == doc.name and r[num] is not None and str(r[num]) == number]
        want_items = [{"문서종류": doc.name, "문서번호": result.number, **row} for row in result.table or []]
        if len(mine) != len(want_items):
            return problems + [f"누적장부 품목 시트에 {number} 줄이 {len(mine)}개입니다(쓴 줄 {len(want_items)}개)"]
        for i, (got, want) in enumerate(zip(mine, want_items)):
            problems += _row_problems(f"누적장부 품목 {i + 1}줄", header, got, want)
        return problems
    finally:
        wb.close()


# --- 실행 한 번 ---

class _StopRun(Exception):
    """처리기록을 쓰지 못해 실행을 멈춰야 할 때. 한 건 오류 처리(except Exception)가 삼키지 않게 다시 던진다."""


class _Run:
    def __init__(self, job, work, now, source, gate, on_success):
        self.job = job
        self.work = work
        self.now = now
        self.today = now.date().isoformat()
        self.source = source
        self.gate = gate
        self.on_success = on_success
        self.report = RunReport()
        self.warnings: list[str] = []
        self.stop_reason = ""
        self.ledger = None
        self.ledger_locked = False
        self.ledger_held = 0  # 누적장부가 열려 있어 보류한 파일 수
        self.master = None
        # 승인 실행(승인대기\ 를 읽음)이면 승인대기 기록을 건너뛰지 않는다
        self.approving = _same_dir(source, Path(job.folder) / WAITING)
        self.waiting_fps: set[str] = set()  # 승인대기에 있는 복사본들의 지문
        self.waiting_unknown = False         # 지문을 읽지 못한 복사본이 있음
        self.checks_index = None             # 승인대기 대조표(waiting_checks), 처음 쓸 때 읽음
        self.remade: list[str] = []          # 승인 실행에서 대조표를 새로 만든 문서

    # 준비

    def _preflight(self) -> str | None:
        job = self.job
        for label, path in (("입력", self.source), ("출력", job.output_dir)):
            try:
                check_work_dir(path)
            except LocalOnlyError as e:
                return f"{label} 폴더를 쓸 수 없습니다. {e}"
            if not Path(path).is_dir():
                return (f"{label} 폴더가 없습니다: {path}\n"
                        f"폴더를 다시 만들거나(외장 디스크면 연결), Code 탭에서 「{job.title} 작업 {label} 폴더 바꿔 줘」라고 말하세요.")
        if is_locked(Path(job.folder) / RECORD_FILE):
            return (f"{RECORD_FILE} 가 다른 프로그램(엑셀 등)에서 열려 있습니다. 닫고 다시 실행하세요.\n"
                    f"파일 위치: {Path(job.folder) / RECORD_FILE}")
        return None

    def _stop(self, reason: str) -> RunReport:
        """시작 전에 멈춤. 매일 자동 실행이 조용히 건너뛰지 않도록 알림 대상(stopped)으로 돌려준다."""
        return RunReport(message=f"중요: {self.job.title} 작업을 실행하지 못했습니다(아무 파일도 처리하지 않음).\n{reason}",
                         stopped=True)

    def run(self) -> RunReport:
        stop = self._preflight()
        if stop:
            return self._stop(stop)
        try:
            self.rules = load_rules(self.job.rules_path)
            self.rules_fp = rules_fingerprint(self.job)
        except (RuleError, OSError) as e:
            return self._stop(f"규칙.yaml 오류: {e}\nCode 탭에서 「{self.job.title} 작업 규칙 고쳐 줘」라고 말하세요.")
        if self.job.output_mode in COMPANY_MODES and self.rules.company_form is None:
            return self._stop(
                f"출력 방식이 '{self.job.output_mode}' 인데 규칙.yaml 에 회사양식이 없습니다. "
                f"Code 탭에서 「{self.job.title} 작업 규칙 고쳐 줘」라고 말해 회사양식을 적거나, 출력 방식을 '누적장부' 로 바꾸세요.")
        try:
            self.state = _current_state(load_records(self.job))
        except (OSError, UnicodeDecodeError, csv.Error) as e:
            return self._stop(
                f"{RECORD_FILE} 를 읽지 못했습니다: {e}\n"
                f"이 파일이 열려 있으면 닫고, 엑셀 등으로 고쳐 저장했다면 원래대로 되돌린 뒤 다시 실행하세요: "
                f"{Path(self.job.folder) / RECORD_FILE}")
        self.success_keys = {(r.get("문서종류"), r.get("번호")) for r in self.state.values() if r.get("상태") == "성공"}
        if not self.approving:
            self._scan_waiting()

        try:
            candidates = self._candidates()
        except OSError as e:
            candidates = []
            self.warnings.append(f"경고: 입력 폴더를 읽지 못함 ({self.source}): {e}")
        try:
            for path in candidates:
                self._one(path)
        except _StopRun as e:
            self.stop_reason = str(e)
            self.report.stopped = True

        if not self.approving:  # 승인은 입력 폴더를 보지 않으므로 그날의 자동 실행(놓친 실행 포함)을 대신하지 않는다
            self.job.last_run = self.today
            try:
                save_job(self.job)
            except OSError as e:
                self.warnings.append(f"경고: 작업설정.yaml 에 마지막 실행 날짜를 쓰지 못함 ({e})")
        try:
            self.report.unmatched_old = unmatched_old(self.job, self.work, self.now)
        except Exception as e:  # 마지막 집계가 실패해도 이번 실행 결과(보고)는 잃지 않는다
            self.warnings.append(f"경고: 입력 폴더에 남은 파일을 세지 못함 ({e})")
        self.report.ledger_locked = self.ledger_locked
        self.report.message = self._message()
        return self.report

    def _candidates(self):
        return _documents(self.source)

    def _scan_waiting(self):
        """승인대기에 있는 복사본들의 지문. 승인하지 않은 문서(복사본을 지움)인지 가릴 때 쓴다."""
        try:
            copies = _documents(Path(self.job.folder) / WAITING)
        except OSError:
            self.waiting_unknown = True
            return
        for path in copies:
            try:
                self.waiting_fps.add(file_fingerprint(path))
            except OSError:
                self.waiting_unknown = True  # 모르면 다시 처리하지 않는다(_2 복사본을 만들지 않게)

    # 한 건

    def _one(self, path: Path):
        try:
            age = self.now.timestamp() - path.stat().st_mtime
        except OSError:
            return
        if age < SETTLE_SECONDS:
            self._done(path, "보류", SETTLING_REASON)
            return
        try:
            fp = file_fingerprint(path)
        except OSError:
            self._done(path, "보류", HELD_REASON)
            return
        if self._skip(fp):  # 이미 처리한 원본은 PDF 뷰어에 열려 있어도 보류로 세지 않는다
            return
        if os.access(path, os.W_OK) and is_locked(path):  # 읽기 전용 파일은 잠김 검사가 늘 실패하므로 건너뜀
            self._done(path, "보류", HELD_REASON)
            return
        try:
            result = read_excel(path, self.rules) if path.suffix.lower() == ".xlsx" else extract(path, self.rules)
        except Exception as e:  # 깨진 파일 한 건이 나머지를 막지 않게
            self._review(path, fp, None, f"파일을 읽지 못함: {e}")
            return
        if result.doc_type is None and result.has_text:
            return  # 판별 실패: 다른 작업 몫일 수 있으므로 기록하지 않고 그대로 둔다
        if result.review_reason:
            self._review(path, fp, result, result.review_reason)
            return
        key = (result.doc_type.name, str(result.number))
        if key in self.success_keys:
            self._record(path, fp, result, "중복")
            self._clear_note(path)
            self._done(path, "중복", f"{key[0]} {key[1]} 이미 처리됨")
            return
        try:
            self._process(path, fp, result, key)
        except _StopRun:
            raise
        except Exception as e:  # 한 건의 오류가 나머지를 막지 않게
            self._review(path, fp, result, f"처리 중 오류: {e}")

    def _skip(self, fp) -> bool:
        row = self.state.get(fp)
        if row is None:
            return False
        status = row.get("상태")
        if status in ("성공", "중복"):
            return True
        if self.approving:
            return False  # 승인은 사람이 누른 실행이므로 승인대기 · 확인필요도 다시 확인한다
        if status == "승인대기":
            # 승인하지 않은 문서(친구가 복사본을 지움)는 규칙을 고친 뒤 다시 처리한다(설계 12-1).
            # 복사본이 아직 승인대기에 있으면 규칙이 바뀌어도 건너뛴다(승인 때 대조표를 새로 만든다)
            return (row.get("규칙지문") == self.rules_fp or self.waiting_unknown or fp in self.waiting_fps)
        if status == "확인필요":
            return row.get("날짜") == self.today and row.get("규칙지문") == self.rules_fp
        return False  # 보류: 다음 실행 때 다시

    def _process(self, path, fp, result, key):
        checks = verify(result, path)
        bad = failed(checks)
        if bad:
            self._review(path, fp, result, "검산 불일치: " + "; ".join(
                f"{c.item} (원본 {c.source}, 결과 {c.output})" for c in bad))
            return
        master_row = None
        form = self.rules.company_form
        if self.job.output_mode in COMPANY_MODES and any(c.source.startswith(MASTER_PREFIX) for c in form.columns):
            customer = result.fields.get("고객명")
            try:
                master_row = lookup(self._master(), customer)
            except MasterError as e:
                self._review(path, fp, result, str(e))
                return
            if master_row is None:
                self._review(path, fp, result, (
                    f"고객마스터에 '{customer}' 가 없습니다. 고객마스터.xlsx 에 이 고객 줄을 더한 뒤 다시 넣으세요"))
                return
        if self.approving:
            reason = self._unseen(path, fp, result)
            if reason:  # 친구가 대조표에서 본 값이 아니므로 내보내지 않는다
                self._remake_check(path, fp, result, checks, reason)
                return
        if self.gate is not None and self.gate(self.job, self.rules, result):
            self._to_waiting(path, fp, result, checks)
            return
        self._write(path, fp, result, key, checks, master_row)

    # 승인대기 대조표

    def _checks(self) -> dict[str, list[tuple[Path, dict]]]:
        if self.checks_index is None:
            self.checks_index = waiting_checks(self.job)
        return self.checks_index

    def _stamps(self, fp, result) -> dict:
        return {STAMP_FORM: form_fingerprint(self.job.rules_path, form_key(result)), STAMP_FILE: fp}

    def _unseen(self, path, fp, result) -> str | None:
        """이 복사본의 대조표가 지금 값과 같은 규칙 · 같은 파일로 만든 것인지. 아니면 그 사유."""
        entries = self._checks().get(path.name.casefold(), [])
        if not entries:
            return REMADE_MISSING
        now = self._stamps(fp, result)
        if any(stamps == now for _, stamps in entries):
            return None
        if any(stamps.get(STAMP_FILE) == fp for _, stamps in entries):
            return REMADE_RULES
        return REMADE_FILE

    def _write_waiting_check(self, copy: Path, fp, result, checks):
        """승인대기 대조표를 (다시) 만든다. 같은 복사본 이름을 가리키던 옛 대조표는 지운다."""
        name = copy.name.casefold()
        for old, _ in self._checks().get(name, []):
            try:
                old.unlink(missing_ok=True)
            except OSError as e:
                self.warnings.append(f"경고: 옛 대조표를 지우지 못함 ({old.name}): {e}")
        stamps = self._stamps(fp, result)
        # 대조표의 원본파일 칸 = 복사본 이름. 승인이 이 이름으로 대조표와 복사본을 짝짓는다
        check = write_check_file(copy.parent / WAITING_CHECKS, result, checks, copy.name, stamps)
        self._checks()[name] = [(check, stamps)]

    def _remake_check(self, path, fp, result, checks, reason):
        try:
            self._write_waiting_check(path, fp, result, checks)
        except Exception as e:
            self._review(path, fp, result, f"대조표를 새로 만들지 못함: {e}")
            return
        self._clear_note(path)
        self.remade.append(f"{path.name}: {reason}")
        self._done(path, "승인대기", reason)

    def _master(self):
        if self.master is None:
            self.master = load_master(self.work / MASTER_FILE)
        return self.master

    def _to_waiting(self, path, fp, result, checks):
        waiting = Path(self.job.folder) / WAITING
        waiting.mkdir(parents=True, exist_ok=True)
        copy = free_name(waiting / path.name)
        try:
            shutil.copy2(path, copy)
            self._write_waiting_check(copy, fp, result, checks)
        except Exception as e:
            if copy.exists():
                copy.unlink()
            self._review(path, fp, result, f"승인대기로 보내지 못함: {e}")
            return
        self._record(path, fp, result, "승인대기")
        self._clear_note(path)
        self._done(path, "승인대기", "")

    def _open_ledger(self):
        if self.ledger is None:
            self.ledger = Ledger.open(ledger_path(self.job), self.rules)
        return self.ledger

    def _hold_for_ledger(self, path, fp, result):
        """누적장부가 엑셀에서 열려 있어 보류. 판별된 파일이라는 표시로 처리기록에 보류를 남긴다
        (「어느 작업에도 맞지 않는 파일」 로 세지 않게). 보류 기록은 건너뛰지 않으므로 다음 실행 때 다시 처리한다."""
        self.ledger_held += 1
        if not self.approving:
            # 승인 실행에서는 승인대기 상태를 그대로 둔다. 보류로 덮으면 다음 자동 실행이 입력 폴더의 원본을
            # 다시 처리해 승인대기에 복사본을 하나 더 만든다. (원본 지문의 승인대기 기록이 「맞지 않는 파일」 도 막는다)
            self._record(path, fp, result, "보류")
        self._done(path, "보류", LEDGER_LOCKED_REASON)

    def _write(self, path, fp, result, key, checks, master_row):
        """출력: 누적장부 · 회사 양식. 저장한 뒤 다시 읽어 다르면 이번 건만 되돌리고 확인필요.

        여기까지 왔으면 처리기록에는 이 (문서종류, 번호)의 성공이 없다. 그런데 결과가 이미 있으면
        지난 실행이 결과를 쓴 뒤 기록 전에 끊긴 것일 수 있으므로, 다시 쓰지 않고 확인필요로 보낸다.
        """
        mode = self.job.output_mode
        form = self.rules.company_form
        final = None
        if mode in COMPANY_MODES:
            final = Path(self.job.output_dir) / company_form_name(result, form)
            if final.exists():
                also = " 누적장부에 이 문서 줄이 있으면 그것도 지우세요." if mode in LEDGER_MODES else ""
                self._review(path, fp, result, (
                    f"출력 폴더에 회사 양식 {final.name} 이 이미 있는데 {INTERRUPTED}. 두 번 입력하지 않도록 새로 "
                    f"만들지 않았습니다. 그 파일이 맞으면 {KEEP_HINT} 다시 만들려면 {final.name} 을 지우세요.{also} "
                    f"{REMAKE_HINT}"))
                return
        if mode in LEDGER_MODES and self.ledger_locked:
            self._hold_for_ledger(path, fp, result)
            return
        ledger, added, saved, company = None, False, False, None
        problems = []
        try:
            if mode in LEDGER_MODES:
                ledger = self._open_ledger()
                if ledger.is_duplicate(*key):
                    also = " 출력 폴더에 이 문서의 회사 양식이 있으면 그것도 지우세요." if mode in COMPANY_MODES else ""
                    self._review(path, fp, result, (
                        f"누적장부에는 {key[0]} {key[1]} 이 있는데 {INTERRUPTED}. 결과가 맞으면 {KEEP_HINT} "
                        f"다시 만들려면 누적장부에서 이 문서 줄(문서 시트와 품목 시트)을 지우세요.{also} {REMAKE_HINT}"))
                    return
                origin_name, origin_path = self._origin(path, fp)
                doc_values = {"처리일시": self.now, "원본파일": origin_name, "정리된경로": str(origin_path),
                              **result.fields}
                ledger.add(result, origin_name, origin_path, processed_at=self.now)
                added = True
                ledger.save()
                saved = True
                try:
                    problems += _ledger_problems(ledger.path, result, doc_values)
                except Exception as e:
                    problems.append(f"저장한 누적장부를 다시 읽지 못함: {e}")
            if mode in COMPANY_MODES and not problems:
                company, more = self._write_company(final, result, master_row, form)
                problems += more
        except _StopRun:
            raise
        except LedgerLocked:
            self.ledger_locked = True
            self._rollback(ledger, added, saved, key, company)
            self._hold_for_ledger(path, fp, result)
            return
        except Exception as e:
            undo = self._rollback(ledger, added, saved, key, company)
            self._review(path, fp, result, f"결과를 쓰는 중 오류: {e}{undo}")
            return
        if problems:
            undo = self._rollback(ledger, added, saved, key, company)
            shown = "; ".join(problems[:5]) + (f" 외 {len(problems) - 5}건" if len(problems) > 5 else "")
            self._review(path, fp, result, f"저장한 파일을 다시 읽어 보니 쓴 값과 다릅니다: {shown}{undo}")
            return

        try:
            write_check_file(Path(self.job.folder) / CHECKS, result, checks, self._origin(path, fp)[0])
        except Exception as e:  # 결과는 이미 나갔으므로 알리기만 한다
            self.warnings.append(f"경고: 검산 파일을 만들지 못함 ({path.name}): {e}")
        if company is not None:
            self.report.output_files.append(company)
        if ledger is not None and ledger.path not in self.report.output_files:
            self.report.output_files.append(ledger.path)
        self._record(path, fp, result, "성공")  # 못 쓰면 _StopRun 으로 실행을 멈춘다
        self.success_keys.add(key)
        self._clear_note(path)
        self._done(path, "성공", "; ".join(result.warnings))
        if self.on_success is not None:
            try:
                self.on_success(self.job, result)
            except Exception as e:
                self.warnings.append(f"경고: 성공 처리 뒤 작업을 하지 못함 ({path.name}): {e}")

    def _origin(self, path, fp) -> tuple[str, Path]:
        """장부 · 검산 파일에 적을 원본 (이름, 경로). 승인 실행이면 승인대기 복사본이 아니라 입력 폴더의 원본:
        처리기록의 그 지문 승인대기 줄의 파일 이름(없으면 복사본 이름)과 입력 폴더 경로."""
        if not self.approving:
            return path.name, path
        row = self.state.get(fp) or {}
        name = row.get("파일") if row.get("상태") == "승인대기" and row.get("파일") else path.name
        return name, Path(self.job.input_dir) / name

    def _write_company(self, final, result, master_row, form):
        """회사 양식을 임시 이름으로 저장 → 다시 읽어 대조 → 같을 때만 최종 이름으로 바꾼다.

        다시 읽기를 통과하지 않은 파일이 최종 이름으로 보이는 일이 없다. (만든 파일 또는 None, 문제 목록)
        """
        tmp = final.with_name(f".{final.stem}.tmp{final.suffix}")
        problems = []
        try:
            company_workbook(result, master_row, form).save(tmp)
            header, rows = company_rows(result, master_row, form)
            try:
                problems += read_back(tmp, form.sheet, header, rows)
            except Exception as e:
                problems.append(f"저장한 회사 양식 파일을 다시 읽지 못함: {e}")
            if problems:
                return None, problems
            if final.exists():
                return None, [f"저장하는 사이 출력 폴더에 {final.name} 이 생겼습니다"]
            os.replace(tmp, final)
            return final, []
        finally:
            if tmp.exists():
                try:
                    tmp.unlink()
                except OSError:
                    pass

    def _rollback(self, ledger, added, saved, key, company) -> str:
        """이번 건이 쓴 것을 되돌린다. 되돌리지 못한 것이 있으면 사유에 붙일 글자를 돌려준다."""
        notes = []
        if added:
            ledger.remove(*key)
            if saved:
                try:
                    ledger.save(backup=False)  # .bak 에는 마지막 좋은 상태를 남긴다
                except Exception as e:
                    notes.append(f"누적장부에서 이 문서 줄을 지우지 못함({e}). 누적장부를 열어 직접 지우세요")
        if company is not None and company.exists():
            try:
                company.unlink()
            except OSError as e:
                notes.append(f"회사 양식 파일을 지우지 못함({company.name}: {e}). 직접 지우세요")
        return "".join(f". {n}" for n in notes)

    # 기록

    def _record(self, path, fp, result, status):
        row = {
            "지문": fp,
            "파일": path.name,
            "문서종류": result.doc_type.name if result is not None and result.doc_type else "",
            "번호": str(result.number) if result is not None and result.number is not None else "",
            "상태": status,
            "규칙지문": self.rules_fp,
            "날짜": self.today,
        }
        try:
            _append_record(self.job, row)
        except OSError as e:
            # 기록 없이 계속하면 다음 실행이 같은 문서를 또 내보낼 수 있다(주문 두 번 입력). 여기서 멈춘다
            raise _StopRun(f"{RECORD_FILE} 에 쓰지 못해 실행을 멈췄습니다({path.name}, {status}: {e}). "
                           f"이 파일의 결과는 이미 만들어졌을 수 있으니 출력 폴더를 확인하세요. "
                           f"남은 파일은 처리하지 않았습니다. {RECORD_FILE} 를 연 프로그램을 닫고 다시 실행하세요.") from e
        self.state[fp] = row

    def _note_path(self, path) -> Path:
        return Path(self.job.folder) / REVIEW / f"{path.name}.txt"

    def _clear_note(self, path):
        """이 파일이 지난번에 남긴 확인필요 사유 파일을 지운다(이제 처리됐으므로)."""
        try:
            self._note_path(path).unlink(missing_ok=True)
        except OSError as e:
            self.warnings.append(f"경고: 지난 확인필요 사유 파일을 지우지 못함 ({path.name}): {e}")

    def _review(self, path, fp, result, reason):
        note = self._note_path(path)
        try:
            note.parent.mkdir(parents=True, exist_ok=True)
            hint = APPROVAL_RETRY_NOTE if self.approving else RETRY_NOTE
            note.write_text(f"원본: {path}\n사유: {reason}\n\n{hint}\n", encoding="utf-8")
        except OSError as e:
            self.warnings.append(f"경고: 확인필요 사유 파일을 쓰지 못함 ({path.name}): {e}")
        if not self.approving:
            # 승인 실행에서는 승인대기 상태를 그대로 둔다. 확인필요로 덮으면 다음 날 자동 실행이 입력 폴더의 원본을
            # 다시 처리해 승인대기에 복사본을 하나 더(_2) 만든다. 복사본은 승인대기에 남아 고친 뒤 다시 승인할 수 있다
            self._record(path, fp, result, "확인필요")
        self._done(path, "확인필요", reason)

    def _done(self, path, outcome, reason):
        self.report.counts[outcome] += 1
        try:
            append_log(Path(self.job.folder) / LOG_FILE, self.now, path.name, outcome, reason)
        except OSError as e:
            self.warnings.append(f"경고: {LOG_FILE} 에 쓰지 못함 ({e})")

    def _message(self) -> str:
        counts = self.report.counts
        parts = [f"{k} {counts[k]}" for k in NEWS_KEYS]
        if counts["보류"]:
            parts.append(f"보류 {counts['보류']}")
        lines = []
        if self.stop_reason:
            lines.append(f"중요: {self.stop_reason}")
        lines.append(f"{self.job.title}: " + " · ".join(parts))
        if self.ledger_locked:
            lines.append(f"누적장부가 엑셀에서 열려 있어 파일 {self.ledger_held}개를 처리하지 못했습니다(보류). "
                         f"엑셀을 닫으면 다음 실행 때 처리합니다: {ledger_path(self.job)}")
        folder = Path(self.job.folder)
        if counts["확인필요"]:
            lines.append(f"확인필요 폴더의 .txt 를 열어 사유를 보세요: {folder / REVIEW}")
        for line in self.remade:
            lines.append(line)
        if counts["승인대기"]:
            lines.append(f"승인대기 폴더의 대조표를 확인한 뒤 승인하세요: {folder / WAITING / WAITING_CHECKS}")
        if counts["보류"]:
            lines.append("보류 파일은 입력 폴더에 그대로 있고 다음 실행 때 처리합니다.")
        if self.report.unmatched_old:
            lines.append(f"어느 작업에도 맞지 않는 파일 {self.report.unmatched_old}개가 입력 폴더에 하루 넘게 "
                         f"남아 있습니다: {self.job.input_dir}")
        return "\n".join(lines + self.warnings)
