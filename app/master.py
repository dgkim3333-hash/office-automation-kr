"""고객마스터.xlsx: 고객명 → 고객코드·결제조건·기본 인도조건. 회사 정보라 사용자가 직접 채운다."""
import hashlib
import io
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.table import Table, TableStyleInfo

NAME_COLUMN = "고객명"
EXAMPLE_COLUMNS = ["고객명", "고객코드", "결제조건", "기본인도조건"]
# 가짜 예시(samples 의 가짜 고객). 실제로 쓸 때는 여러분 고객으로 바꾼다.
EXAMPLE_ROWS = [
    ["ACME Semiconductor Korea Co., Ltd.", "C-KR-0417", "Net 60", "DAP Cheongun"],
    ["Hanil DataCenter Co., Ltd.", "C-KR-0522", "Net 30", "CIF Haeun"],
]


class MasterError(Exception):
    """고객마스터.xlsx 를 읽을 수 없을 때."""


class DuplicateCustomer(MasterError):
    """같은 고객 줄이 두 개 이상이라 어느 줄이 맞는지 모를 때. 그 고객 문서만 확인필요로 간다."""


class Master(dict):
    """고객명 열쇠(_key) → 줄. duplicates = 두 줄 이상인 고객의 열쇠 → 엑셀 줄 번호들(머리글이 1번째 줄)."""

    def __init__(self):
        super().__init__()
        self.duplicates: dict[str, list[int]] = {}


def _key(name) -> str:
    return " ".join(str(name or "").split()).lower()


def load_master(path) -> dict[str, dict]:
    path = Path(path)
    if not path.exists():
        raise MasterError(f"고객마스터.xlsx 가 없습니다: {path}")
    return _master_from(path)


def master_from_bytes(data: bytes) -> dict[str, dict]:
    """한 번 읽은 파일 바이트에서 고객마스터를 만든다(실행 중 파일이 바뀌어도 지문과 같은 내용을 쓰게)."""
    return _master_from(io.BytesIO(data))


def _master_from(source) -> dict[str, dict]:
    try:
        wb = openpyxl.load_workbook(source, data_only=True)
    except Exception as e:
        raise MasterError(f"고객마스터.xlsx 를 열지 못했습니다(엑셀 통합 문서 .xlsx 인지 확인): {e}") from None
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header = [str(h).strip() if h is not None else "" for h in (rows[0] if rows else [])]
    if NAME_COLUMN not in header:
        raise MasterError(f"고객마스터.xlsx 첫 줄에 '{NAME_COLUMN}' 열이 있어야 합니다.")
    master = Master()
    lines: dict[str, list[int]] = {}
    for line, values in enumerate(rows[1:], 2):
        row = {h: v for h, v in zip(header, values) if h}
        key = _key(row.get(NAME_COLUMN))
        if key:
            master.setdefault(key, row)
            lines.setdefault(key, []).append(line)
    master.duplicates = {k: v for k, v in lines.items() if len(v) > 1}
    return master


def lookup(master, customer_name):
    """고객명으로 찾는다(대소문자·띄어쓰기 차이는 무시). 없으면 None.
    같은 고객 줄이 두 개 이상이면 DuplicateCustomer(어느 줄의 고객코드 · 결제조건이 맞는지 모르므로)."""
    key = _key(customer_name)
    lines = getattr(master, "duplicates", {}).get(key)
    if lines:
        name = master[key].get(NAME_COLUMN)
        raise DuplicateCustomer(
            f"고객마스터에 '{name}' 줄이 {len(lines)}개 있습니다({' · '.join(map(str, lines))}번째 줄). "
            "어느 줄이 맞는지 알 수 없어 회사 양식을 만들지 않았습니다. "
            "고객마스터.xlsx 에서 한 줄만 남기고 저장한 뒤 지금 실행을 누르세요")
    return master.get(key)


def master_stamp(row) -> str:
    """고객마스터 한 고객 줄의 지문. 승인대기 대조표에 남겨, 친구가 본 뒤 그 고객 줄이 바뀌면 대조표를 새로 만든다
    (다른 고객 줄을 더하거나 고친 것은 상관없음). 줄을 쓰지 않으면 빈 글자."""
    if row is None:
        return ""
    text = "\n".join(f"{k}\t{type(v).__name__}\t{v}" for k, v in sorted(row.items()))
    return hashlib.sha256(text.encode("utf-8")).hexdigest()[:16]


def write_example(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "고객마스터"
    ws.append(EXAMPLE_COLUMNS)
    for row in EXAMPLE_ROWS:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for letter, width in zip("ABCD", (40, 14, 14, 18)):
        ws.column_dimensions[letter].width = width
    table = Table(displayName="Customer_Master", ref=f"A1:D{len(EXAMPLE_ROWS) + 1}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A2"
    wb.save(path)
