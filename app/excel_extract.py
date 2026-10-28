"""고객이 보낸 엑셀 PO 를 규칙의 엑셀양식대로 읽는다. 파일을 옮기거나 쓰지 않는다."""
import math
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from openpyxl.utils import column_index_from_string

from app.extract import HEAD_LINES, ExtractResult, to_amount, to_date
from app.rules import CELL, RIGHT_OF

DETECT_ROWS = 20  # 판별키워드를 찾는 줄 수
EXCEL_DIGITS = 15  # 엑셀이 숫자를 보관 · 계산하는 유효 자리 수


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(" ")
    return str(value).strip()


def _sheet_lines(ws, rows) -> list[str]:
    lines = []
    for row in ws.iter_rows(min_row=1, max_row=rows, values_only=True):
        cells = [_text(v) for v in row if v is not None]
        if cells:
            lines.append("\t".join(cells))
    return lines


def _find(ws, label):
    """글자가 정확히 같은 첫 칸."""
    for row in ws.iter_rows():
        for cell in row:
            if _text(cell.value) == label:
                return cell
    return None


def _head_value(ws, where):
    if CELL.match(where):
        return ws[where].value, True
    found = _find(ws, where[len(RIGHT_OF):].strip())
    if found is None:
        return None, False
    return ws.cell(row=found.row, column=found.column + 1).value, True


def _as_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return to_date(str(value)) if value is not None else None


def excel_number(value) -> Decimal | None:
    """엑셀 칸의 숫자를 Decimal 로. float 은 엑셀처럼 유효 15자리로 맞추고 뒤의 0 은 뺀다.

    수식 =3*0.1 은 0.30000000000000004 로 저장되지만 엑셀은 0.3 으로 보고 계산한다.
    """
    if value is None or isinstance(value, bool):
        return None
    try:
        if isinstance(value, Decimal):
            return value if value.is_finite() else None
        if isinstance(value, int):
            return Decimal(value)
        if isinstance(value, float):
            if not math.isfinite(value):
                return None
            number = Decimal(format(value, f".{EXCEL_DIGITS}g"))
            if number.as_tuple().exponent > 0:  # 1.5e+20 같은 표기를 정수로
                number = number.quantize(Decimal(1))
            return number
    except InvalidOperation:
        return None
    return None


def _as_amount(value):
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        return excel_number(value)
    return to_amount(str(value))


def _plain(value):
    if isinstance(value, datetime):
        return value.date() if value.time() == datetime.min.time() else value
    if isinstance(value, str):
        return value.strip()
    return value


def _convert(name, value, fmt, warnings):
    if value is None or _text(value) == "":
        return None
    if fmt == "날짜":
        converted = _as_date(value)
    elif fmt == "금액":
        converted = _as_amount(value)
    else:
        return _plain(value)
    if converted is None:
        warnings.append(f"'{name}' 값 '{_text(value)}' 를 {fmt}로 바꾸지 못함")
    return converted


def _items(ws, spec, warnings):
    header = _find(ws, spec.header_text)
    if header is None:
        warnings.append(f"품목 머리글 '{spec.header_text}' 를 못 찾음, 품목 건너뜀")
        return None
    stop_key = "품번" if "품번" in spec.columns else next(iter(spec.columns))
    stop_col = column_index_from_string(spec.columns[stop_key])
    rows = []
    for r in range(header.row + 1, ws.max_row + 1):
        if _text(ws.cell(row=r, column=stop_col).value) == "":
            break
        row = {}
        for name, letter in spec.columns.items():
            value = ws.cell(row=r, column=column_index_from_string(letter)).value
            if name in spec.amount_columns:
                value = _as_amount(value)
            elif name in spec.quantity_columns:
                amount = _as_amount(value)
                value = int(amount) if amount is not None and amount == amount.to_integral_value() else None
            else:
                value = _plain(value) if value is not None else ""
            if value is None:
                warnings.append(f"품목 {r}행 '{name}' 값이 숫자가 아님, 품목 건너뜀")
                return None
            row[name] = value
        rows.append(row)
    return rows


def read_excel(path, rules) -> ExtractResult:
    wb = openpyxl.load_workbook(Path(path), data_only=True, read_only=False)
    ws = wb.active
    lines = _sheet_lines(ws, max(DETECT_ROWS, HEAD_LINES))
    head_text = "\n".join(lines[:HEAD_LINES])
    if not lines:
        return ExtractResult(doc_type=None, has_text=False, is_excel=True)

    detect = "\n".join(lines[:DETECT_ROWS]).lower()
    form = next((f for f in rules.excel_forms if any(k.lower() in detect for k in f.keywords)), None)
    if form is None:
        return ExtractResult(doc_type=None, text_head=head_text, is_excel=True)

    doc_type = rules.get(form.doc_type)
    result = ExtractResult(doc_type=doc_type, text_head=head_text, source_form=form.name, is_excel=True)
    for name, cell in form.head.items():
        value, found = _head_value(ws, cell.where)
        if not found:
            result.warnings.append(f"머리 '{name}' 칸({cell.where})을 못 찾음")
        result.fields[name] = _convert(name, value, cell.fmt, result.warnings)
    result.missing_required = result.fields.get(doc_type.number_field) is None
    if form.items is not None:
        result.table = _items(ws, form.items, result.warnings)
    return result
