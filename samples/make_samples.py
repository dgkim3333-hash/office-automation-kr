# -*- coding: utf-8 -*-
"""office-automation-kr 샘플 생성기.
가짜 회사·가짜 고객·가짜 금액으로 실제 거래 문서와 비슷한 예제를 만든다.
실행: python make_samples.py <출력폴더>
"""
import sys, os, csv
from datetime import date, datetime
from decimal import Decimal
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib.units import inch
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

OUT = sys.argv[1] if len(sys.argv) > 1 else "samples"
os.makedirs(OUT, exist_ok=True)
for sub in ["01_고객오더_엑셀", "02_회사양식_변환", "03_Invoice_DM_PDF", "04_장부_엑셀"]:
    os.makedirs(os.path.join(OUT, sub), exist_ok=True)

# ---------- 가짜 당사자 ----------
SUPPLIER = {
    "name": "Vertex Thermal Systems, Inc.",
    "addr1": "1200 Thermal Way, Building C",
    "addr2": "Prairie Bend, TX 76999, USA",
    "tel": "+1 (512) 555-0140",
    "branch": "Vertex Thermal Systems Korea (Branch)",
    "branch_addr": "Haneul-si, Korea",
}
CUSTOMERS = {
    "ACME": {
        "code": "C-KR-0417",
        "name": "ACME Semiconductor Korea Co., Ltd.",
        "bill": "ACME Semiconductor Korea Co., Ltd.\n100 Gaesang-ro, Cheongun-si, Korea",
        "ship": "ACME Fab 1, Dock 7\n100 Gaesang-ro, Cheongun-si, Korea",
        "contact": "Mr. J. Park (Procurement)",
    },
    "HANIL": {
        "code": "C-KR-0522",
        "name": "Hanil DataCenter Co., Ltd.",
        "bill": "Hanil DataCenter Co., Ltd.\n250 Byeolbit-ro, Hanbit-si, Korea",
        "ship": "Hanil IDC Campus 2, Loading Bay B\n250 Byeolbit-ro, Hanbit-si, Korea",
        "contact": "Ms. S. Kim (Facilities)",
    },
}
# 품목 마스터 (가짜 부품번호)
PARTS = {
    "VTS-CDU-300": ("Coolant Distribution Unit, 300 kW, N+1 pump", Decimal("48500.00")),
    "VTS-RDHX-42U": ("Rear Door Heat Exchanger, 42U, 35 kW", Decimal("6900.00")),
    "VTS-MFD-12": ("Coolant Manifold, 12-port, SS316", Decimal("1850.00")),
    "VTS-QD-DN25": ("Quick Disconnect Coupling DN25 (pair)", Decimal("260.00")),
    "VTS-CHL-500RT": ("Air-cooled Chiller, 500 RT, VFD", Decimal("212000.00")),
    "VTS-PMP-CDU-K": ("CDU Pump Cartridge Kit", Decimal("3100.00")),
    "VTS-SENS-LEAK": ("Leak Detection Sensor Cable, 10 m", Decimal("95.00")),
    "VTS-FLT-50": ("Inline Particle Filter 50 um", Decimal("140.00")),
}

# ---------- 거래 2건 ----------
ORDERS = [
    {
        "cust": "ACME", "po": "PO-ACME-2026-0417", "po_date": date(2026, 9, 18),
        "so": "SO-26-11892", "req_date": date(2026, 11, 20),
        "incoterms": "DAP Cheongun", "terms": "Net 60", "currency": "USD",
        "lines": [("VTS-CDU-300", 2), ("VTS-MFD-12", 4), ("VTS-QD-DN25", 24), ("VTS-SENS-LEAK", 10)],
        "inv": {"no": "INV-2026-10231", "date": date(2026, 10, 14), "ship_date": date(2026, 10, 13),
                "carrier": "DHL Global Forwarding", "awb": "DHLGF-7731-884210", "freight": Decimal("3850.00"),
                "vessel": "", "eta": date(2026, 10, 17)},
        "dm": {"no": "DM-2026-0092", "date": date(2026, 10, 21),
               "reason": "Freight surcharge - expedited air (customer request 2026-10-10)",
               "lines": [("FRT-EXP", "Expedited air freight surcharge", 1, Decimal("1420.00"))]},
    },
    {
        "cust": "HANIL", "po": "HDC-PUR-26-0908", "po_date": date(2026, 9, 25),
        "so": "SO-26-11930", "req_date": date(2026, 12, 5),
        "incoterms": "CIF Haeun", "terms": "Net 30", "currency": "USD",
        "lines": [("VTS-CHL-500RT", 1), ("VTS-PMP-CDU-K", 3), ("VTS-FLT-50", 20)],
        "inv": {"no": "INV-2026-10278", "date": date(2026, 10, 22), "ship_date": date(2026, 10, 20),
                "carrier": "Maersk Line", "awb": "MAEU-2261-907733 / Cont. MSKU-884120-7", "freight": Decimal("9600.00"),
                "vessel": "Maersk Edmonton V.041E", "eta": date(2026, 11, 14)},
        "dm": {"no": "DM-2026-0097", "date": date(2026, 10, 27),
               "reason": "Price adjustment - VTS-FLT-50 list price change effective 2026-10-01",
               "lines": [("VTS-FLT-50", "Price difference 20 x (USD 155.00 - 140.00)", 20, Decimal("15.00"))]},
    },
]

def money(d): return f"${d:,.2f}"
def usd(d): return f"{d:,.2f}"

# ---------- PDF 공통 ----------
def header(c, title, no, dt):
    c.setFont("Helvetica-Bold", 16); c.drawString(0.8*inch, 10.3*inch, SUPPLIER["name"])
    c.setFont("Helvetica", 9)
    c.drawString(0.8*inch, 10.1*inch, SUPPLIER["addr1"]); c.drawString(0.8*inch, 9.95*inch, SUPPLIER["addr2"])
    c.drawString(0.8*inch, 9.8*inch, "Tel " + SUPPLIER["tel"])
    c.setFont("Helvetica-Bold", 22); c.drawRightString(7.7*inch, 10.3*inch, title)
    c.setFont("Helvetica", 10)
    c.drawRightString(7.7*inch, 10.0*inch, f"{title.title()} No: {no}")
    c.drawRightString(7.7*inch, 9.82*inch, f"Date: {dt.strftime('%m/%d/%Y')}")
    c.line(0.8*inch, 9.65*inch, 7.7*inch, 9.65*inch)

def block(c, x, y, label, text):
    c.setFont("Helvetica-Bold", 9); c.drawString(x, y, label)
    c.setFont("Helvetica", 9)
    for i, ln in enumerate(text.split("\n")):
        c.drawString(x, y - 0.17*inch*(i+1), ln)

def table(c, y, cols, widths, rows):
    x0 = 0.8*inch
    c.setFont("Helvetica-Bold", 9)
    x = x0
    for col, w in zip(cols, widths):
        c.drawString(x+2, y, col); x += w
    c.line(x0, y-3, 7.7*inch, y-3)
    c.setFont("Helvetica", 9); y -= 0.22*inch
    for r in rows:
        x = x0
        for val, w, col in zip(r, widths, cols):
            if col in ("Qty", "Unit Price", "Amount"):
                c.drawRightString(x+w-4, y, str(val))
            else:
                c.drawString(x+2, y, str(val))
            x += w
        y -= 0.2*inch
    c.line(x0, y+0.1*inch, 7.7*inch, y+0.1*inch)
    return y

def make_invoice(o):
    cu = CUSTOMERS[o["cust"]]; inv = o["inv"]
    path = os.path.join(OUT, "03_Invoice_DM_PDF", f"{inv['no']}.pdf")
    c = canvas.Canvas(path, pagesize=letter)
    header(c, "INVOICE", inv["no"], inv["date"])
    block(c, 0.8*inch, 9.4*inch, "Bill To:", cu["bill"])
    block(c, 4.4*inch, 9.4*inch, "Ship To:", cu["ship"])
    y = 8.6*inch
    c.setFont("Helvetica", 9)
    meta = [("Customer PO", o["po"]), ("Sales Order", o["so"]), ("Payment Terms", o["terms"]),
            ("Incoterms", o["incoterms"]), ("Ship Date", inv["ship_date"].strftime("%m/%d/%Y")),
            ("Carrier", inv["carrier"]), ("Tracking / AWB", inv["awb"]), ("ETA", inv["eta"].strftime("%m/%d/%Y"))]
    if inv["vessel"]: meta.insert(6, ("Vessel / Voyage", inv["vessel"]))
    for i, (k, v) in enumerate(meta):
        col = i % 2; row = i // 2
        x = 0.8*inch + col*3.6*inch; yy = y - row*0.18*inch
        c.setFont("Helvetica-Bold", 9); c.drawString(x, yy, k + ":")
        c.setFont("Helvetica", 9); c.drawString(x + 1.15*inch, yy, v)
    y = 7.5*inch
    rows = []; subtotal = Decimal("0")
    for i, (pn, q) in enumerate(o["lines"], 1):
        desc, up = PARTS[pn]; amt = up * q; subtotal += amt
        rows.append((i, pn, desc, q, usd(up), usd(amt)))
    y = table(c, y, ["Line", "Item", "Description", "Qty", "Unit Price", "Amount"],
              [0.45*inch, 1.15*inch, 3.1*inch, 0.5*inch, 0.9*inch, 0.8*inch], rows)
    y -= 0.15*inch
    c.setFont("Helvetica", 10)
    for k, v in [("Subtotal:", money(subtotal)), ("Freight:", money(inv["freight"])), ("Tax:", money(Decimal("0")))]:
        c.drawRightString(6.6*inch, y, k); c.drawRightString(7.7*inch, y, v); y -= 0.2*inch
    total = subtotal + inv["freight"]
    c.setFont("Helvetica-Bold", 11)
    c.drawRightString(6.6*inch, y, "Total:"); c.drawRightString(7.7*inch, y, money(total))
    c.setFont("Helvetica", 8)
    c.drawString(0.8*inch, 1.4*inch, "Remit to: Vertex Thermal Systems, Inc. - Wire details on file. Please reference the invoice number on payment.")
    c.drawString(0.8*inch, 1.25*inch, f"Questions: ar@vertex-thermal.example  |  Korea branch: {SUPPLIER['branch']}, {SUPPLIER['branch_addr']}")
    c.showPage(); c.save()
    return path, subtotal, total

def make_dm(o):
    cu = CUSTOMERS[o["cust"]]; dm = o["dm"]; inv = o["inv"]
    path = os.path.join(OUT, "03_Invoice_DM_PDF", f"{dm['no']}.pdf")
    c = canvas.Canvas(path, pagesize=letter)
    header(c, "DEBIT MEMO", dm["no"], dm["date"])
    block(c, 0.8*inch, 9.4*inch, "Bill To:", cu["bill"])
    y = 8.6*inch
    for i, (k, v) in enumerate([("Reference Invoice", inv["no"]), ("Customer PO", o["po"]),
                                ("Sales Order", o["so"]), ("Payment Terms", o["terms"])]):
        col = i % 2; row = i // 2
        x = 0.8*inch + col*3.6*inch; yy = y - row*0.18*inch
        c.setFont("Helvetica-Bold", 9); c.drawString(x, yy, k + ":")
        c.setFont("Helvetica", 9); c.drawString(x + 1.25*inch, yy, v)
    c.setFont("Helvetica-Bold", 9); c.drawString(0.8*inch, 8.0*inch, "Reason:")
    c.setFont("Helvetica", 9); c.drawString(1.4*inch, 8.0*inch, dm["reason"])
    rows = []; total = Decimal("0")
    for i, (code, desc, q, up) in enumerate(dm["lines"], 1):
        amt = up * q; total += amt
        rows.append((i, code, desc, q, usd(up), usd(amt)))
    y = table(c, 7.5*inch, ["Line", "Item", "Description", "Qty", "Unit Price", "Amount"],
              [0.45*inch, 1.15*inch, 3.1*inch, 0.5*inch, 0.9*inch, 0.8*inch], rows)
    y -= 0.15*inch
    c.setFont("Helvetica-Bold", 11)
    c.drawRightString(6.6*inch, y, "Total Debit:"); c.drawRightString(7.7*inch, y, money(total))
    c.setFont("Helvetica", 8)
    c.drawString(0.8*inch, 1.4*inch, "This debit memo is payable with the referenced invoice. Please reference the memo number on payment.")
    c.showPage(); c.save()
    return path, total

def make_unknown():
    path = os.path.join(OUT, "03_Invoice_DM_PDF", "Letter_warranty_notice.pdf")
    c = canvas.Canvas(path, pagesize=letter)
    c.setFont("Helvetica-Bold", 14); c.drawString(1*inch, 10*inch, SUPPLIER["name"])
    c.setFont("Helvetica", 10)
    for i, ln in enumerate([
        "October 20, 2026", "",
        "To: ACME Semiconductor Korea Co., Ltd.", "",
        "Subject: Extended warranty registration - CDU units shipped under SO-26-11892", "",
        "Dear Mr. Park,", "",
        "Thank you for your order. The two Coolant Distribution Units shipped on October 13 are",
        "eligible for the 24-month extended warranty at no additional charge. Please register the",
        "serial numbers at the customer portal within 30 days of commissioning.", "",
        "Sincerely,", "Vertex Thermal Systems - Customer Care"]):
        c.drawString(1*inch, 9.6*inch - i*0.2*inch, ln)
    c.showPage(); c.save()
    return path

def make_blank():
    path = os.path.join(OUT, "03_Invoice_DM_PDF", "Scan_signed_packing_list.pdf")
    c = canvas.Canvas(path, pagesize=letter)
    c.setFillGray(0.85); c.rect(1*inch, 2*inch, 6.5*inch, 8*inch, fill=1, stroke=0)
    c.showPage(); c.save()
    return path

# ---------- 엑셀 공통 ----------
thin = Side(style="thin", color="999999")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")

def style_header(ws, row, ncol):
    for cidx in range(1, ncol+1):
        cell = ws.cell(row=row, column=cidx)
        cell.font = Font(bold=True); cell.fill = HEAD_FILL; cell.border = BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")

def autowidth(ws, minw=8, maxw=48):
    for col in ws.columns:
        w = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[get_column_letter(col[0].column)].width = max(minw, min(maxw, w + 2))

# ---------- 1. 고객 오더 엑셀 (고객 양식 — 고객마다 다름) ----------
def make_customer_po(o):
    cu = CUSTOMERS[o["cust"]]
    wb = Workbook(); ws = wb.active; ws.title = "Purchase Order"
    if o["cust"] == "ACME":
        ws["A1"] = "PURCHASE ORDER"; ws["A1"].font = Font(bold=True, size=16)
        ws["A3"] = "Buyer:"; ws["B3"] = cu["name"]
        ws["A4"] = "PO Number:"; ws["B4"] = o["po"]
        ws["A5"] = "PO Date:"; ws["B5"] = o["po_date"]; ws["B5"].number_format = "yyyy-mm-dd"
        ws["A6"] = "Supplier:"; ws["B6"] = SUPPLIER["name"]
        ws["A7"] = "Deliver To:"; ws["B7"] = cu["ship"].replace("\n", ", ")
        ws["A8"] = "Requested Delivery:"; ws["B8"] = o["req_date"]; ws["B8"].number_format = "yyyy-mm-dd"
        ws["A9"] = "Incoterms:"; ws["B9"] = o["incoterms"]
        ws["A10"] = "Currency:"; ws["B10"] = o["currency"]
        ws["A11"] = "Buyer Contact:"; ws["B11"] = cu["contact"]
        hdr = ["No", "Supplier Part No", "Description", "Qty", "UoM", "Unit Price (USD)", "Amount (USD)", "Need-by Date"]
        r0 = 13
    else:
        ws["A1"] = cu["name"]; ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = "구매발주서 (Purchase Order)"
        ws["A4"] = "발주번호"; ws["B4"] = o["po"]
        ws["C4"] = "발주일"; ws["D4"] = o["po_date"]; ws["D4"].number_format = "yyyy-mm-dd"
        ws["A5"] = "공급사"; ws["B5"] = SUPPLIER["name"]
        ws["C5"] = "납기요청일"; ws["D5"] = o["req_date"]; ws["D5"].number_format = "yyyy-mm-dd"
        ws["A6"] = "납품처"; ws["B6"] = cu["ship"].replace("\n", ", ")
        ws["A7"] = "인도조건"; ws["B7"] = o["incoterms"]
        ws["C7"] = "통화"; ws["D7"] = o["currency"]
        ws["A8"] = "담당"; ws["B8"] = cu["contact"]
        hdr = ["순번", "품번", "품명", "수량", "단위", "단가(USD)", "금액(USD)", "납기"]
        r0 = 10
    for j, h in enumerate(hdr, 1): ws.cell(row=r0, column=j, value=h)
    style_header(ws, r0, len(hdr))
    total = Decimal("0")
    for i, (pn, q) in enumerate(o["lines"], 1):
        desc, up = PARTS[pn]; amt = up * q; total += amt
        r = r0 + i
        vals = [i, pn, desc, q, "EA", float(up), float(amt), o["req_date"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=r, column=j, value=v); c.border = BORDER
        ws.cell(row=r, column=6).number_format = "#,##0.00"
        ws.cell(row=r, column=7).number_format = "#,##0.00"
        ws.cell(row=r, column=8).number_format = "yyyy-mm-dd"
    r = r0 + len(o["lines"]) + 1
    ws.cell(row=r, column=6, value="Total" if o["cust"] == "ACME" else "합계").font = Font(bold=True)
    ws.cell(row=r, column=7, value=float(total)).number_format = "#,##0.00"
    ws.cell(row=r, column=7).font = Font(bold=True)
    autowidth(ws)
    path = os.path.join(OUT, "01_고객오더_엑셀", f"고객오더_{o['cust']}_{o['po']}.xlsx")
    wb.save(path); return path, total

# ---------- 2. 회사 입력 양식 (본사 시스템 업로드용 — 열 고정) ----------
SO_COLS = ["Order Date", "Customer Code", "Customer Name", "Customer PO No", "Line No", "Part No",
           "Description", "Qty", "UoM", "Unit Price", "Ext Price", "Currency", "Requested Ship Date",
           "Incoterms", "Payment Terms", "Ship-To Address", "Customer Contact", "Entered By", "Remarks"]

def make_company_so(o):
    cu = CUSTOMERS[o["cust"]]
    wb = Workbook(); ws = wb.active; ws.title = "SO_Upload"
    for j, h in enumerate(SO_COLS, 1): ws.cell(row=1, column=j, value=h)
    style_header(ws, 1, len(SO_COLS))
    for i, (pn, q) in enumerate(o["lines"], 1):
        desc, up = PARTS[pn]
        vals = [o["po_date"], cu["code"], cu["name"], o["po"], i*10, pn, desc, q, "EA", float(up), float(up*q),
                o["currency"], o["req_date"], o["incoterms"], o["terms"], cu["ship"].replace("\n", ", "),
                cu["contact"], "KR-Branch / D. Lee", ""]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i+1, column=j, value=v); c.border = BORDER
        ws.cell(row=i+1, column=1).number_format = "yyyy-mm-dd"
        ws.cell(row=i+1, column=13).number_format = "yyyy-mm-dd"
        ws.cell(row=i+1, column=10).number_format = "#,##0.00"
        ws.cell(row=i+1, column=11).number_format = "#,##0.00"
    n = len(o["lines"]) + 1
    t = Table(displayName="SO_Upload", ref=f"A1:{get_column_letter(len(SO_COLS))}{n}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t); ws.freeze_panes = "A2"; autowidth(ws)
    ws2 = wb.create_sheet("열설명")
    for r, (k, v) in enumerate([
        ("열", "설명"),
        ("Order Date", "고객 PO 발주일"), ("Customer Code", "본사 시스템 고객코드 (고객마스터 참조)"),
        ("Customer PO No", "고객 발주번호 — 그대로 복사"), ("Line No", "10, 20, 30 … 10단위"),
        ("Part No", "당사 부품번호. 고객 품번이 다르면 품번 매핑표 참조"), ("Qty/UoM", "수량과 단위(EA)"),
        ("Unit Price / Ext Price", "달러. Ext = 단가×수량"), ("Requested Ship Date", "고객 납기요청일"),
        ("Incoterms / Payment Terms", "고객 PO 또는 계약조건"), ("Entered By", "지사 입력자")], 1):
        ws2.cell(row=r, column=1, value=k); ws2.cell(row=r, column=2, value=v)
    style_header(ws2, 1, 2); autowidth(ws2)
    path = os.path.join(OUT, "02_회사양식_변환", f"회사양식_SO입력_{o['po']}.xlsx")
    wb.save(path); return path

# ---------- 4-a. 단건 추출결과 (check 출력과 같은 내용) ----------
def make_extract_result(o, inv_path, subtotal, total):
    inv = o["inv"]; cu = CUSTOMERS[o["cust"]]
    wb = Workbook(); ws = wb.active; ws.title = "추출결과"
    ws.append(["필드", "값", "비고"])
    style_header(ws, 1, 3)
    for k, v in [("문서종류", "Invoice"), ("원본파일", os.path.basename(inv_path)),
                 ("인보이스번호", inv["no"]), ("발행일", inv["date"]), ("고객명", cu["name"]),
                 ("고객PO번호", o["po"]), ("판매오더번호", o["so"]), ("결제조건", o["terms"]),
                 ("인도조건", o["incoterms"]), ("선적일", inv["ship_date"]), ("운송사", inv["carrier"]),
                 ("추적번호", inv["awb"]), ("도착예정일", inv["eta"]), ("소계", float(subtotal)),
                 ("운임", float(inv["freight"])), ("합계금액", float(total))]:
        ws.append([k, v, ""])
    for r in range(2, ws.max_row+1):
        c = ws.cell(row=r, column=2)
        if isinstance(c.value, (date, datetime)): c.number_format = "yyyy-mm-dd"
        if isinstance(c.value, float): c.number_format = "#,##0.00"
    autowidth(ws)
    ws2 = wb.create_sheet("품목")
    ws2.append(["줄", "품번", "품명", "수량", "단가", "금액"]); style_header(ws2, 1, 6)
    for i, (pn, q) in enumerate(o["lines"], 1):
        desc, up = PARTS[pn]; ws2.append([i, pn, desc, q, float(up), float(up*q)])
        ws2.cell(row=i+1, column=5).number_format = "#,##0.00"; ws2.cell(row=i+1, column=6).number_format = "#,##0.00"
    autowidth(ws2)
    path = os.path.join(OUT, "04_장부_엑셀", f"추출결과_단건_{inv['no']}.xlsx")
    wb.save(path); return path

# ---------- 4-b. 누적장부 (설계 7절 그대로, 샘플 처리 후 기대 상태) ----------
INV_FIELDS = ["인보이스번호", "발행일", "고객명", "고객PO번호", "판매오더번호", "결제조건", "인도조건",
              "선적일", "운송사", "추적번호", "도착예정일", "소계", "운임", "합계금액"]
PO_FIELDS = ["고객PO번호", "발주일", "고객명", "납기요청일", "인도조건", "통화", "합계금액"]
DM_FIELDS = ["DM번호", "발행일", "고객명", "참조인보이스", "고객PO번호", "사유", "합계금액"]
FIXED = ["처리일시", "원본파일"]

def add_table(ws, name):
    ref = f"A1:{get_column_letter(ws.max_column)}{max(ws.max_row, 2)}"
    t = Table(displayName=name, ref=ref); t.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(t); ws.freeze_panes = "A2"

def fmt_row(ws, r):
    for c in ws[r]:
        if isinstance(c.value, datetime): c.number_format = "yyyy-mm-dd hh:mm"
        elif isinstance(c.value, date): c.number_format = "yyyy-mm-dd"
        elif isinstance(c.value, float): c.number_format = "#,##0.00"

def make_ledger(results):
    wb = Workbook(); now = datetime(2026, 10, 28, 9, 40)
    ws = wb.active; ws.title = "Invoice"
    ws.append(FIXED + INV_FIELDS + ["정리된경로"]); style_header(ws, 1, ws.max_column)
    for o, inv_path, subtotal, total, dm_path, dm_total in results:
        inv = o["inv"]; cu = CUSTOMERS[o["cust"]]
        ws.append([now, os.path.basename(inv_path), inv["no"], inv["date"], cu["name"], o["po"], o["so"], o["terms"],
                   o["incoterms"], inv["ship_date"], inv["carrier"], inv["awb"], inv["eta"], float(subtotal),
                   float(inv["freight"]), float(total),
                   f"정리문서/{cu['name'].split(' ')[0]}/{inv['date']:%Y-%m}/Invoice_{inv['no']}.pdf"])
        fmt_row(ws, ws.max_row)
    add_table(ws, "T_Invoice"); autowidth(ws)

    ws = wb.create_sheet("PO"); ws.append(FIXED + PO_FIELDS + ["정리된경로"]); style_header(ws, 1, ws.max_column)
    for o, *_ in results:
        cu = CUSTOMERS[o["cust"]]; tot = sum(PARTS[pn][1]*q for pn, q in o["lines"])
        ws.append([now, f"고객오더_{o['cust']}_{o['po']}.xlsx", o["po"], o["po_date"], cu["name"], o["req_date"],
                   o["incoterms"], o["currency"], float(tot),
                   f"정리문서/{cu['name'].split(' ')[0]}/{o['po_date']:%Y-%m}/PO_{o['po']}.xlsx"])
        fmt_row(ws, ws.max_row)
    add_table(ws, "T_PO"); autowidth(ws)

    ws = wb.create_sheet("DM"); ws.append(FIXED + DM_FIELDS + ["정리된경로"]); style_header(ws, 1, ws.max_column)
    for o, inv_path, subtotal, total, dm_path, dm_total in results:
        dm = o["dm"]; cu = CUSTOMERS[o["cust"]]
        ws.append([now, os.path.basename(dm_path), dm["no"], dm["date"], cu["name"], o["inv"]["no"], o["po"],
                   dm["reason"], float(dm_total),
                   f"정리문서/{cu['name'].split(' ')[0]}/{dm['date']:%Y-%m}/DM_{dm['no']}.pdf"])
        fmt_row(ws, ws.max_row)
    add_table(ws, "T_DM"); autowidth(ws)

    ws = wb.create_sheet("품목"); ws.append(["문서종류", "문서번호", "줄", "품번", "품명", "수량", "단가", "금액"]); style_header(ws, 1, 8)
    for o, inv_path, subtotal, total, dm_path, dm_total in results:
        for i, (pn, q) in enumerate(o["lines"], 1):
            desc, up = PARTS[pn]; ws.append(["Invoice", o["inv"]["no"], i, pn, desc, q, float(up), float(up*q)]); fmt_row(ws, ws.max_row)
        for i, (pn, q) in enumerate(o["lines"], 1):
            desc, up = PARTS[pn]; ws.append(["PO", o["po"], i, pn, desc, q, float(up), float(up*q)]); fmt_row(ws, ws.max_row)
        for i, (code, desc, q, up) in enumerate(o["dm"]["lines"], 1):
            ws.append(["DM", o["dm"]["no"], i, code, desc, q, float(up), float(up*q)]); fmt_row(ws, ws.max_row)
    add_table(ws, "T_Items"); autowidth(ws)

    ws = wb.create_sheet("처리로그"); ws.append(["일시", "파일", "결과", "사유"]); style_header(ws, 1, 4)
    for o, inv_path, subtotal, total, dm_path, dm_total in results:
        ws.append([now, os.path.basename(inv_path), "성공", ""]); fmt_row(ws, ws.max_row)
        ws.append([now, os.path.basename(dm_path), "성공", ""]); fmt_row(ws, ws.max_row)
        ws.append([now, f"고객오더_{o['cust']}_{o['po']}.xlsx", "성공", ""]); fmt_row(ws, ws.max_row)
    ws.append([now, "Letter_warranty_notice.pdf", "확인필요", "문서종류 판별 실패 (판별키워드 없음)"]); fmt_row(ws, ws.max_row)
    ws.append([now, "Scan_signed_packing_list.pdf", "확인필요", "텍스트 없음 — 스캔본, OCR 범위 밖"]); fmt_row(ws, ws.max_row)
    add_table(ws, "T_Log"); autowidth(ws)
    path = os.path.join(OUT, "04_장부_엑셀", "누적장부_기대결과.xlsx"); wb.save(path); return path

def make_empty_ledger():
    wb = Workbook(); ws = wb.active; ws.title = "Invoice"; ws.append(FIXED + INV_FIELDS + ["정리된경로"]); style_header(ws, 1, ws.max_column); autowidth(ws)
    ws = wb.create_sheet("PO"); ws.append(FIXED + PO_FIELDS + ["정리된경로"]); style_header(ws, 1, ws.max_column); autowidth(ws)
    ws = wb.create_sheet("DM"); ws.append(FIXED + DM_FIELDS + ["정리된경로"]); style_header(ws, 1, ws.max_column); autowidth(ws)
    ws = wb.create_sheet("품목"); ws.append(["문서종류", "문서번호", "줄", "품번", "품명", "수량", "단가", "금액"]); style_header(ws, 1, 8); autowidth(ws)
    ws = wb.create_sheet("처리로그"); ws.append(["일시", "파일", "결과", "사유"]); style_header(ws, 1, 4); autowidth(ws)
    path = os.path.join(OUT, "04_장부_엑셀", "누적장부_빈양식.xlsx"); wb.save(path); return path

# ---------- 실행 ----------
made = []
results = []
for o in ORDERS:
    p1, _ = make_customer_po(o); made.append(p1)
    p2 = make_company_so(o); made.append(p2)
    inv_path, subtotal, total = make_invoice(o); made.append(inv_path)
    dm_path, dm_total = make_dm(o); made.append(dm_path)
    results.append((o, inv_path, subtotal, total, dm_path, dm_total))
made.append(make_unknown()); made.append(make_blank())
made.append(make_extract_result(*results[0][:4]))
made.append(make_ledger(results)); made.append(make_empty_ledger())

# 자가 확인: PDF 텍스트 추출
import pdfplumber
print("---- 자가 확인 (pdfplumber) ----")
for p in made:
    if p.endswith(".pdf"):
        with pdfplumber.open(p) as pdf:
            t = pdf.pages[0].extract_text() or ""
        print(f"{os.path.basename(p)}: {len(t)}자, 1행: {t.splitlines()[0] if t else '(없음)'}")
print("---- 생성 파일 ----")
for p in made: print(p)
