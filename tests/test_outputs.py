"""문서별 결과 엑셀 — 엑셀 PO → 회사 SO 양식, PDF → 추출결과. 정답은 samples/02·04."""
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
import pytest

from app.excel_extract import read_excel
from app.extract import extract
from app.markup import load_markup, write_example_markup
from app.master import load_master, lookup, write_example
from app.outputs import (company_rows, extract_result_rows, read_back, write_check_file, write_company_form,
                         write_extract_result)
from app.rules import load_rules
from app.verify import Check

REPO = Path(__file__).resolve().parent.parent
SAMPLES = REPO / "samples"
CASES = [
    ("고객오더_ACME_PO-ACME-2026-0417.xlsx", "회사양식_SO입력_PO-ACME-2026-0417.xlsx"),
    ("고객오더_HANIL_HDC-PUR-26-0908.xlsx", "회사양식_SO입력_HDC-PUR-26-0908.xlsx"),
]


@pytest.fixture(scope="module")
def rules():
    return load_rules(REPO / "rules" / "규칙.예시.yaml")


@pytest.fixture(scope="module")
def master(tmp_path_factory):
    path = tmp_path_factory.mktemp("m") / "고객마스터.xlsx"
    write_example(path)
    return load_master(path)


@pytest.fixture(scope="module")
def markup(tmp_path_factory):
    path = tmp_path_factory.mktemp("k") / "마크업표.xlsx"
    write_example_markup(path)
    return load_markup(path)


def _same(a, b):
    if hasattr(a, "year") and hasattr(b, "year"):
        return (a.year, a.month, a.day) == (b.year, b.month, b.day)
    try:
        return Decimal(str(a)) == Decimal(str(b))
    except (InvalidOperation, ValueError):
        return a == b


def _sheet(path, name):
    rows = [r for r in openpyxl.load_workbook(path)[name].iter_rows(values_only=True)
            if any(v is not None for v in r)]
    return rows[0], rows[1:]


@pytest.mark.parametrize("order, expected", CASES)
def test_company_rows_match_expected_form(rules, master, markup, order, expected):
    result = read_excel(SAMPLES / "01_고객오더_엑셀" / order, rules)
    header, rows = company_rows(result, lookup(master, result.fields["고객명"]), rules.company_form, markup)
    exp_header, exp_rows = _sheet(SAMPLES / "02_회사양식_변환" / expected, "SO_Upload")
    assert header == list(exp_header)
    assert len(rows) == len(exp_rows)
    for got, exp in zip(rows, exp_rows):
        for name, g, e in zip(header, got, exp):
            assert _same(g, e), f"{order} {name}: {g!r} != {e!r}"


@pytest.mark.parametrize("order, expected", CASES)
def test_written_company_form_file(rules, master, markup, tmp_path, order, expected):
    result = read_excel(SAMPLES / "01_고객오더_엑셀" / order, rules)
    path = write_company_form(tmp_path, result, lookup(master, result.fields["고객명"]), rules.company_form, markup)
    assert path.name == expected
    got_header, got_rows = _sheet(path, "SO_Upload")
    exp_header, exp_rows = _sheet(SAMPLES / "02_회사양식_변환" / expected, "SO_Upload")
    assert got_header == exp_header and len(got_rows) == len(exp_rows)
    assert all(_same(g, e) for gr, er in zip(got_rows, exp_rows) for g, e in zip(gr, er))


def test_customer_missing_from_master_leaves_blank(rules):
    result = read_excel(SAMPLES / "01_고객오더_엑셀" / CASES[0][0], rules)
    header, rows = company_rows(result, None, rules.company_form)
    assert rows[0][header.index("Customer Code")] is None


def test_extract_result_file_matches_sample(rules, tmp_path):
    pdf = SAMPLES / "03_Invoice_DM_PDF" / "INV-2026-10231.pdf"
    path = write_extract_result(tmp_path, extract(pdf, rules), pdf.name)
    assert path.name == "추출결과_INV-2026-10231.xlsx"
    sample = SAMPLES / "04_장부_엑셀" / "추출결과_단건_INV-2026-10231.xlsx"
    for sheet in ("추출결과", "품목"):
        got_header, got_rows = _sheet(path, sheet)
        exp_header, exp_rows = _sheet(sample, sheet)
        assert got_header[:2] == exp_header[:2]
        assert len(got_rows) == len(exp_rows), sheet
        for gr, er in zip(got_rows, exp_rows):
            assert all(_same(g, e) for g, e in zip(gr[:2] if sheet == "추출결과" else gr, er)), (sheet, gr, er)


def test_output_name_collision_adds_suffix(rules, master, tmp_path):
    result = read_excel(SAMPLES / "01_고객오더_엑셀" / CASES[1][0], rules)
    row = lookup(master, result.fields["고객명"])
    first = write_company_form(tmp_path, result, row, rules.company_form)
    second = write_company_form(tmp_path, result, row, rules.company_form)
    assert first.exists() and second.name == first.stem + "_2.xlsx"


def test_check_file_has_one_row_per_check(rules, tmp_path):
    pdf = SAMPLES / "03_Invoice_DM_PDF" / "INV-2026-10231.pdf"
    result = extract(pdf, rules)
    checks = [Check("합계", "1,000.00", "1000.00", True), Check("1줄 금액", "2 x 3 = 6", "7", False),
              Check("2줄 금액", "", "", None)]
    path = write_check_file(tmp_path, result, checks, pdf.name)
    assert path.name == "검산_INV-2026-10231.xlsx"
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames == ["검산", "추출값", "품목"]  # 회사 양식 값을 주지 않으면 회사양식 시트는 없음
    rows = list(wb["검산"].iter_rows(values_only=True))
    assert rows[0][:2] == ("원본파일", pdf.name)
    assert rows[1] == ("항목", "원본 값", "결과 값", "일치")
    assert len(rows) == 2 + 3
    assert [r[3] for r in rows[2:]] == ["일치", "불일치", "해당 없음"]
    assert rows[3][:3] == ("1줄 금액", "2 x 3 = 6", "7")


def test_read_back_matches_company_form(rules, master, markup, tmp_path):
    result = read_excel(SAMPLES / "01_고객오더_엑셀" / CASES[0][0], rules)
    row = lookup(master, result.fields["고객명"])
    path = write_company_form(tmp_path, result, row, rules.company_form, markup)
    header, rows = company_rows(result, row, rules.company_form, markup)
    assert read_back(path, rules.company_form.sheet, header, rows) == []


def test_read_back_matches_extract_result(rules, tmp_path):
    pdf = SAMPLES / "03_Invoice_DM_PDF" / "INV-2026-10231.pdf"
    result = extract(pdf, rules)
    path = write_extract_result(tmp_path, result, pdf.name)
    header, rows = extract_result_rows(result, pdf.name)
    assert header == ["필드", "값", "비고"]
    assert rows[0][:2] == ["문서종류", result.doc_type.name]
    assert read_back(path, "추출결과", header, rows) == []


def test_read_back_reports_changed_cell(rules, master, markup, tmp_path):
    result = read_excel(SAMPLES / "01_고객오더_엑셀" / CASES[0][0], rules)
    row = lookup(master, result.fields["고객명"])
    path = write_company_form(tmp_path, result, row, rules.company_form, markup)
    header, rows = company_rows(result, row, rules.company_form, markup)
    wb = openpyxl.load_workbook(path)
    ws = wb[rules.company_form.sheet]
    column = header.index("Unit Price") + 1
    written = ws.cell(row=3, column=column).value
    ws.cell(row=3, column=column).value = 999
    wb.save(path)
    problems = read_back(path, rules.company_form.sheet, header, rows)
    assert len(problems) == 1
    assert problems[0].startswith("3행 Unit Price: 쓴 값 ") and "읽은 값 999" in problems[0]
    assert written != 999


def test_company_form_from_pdf_result(rules):
    result = extract(SAMPLES / "03_Invoice_DM_PDF" / "INV-2026-10231.pdf", rules)
    assert result.table
    header, rows = company_rows(result, None, rules.company_form)
    assert len(rows) == len(result.table)
    assert header[0] == "Order Date" and rows[0][header.index("Line No")] == 10


def _roundtrip(tmp_path, written, stored):
    """written 을 header/rows 로 넘기고, 파일에는 stored 가 들어 있을 때 read_back 결과."""
    path = tmp_path / "t.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "S"
    wb.active.append(["Unit Price"])
    wb.active.append([stored])
    wb.save(path)
    return read_back(path, "S", ["Unit Price"], [[written]])


def test_read_back_decimal_equal_by_value(tmp_path):
    assert _roundtrip(tmp_path, Decimal("12.50"), 12.5) == []
    assert _roundtrip(tmp_path, Decimal("12.50"), 12) == ["2행 Unit Price: 쓴 값 12.50, 읽은 값 12"]


def test_read_back_large_integers_must_match_exactly(tmp_path):
    assert _roundtrip(tmp_path, 9007199254740993, 9007199254740992) != []
    assert _roundtrip(tmp_path, 9007199254740992, 9007199254740992) == []  # 엑셀은 정수도 double 로 저장
    assert _roundtrip(tmp_path, Decimal("9007199254740993"), 9007199254740992) != []


def test_read_back_date_and_blank(tmp_path):
    assert _roundtrip(tmp_path, date(2026, 10, 28), datetime(2026, 10, 28)) == []
    assert _roundtrip(tmp_path, date(2026, 10, 28), datetime(2026, 10, 29)) != []
    assert _roundtrip(tmp_path, date(2026, 10, 28), datetime(2026, 10, 28, 15, 30)) != []
    assert _roundtrip(tmp_path, None, None) == []
    assert _roundtrip(tmp_path, "", None) == []
    assert _roundtrip(tmp_path, None, "x") == ["2행 Unit Price: 쓴 값 (빈칸), 읽은 값 x"]
    assert _roundtrip(tmp_path, "ab", "ab ") != []


def test_read_back_reports_header_row_count_and_missing_sheet(tmp_path):
    path = tmp_path / "t.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "S"
    wb.active.append(["A", "B"])
    wb.active.append([1, 2])
    wb.save(path)
    assert read_back(path, "S", ["A", "B"], [[1, 2]]) == []
    assert read_back(path, "S", ["A", "C"], [[1, 2]]) == ["머리글 2열: 쓴 값 C, 읽은 값 B"]
    assert len(read_back(path, "S", ["A", "B"], [[1, 2], [3, 4]])) == 1
    assert len(read_back(path, "S", ["A", "B"], [])) == 1
    assert read_back(path, "없는시트", ["A"], []) == ["시트 '없는시트' 가 없습니다"]
