"""PDF 6개 + 고객 엑셀 PO 2개를 함께 처리 → 장부(PO 포함)·문서별 결과 엑셀·정리 폴더 검증.

정답: samples/04_장부_엑셀/누적장부_기대결과.xlsx, samples/02_회사양식_변환/*.xlsx
"""
import shutil
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
import pytest

from app.__main__ import main
from app.init import init

REPO = Path(__file__).resolve().parent.parent
SAMPLES = REPO / "samples"
EXPECTED = SAMPLES / "04_장부_엑셀" / "누적장부_기대결과.xlsx"
SKIP = {"처리일시", "정리된경로"}


@pytest.fixture
def env(tmp_path):
    repo = tmp_path / "repo"
    (repo / "rules").mkdir(parents=True)
    shutil.copy(REPO / "rules" / "규칙.예시.yaml", repo / "rules" / "규칙.예시.yaml")
    work = init(work_dir=tmp_path / "work", repo_root=repo)
    return repo, work


def _drop_all(work):
    for f in list((SAMPLES / "03_Invoice_DM_PDF").glob("*.pdf")) + list((SAMPLES / "01_고객오더_엑셀").glob("*.xlsx")):
        shutil.copy(f, work / "받은문서" / f.name)


def _records(path, sheet):
    rows = list(openpyxl.load_workbook(path)[sheet].iter_rows(values_only=True))
    header = [h for h in rows[0] if h is not None]
    return [dict(zip(header, r)) for r in rows[1:] if any(v is not None for v in r)]


def _same(a, b):
    if hasattr(a, "year") and hasattr(b, "year"):
        return (a.year, a.month, a.day) == (b.year, b.month, b.day)
    try:
        return Decimal(str(a)) == Decimal(str(b))
    except (InvalidOperation, ValueError):
        return a == b


def test_pdf_and_excel_together(env, capsys):
    repo, work = env
    _drop_all(work)
    assert main([], repo_root=repo) == 0
    out = capsys.readouterr().out
    assert "성공 6 · 중복 0 · 확인필요 2" in out
    assert not list((work / "받은문서").glob("*.*"))
    acme_po = work / "정리문서" / "ACME Semiconductor Korea Co., Ltd" / "2026-09" / "PO_PO-ACME-2026-0417.xlsx"
    assert acme_po.exists()  # 정리 파일은 원래 확장자(.xlsx)
    outputs = sorted(p.name for p in (work / "출력").glob("*.xlsx"))
    assert outputs == sorted([
        "회사양식_SO입력_PO-ACME-2026-0417.xlsx", "회사양식_SO입력_HDC-PUR-26-0908.xlsx",
        "추출결과_INV-2026-10231.xlsx", "추출결과_INV-2026-10278.xlsx",
        "추출결과_DM-2026-0092.xlsx", "추출결과_DM-2026-0097.xlsx",
    ])


@pytest.mark.parametrize("sheet", ["Invoice", "DM", "PO"])
def test_ledger_matches_expected(env, sheet):
    repo, work = env
    _drop_all(work)
    main([], repo_root=repo)
    got = {r["원본파일"]: r for r in _records(work / "누적장부.xlsx", sheet)}
    expected = {r["원본파일"]: r for r in _records(EXPECTED, sheet)}
    assert sorted(got) == sorted(expected)
    for name, exp in expected.items():
        for column, value in exp.items():
            if column not in SKIP:
                assert _same(got[name][column], value), f"{sheet} {name} {column}: {got[name][column]!r} != {value!r}"
        assert Path(got[name]["정리된경로"]).name == Path(exp["정리된경로"]).name


def test_items_and_log_match_expected(env):
    repo, work = env
    _drop_all(work)
    main([], repo_root=repo)
    key = lambda r: (r["문서종류"], r["문서번호"], str(r["줄"]))  # noqa: E731
    got = sorted(_records(work / "누적장부.xlsx", "품목"), key=key)
    exp = sorted(_records(EXPECTED, "품목"), key=key)
    assert len(got) == len(exp) == 16
    for g, e in zip(got, exp):
        assert all(_same(g[c], v) for c, v in e.items()), (g, e)
    got_log = {r["파일"]: r["결과"] for r in _records(work / "누적장부.xlsx", "처리로그")}
    exp_log = {r["파일"]: r["결과"] for r in _records(EXPECTED, "처리로그")}
    assert got_log == exp_log


@pytest.mark.parametrize("order, form", [
    ("고객오더_ACME_PO-ACME-2026-0417.xlsx", "회사양식_SO입력_PO-ACME-2026-0417.xlsx"),
    ("고객오더_HANIL_HDC-PUR-26-0908.xlsx", "회사양식_SO입력_HDC-PUR-26-0908.xlsx"),
])
def test_company_form_output_matches_expected(env, order, form):
    repo, work = env
    shutil.copy(SAMPLES / "01_고객오더_엑셀" / order, work / "받은문서" / order)
    main([], repo_root=repo)
    got = list(openpyxl.load_workbook(work / "출력" / form)["SO_Upload"].iter_rows(values_only=True))
    exp = list(openpyxl.load_workbook(SAMPLES / "02_회사양식_변환" / form)["SO_Upload"].iter_rows(values_only=True))
    assert got[0] == exp[0] and len(got) == len(exp)
    assert all(_same(g, e) for gr, er in zip(got[1:], exp[1:]) for g, e in zip(gr, er))


def test_customer_not_in_master_goes_to_review(env, capsys):
    repo, work = env
    wb = openpyxl.load_workbook(work / "고객마스터.xlsx")
    ws = wb.active
    ws.delete_rows(2)  # ACME 줄을 지움
    wb.save(work / "고객마스터.xlsx")
    order = "고객오더_ACME_PO-ACME-2026-0417.xlsx"
    shutil.copy(SAMPLES / "01_고객오더_엑셀" / order, work / "받은문서" / order)
    main([], repo_root=repo)
    assert "확인필요 1" in capsys.readouterr().out
    note = (work / "확인필요" / "고객오더_ACME_PO-ACME-2026-0417.txt").read_text(encoding="utf-8")
    assert "고객마스터" in note and "ACME" in note
    assert not (work / "출력" / "회사양식_SO입력_PO-ACME-2026-0417.xlsx").exists()


def test_old_xls_goes_to_review_with_hint(env, capsys):
    repo, work = env
    (work / "받은문서" / "옛주문.xls").write_bytes(b"\xd0\xcf\x11\xe0 old excel")
    main([], repo_root=repo)
    assert "확인필요 1" in capsys.readouterr().out
    assert ".xlsx" in (work / "확인필요" / "옛주문.txt").read_text(encoding="utf-8")


def test_excel_lock_file_is_ignored(env, capsys):
    repo, work = env
    (work / "받은문서" / "~$주문.xlsx").write_bytes(b"lock")
    assert main([], repo_root=repo) == 0
    assert (work / "받은문서" / "~$주문.xlsx").exists()
    assert "없습니다" in capsys.readouterr().out
