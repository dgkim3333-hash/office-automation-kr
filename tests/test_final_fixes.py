"""최종 검토 수정의 회귀 테스트: 잠금 재시도 알림 · 검산 열 이름 · 출력방식 검사 · 장부 소수 표시 · 방식 변경 승인 ·
run_job 기본 승인 · 대조표 글자 값 · 문서 안내 · outputs 공백."""
import os
import re
import shutil
import time
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
import yaml

import app.__main__ as cli
import app.job_run as job_run
from app import config, schedule
from app.__main__ import main
from app.approval import approval_gate, approve
from app.excel_extract import read_excel
from app.job_run import job_lock, load_records, run_job
from app.jobs import JOB_FILE, JobError, create_job, find_job, load_job, save_job
from app.ledger import _append
from app.markup import load_markup, write_example_markup
from app.master import load_master, lookup, write_example
from app.outputs import company_rows
from app.rules import RuleError, load_rules, rules_from_text
from app.verify import verify

REPO = Path(__file__).resolve().parent.parent
EXAMPLE_RULES = REPO / "rules" / "규칙.예시.yaml"
ORDERS = REPO / "samples" / "01_고객오더_엑셀"
PDFS = REPO / "samples" / "03_Invoice_DM_PDF"
ACME_PO = ORDERS / "고객오더_ACME_PO-ACME-2026-0417.xlsx"
HANIL_PO = ORDERS / "고객오더_HANIL_HDC-PUR-26-0908.xlsx"
INV = PDFS / "INV-2026-10231.pdf"
COMPANY_ACME = "회사양식_SO입력_PO-ACME-2026-0417.xlsx"
TITLE = "일일 정리"
LOCKED_TEXT = "다른 일로 잠겨 있어 오늘 자동 실행을 못 했습니다. 지금 실행을 눌러 주세요"


def _drop(src, folder, name=None, age=3600):
    dst = Path(folder) / (name or Path(src).name)
    shutil.copy(src, dst)
    t = time.time() - age
    os.utime(dst, (t, t))
    return dst


def _files(folder):
    folder = Path(folder)
    if not folder.is_dir():
        return []
    return sorted(p.name for p in folder.iterdir() if p.is_file() and not p.name.endswith(".bak.xlsx"))


# --- 1. 자동 실행 · 놓친 실행이 잠금을 만나면 기다렸다가, 그래도 잠겨 있으면 알린다 ---

@pytest.fixture
def cli_env(tmp_path, capsys):
    repo = tmp_path / "repo"
    shutil.copytree(REPO / "rules", repo / "rules")
    work, inp, out = tmp_path / "work", tmp_path / "in", tmp_path / "out"
    inp.mkdir()
    out.mkdir()
    assert main(["init", str(work)], repo_root=repo) == 0
    capsys.readouterr()
    assert main(["job", "new", "--제목", TITLE, "--입력", str(inp), "--출력", str(out), "--방식", "누적장부"],
                repo_root=repo) == 0
    job = find_job(config.work_dir(repo), TITLE)
    assert main(["job", "resume", str(job.number)], repo_root=repo) == 0
    shutil.copy(EXAMPLE_RULES, job.rules_path)
    job = load_job(job.folder)
    job.run_time = "00:00"
    save_job(job)
    capsys.readouterr()
    return repo, work.resolve(), job


@pytest.fixture
def failures(monkeypatch):
    calls = []
    monkeypatch.setattr(cli, "notify_failure",
                        lambda what, message, folder=None: calls.append((what, message, folder)))
    return calls


class _Holder:
    """다른 일(승인 · job set 등)이 쥔 잠금 흉내. release_after 번째 기다림 뒤에 놓는다(None 이면 안 놓음)."""

    def __init__(self, job, release_after=None, then=None):
        self.cm = job_lock(job)
        self.cm.__enter__()
        self.held = True
        self.release_after = release_after
        self.then = then
        self.sleeps = []

    def sleep(self, seconds):
        self.sleeps.append(seconds)
        if self.held and self.release_after is not None and len(self.sleeps) >= self.release_after:
            if self.then:
                self.then()
            self.release()

    def release(self):
        if self.held:
            self.held = False
            self.cm.__exit__(None, None, None)


def _today():
    return datetime.now().date().isoformat()


@pytest.mark.parametrize("command", [["run", "1", "--자동"], ["catch-up"]])
def test_auto_run_waits_for_lock_then_runs(cli_env, failures, monkeypatch, command):
    repo, work, job = cli_env
    holder = _Holder(job, release_after=2)
    monkeypatch.setattr(schedule, "_sleep", holder.sleep)
    try:
        assert main(command, repo_root=repo) == 0
    finally:
        holder.release()
    assert holder.sleeps == [schedule.LOCK_RETRY_SECONDS] * 2
    assert failures == []
    assert load_job(job.folder).last_run == _today()


@pytest.mark.parametrize("command", [["run", "1", "--자동"], ["catch-up"]])
def test_auto_run_still_locked_after_limit_notifies(cli_env, failures, monkeypatch, command):
    repo, work, job = cli_env
    monkeypatch.setattr(schedule, "LOCK_RETRY_SECONDS", 30)
    monkeypatch.setattr(schedule, "LOCK_RETRY_LIMIT_SECONDS", 15 * 60)
    holder = _Holder(job)
    monkeypatch.setattr(schedule, "_sleep", holder.sleep)
    try:
        assert main(command, repo_root=repo) != 0
    finally:
        holder.release()
    assert holder.sleeps == [30] * 30  # 30초마다 15분
    assert len(failures) == 1 and LOCKED_TEXT in failures[0][1] and TITLE in failures[0][1]
    assert LOCKED_TEXT in (job.folder / "실행기록.txt").read_text(encoding="utf-8")
    assert load_job(job.folder).last_run != _today()


def test_auto_run_does_not_run_if_paused_while_waiting(cli_env, failures, monkeypatch):
    """기다리는 동안 job pause 가 잠금을 쥐고 멈춤으로 바꿨으면 실행하지 않는다."""
    repo, work, job = cli_env

    def pause():
        paused = load_job(job.folder)
        paused.paused = True
        save_job(paused)

    holder = _Holder(job, release_after=1, then=pause)
    monkeypatch.setattr(schedule, "_sleep", holder.sleep)
    try:
        assert main(["run", "1", "--자동"], repo_root=repo) == 0
    finally:
        holder.release()
    assert failures == []
    assert load_job(job.folder).last_run != _today()


# --- 2. 표의 열 이름이 수량 · 단가 · 금액 이 아니면 검산 열 이름을 적게 한다 ---

TABLE_RULES = r'''
문서종류:
  Invoice:
    판별키워드: ["Invoice No"]
    번호필드: 번호
    필드:
      번호: {패턴: "Invoice No[.:]?\\s*([A-Z0-9-]+)"}
      합계금액: {패턴: "Subtotal[.:]?\\s*\\$?([\\d,]+\\.\\d{2})", 형식: 금액}
    표:
      시작키워드: "Line Item Description"
      끝키워드: "Subtotal"
      행패턴: "^(\\d+)\\s+(\\S+)\\s+(.+?)\\s+(\\d+)\\s+([\\d,]+\\.\\d{2})\\s+([\\d,]+\\.\\d{2})$"
      열: [줄, 품번, 품명, COLUMNS]
      금액열: MONEY
      수량열: [QTY]
CHECK
    파일명: "Invoice_{번호}.pdf"
'''


def _table_rules(columns="Qty, Price, Amount", money="[Price, Amount]", qty="Qty", check=""):
    return (TABLE_RULES.replace("COLUMNS", columns).replace("MONEY", money).replace("QTY", qty)
            .replace("CHECK", check))


LINE_RULE_ERROR = "검산: 에 수량열 · 단가열 · 금액열 을 적으세요(표의 열 이름이 수량 · 단가 · 금액 이 아닙니다)"


def test_rules_require_check_columns_when_table_names_differ():
    with pytest.raises(RuleError, match=re.escape(LINE_RULE_ERROR)):
        rules_from_text(_table_rules())
    with pytest.raises(RuleError, match="검산.*금액열"):  # 합계만 적고 열 이름은 안 적음(있던 검사가 먼저 잡음)
        rules_from_text(_table_rules(check="    검산: {합계: 합계금액}"))
    with pytest.raises(RuleError, match=re.escape(LINE_RULE_ERROR)):  # 금액열만 적음: 수량 · 단가는 기본 이름
        rules_from_text(_table_rules(check="    검산: {합계: 합계금액, 금액열: Amount}"))
    named = rules_from_text(_table_rules(check="    검산: {수량열: Qty, 단가열: Price, 금액열: Amount}"))
    spec = named.doc_types[0].check
    assert (spec.quantity, spec.price, spec.amount) == ("Qty", "Price", "Amount")


def test_rules_allow_default_names_and_tables_without_unit_price():
    rules_from_text(_table_rules(columns="수량, 단가, 금액", money="[단가, 금액]", qty="수량"))
    # 단가가 없는 문서(금액 열 하나): 줄 검산은 할 수 없고 합계 검산만 한다
    text = _table_rules(columns="수량, 비고, 금액", money="[금액]", qty="수량")
    rules_from_text(text)


def test_excel_items_require_check_columns_when_names_differ():
    data = yaml.safe_load(EXAMPLE_RULES.read_text(encoding="utf-8"))
    items = data["엑셀양식"]["ACME"]["품목"]
    items["열"] = {"줄": "A", "품번": "B", "품명": "C", "Qty": "D", "단위": "E", "Price": "F", "Amount": "G"}
    items["금액열"], items["수량열"] = ["Price", "Amount"], ["Qty"]
    with pytest.raises(RuleError, match="검산.*금액열"):  # PO 의 합계 검산: 있던 검사가 금액열 이름을 묻는다
        rules_from_text(yaml.safe_dump(data, allow_unicode=True))
    del data["문서종류"]["PO"]["검산"]
    with pytest.raises(RuleError, match=re.escape(LINE_RULE_ERROR)):
        rules_from_text(yaml.safe_dump(data, allow_unicode=True))
    data["문서종류"]["PO"]["검산"] = {"합계": "합계금액", "수량열": "Qty", "단가열": "Price", "금액열": "Amount"}
    data["회사양식"]["마크업"] = {"품번열": "품번", "단가열": "Price", "수량열": "Qty"}
    data["엑셀양식"]["HANIL"]["품목"].update({k: items[k] for k in ("열", "금액열", "수량열")})  # 같은 문서종류
    rules_from_text(yaml.safe_dump(data, allow_unicode=True))  # 열 이름을 적으면 됨


def test_example_rules_still_load_and_po_checks_total():
    rules = load_rules(EXAMPLE_RULES)
    assert rules.get("PO").check.total == "합계금액"
    for po in (ACME_PO, HANIL_PO):
        result = read_excel(po, rules)
        checks = {c.item: c.ok for c in verify(result, po)}
        assert checks["품목 합 = 합계"] is True
        assert all(ok is not False for ok in checks.values())


def test_check_warns_when_every_line_check_is_not_applicable(tmp_path, capsys):
    rules = rules_from_text(_table_rules(columns="수량, 단가표시, 금액", money="[금액]", qty="수량"))
    pdf = tmp_path / INV.name
    shutil.copy(INV, pdf)
    assert cli.check(pdf, rules) == 0
    shown = capsys.readouterr().out
    assert "해당 없음" in shown
    assert "줄 검산(수량 x 단가 = 금액)이 모두 해당 없음" in shown
    assert "단가가 없는 문서가 아니면" in shown


def test_check_does_not_warn_when_line_checks_ran(tmp_path, capsys):
    rules = load_rules(EXAMPLE_RULES)
    pdf = tmp_path / INV.name
    shutil.copy(INV, pdf)
    assert cli.check(pdf, rules) == 0
    assert "모두 해당 없음" not in capsys.readouterr().out


def test_work_order_says_line_not_applicable_is_not_ok():
    text = (REPO / "작업지시서.md").read_text(encoding="utf-8")
    assert "단가가 없는 문서가 아니면 n줄 금액 해당 없음은 안 됩니다" in text


# --- 공통 작업 환경 ---

@pytest.fixture
def env(tmp_path):
    work, inp, out = tmp_path / "work", tmp_path / "in", tmp_path / "out"
    for d in (work, inp, out):
        d.mkdir()
    write_example(work / "고객마스터.xlsx")
    write_example_markup(work / "마크업표.xlsx")
    return work, inp, out


def _job(work, inp, out, mode):
    job = create_job(work, TITLE, inp, out, mode)
    shutil.copy(EXAMPLE_RULES, job.rules_path)
    return job


# --- 3. 작업설정.yaml 의 출력방식이 잘못되면 성공으로 기록하지 않는다 ---

def test_load_job_rejects_bad_output_mode(env):
    work, inp, out = env
    job = _job(work, inp, out, "누적장부")
    data = yaml.safe_load((job.folder / JOB_FILE).read_text(encoding="utf-8"))
    data["출력방식"] = "회사 양식"
    (job.folder / JOB_FILE).write_text(yaml.safe_dump(data, allow_unicode=True), encoding="utf-8")
    with pytest.raises(JobError, match="출력 방식 '회사 양식'"):
        load_job(job.folder)


def test_auto_run_with_bad_output_mode_notifies(cli_env, failures):
    repo, work, job = cli_env
    data = yaml.safe_load((job.folder / JOB_FILE).read_text(encoding="utf-8"))
    data["출력방식"] = "장부"
    (job.folder / JOB_FILE).write_text(yaml.safe_dump(data, allow_unicode=True), encoding="utf-8")
    assert main(["run", "1", "--자동"], repo_root=repo) != 0
    assert len(failures) == 1 and "출력 방식 '장부'" in failures[0][1]


def test_write_with_bad_mode_is_not_success(env):
    """잠금 뒤 다시 읽기를 거치지 않고 잘못된 방식이 들어와도 결과 없이 성공으로 기록하지 않는다."""
    work, inp, out = env
    job = _job(work, inp, out, "누적장부")
    _drop(INV, inp)
    job.output_mode = "엉뚱"
    report = job_run._Run(job, work, datetime.now().replace(microsecond=0), inp, None, None).run()
    assert report.counts["성공"] == 0 and report.counts["확인필요"] == 1
    assert _files(out) == []
    assert [r["상태"] for r in load_records(job)] == ["확인필요"]


# --- 4. 누적장부의 소수 칸은 소수 자리를 다 보인다 ---

def test_ledger_decimal_cells_show_all_places():
    ws = openpyxl.Workbook().active
    ws.append(["a", "b", "c"])
    _append(ws, {"a": Decimal("0.3745"), "b": Decimal("12.50"), "c": Decimal("3")})
    assert ws.cell(2, 1).number_format == "#,##0.0000"
    assert ws.cell(2, 2).number_format == "#,##0.00"
    assert ws.cell(2, 3).number_format == "#,##0.00"


# --- 5. 방식을 바꿔 회사 양식을 새로 내면 승인 수를 처음부터 센다 ---

def _approve_ten_in_current_mode(job, rules_path, po=ACME_PO):
    """지금 방식으로 그 양식 승인 수를 10 으로 둔다(지문은 승인 게이트가 저장한 그대로)."""
    rules = load_rules(rules_path)
    result = read_excel(po, rules)
    approval_gate(job, rules, result)
    job.approvals["ACME"] = 10
    save_job(job)


def test_ledger_mode_with_ten_approvals_outputs(env):
    work, inp, out = env
    job = _job(work, inp, out, "누적장부")
    _approve_ten_in_current_mode(job, job.rules_path)
    _drop(ACME_PO, inp)
    assert run_job(job, work, gate=approval_gate).counts["성공"] == 1


@pytest.mark.parametrize("new_mode", ["둘다", "회사양식"])
def test_mode_change_to_company_form_needs_approval_again(env, new_mode):
    work, inp, out = env
    job = _job(work, inp, out, "누적장부")
    _approve_ten_in_current_mode(job, job.rules_path)
    job.output_mode = new_mode  # job set --방식 과 같은 저장
    save_job(job)
    _drop(ACME_PO, inp)
    report = run_job(job, work, gate=approval_gate)
    assert report.counts["승인대기"] == 1 and report.counts["성공"] == 0
    assert COMPANY_ACME not in _files(out)
    assert load_job(job.folder).approvals["ACME"] == 0


def test_job_set_mode_change_then_run_goes_to_waiting(cli_env, capsys):
    repo, work, job = cli_env
    write_example(work / "고객마스터.xlsx")
    write_example_markup(work / "마크업표.xlsx")
    _approve_ten_in_current_mode(job, job.rules_path)
    assert main(["job", "set", "1", "--방식", "둘다"], repo_root=repo) == 0
    _drop(ACME_PO, job.input_dir)
    report = run_job(load_job(job.folder), work, gate=approval_gate)
    assert report.counts["승인대기"] == 1
    assert COMPANY_ACME not in _files(job.output_dir)


def _without_markup_columns():
    """예시 규칙에서 회사양식의 마크업.* 열을 뺀 글자(마크업 지문이 대조표를 대신 새로 만들지 않게)."""
    data = yaml.safe_load(EXAMPLE_RULES.read_text(encoding="utf-8"))
    form = data["회사양식"]
    form["열"] = [c for c in form["열"] if not str(c["값"]).startswith("마크업.")]
    return yaml.safe_dump(data, allow_unicode=True, sort_keys=False)


def test_waiting_copy_made_before_mode_change_is_not_output_on_approve(env):
    """누적장부 방식에서 만든 대조표(회사 양식 값 없음)로 회사 양식을 내보내지 않는다."""
    work, inp, out = env
    job = _job(work, inp, out, "누적장부")
    job.rules_path.write_text(_without_markup_columns(), encoding="utf-8")
    _drop(ACME_PO, inp)
    assert run_job(job, work, gate=approval_gate).counts["승인대기"] == 1
    job = load_job(job.folder)
    job.output_mode = "둘다"
    save_job(job)
    report = approve(job, work)
    assert report.counts["성공"] == 0 and report.counts["승인대기"] == 1
    assert "출력 방식이 바뀌어 대조표를 새로 만들었습니다" in report.message
    assert COMPANY_ACME not in _files(out)
    report = approve(load_job(job.folder), work)  # 새 대조표를 본 뒤 승인
    assert report.counts["성공"] == 1
    assert COMPANY_ACME in _files(out)


# --- 6. run_job 의 기본은 승인 게이트 ---

def test_run_job_default_gate_is_approval(env):
    work, inp, out = env
    job = _job(work, inp, out, "누적장부")
    _drop(INV, inp)
    report = run_job(job, work)
    assert report.counts["승인대기"] == 1 and report.counts["성공"] == 0


def test_run_job_gate_none_means_no_approval(env):
    work, inp, out = env
    job = _job(work, inp, out, "누적장부")
    _drop(INV, inp)
    assert run_job(job, work, gate=None).counts["성공"] == 1


# --- 7. 대조표에 글자 값 · 회사 양식 값까지 ---

def _sheet(path, name):
    return [list(r) for r in openpyxl.load_workbook(path)[name].iter_rows(values_only=True)]


def _expected_company(work, po):
    rules = load_rules(EXAMPLE_RULES)
    result = read_excel(po, rules)
    master_row = lookup(load_master(work / "고객마스터.xlsx"), result.fields["고객명"])
    return result, company_rows(result, master_row, rules.company_form, load_markup(work / "마크업표.xlsx"))


def test_waiting_check_shows_text_fields_items_and_company_rows(env):
    work, inp, out = env
    job = _job(work, inp, out, "둘다")
    _drop(ACME_PO, inp)
    assert run_job(job, work, gate=approval_gate).counts["승인대기"] == 1
    [check] = sorted((job.folder / "승인대기" / "대조표").glob("*.xlsx"))
    wb = openpyxl.load_workbook(check)
    assert wb.sheetnames[0] == "검산"
    assert wb["검산"]["A1"].value == "원본파일"  # 승인이 짝짓는 칸은 그대로
    result, (header, rows) = _expected_company(work, ACME_PO)

    fields = {r[0]: r[1] for r in _sheet(check, "추출값")[1:]}
    assert fields["납품처"] == result.fields["납품처"] and fields["인도조건"] == result.fields["인도조건"]
    items = _sheet(check, "품목")
    assert "품명" in items[0]
    assert [r[items[0].index("품명")] for r in items[1:]] == [row["품명"] for row in result.table]

    company = _sheet(check, "회사양식")
    assert company[0] == header
    code = header.index("Customer Code")
    assert [r[code] for r in company[1:]] == [r[code] for r in rows] and rows[0][code]
    terms = header.index("Payment Terms")
    assert [r[terms] for r in company[1:]] == [r[terms] for r in rows] and rows[0][terms]


def test_real_run_check_file_has_same_sheets(env):
    work, inp, out = env
    job = _job(work, inp, out, "둘다")
    _drop(ACME_PO, inp)
    assert run_job(job, work, gate=None).counts["성공"] == 1
    [check] = sorted((job.folder / "검산").glob("*.xlsx"))
    assert openpyxl.load_workbook(check).sheetnames == ["검산", "추출값", "품목", "회사양식"]
    _, (header, _) = _expected_company(work, ACME_PO)
    assert _sheet(check, "회사양식")[0] == header


def test_ledger_only_check_file_has_no_company_sheet(env):
    work, inp, out = env
    job = _job(work, inp, out, "누적장부")
    _drop(INV, inp)
    assert run_job(job, work, gate=approval_gate).counts["승인대기"] == 1
    [check] = sorted((job.folder / "승인대기" / "대조표").glob("*.xlsx"))
    names = openpyxl.load_workbook(check).sheetnames
    assert names == ["검산", "추출값", "품목"]
    fields = {r[0]: r[1] for r in _sheet(check, "추출값")[1:]}
    assert fields["인도조건"] and fields["운송사"]


# --- 8. 마크업 사유는 Claude 에게 붙여 넣지 않게 ---

MARKUP_WARNING = "마크업 검산 불일치 · 마크업표 관련 사유는 붙이지 마세요(마크업% · 본사 단가가 들어 있습니다). 「마크업 불일치」라고 말로만 알려 주세요"


@pytest.mark.parametrize("doc", ["05_내_양식_가르치기.md", "99_막혔을_때.md"])
def test_docs_say_not_to_paste_markup_reasons(doc):
    assert MARKUP_WARNING in (REPO / doc).read_text(encoding="utf-8")


# --- 9. outputs.py 의 긴 공백 ---

def test_outputs_has_no_stray_long_whitespace_inside_lines():
    for n, line in enumerate((REPO / "app" / "outputs.py").read_text(encoding="utf-8").splitlines(), 1):
        assert not re.search(r"\S {8,}\S", line), f"outputs.py {n}줄"
