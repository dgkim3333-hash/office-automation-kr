"""고객마스터에 같은 고객 줄이 두 개 이상이면 어느 줄이 맞는지 모르므로 그 고객 문서를 확인필요로 보내고 알린다."""
import os
import shutil
import time
from pathlib import Path

import openpyxl
import pytest

import app.notify as notify
from app.__main__ import main
from app.approval import approval_gate, approve
from app.job_run import run_job
from app.jobs import create_job, load_job
from app.markup import write_example_markup
from app.master import MasterError, load_master, lookup, write_example

REPO = Path(__file__).resolve().parent.parent
EXAMPLE_RULES = REPO / "rules" / "규칙.예시.yaml"
ORDERS = REPO / "samples" / "01_고객오더_엑셀"
ACME_PO = ORDERS / "고객오더_ACME_PO-ACME-2026-0417.xlsx"
HANIL_PO = ORDERS / "고객오더_HANIL_HDC-PUR-26-0908.xlsx"
ACME = "ACME Semiconductor Korea Co., Ltd."
COMPANY_ACME = "회사양식_SO입력_PO-ACME-2026-0417.xlsx"
COMPANY_HANIL = "회사양식_SO입력_HDC-PUR-26-0908.xlsx"
WARNING = "고객마스터에 'ACME Semiconductor Korea Co., Ltd.' 줄이 2개 있습니다(2 · 4번째 줄)"


def _add_master_row(path, row):
    wb = openpyxl.load_workbook(path)
    wb.worksheets[0].append(row)
    wb.save(path)


def _drop(src, folder):
    dst = Path(folder) / Path(src).name
    shutil.copy(src, dst)
    t = time.time() - 3600
    os.utime(dst, (t, t))
    return dst


def _files(folder):
    return sorted(p.name for p in Path(folder).iterdir() if p.is_file()) if Path(folder).is_dir() else []


@pytest.fixture
def env(tmp_path):
    work, inp, out = tmp_path / "work", tmp_path / "in", tmp_path / "out"
    for d in (work, inp, out):
        d.mkdir()
    write_example(work / "고객마스터.xlsx")
    write_example_markup(work / "마크업표.xlsx")
    # 대소문자 · 띄어쓰기만 다른 같은 고객 줄(고객코드가 다름)
    _add_master_row(work / "고객마스터.xlsx", ["acme  semiconductor korea co., ltd.", "TEST-001", "Net 30", None])
    job = create_job(work, "일일 정리", inp, out, "둘다")
    shutil.copy(EXAMPLE_RULES, job.rules_path)
    return work, inp, out, job


def test_lookup_of_duplicated_customer_names_both_lines(env):
    work, *_ = env
    master = load_master(work / "고객마스터.xlsx")
    with pytest.raises(MasterError) as e:
        lookup(master, ACME)
    assert WARNING in str(e.value)
    assert lookup(master, "Hanil DataCenter Co., Ltd.")["고객코드"] == "C-KR-0522"  # 다른 고객은 그대로


def test_duplicated_customer_goes_to_review_with_warning(env):
    work, inp, out, job = env
    _drop(ACME_PO, inp)
    report = run_job(job, work, gate=None)
    assert report.counts["확인필요"] == 1 and report.counts["성공"] == 0
    assert COMPANY_ACME not in _files(out)
    assert WARNING in report.message
    [note] = (Path(job.folder) / "확인필요").glob("*.txt")
    assert WARNING in note.read_text(encoding="utf-8")


def test_notification_carries_duplicate_warning(env):
    work, inp, out, job = env
    _drop(ACME_PO, inp)
    report = run_job(job, work, gate=None)
    text = notify.report_text(job, report)
    assert "확인필요 1건" in text
    assert WARNING in text


def test_other_customer_still_outputs(env):
    work, inp, out, job = env
    _drop(HANIL_PO, inp)
    report = run_job(job, work, gate=None)
    assert report.counts["성공"] == 1 and COMPANY_HANIL in _files(out)
    assert "고객마스터에" not in notify.report_text(job, report)


def test_fixing_master_reprocesses_same_day(env):
    work, inp, out, job = env
    _drop(ACME_PO, inp)
    assert run_job(job, work, gate=None).counts["확인필요"] == 1
    wb = openpyxl.load_workbook(work / "고객마스터.xlsx")
    wb.worksheets[0].delete_rows(4)  # 겹친 줄을 지움
    wb.save(work / "고객마스터.xlsx")
    report = run_job(load_job(job.folder), work, gate=None)
    assert report.counts["성공"] == 1 and COMPANY_ACME in _files(out)


def test_approve_with_duplicated_customer_keeps_waiting(env):
    work, inp, out, job = env
    _drop(ACME_PO, inp)
    master = work / "고객마스터.xlsx"
    keep = master.read_bytes()
    wb = openpyxl.load_workbook(master)
    wb.worksheets[0].delete_rows(4)
    wb.save(master)
    assert run_job(job, work, gate=approval_gate).counts["승인대기"] == 1
    master.write_bytes(keep)  # 승인 전에 같은 고객 줄이 하나 더 생김
    report = approve(load_job(job.folder), work)
    assert report.counts["성공"] == 0 and COMPANY_ACME not in _files(out)
    assert WARNING in notify.report_text(job, report)


def test_check_shows_duplicate_warning_not_missing(env, capsys, tmp_path):
    work, inp, out, job = env
    repo = tmp_path / "repo"
    shutil.copytree(REPO / "rules", repo / "rules")
    (repo / "작업폴더.txt").write_text(str(work) + "\n", encoding="utf-8")
    main(["check", str(ACME_PO), "--작업", str(job.number)], repo_root=repo)
    printed = capsys.readouterr().out
    assert f"고객마스터: {WARNING}" in printed
    assert "고객마스터: 없음" not in printed
    assert "TEST-001" not in printed and "C-KR-0417" not in printed  # 고객코드는 보여 주지 않는다
