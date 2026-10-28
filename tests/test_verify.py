"""검산: 원문 대조 · 소수 자리 · 줄 검산 · 합계 검산."""
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
import yaml
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

from app.excel_extract import read_excel
from app.extract import ExtractResult, extract
from app.rules import load_rules
from app.verify import Check, failed, verify

REPO = Path(__file__).resolve().parent.parent
EXAMPLE = REPO / "rules" / "규칙.예시.yaml"
PDFS = REPO / "samples" / "03_Invoice_DM_PDF"
ORDERS = REPO / "samples" / "01_고객오더_엑셀"

AMOUNT = "\\$?([\\d,]+\\.\\d+)"
TEST_RULES = {
    "문서종류": {
        # 표·검산 없는 짧은 문서
        "Memo": {
            "판별키워드": ["Memo No"],
            "번호필드": "메모번호",
            "필드": {
                "메모번호": {"패턴": "Memo No[.:]?\\s*([A-Z0-9-]+)"},
                "발행일": {"패턴": "Date[.:]?\\s*(\\d{1,2}/\\d{1,2}/\\d{4})", "형식": "날짜"},
                "합계금액": {"패턴": "Total[.:]?\\s*" + AMOUNT, "형식": "금액"},
            },
            "파일명": "Memo_{메모번호}.pdf",
        },
        # 품목 표 + 소계 · 운임 · 합계
        "Bill": {
            "판별키워드": ["Bill No"],
            "번호필드": "청구번호",
            "필드": {
                "청구번호": {"패턴": "Bill No[.:]?\\s*([A-Z0-9-]+)"},
                "소계": {"패턴": "Subtotal[.:]?\\s*" + AMOUNT, "형식": "금액"},
                "운임": {"패턴": "Freight[.:]?\\s*" + AMOUNT, "형식": "금액"},
                "합계금액": {"패턴": "\\nTotal[.:]?\\s*" + AMOUNT, "형식": "금액"},
            },
            "표": {
                "시작키워드": "Line Item",
                "끝키워드": "Subtotal",
                "행패턴": "^(\\d+)\\s+(\\S+)\\s+(\\d+)\\s+([\\d,]+\\.\\d+)\\s+([\\d,]+\\.\\d+)$",
                "열": ["줄", "품번", "수량", "단가", "금액"],
                "금액열": ["단가", "금액"],
                "수량열": ["수량"],
            },
            "검산": {"소계": "소계", "부대비용": ["운임"], "합계": "합계금액"},
            "파일명": "Bill_{청구번호}.pdf",
        },
    }
}


@pytest.fixture(scope="module")
def rules(tmp_path_factory):
    path = tmp_path_factory.mktemp("rules") / "규칙.yaml"
    path.write_text(yaml.safe_dump(TEST_RULES, allow_unicode=True), encoding="utf-8")
    return load_rules(path)


@pytest.fixture(scope="module")
def example_rules():
    return load_rules(EXAMPLE)


def _pdf(path, lines):
    c = canvas.Canvas(str(path), pagesize=letter, invariant=1)
    c.setFont("Helvetica", 11)
    y = 720
    for line in lines:
        c.drawString(72, y, line)
        y -= 16
    c.showPage()
    c.save()
    return path


def _memo(tmp_path, total="$1,234.50"):
    return _pdf(tmp_path / "memo.pdf", ["MEMO", "Memo No: M-2026-01", "Date: 10/15/2026", f"Total: {total}"])


def _bill(tmp_path, items, subtotal, freight, total):
    lines = ["BILL", "Bill No: B-2026-07", "Line Item Part Qty Price Amount"]
    lines += [f"{n} P-{n} {q} {p} {a}" for n, (q, p, a) in enumerate(items, 1)]
    lines += [f"Subtotal: ${subtotal}", f"Freight: ${freight}", f"Total: ${total}"]
    return _pdf(tmp_path / "bill.pdf", lines)


def _items(checks):
    return [c.item for c in failed(checks)]


# --- 예시 샘플은 깨끗이 통과 ---

@pytest.mark.parametrize("name", ["INV-2026-10231.pdf", "INV-2026-10278.pdf", "DM-2026-0092.pdf", "DM-2026-0097.pdf"])
def test_sample_pdfs_pass_all_checks(example_rules, name):
    pdf = PDFS / name
    checks = verify(extract(pdf, example_rules), pdf)
    assert failed(checks) == []
    assert any(c.item == "합계금액" and c.ok for c in checks)


def test_sample_invoice_passes_all_checks(example_rules):
    pdf = PDFS / "INV-2026-10231.pdf"
    checks = verify(extract(pdf, example_rules), pdf)
    assert failed(checks) == []
    by_item = {c.item: c for c in checks}
    assert by_item["품목 합 = 소계"].ok is True
    assert by_item["소계 + 부대비용 = 합계"].ok is True
    assert by_item["1줄 금액"].ok is True
    assert by_item["인보이스번호"].ok is True
    assert by_item["발행일"].ok is True
    assert by_item["합계금액"].source == "115,440.00" and by_item["합계금액"].output == "115440.00"


@pytest.mark.parametrize("name", ["고객오더_ACME_PO-ACME-2026-0417.xlsx", "고객오더_HANIL_HDC-PUR-26-0908.xlsx"])
def test_sample_excel_orders_pass(example_rules, name):
    path = ORDERS / name
    checks = verify(read_excel(path, example_rules), path)
    assert failed(checks) == []
    assert any(c.item == "1줄 금액" and c.ok for c in checks)


# --- 원문 대조 ---

def test_wrong_digit_or_decimal_fails_number_match(tmp_path, rules):
    pdf = _memo(tmp_path)
    result = extract(pdf, rules)
    assert failed(verify(result, pdf)) == []
    result.fields["합계금액"] = Decimal("12345.00")  # 원본 1234.50
    assert _items(verify(result, pdf)) == ["합계금액"]


def test_shifted_decimal_point_fails(tmp_path, rules):
    pdf = _memo(tmp_path)
    result = extract(pdf, rules)
    result.fields["합계금액"] = Decimal("123.450")
    assert _items(verify(result, pdf)) == ["합계금액"]


def test_same_value_different_notation_passes(tmp_path, rules):
    pdf = _memo(tmp_path, "$1,234.50")
    result = extract(pdf, rules)
    assert result.fields["합계금액"] == Decimal("1234.50")
    checks = verify(result, pdf)
    assert failed(checks) == []
    total = next(c for c in checks if c.item == "합계금액")
    assert total.ok is True and total.source == "1,234.50"


def test_decimal_places_must_match_pdf_token(tmp_path, rules):
    pdf = _memo(tmp_path, "$12.50")
    result = extract(pdf, rules)
    result.fields["합계금액"] = Decimal("12.5")
    assert _items(verify(result, pdf)) == ["합계금액 소수 자리"]


def test_wrong_date_fails(tmp_path, rules):
    pdf = _memo(tmp_path)
    result = extract(pdf, rules)
    result.fields["발행일"] = date(2026, 10, 16)
    assert _items(verify(result, pdf)) == ["발행일"]


def test_number_field_must_be_in_source_text(tmp_path, rules):
    pdf = _memo(tmp_path)
    result = extract(pdf, rules)
    result.fields["메모번호"] = "M-2026-07"
    assert _items(verify(result, pdf)) == ["메모번호"]


@pytest.mark.parametrize("wrong", ["INV-2026-1023", "2026", "INV-2026", "10231"])
def test_number_field_part_of_longer_token_fails(example_rules, wrong):
    pdf = PDFS / "INV-2026-10231.pdf"
    result = extract(pdf, example_rules)
    result.fields["인보이스번호"] = wrong
    assert _items(verify(result, pdf)) == ["인보이스번호"]


def test_number_field_before_sentence_period_passes(tmp_path, rules):
    pdf = _pdf(tmp_path / "dot.pdf", ["Memo No: M-2026-01.", "Date: 10/15/2026", "Total: $1,234.50"])
    result = extract(pdf, rules)
    assert result.fields["메모번호"] == "M-2026-01"
    assert failed(verify(result, pdf)) == []


def test_number_field_right_after_label_period_passes(tmp_path, rules):
    pdf = _pdf(tmp_path / "nodot.pdf", ["Memo No.M-2026-01", "Date: 10/15/2026", "Total: $1,234.50"])
    result = extract(pdf, rules)
    assert result.fields["메모번호"] == "M-2026-01"
    assert failed(verify(result, pdf)) == []


def test_number_field_after_digit_period_is_partial(tmp_path, rules):
    pdf = _pdf(tmp_path / "part.pdf", ["Memo No: M-1", "Ref 1.2026", "Date: 10/15/2026", "Total: $1,234.50"])
    result = extract(pdf, rules)
    result.fields["메모번호"] = "2026"  # 1.2026 의 일부, 10/15/2026 의 일부
    assert _items(verify(result, pdf)) == ["메모번호"]


def test_line_check_without_price_is_not_applicable_with_total_only(tmp_path):
    data = yaml.safe_load(yaml.safe_dump(TEST_RULES, allow_unicode=True))
    bill = data["문서종류"]["Bill"]
    bill["표"].update(열=["줄", "품번", "수량", "비고", "금액"], 금액열=["금액"])
    bill["검산"] = {"합계": "합계금액"}
    path = tmp_path / "규칙.yaml"
    path.write_text(yaml.safe_dump(data, allow_unicode=True), encoding="utf-8")
    pdf = _bill(tmp_path, [("2", "50.00", "100.00")], "100.00", "0.00", "100.00")
    checks = verify(extract(pdf, load_rules(path)), pdf)
    assert next(c for c in checks if c.item == "1줄 금액").ok is None
    assert next(c for c in checks if c.item == "품목 합 = 합계").ok is True
    assert failed(checks) == []


def test_excel_number_field_must_equal_a_whole_cell(tmp_path, example_rules):
    wb = openpyxl.Workbook()
    wb.active["A1"], wb.active["B1"] = "PO Number:", "  PO-T-12 "
    path = tmp_path / "po.xlsx"
    wb.save(path)
    result = ExtractResult(doc_type=example_rules.get("PO"), is_excel=True, fields={"고객PO번호": "PO-T-12"})
    assert failed(verify(result, path)) == []
    result.fields["고객PO번호"] = "PO-T-1"
    assert _items(verify(result, path)) == ["고객PO번호"]


def test_line_and_total_checks_use_named_columns(tmp_path):
    data = yaml.safe_load(yaml.safe_dump(TEST_RULES, allow_unicode=True))
    bill = data["문서종류"]["Bill"]
    bill["표"].update(열=["줄", "품번", "Qty", "Price", "Amount"], 금액열=["Price", "Amount"], 수량열=["Qty"])
    bill["검산"].update(수량열="Qty", 단가열="Price", 금액열="Amount")
    path = tmp_path / "규칙.yaml"
    path.write_text(yaml.safe_dump(data, allow_unicode=True), encoding="utf-8")
    named = load_rules(path)

    pdf = _bill(tmp_path, [("2", "50.00", "100.00")], "100.00", "15.00", "115.00")
    checks = verify(extract(pdf, named), pdf)
    assert failed(checks) == []
    assert next(c for c in checks if c.item == "1줄 금액").ok is True
    assert next(c for c in checks if c.item == "품목 합 = 소계").ok is True

    pdf = _bill(tmp_path, [("2", "50.00", "100.01")], "100.01", "15.00", "115.01")
    assert _items(verify(extract(pdf, named), pdf)) == ["1줄 금액"]


def test_wrong_table_value_fails(tmp_path, rules):
    pdf = _bill(tmp_path, [("2", "50.00", "100.00")], "100.00", "15.00", "115.00")
    result = extract(pdf, rules)
    assert failed(verify(result, pdf)) == []
    result.table[0]["단가"] = Decimal("5.00")
    assert "품목 1줄 단가" in _items(verify(result, pdf))


def test_negative_amount_matches_signed_token(tmp_path, rules):
    pdf = _pdf(tmp_path / "neg.pdf", ["Memo No: M-1", "Date: 10/15/2026", "Total: -12.50"])
    result = ExtractResult(doc_type=rules.get("Memo"), fields={
        "메모번호": "M-1", "발행일": date(2026, 10, 15), "합계금액": Decimal("-12.50")})
    assert failed(verify(result, pdf)) == []


# --- 줄 검산 ---

def test_line_check_allows_rounding(tmp_path, rules):
    pdf = _bill(tmp_path, [("3", "0.333", "1.00")], "1.00", "0.00", "1.00")
    assert failed(verify(extract(pdf, rules), pdf)) == []

    pdf = _bill(tmp_path, [("3", "0.333", "1.10")], "1.10", "0.00", "1.10")
    assert _items(verify(extract(pdf, rules), pdf)) == ["1줄 금액"]


def test_line_check_half_even_also_passes(tmp_path, rules):
    # 1 x 0.125 = 0.125 → 사사오입 0.13, 은행원식 0.12. 둘 중 하나라도 같으면 통과
    pdf = _bill(tmp_path, [("1", "0.125", "0.12")], "0.12", "0.00", "0.12")
    assert failed(verify(extract(pdf, rules), pdf)) == []


def test_missing_values_are_not_applicable(tmp_path, rules):
    pdf = _bill(tmp_path, [("2", "50.00", "100.00")], "100.00", "15.00", "115.00")
    result = extract(pdf, rules)
    for row in result.table:
        del row["단가"]  # 단가 없는 표
    checks = verify(result, pdf)
    line = next(c for c in checks if c.item == "1줄 금액")
    assert line.ok is None
    assert failed(checks) == []


# --- 합계 검산 ---

def test_total_check_with_freight(tmp_path, rules):
    pdf = _bill(tmp_path, [("2", "50.00", "100.00")], "100.00", "15.00", "115.00")
    checks = verify(extract(pdf, rules), pdf)
    assert failed(checks) == []
    assert next(c for c in checks if c.item == "소계 + 부대비용 = 합계").ok is True

    pdf = _bill(tmp_path, [("2", "50.00", "100.00")], "100.00", "15.00", "116.00")
    assert _items(verify(extract(pdf, rules), pdf)) == ["소계 + 부대비용 = 합계"]


def test_item_sum_must_equal_subtotal(tmp_path, rules):
    pdf = _bill(tmp_path, [("2", "50.00", "100.00"), ("1", "5.00", "5.00")], "100.00", "15.00", "115.00")
    assert _items(verify(extract(pdf, rules), pdf)) == ["품목 합 = 소계"]


def test_missing_total_is_not_applicable(tmp_path, rules):
    pdf = _bill(tmp_path, [("2", "50.00", "100.00")], "100.00", "15.00", "115.00")
    result = extract(pdf, rules)
    result.fields["합계금액"] = None
    checks = verify(result, pdf)
    assert next(c for c in checks if c.item == "소계 + 부대비용 = 합계").ok is None
    assert failed(checks) == []


def test_missing_subtotal_falls_back_to_item_sum(tmp_path, rules):
    pdf = _bill(tmp_path, [("2", "50.00", "100.00")], "100.00", "15.00", "115.00")
    result = extract(pdf, rules)
    result.fields["소계"] = None
    checks = verify(result, pdf)
    assert next(c for c in checks if c.item == "품목 합 = 소계").ok is None
    assert next(c for c in checks if c.item == "소계 + 부대비용 = 합계").ok is True
    result.table[0]["금액"] = Decimal("99.00")
    assert "소계 + 부대비용 = 합계" in _items(verify(result, pdf))


def test_missing_extra_counts_as_zero(tmp_path, rules):
    pdf = _bill(tmp_path, [("2", "50.00", "100.00")], "100.00", "15.00", "115.00")
    result = extract(pdf, rules)
    result.fields["운임"] = None  # 운임을 못 읽었으면 0 으로 보고, 합계가 안 맞아 잡힌다
    assert _items(verify(result, pdf)) == ["소계 + 부대비용 = 합계"]


def test_total_only_compares_item_sum(tmp_path, example_rules):
    # 예시 DM: 검산 {합계: 합계금액} → 품목 합 = 합계
    pdf = PDFS / "DM-2026-0092.pdf"
    result = extract(pdf, example_rules)
    checks = verify(result, pdf)
    assert next(c for c in checks if c.item == "품목 합 = 합계").ok is True
    result.table[0]["금액"] = Decimal("1240.00")
    assert "품목 합 = 합계" in _items(verify(result, pdf))


# --- 엑셀 ---

def test_excel_float_values_pass_and_wrong_scale_fails(tmp_path, example_rules):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"], ws["B1"] = "PO Number:", "PO-T-1"
    ws["F9"] = 12.3
    ws["F9"].number_format = "0.00"
    path = tmp_path / "po.xlsx"
    wb.save(path)

    result = ExtractResult(doc_type=example_rules.get("PO"), is_excel=True, source_form="ACME",
                           fields={"고객PO번호": "PO-T-1", "합계금액": Decimal("12.30")})
    assert failed(verify(result, path)) == []
    result.fields["합계금액"] = Decimal("123.0")
    assert _items(verify(result, path)) == ["합계금액"]


def test_excel_values_on_other_sheets_and_string_numbers_count(tmp_path, example_rules):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "PO-T-2"
    other = wb.create_sheet("other")
    other["B2"] = "1,500.25"
    other["C3"] = date(2026, 9, 18)
    path = tmp_path / "po.xlsx"
    wb.save(path)
    result = ExtractResult(doc_type=example_rules.get("PO"), is_excel=True, fields={
        "고객PO번호": "PO-T-2", "합계금액": Decimal("1500.25"), "발주일": date(2026, 9, 18)})
    assert failed(verify(result, path)) == []
    result.fields["발주일"] = date(2026, 9, 19)
    assert _items(verify(result, path)) == ["발주일"]


def test_check_dataclass_shape():
    c = Check("합계금액", "1,234.50", "1234.50", True)
    assert (c.item, c.source, c.output, c.ok) == ("합계금액", "1,234.50", "1234.50", True)
    assert failed([c, Check("x", "", "", None), Check("y", "1", "2", False)]) == [Check("y", "1", "2", False)]


def test_excel_formula_float_passes_line_check(tmp_path, example_rules):
    wb = openpyxl.load_workbook(ORDERS / "고객오더_ACME_PO-ACME-2026-0417.xlsx")
    ws = wb.active
    ws["D14"], ws["F14"], ws["G14"] = 3, 0.1, 0.375  # =3*0.1 의 저장값
    path = tmp_path / "po.xlsx"
    wb.save(path)
    _raw_float(path, "0.375", "0.30000000000000004")
    checks = verify(read_excel(path, example_rules), path)
    assert failed(checks) == []
    line = next(c for c in checks if c.item == "1줄 금액")
    assert line.ok is True and line.output == "0.3"


def _raw_float(path, placeholder, raw):
    """openpyxl 은 float 을 짧게 쓰므로, 엑셀이 저장하는 긴 값(=3*0.1 → 0.30000000000000004)을 XML 에 직접 넣는다."""
    import zipfile
    with zipfile.ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
    sheet = "xl/worksheets/sheet1.xml"
    assert files[sheet].count(f"<v>{placeholder}</v>".encode()) == 1
    files[sheet] = files[sheet].replace(f"<v>{placeholder}</v>".encode(), f"<v>{raw}</v>".encode())
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in files.items():
            z.writestr(n, data)


def test_excel_int_amount_rounds_to_two_places(tmp_path, example_rules):
    # 엑셀은 100.00 을 정수 100 으로 저장한다. 2 x 50.2 = 100.4 가 100 으로 통과하면 안 된다
    wb = openpyxl.load_workbook(ORDERS / "고객오더_ACME_PO-ACME-2026-0417.xlsx")
    ws = wb.active
    ws["D14"], ws["F14"], ws["G14"] = 2, 50.2, 100
    path = tmp_path / "po.xlsx"
    wb.save(path)
    assert _items(verify(read_excel(path, example_rules), path)) == ["1줄 금액"]
