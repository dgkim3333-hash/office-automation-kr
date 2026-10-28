"""누적장부.xlsx — 생성·추가·중복·열 동기화·덮어쓰기·잠김."""
import copy
import re
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from app.extract import extract
from app.ledger import Ledger, LedgerLocked
from app.rules import load_rules

REPO = Path(__file__).resolve().parent.parent
SAMPLES = REPO / "samples"
INVOICE = SAMPLES / "ACME_Invoice_INV-2026-0001.pdf"
PO = SAMPLES / "ACME_PO_PO-77001.pdf"


@pytest.fixture(scope="module")
def rules():
    return load_rules(REPO / "rules" / "규칙.예시.yaml")


@pytest.fixture
def book(tmp_path):
    return tmp_path / "누적장부.xlsx"


def _header(ws):
    return [c.value for c in ws[1] if c.value is not None]


def _data_rows(ws):
    return [r for r in ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]


def _add(ledger, pdf, rules, when=datetime(2026, 10, 28, 9, 0)):
    ledger.add(extract(pdf, rules), pdf.name, f"정리문서/x/{pdf.name}", processed_at=when)


def test_new_ledger_has_five_sheets_and_headers(book, rules):
    Ledger.open(book, rules).save()
    wb = openpyxl.load_workbook(book)
    assert wb.sheetnames == ["Invoice", "PO", "DM", "품목", "처리로그"]
    assert _header(wb["Invoice"]) == ["처리일시", "원본파일", "인보이스번호", "발행일", "고객명", "합계금액", "정리된경로"]
    assert _header(wb["품목"]) == ["문서종류", "문서번호", "품번", "품명", "수량", "단가", "금액"]
    assert _header(wb["처리로그"]) == ["일시", "파일", "결과", "사유"]
    assert wb["Invoice"].tables  # 엑셀 표(필터)
    assert wb["Invoice"].freeze_panes == "A2"


def test_add_row_persists_with_formats(book, rules):
    ledger = Ledger.open(book, rules)
    _add(ledger, INVOICE, rules)
    ledger.save()

    ws = openpyxl.load_workbook(book)["Invoice"]
    header = _header(ws)
    rows = _data_rows(ws)
    assert len(rows) == 1
    row = dict(zip(header, rows[0]))
    assert row["인보이스번호"] == "INV-2026-0001"
    assert row["발행일"].date() == date(2026, 10, 15)
    assert row["합계금액"] == pytest.approx(12345.67)
    assert row["원본파일"] == INVOICE.name
    date_cell = ws.cell(row=2, column=header.index("발행일") + 1)
    amount_cell = ws.cell(row=2, column=header.index("합계금액") + 1)
    assert "yy" in date_cell.number_format
    assert amount_cell.number_format == "#,##0.00"
    assert len(_data_rows(openpyxl.load_workbook(book)["품목"])) == 3
    # 표 범위가 데이터 행까지 늘어남
    ref = ws.tables[next(iter(ws.tables))].ref
    assert re.fullmatch(r"A1:[A-Z]+2", ref)


def test_duplicate_detected_after_reopen(book, rules):
    ledger = Ledger.open(book, rules)
    _add(ledger, INVOICE, rules)
    ledger.save()

    again = Ledger.open(book, rules)
    assert again.is_duplicate("Invoice", "INV-2026-0001") is True
    assert again.is_duplicate("Invoice", "INV-0000") is False
    assert again.is_duplicate("PO", "INV-2026-0001") is False
    again.save()
    assert len(_data_rows(openpyxl.load_workbook(book)["Invoice"])) == 1


def test_new_field_column_appended_at_end(book, rules):
    ledger = Ledger.open(book, rules)
    _add(ledger, INVOICE, rules)
    ledger.save()

    more = copy.deepcopy(rules)
    invoice = more.get("Invoice")
    invoice.fields = {"통화": copy.deepcopy(invoice.fields["고객명"]), **invoice.fields}
    invoice.fields["통화"].name = "통화"
    Ledger.open(book, more).save()

    ws = openpyxl.load_workbook(book)["Invoice"]
    assert _header(ws) == ["처리일시", "원본파일", "인보이스번호", "발행일", "고객명", "합계금액", "정리된경로", "통화"]
    assert _data_rows(ws)[0][2] == "INV-2026-0001"


def test_overwrite_replaces_document_and_items(book, rules):
    ledger = Ledger.open(book, rules)
    _add(ledger, INVOICE, rules, when=datetime(2026, 10, 1, 9, 0))
    _add(ledger, PO, rules)
    ledger.save()

    ledger = Ledger.open(book, rules)
    ledger.remove("Invoice", "INV-2026-0001")
    _add(ledger, INVOICE, rules, when=datetime(2026, 10, 28, 15, 30))
    ledger.save()

    wb = openpyxl.load_workbook(book)
    rows = _data_rows(wb["Invoice"])
    assert len(rows) == 1
    assert rows[0][0] == datetime(2026, 10, 28, 15, 30)
    items = _data_rows(wb["품목"])
    assert len(items) == 3 + 2  # Invoice 3줄 + PO 2줄, 중복 없음
    assert sum(1 for r in items if r[0] == "Invoice") == 3


def test_log_rows(book, rules):
    ledger = Ledger.open(book, rules)
    ledger.add_log(datetime(2026, 10, 28, 9, 0), "a.pdf", "성공", "")
    ledger.add_log(datetime(2026, 10, 28, 9, 0), "b.pdf", "확인필요", "판별 실패")
    ledger.save()
    rows = _data_rows(openpyxl.load_workbook(book)["처리로그"])
    assert [r[1:] for r in rows] == [("a.pdf", "성공", None), ("b.pdf", "확인필요", "판별 실패")]


def test_locked_on_open(book, rules):
    Ledger.open(book, rules).save()
    with open(book, "rb"):  # 엑셀이 열어 둔 상태 흉내
        with pytest.raises(LedgerLocked):
            Ledger.open(book, rules)


def test_locked_on_save(book, rules):
    Ledger.open(book, rules).save()
    ledger = Ledger.open(book, rules)
    _add(ledger, INVOICE, rules)
    with open(book, "rb"):
        with pytest.raises(LedgerLocked):
            ledger.save()
    assert len(_data_rows(openpyxl.load_workbook(book)["Invoice"])) == 0
