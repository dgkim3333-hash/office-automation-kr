"""고객 엑셀 PO 읽기 — samples/01_고객오더_엑셀 기준."""
from datetime import date
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from app.excel_extract import read_excel
from app.rules import load_rules

REPO = Path(__file__).resolve().parent.parent
ORDERS = REPO / "samples" / "01_고객오더_엑셀"
ACME = ORDERS / "고객오더_ACME_PO-ACME-2026-0417.xlsx"
HANIL = ORDERS / "고객오더_HANIL_HDC-PUR-26-0908.xlsx"


@pytest.fixture(scope="module")
def rules():
    return load_rules(REPO / "rules" / "규칙.예시.yaml")


def test_acme_head(rules):
    r = read_excel(ACME, rules)
    assert r.doc_type.name == "PO" and r.source_form == "ACME"
    assert r.fields["고객PO번호"] == "PO-ACME-2026-0417"
    assert r.fields["발주일"] == date(2026, 9, 18) and type(r.fields["발주일"]) is date
    assert r.fields["고객명"] == "ACME Semiconductor Korea Co., Ltd."
    assert r.fields["납기요청일"] == date(2026, 11, 20)
    assert r.fields["인도조건"] == "DAP Cheongun"
    assert r.fields["통화"] == "USD"
    assert r.fields["합계금액"] == Decimal("111590")
    assert r.fields["담당"] == "Mr. J. Park (Procurement)"
    assert r.number == "PO-ACME-2026-0417"
    assert r.warnings == [] and r.review_reason is None


def test_acme_items_stop_before_total_row(rules):
    rows = read_excel(ACME, rules).table
    assert len(rows) == 4
    assert rows[0] == {"줄": 1, "품번": "VTS-CDU-300", "품명": "Coolant Distribution Unit, 300 kW, N+1 pump",
                       "수량": 2, "단위": "EA", "단가": Decimal("48500"), "금액": Decimal("97000")}
    assert rows[3]["품번"] == "VTS-SENS-LEAK"


def test_hanil_korean_form(rules):
    r = read_excel(HANIL, rules)
    assert r.source_form == "HANIL"
    assert r.fields["고객명"] == "Hanil DataCenter Co., Ltd."  # A1 셀
    assert r.fields["고객PO번호"] == "HDC-PUR-26-0908"
    assert r.fields["발주일"] == date(2026, 9, 25)
    assert r.fields["합계금액"] == Decimal("224100")
    assert [row["품번"] for row in r.table] == ["VTS-CHL-500RT", "VTS-PMP-CDU-K", "VTS-FLT-50"]


def test_unknown_excel_is_review(rules):
    r = read_excel(REPO / "samples" / "02_회사양식_변환" / "회사양식_SO입력_HDC-PUR-26-0908.xlsx", rules)
    assert r.doc_type is None
    assert r.review_reason.startswith("판별 실패")
    assert "Order Date" in r.text_head  # 무엇이 들어 있었는지 확인필요 메모에 남김


def test_missing_label_is_warning_and_missing_number_is_required(rules, tmp_path):
    wb = openpyxl.load_workbook(ACME)
    ws = wb.active
    ws["A4"] = "Order No:"   # 'PO Number:' 글자가 없어짐 → 번호필드 못 찾음
    ws["A11"] = None         # 'Buyer Contact:' 글자가 없어짐 → 담당 못 찾음
    broken = tmp_path / "broken.xlsx"
    wb.save(broken)
    r = read_excel(broken, rules)
    assert r.fields["고객PO번호"] is None and r.missing_required is True
    assert r.fields["담당"] is None
    assert any("담당" in w for w in r.warnings)
