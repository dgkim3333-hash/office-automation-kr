"""처음 10건 승인: 양식(엑셀양식 · 문서종류) 단위로 승인 수를 세고, 규칙이 바뀐 양식만 0 으로 되돌린다."""
import os
import shutil
import time
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
import pytest

import app.job_run as job_run
from app.approval import REQUIRED, approval_gate, approve, form_fingerprint, form_key
from app.job_run import JobLocked, check_file_info, job_lock, load_records, run_job
from app.jobs import create_job, load_job, save_job
from app.master import write_example
from app.rules import RuleError, load_rules
from app.verify import Check

REPO = Path(__file__).resolve().parent.parent
ORDERS = REPO / "samples" / "01_고객오더_엑셀"
PDFS = REPO / "samples" / "03_Invoice_DM_PDF"
EXAMPLE_RULES = REPO / "rules" / "규칙.예시.yaml"
ACME_PO = ORDERS / "고객오더_ACME_PO-ACME-2026-0417.xlsx"
HANIL_PO = ORDERS / "고객오더_HANIL_HDC-PUR-26-0908.xlsx"
INV = PDFS / "INV-2026-10231.pdf"
TITLE = "일일 정리"
COMPANY_ACME = "회사양식_SO입력_PO-ACME-2026-0417.xlsx"
COMPANY_HANIL = "회사양식_SO입력_HDC-PUR-26-0908.xlsx"
CHECK_ACME = "검산_PO-ACME-2026-0417.xlsx"


@pytest.fixture
def env(tmp_path):
    work, inp, out = tmp_path / "work", tmp_path / "in", tmp_path / "out"
    for d in (work, inp, out):
        d.mkdir()
    write_example(work / "고객마스터.xlsx")
    job = create_job(work, TITLE, inp, out, "둘다")
    shutil.copy(EXAMPLE_RULES, job.rules_path)
    return work, inp, out, job


def _drop(src, folder, name=None, age=3600):
    dst = Path(folder) / (name or Path(src).name)
    shutil.copy(src, dst)
    _age(dst, age)
    return dst


def _age(path, age=3600):
    t = time.time() - age
    os.utime(path, (t, t))


def _files(folder):
    folder = Path(folder)
    if not folder.is_dir():
        return []
    return sorted(p.name for p in folder.iterdir() if p.is_file() and not p.name.endswith(".bak.xlsx"))


def _approved(job, key, count):
    """그 양식이 지금 규칙으로 이미 count 건 승인된 상태로 둔다."""
    job.approvals[key] = count
    job.fingerprints[key] = form_fingerprint(job.rules_path, key)


def test_form_key_and_fingerprint(env):
    work, inp, out, job = env
    from app.excel_extract import read_excel
    from app.rules import load_rules

    rules = load_rules(job.rules_path)
    assert form_key(read_excel(ACME_PO, rules)) == "ACME"
    acme = form_fingerprint(job.rules_path, "ACME")
    assert len(acme) == 16 and acme != form_fingerprint(job.rules_path, "HANIL")
    assert form_fingerprint(job.rules_path, "Invoice") != form_fingerprint(job.rules_path, "DM")


def test_new_form_waits_for_approval(env):
    work, inp, out, job = env
    _drop(ACME_PO, inp)

    report = run_job(job, work, gate=approval_gate)

    assert report.counts["승인대기"] == 1 and report.counts["성공"] == 0
    assert _files(out) == []
    assert _files(job.folder / "승인대기") == [ACME_PO.name]
    assert _files(job.folder / "승인대기" / "대조표") == [CHECK_ACME]
    saved = load_job(job.folder)
    assert saved.approvals == {"ACME": 0}
    assert saved.fingerprints == {"ACME": form_fingerprint(job.rules_path, "ACME")}


def test_approve_moves_to_output_and_counts(env):
    work, inp, out, job = env
    _drop(ACME_PO, inp)
    run_job(job, work, gate=approval_gate)
    job.last_run = ""  # 자동 실행 전 아침에 승인하는 경우
    save_job(job)

    report = approve(job, work)

    assert report.counts["성공"] == 1
    assert COMPANY_ACME in _files(out) and f"누적장부_{TITLE}.xlsx" in _files(out)
    assert job.approvals["ACME"] == 1 and load_job(job.folder).approvals["ACME"] == 1
    assert _files(job.folder / "승인대기") == []
    assert _files(job.folder / "승인대기" / "대조표") == []
    assert load_records(job)[-1]["상태"] == "성공"
    assert "ACME 1/10" in report.message
    assert load_job(job.folder).last_run == ""  # 승인은 그날의 자동 실행을 대신하지 않는다(놓친 실행 판단)
    # 입력 폴더의 원본은 이미 처리됨(같은 지문): 다음 자동 실행이 다시 내보내거나 승인대기로 보내지 않는다
    again = run_job(job, work, gate=approval_gate)
    assert again.counts["성공"] == 0 and again.counts["승인대기"] == 0


def test_after_ten_approvals_goes_straight_out(env):
    work, inp, out, job = env
    _approved(job, "ACME", REQUIRED)
    _drop(ACME_PO, inp)

    report = run_job(job, work, gate=approval_gate)

    assert report.counts["승인대기"] == 0 and report.counts["성공"] == 1
    assert COMPANY_ACME in _files(out)
    assert job.approvals["ACME"] == REQUIRED  # 자동으로 나간 건은 세지 않는다


def test_rule_change_resets_only_that_form(env):
    work, inp, out, job = env
    _approved(job, "ACME", REQUIRED)
    _approved(job, "HANIL", REQUIRED)
    text = job.rules_path.read_text(encoding="utf-8")
    changed = text.replace('납품처:     "오른쪽:Deliver To:"', '납품처:     "B7"')  # ACME 머리 칸 하나(같은 값)
    assert changed != text
    job.rules_path.write_text(changed, encoding="utf-8")
    _drop(ACME_PO, inp)
    _drop(HANIL_PO, inp)

    report = run_job(job, work, gate=approval_gate)

    assert report.counts["승인대기"] == 1 and report.counts["성공"] == 1
    assert job.approvals == {"ACME": 0, "HANIL": REQUIRED}
    assert job.fingerprints["ACME"] == form_fingerprint(job.rules_path, "ACME")
    assert _files(job.folder / "승인대기") == [ACME_PO.name]
    assert COMPANY_HANIL in _files(out) and COMPANY_ACME not in _files(out)


def test_approve_with_empty_waiting(env):
    work, inp, out, job = env

    report = approve(job, work)

    assert report.message == "승인할 문서가 없습니다."
    assert report.counts == dict.fromkeys(report.counts, 0) and not report.has_news


def test_approve_rechecks_and_rejects_bad_file(env):
    """승인대기의 복사본을 지금 규칙으로 다시 검산한다. 틀린 파일은 확인필요, 승인 수 그대로, 결과 없음."""
    work, inp, out, job = env
    _drop(ACME_PO, inp)
    run_job(job, work, gate=approval_gate)
    copy = job.folder / "승인대기" / ACME_PO.name
    wb = openpyxl.load_workbook(copy)
    ws = wb.active
    ws["G15"] = 1  # 둘째 품목 금액을 틀리게(수량 x 단가 와 다름)
    wb.save(copy)
    _age(copy)

    report = approve(job, work)

    assert report.counts["확인필요"] == 1 and report.counts["성공"] == 0
    assert job.approvals["ACME"] == 0
    assert _files(out) == []
    assert copy.is_file()  # 고쳐서 다시 승인하거나, 승인하지 않으려면 지우면 된다
    assert _files(job.folder / "승인대기" / "대조표") == [CHECK_ACME]
    note = (job.folder / "확인필요" / f"{ACME_PO.name}.txt").read_text(encoding="utf-8")
    assert "검산 불일치" in note and "다시 승인" in note


# --- W4 리뷰에서 넘어온 것: 승인 실행의 확인필요와 자동 실행 ---

def test_approval_review_keeps_waiting_so_next_day_run_makes_no_second_copy(env, monkeypatch):
    """승인 실행에서 확인필요(고객마스터에 없음)가 돼도 승인대기 상태를 지킨다. 다음 날 자동 실행이 _2 복사본을 만들지 않는다."""
    work, inp, out, job = env
    _drop(ACME_PO, inp)
    run_job(job, work, gate=approval_gate)
    monkeypatch.setattr(job_run, "lookup", lambda master, customer: None)

    first = approve(job, work)

    assert first.counts["확인필요"] == 1 and first.counts["성공"] == 0
    assert load_records(job)[-1]["상태"] == "승인대기"
    monkeypatch.undo()

    tomorrow = datetime.now() + timedelta(days=1)
    scheduled = run_job(job, work, now=tomorrow, gate=approval_gate)

    assert scheduled.counts["승인대기"] == 0 and scheduled.counts["확인필요"] == 0
    assert _files(job.folder / "승인대기") == [ACME_PO.name]
    assert _files(job.folder / "승인대기" / "대조표") == [CHECK_ACME]

    second = approve(job, work)  # 고객마스터를 고친 뒤 같은 날 다시 승인하면 다시 확인한다

    assert second.counts["성공"] == 1 and job.approvals["ACME"] == 1
    assert _files(job.folder / "승인대기") == []
    assert not (job.folder / "확인필요" / f"{ACME_PO.name}.txt").exists()


def test_approval_verify_failure_then_scheduled_run_same_and_next_day(env, monkeypatch):
    work, inp, out, job = env
    _drop(INV, inp)
    run_job(job, work, gate=approval_gate)
    monkeypatch.setattr(job_run, "verify", lambda result, path: [Check("합계금액", "100.00", "1000.00", False)])

    assert approve(job, work).counts["확인필요"] == 1
    monkeypatch.undo()

    for day in (0, 1, 2):
        report = run_job(job, work, now=datetime.now() + timedelta(days=day), gate=approval_gate)
        assert report.counts["승인대기"] == 0
    assert _files(job.folder / "승인대기") == [INV.name]


def test_approve_clears_duplicate_copy_and_orphan_check_files(env):
    """같은 번호의 다른 파일 두 개가 승인대기에 있으면 하나는 성공, 하나는 중복. 둘 다 승인대기에서 치운다.
    친구가 복사본을 지워 승인하지 않은 문서의 대조표도 승인 때 치운다."""
    work, inp, out, job = env
    _drop(ACME_PO, inp)
    other = inp / "고객오더_ACME_다시보냄.xlsx"
    wb = openpyxl.load_workbook(ACME_PO)
    wb.active["J1"] = "resent"
    wb.save(other)
    _age(other)
    _drop(INV, inp)
    run_job(job, work, gate=approval_gate)
    waiting = job.folder / "승인대기"
    assert len(_files(waiting)) == 3 and len(_files(waiting / "대조표")) == 3
    (waiting / INV.name).unlink()  # 승인하지 않음

    report = approve(job, work)

    assert report.counts["성공"] == 1 and report.counts["중복"] == 1
    assert job.approvals["ACME"] == 1
    assert _files(waiting) == []
    assert _files(waiting / "대조표") == []


def test_check_file_is_paired_with_copy_name_not_original_name(env, monkeypatch):
    """입력 폴더에서 같은 이름으로 들어온 다른 문서는 복사본이 _2 가 된다. 대조표는 복사본 이름으로 짝지어
    한쪽을 승인해도 아직 기다리는 쪽의 대조표를 지우지 않는다."""
    work, inp, out, job = env
    _drop(ACME_PO, inp, name="PO.xlsx")
    run_job(job, work, gate=approval_gate)
    (inp / "PO.xlsx").unlink()
    _drop(HANIL_PO, inp, name="PO.xlsx")
    run_job(job, work, gate=approval_gate)
    waiting = job.folder / "승인대기"
    assert _files(waiting) == ["PO.xlsx", "PO_2.xlsx"]
    real = job_run.lookup
    monkeypatch.setattr(job_run, "lookup", lambda master, name: None if "Hanil" in str(name) else real(master, name))

    report = approve(job, work)

    assert report.counts["성공"] == 1 and report.counts["확인필요"] == 1
    assert _files(waiting) == ["PO_2.xlsx"]
    assert _files(waiting / "대조표") == ["검산_HDC-PUR-26-0908.xlsx"]


def _edit_rules(job, old, new):
    text = job.rules_path.read_text(encoding="utf-8")
    assert text.count(old) == 1
    job.rules_path.write_text(text.replace(old, new), encoding="utf-8")


def test_fingerprint_covers_linked_doc_type_and_company_form(env):
    work, inp, out, job = env
    keys = ("ACME", "HANIL", "PO", "Invoice", "DM")
    before = {k: form_fingerprint(job.rules_path, k) for k in keys}

    _edit_rules(job, '통화:       {패턴: "Currency', '통화:       {패턴: "(?i)Currency')  # 문서종류 PO
    after_po = {k: form_fingerprint(job.rules_path, k) for k in keys}
    assert [k for k in keys if after_po[k] != before[k]] == ["ACME", "HANIL", "PO"]

    _edit_rules(job, '"고정:KR-Branch / D. Lee"', '"고정:KR-Branch / D. Kim"')  # 회사양식
    after_form = {k: form_fingerprint(job.rules_path, k) for k in keys}
    assert all(after_form[k] != after_po[k] for k in keys)


def test_rule_change_while_waiting_remakes_check_file_before_output(env):
    """승인대기에 있는 사이 그 양식의 규칙이 바뀌면 친구가 본 값이 아니므로 내보내지 않고 대조표를 새로 만든다.
    새 대조표를 보고 다시 승인하면 결과가 나가고, 새 규칙의 첫 승인으로 센다."""
    work, inp, out, job = env
    _approved(job, "ACME", 5)
    _drop(ACME_PO, inp)
    run_job(job, work, gate=approval_gate)
    _edit_rules(job, '납품처:     "오른쪽:Deliver To:"', '납품처:     "B7"')

    report = approve(job, work)

    assert report.counts["승인대기"] == 1 and report.counts["성공"] == 0
    assert "규칙이 바뀌어 대조표를 새로 만들었습니다. 다시 보고 승인하세요" in report.message
    assert _files(out) == [] and job.approvals["ACME"] == 5
    waiting = job.folder / "승인대기"
    assert _files(waiting) == [ACME_PO.name]
    assert _files(waiting / "대조표") == [CHECK_ACME]
    _, stamps = check_file_info(waiting / "대조표" / CHECK_ACME)
    assert stamps["form_fingerprint"] == form_fingerprint(job.rules_path, "ACME")

    again = approve(job, work)

    assert again.counts["성공"] == 1 and COMPANY_ACME in _files(out)
    assert job.approvals["ACME"] == 1
    assert job.fingerprints["ACME"] == form_fingerprint(job.rules_path, "ACME")


def test_copy_without_check_file_gets_one_before_output(env, monkeypatch):
    """대조표가 없는 복사본(친구가 직접 넣음)은 대조표부터 만든다. 같은 날 · 같은 규칙의 확인필요 기록이 있어도
    승인은 사람이 누른 실행이므로 건너뛰지 않는다."""
    work, inp, out, job = env
    _drop(ACME_PO, inp)
    monkeypatch.setattr(job_run, "lookup", lambda master, customer: None)
    assert run_job(job, work, gate=approval_gate).counts["확인필요"] == 1
    monkeypatch.undo()
    shutil.copy2(inp / ACME_PO.name, job.folder / "승인대기" / ACME_PO.name)  # 친구가 고객마스터를 고친 뒤 직접 넣음

    report = approve(job, work)

    assert report.counts["승인대기"] == 1 and report.counts["성공"] == 0
    assert "대조표가 없어 새로 만들었습니다" in report.message
    assert _files(out) == []
    assert _files(job.folder / "승인대기" / "대조표") == [CHECK_ACME]

    again = approve(job, work)

    assert again.counts["성공"] == 1 and job.approvals["ACME"] == 1
    assert _files(job.folder / "승인대기") == []


def test_changed_copy_gets_new_check_file_before_output(env):
    """대조표를 만든 뒤 승인대기의 파일이 바뀌면(검산은 통과해도) 내보내지 않고 대조표를 새로 만든다."""
    work, inp, out, job = env
    _drop(ACME_PO, inp)
    run_job(job, work, gate=approval_gate)
    copy = job.folder / "승인대기" / ACME_PO.name
    wb = openpyxl.load_workbook(copy)
    wb.active["J1"] = "memo"
    wb.save(copy)
    _age(copy)

    report = approve(job, work)

    assert report.counts["승인대기"] == 1 and report.counts["성공"] == 0
    assert "파일이 대조표를 만든 뒤 바뀌어" in report.message
    assert approve(job, work).counts["성공"] == 1


# --- 수정 1차: 거절 뒤 재처리 · 잠금 안 정리 · 장부 출처 · 판별 순서 · 이름 겹침 · 사유 파일 ---

def test_rejected_document_is_reprocessed_only_after_rules_change(env):
    """친구가 복사본을 지워 승인하지 않은 문서는 그대로 건너뛰다가, 규칙을 고치면 다시 처리한다(설계 12-1)."""
    work, inp, out, job = env
    _drop(ACME_PO, inp)
    run_job(job, work, gate=approval_gate)
    waiting = job.folder / "승인대기"
    (waiting / ACME_PO.name).unlink()

    tomorrow = run_job(job, work, now=datetime.now() + timedelta(days=1), gate=approval_gate)

    assert tomorrow.counts == dict.fromkeys(tomorrow.counts, 0)
    assert _files(waiting) == [] and _files(out) == []

    _edit_rules(job, '납품처:     "오른쪽:Deliver To:"', '납품처:     "B7"')
    after_fix = run_job(job, work, now=datetime.now() + timedelta(days=1), gate=approval_gate)

    assert after_fix.counts["승인대기"] == 1
    assert _files(waiting) == [ACME_PO.name]


def test_rules_change_with_copy_present_makes_no_second_copy(env):
    work, inp, out, job = env
    _drop(ACME_PO, inp)
    run_job(job, work, gate=approval_gate)
    _edit_rules(job, '납품처:     "오른쪽:Deliver To:"', '납품처:     "B7"')

    report = run_job(job, work, now=datetime.now() + timedelta(days=1), gate=approval_gate)

    assert report.counts["승인대기"] == 0
    assert _files(job.folder / "승인대기") == [ACME_PO.name]


def test_approval_cleanup_runs_inside_job_lock(env, monkeypatch):
    import app.approval as approval

    work, inp, out, job = env
    _drop(ACME_PO, inp)
    run_job(job, work, gate=approval_gate)
    seen = []
    real = approval._clear_orphan_checks

    def spy(job_, waiting, warnings):
        with pytest.raises(JobLocked):
            with job_lock(job_):
                pass
        seen.append(True)
        real(job_, waiting, warnings)

    monkeypatch.setattr(approval, "_clear_orphan_checks", spy)

    assert approve(job, work).counts["성공"] == 1
    assert approve(job, work).message == "승인할 문서가 없습니다."
    assert seen == [True, True]


def test_ledger_shows_original_file_not_waiting_copy(env):
    work, inp, out, job = env
    _drop(ACME_PO, inp, name="PO.xlsx")
    run_job(job, work, gate=approval_gate)
    (inp / "PO.xlsx").unlink()
    _drop(HANIL_PO, inp, name="PO.xlsx")
    run_job(job, work, gate=approval_gate)
    assert _files(job.folder / "승인대기") == ["PO.xlsx", "PO_2.xlsx"]

    assert approve(job, work).counts["성공"] == 2

    rows = list(openpyxl.load_workbook(out / f"누적장부_{TITLE}.xlsx")["PO"].iter_rows(values_only=True))
    header = rows[0]
    got = [(r[header.index("원본파일")], r[header.index("정리된경로")]) for r in rows[1:] if any(r)]
    assert got == [("PO.xlsx", str(inp / "PO.xlsx"))] * 2
    checks = sorted((job.folder / "검산").glob("*.xlsx"))
    assert len(checks) == 2 and [check_file_info(c)[0] for c in checks] == ["PO.xlsx", "PO.xlsx"]


def test_detection_keywords_of_other_forms_change_fingerprint(env):
    """다른 양식의 판별키워드 · 순서를 고치면 이 양식으로 판별되는 문서가 바뀔 수 있으므로 지문이 바뀐다."""
    work, inp, out, job = env
    keys = ("ACME", "HANIL", "Invoice", "DM")
    before = {k: form_fingerprint(job.rules_path, k) for k in keys}

    _edit_rules(job, '판별키워드: ["Buyer:"]', '판별키워드: ["Buyer:", "PO Number:"]')  # ACME
    after_excel = {k: form_fingerprint(job.rules_path, k) for k in keys}
    assert [k for k in keys if after_excel[k] != before[k]] == ["ACME", "HANIL"]

    _edit_rules(job, '판별키워드: ["Invoice No"]', '판별키워드: ["Invoice No", "INVOICE"]')
    after_pdf = {k: form_fingerprint(job.rules_path, k) for k in keys}
    assert [k for k in keys if after_pdf[k] != after_excel[k]] == ["Invoice", "DM"]


def test_excel_form_named_like_doc_type_is_rejected(env):
    work, inp, out, job = env
    _edit_rules(job, "  HANIL:\n", "  PO:\n")

    with pytest.raises(RuleError, match="문서종류 이름과 같습니다"):
        load_rules(job.rules_path)


def test_rejected_copy_note_is_removed_with_its_check_file(env, monkeypatch):
    """승인 실행에서 확인필요가 된 복사본을 친구가 지우면, 다음 승인 때 대조표와 그 사유 파일을 치운다.
    같은 이름이어도 입력 폴더 파일의 사유 파일은 둔다."""
    from app.excel_extract import read_excel
    from app.outputs import write_check_file

    work, inp, out, job = env
    _drop(ACME_PO, inp)
    run_job(job, work, gate=approval_gate)
    monkeypatch.setattr(job_run, "lookup", lambda master, customer: None)
    assert approve(job, work).counts["확인필요"] == 1
    monkeypatch.undo()
    waiting = job.folder / "승인대기"
    note = job.folder / "확인필요" / f"{ACME_PO.name}.txt"
    assert note.is_file()
    (waiting / ACME_PO.name).unlink()

    assert approve(job, work).message == "승인할 문서가 없습니다."
    assert _files(waiting / "대조표") == [] and not note.exists()

    # 같은 이름의 입력 폴더 파일이 남긴 사유 파일은 그대로 둔다
    note.write_text(f"원본: {inp / ACME_PO.name}\n사유: 시험\n", encoding="utf-8")
    write_check_file(waiting / "대조표", read_excel(ACME_PO, load_rules(job.rules_path)), [], ACME_PO.name)
    approve(job, work)
    assert _files(waiting / "대조표") == [] and note.is_file()


# --- 수정 2차: 거절 문서는 늘 승인대기로 · 규칙 한 번 읽기 · 판별 안 되는 복사본 ---

def _reject_acme_with_full_approvals(env):
    """ACME 문서를 승인대기로 보낸 뒤 친구가 복사본을 지워 거절하고, 그 뒤 다른 문서로 ACME 승인 수가 10 이 된 상태."""
    work, inp, out, job = env
    _drop(ACME_PO, inp)
    run_job(job, work, gate=approval_gate)
    (job.folder / "승인대기" / ACME_PO.name).unlink()
    _approved(job, "ACME", REQUIRED)
    save_job(job)
    return work, inp, out, job


@pytest.mark.parametrize("edit", [
    ('합계금액:    {패턴: "Total Debit', '합계금액:    {패턴: "(?i)Total Debit'),  # 다른 양식(DM)만 고침
    ("시트순서: [Invoice, DM, PO]", "# 메모\n시트순서: [Invoice, DM, PO]"),  # 설명 줄만 더함
])
def test_rejected_document_stays_rejected_when_its_form_did_not_change(env, edit):
    work, inp, out, job = _reject_acme_with_full_approvals(env)
    _edit_rules(job, *edit)

    report = run_job(job, work, now=datetime.now() + timedelta(days=1), gate=approval_gate)

    assert report.counts == dict.fromkeys(report.counts, 0)
    assert _files(out) == [] and _files(job.folder / "승인대기") == []


def test_rejected_document_always_goes_back_to_waiting_never_out(env):
    """그 양식 규칙을 고쳐 다시 처리하는 거절 문서는 승인 수 · 게이트와 상관없이 늘 승인대기로 간다."""
    work, inp, out, job = _reject_acme_with_full_approvals(env)
    _edit_rules(job, '납품처:     "오른쪽:Deliver To:"', '납품처:     "B7"')

    report = run_job(job, work, now=datetime.now() + timedelta(days=1), gate=lambda *a: False)

    assert report.counts["성공"] == 0 and report.counts["승인대기"] == 1
    assert _files(out) == [] and _files(job.folder / "승인대기") == [ACME_PO.name]
    assert load_records(job)[-1]["양식지문"] == form_fingerprint(job.rules_path, "ACME")


def test_rejected_document_review_keeps_waiting_state(env, monkeypatch):
    """다시 처리한 거절 문서가 확인필요가 돼도 승인대기 기록을 지켜, 고친 뒤에도 내보내지 않고 승인대기로 보낸다."""
    work, inp, out, job = _reject_acme_with_full_approvals(env)
    _edit_rules(job, '납품처:     "오른쪽:Deliver To:"', '납품처:     "B7"')
    monkeypatch.setattr(job_run, "lookup", lambda master, customer: None)

    first = run_job(job, work, now=datetime.now() + timedelta(days=1), gate=lambda *a: False)

    assert first.counts["확인필요"] == 1 and load_records(job)[-1]["상태"] == "승인대기"
    monkeypatch.undo()

    second = run_job(job, work, now=datetime.now() + timedelta(days=2), gate=lambda *a: False)

    assert second.counts["성공"] == 0 and second.counts["승인대기"] == 1
    assert _files(out) == []


def test_rules_are_read_once_per_run(env, monkeypatch):
    """추출 뒤 대조표를 만들기 전에 규칙.yaml 이 바뀌어도, 대조표 지문 · 승인 지문 · 처리기록은 추출에 쓴 규칙의 것."""
    import yaml as _yaml

    from app.forms import form_fingerprint_data

    work, inp, out, job = env
    original = job.rules_path.read_text(encoding="utf-8")
    want = form_fingerprint_data(_yaml.safe_load(original), "ACME")
    real = job_run.verify

    def verify_then_edit(result, path):
        checks = real(result, path)
        job.rules_path.write_text(original.replace('납품처:     "오른쪽:Deliver To:"', '납품처:     "B7"'),
                                  encoding="utf-8")
        return checks

    monkeypatch.setattr(job_run, "verify", verify_then_edit)
    _drop(ACME_PO, inp)

    assert run_job(job, work, gate=approval_gate).counts["승인대기"] == 1

    assert form_fingerprint(job.rules_path, "ACME") != want  # 파일은 바뀌었다
    _, stamps = check_file_info(job.folder / "승인대기" / "대조표" / CHECK_ACME)
    assert stamps["form_fingerprint"] == want
    assert job.fingerprints["ACME"] == want
    assert load_records(job)[-1]["양식지문"] == want


def test_waiting_copy_not_detected_any_more_is_reported(env):
    """승인 실행에서 지금 규칙으로 판별되지 않는 복사본은 조용히 두지 않고 확인필요로 알린다. 복사본은 남긴다."""
    work, inp, out, job = env
    _drop(ACME_PO, inp)
    run_job(job, work, gate=approval_gate)
    _edit_rules(job, '판별키워드: ["Buyer:"]', '판별키워드: ["NoSuchWord"]')

    report = approve(job, work)

    assert report.counts["확인필요"] == 1 and report.counts["성공"] == 0
    assert _files(job.folder / "승인대기") == [ACME_PO.name]
    assert load_records(job)[-1]["상태"] == "승인대기"
    note = (job.folder / "확인필요" / f"{ACME_PO.name}.txt").read_text(encoding="utf-8")
    assert "승인대기의 이 파일을 지금 규칙으로 알아보지 못했습니다" in note


def test_approval_counts_under_rules_read_at_start(env, monkeypatch):
    """승인 실행 중 규칙.yaml 이 바뀌어도 승인 수는 추출 · 대조에 쓴 규칙의 지문으로 센다."""
    work, inp, out, job = env
    _drop(ACME_PO, inp)
    run_job(job, work, gate=approval_gate)
    want = form_fingerprint(job.rules_path, "ACME")
    original = job.rules_path.read_text(encoding="utf-8")
    real = job_run.verify

    def verify_then_edit(result, path):
        checks = real(result, path)
        job.rules_path.write_text(original.replace('납품처:     "오른쪽:Deliver To:"', '납품처:     "B7"'),
                                  encoding="utf-8")
        return checks

    monkeypatch.setattr(job_run, "verify", verify_then_edit)

    assert approve(job, work).counts["성공"] == 1
    assert job.fingerprints["ACME"] == want and job.approvals["ACME"] == 1
