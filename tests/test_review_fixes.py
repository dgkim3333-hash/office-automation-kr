"""코드 검토(2026-10-28)의 치명 C1~C4 · 중요 I3 재현 테스트.

받은문서 흐름을 전제로 한 C3(옮기지 못한 파일 · 처리로그.csv 잠김)와 I3(한 건 오류)의 실행 테스트는 작업 방식으로
옮겨졌다(tests/test_job_run.py: 열린 파일 보류 · 처리기록 잠김 · 한 문서만 되돌리기). 작업 방식은 원본을 옮기지 않는다.
"""
import copy
import os
import shutil
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

from app.init import init
from app.jobs import create_job
from app.ledger import Ledger
from app.rules import RuleError, load_rules

REPO = Path(__file__).resolve().parent.parent
SAMPLES = REPO / "samples" / "03_Invoice_DM_PDF"
EXAMPLE = REPO / "rules" / "규칙.예시.yaml"


@pytest.fixture
def env(tmp_path):
    repo = tmp_path / "repo"
    (repo / "rules").mkdir(parents=True)
    shutil.copy(EXAMPLE, repo / "rules" / "규칙.예시.yaml")
    work = init(work_dir=tmp_path / "work", repo_root=repo)
    (tmp_path / "in").mkdir()
    (tmp_path / "out").mkdir()
    job = create_job(work, "일일 정리", tmp_path / "in", tmp_path / "out", "누적장부")
    shutil.copy(EXAMPLE, job.rules_path)
    return repo, work


def _records(ws):
    rows = list(ws.iter_rows(values_only=True))
    header = [h for h in rows[0] if h is not None]
    return [dict(zip(header, r)) for r in rows[1:] if any(v is not None for v in r)]


# --- C1: 문서종류를 추가해도 장부 표 이름이 충돌하지 않는다 ---

def test_c1_adding_doc_type_to_existing_ledger(tmp_path):
    rules = load_rules(EXAMPLE)
    book = tmp_path / "누적장부.xlsx"
    Ledger.open(book, rules).save()

    more = copy.deepcopy(rules)
    extra = copy.deepcopy(more.get("Invoice"))
    extra.name = "CreditNote"
    more.doc_types.append(extra)
    Ledger.open(book, more).save()
    Ledger.open(book, more).save()  # 두 번째 실행도 죽지 않음

    wb = openpyxl.load_workbook(book)
    names = [t for ws in wb.worksheets for t in ws.tables]
    assert len(names) == len(set(names)) == 6


# --- C2: 사용자가 만든 시트·표는 건드리지 않고, 저장 전 백업을 남긴다 ---

def test_c2_user_sheets_are_left_alone(tmp_path):
    rules = load_rules(EXAMPLE)
    book = tmp_path / "누적장부.xlsx"
    Ledger.open(book, rules).save()

    wb = openpyxl.load_workbook(book)
    wb.create_sheet("빈시트")
    calc = wb.create_sheet("계산")
    calc["B2"] = "=1+1"
    mine = wb.create_sheet("내표")
    mine.append(["고객", "메모"])
    mine.append(["A", "B"])
    from openpyxl.worksheet.table import Table
    mine.add_table(Table(displayName="내표", ref="A1:B2"))
    wb.save(book)

    Ledger.open(book, rules).save()
    Ledger.open(book, rules).save()
    wb = openpyxl.load_workbook(book)
    assert list(wb["내표"].tables) == ["내표"]
    assert not wb["빈시트"].tables and not wb["계산"].tables
    assert wb["계산"]["B2"].value == "=1+1"


def test_c2_backup_before_replace(tmp_path):
    rules = load_rules(EXAMPLE)
    book = tmp_path / "누적장부.xlsx"
    Ledger.open(book, rules).save()
    first = book.read_bytes()
    ledger = Ledger.open(book, rules)
    ledger.add_log(__import__("datetime").datetime(2026, 10, 28), "a.pdf", "성공", "")
    ledger.save()
    assert (tmp_path / "누적장부.bak.xlsx").read_bytes() == first


# --- C4: 출력이 파이프여도 utf-8 로 나오고 죽지 않는다 ---

def test_c4_piped_output_is_utf8(env):
    repo, _ = env
    code = (
        "import sys; from app.__main__ import main; "
        f"sys.exit(main(['check', r'{SAMPLES / 'Letter_warranty_notice.pdf'}'], repo_root=r'{repo}'))"
    )
    clean_env = {k: v for k, v in os.environ.items() if k not in ("PYTHONIOENCODING", "PYTHONUTF8")}
    proc = subprocess.run([sys.executable, "-c", code], cwd=REPO, env=clean_env, capture_output=True)
    assert proc.returncode == 0, proc.stderr.decode("utf-8", "replace")
    assert "판별 실패" in proc.stdout.decode("utf-8")


def test_c4_no_em_dash_in_program_messages():
    for path in (REPO / "app").glob("*.py"):
        assert "—" not in path.read_text(encoding="utf-8"), path.name


# --- I3: 한 건의 오류가 실행 전체를 멈추지 않는다 ---

def test_i3_format_spec_only_on_date_fields(tmp_path):
    import yaml
    data = yaml.safe_load(EXAMPLE.read_text(encoding="utf-8"))
    data["문서종류"]["Invoice"]["파일명"] = "{고객명:%Y-%m}/Invoice_{인보이스번호}.pdf"
    path = tmp_path / "규칙.yaml"
    path.write_text(yaml.safe_dump(data, allow_unicode=True), encoding="utf-8")
    with pytest.raises(RuleError, match="고객명"):
        load_rules(path)


def test_i3_control_characters_do_not_break_ledger(tmp_path):
    from app.extract import extract
    rules = load_rules(EXAMPLE)
    result = extract(SAMPLES / "INV-2026-10231.pdf", rules)
    result.fields["운송사"] = "DHL\x01\x02 Global"
    ledger = Ledger.open(tmp_path / "누적장부.xlsx", rules)
    ledger.add(result, "x.pdf", "정리문서/x.pdf")
    ledger.save()
    row = _records(openpyxl.load_workbook(tmp_path / "누적장부.xlsx")["Invoice"])[0]
    assert row["운송사"] == "DHL Global"


def test_i3_deleted_items_sheet_is_recreated(tmp_path):
    rules = load_rules(EXAMPLE)
    book = tmp_path / "누적장부.xlsx"
    Ledger.open(book, rules).save()
    wb = openpyxl.load_workbook(book)
    del wb["품목"]
    wb.save(book)
    Ledger.open(book, rules).save()
    assert "품목" in openpyxl.load_workbook(book).sheetnames
