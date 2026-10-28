"""작업 방식 끝까지: samples 의 PDF · 엑셀 PO 를 작업으로 처리 → 장부 · 회사 양식이 정답과 같은지.

정답: samples/04_장부_엑셀/누적장부_기대결과.xlsx, samples/02_회사양식_변환/*.xlsx
(옛 test_e2e.py · test_e2e_excel.py 의 시나리오를 작업 방식으로 옮긴 것. 원본은 입력 폴더에 그대로 남는다)
"""
import hashlib
import os
import shutil
import time
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
import pytest

from app.job_run import load_records, run_job
from app.jobs import create_job
from app.markup import write_example_markup
from app.master import write_example

REPO = Path(__file__).resolve().parent.parent
SAMPLES = REPO / "samples"
PDFS = SAMPLES / "03_Invoice_DM_PDF"
ORDERS = SAMPLES / "01_고객오더_엑셀"
FORMS = SAMPLES / "02_회사양식_변환"
EXPECTED = SAMPLES / "04_장부_엑셀" / "누적장부_기대결과.xlsx"
SKIP = {"처리일시", "정리된경로"}  # 실행 시각 · 경로는 정답과 비교하지 않음


@pytest.fixture
def env(tmp_path):
    work, out = tmp_path / "work", tmp_path / "out"
    pdf_in, excel_in = tmp_path / "pdf_in", tmp_path / "excel_in"
    for d in (work, out, pdf_in, excel_in):
        d.mkdir()
    write_example(work / "고객마스터.xlsx")
    write_example_markup(work / "마크업표.xlsx")
    for folder, files in ((pdf_in, PDFS.glob("*.pdf")), (excel_in, ORDERS.glob("*.xlsx"))):
        for f in files:
            dst = folder / f.name
            shutil.copy(f, dst)
            t = time.time() - 3600
            os.utime(dst, (t, t))
    pdf_job = create_job(work, "PDF 장부", pdf_in, out, "누적장부")
    excel_job = create_job(work, "PO 입력", excel_in, out, "둘다")
    for job in (pdf_job, excel_job):
        shutil.copy(REPO / "rules" / "규칙.예시.yaml", job.rules_path)
    return work, out, pdf_job, excel_job


def _snapshot(folder):
    return {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in Path(folder).iterdir() if p.is_file()}


def _records(path, sheet):
    rows = list(openpyxl.load_workbook(path)[sheet].iter_rows(values_only=True))
    header = [h for h in rows[0] if h is not None]
    return [dict(zip(header, r)) for r in rows[1:] if any(v is not None for v in r)]


def _same(a, b):
    if hasattr(a, "year") and hasattr(b, "year"):
        return (a.year, a.month, a.day) == (b.year, b.month, b.day)
    try:
        return Decimal(str(a)) == Decimal(str(b))
    except (InvalidOperation, ValueError):
        return a == b


def _compare_sheet(got_rows, expected_rows, key, where):
    got = {r[key]: r for r in got_rows}
    expected = {r[key]: r for r in expected_rows}
    assert sorted(got) == sorted(expected)
    for name, exp in expected.items():
        for column, value in exp.items():
            if column not in SKIP:
                assert _same(got[name][column], value), f"{where} {name} {column}: {got[name][column]!r} != {value!r}"


def test_pdf_job_counts_and_leaves_originals(env):
    work, out, pdf_job, _ = env
    before = _snapshot(pdf_job.input_dir)

    report = run_job(pdf_job, work, gate=None)

    assert report.counts == {"성공": 4, "중복": 0, "확인필요": 1, "승인대기": 0, "보류": 0}
    assert _snapshot(pdf_job.input_dir) == before
    states = {r["파일"]: r["상태"] for r in load_records(pdf_job)}
    assert states == {
        "INV-2026-10231.pdf": "성공", "INV-2026-10278.pdf": "성공",
        "DM-2026-0092.pdf": "성공", "DM-2026-0097.pdf": "성공",
        "Scan_signed_packing_list.pdf": "확인필요",
    }  # 판별되지 않은 서신은 기록하지 않고 그대로 둔다
    assert sorted(p.name for p in (pdf_job.folder / "확인필요").iterdir()) == ["Scan_signed_packing_list.pdf.txt"]
    assert len(list((pdf_job.folder / "검산").glob("검산_*.xlsx"))) == 4
    assert sorted(p.name for p in out.iterdir() if not p.name.endswith(".bak.xlsx")) == ["누적장부_PDF 장부.xlsx"]


@pytest.mark.parametrize("sheet", ["Invoice", "DM"])
def test_pdf_job_ledger_matches_expected(env, sheet):
    work, out, pdf_job, _ = env
    run_job(pdf_job, work, gate=None)
    _compare_sheet(_records(out / "누적장부_PDF 장부.xlsx", sheet), _records(EXPECTED, sheet), "원본파일", sheet)


def test_pdf_job_items_match_expected(env):
    work, out, pdf_job, _ = env
    run_job(pdf_job, work, gate=None)
    key = lambda r: (r["문서종류"], r["문서번호"], str(r["줄"]))  # noqa: E731
    got = sorted(_records(out / "누적장부_PDF 장부.xlsx", "품목"), key=key)
    expected = sorted((r for r in _records(EXPECTED, "품목") if r["문서종류"] != "PO"), key=key)
    assert len(got) == len(expected) == 9
    for g, e in zip(got, expected):
        for column, value in e.items():
            assert _same(g[column], value), f"품목 {key(e)} {column}: {g[column]!r} != {value!r}"


def test_excel_job_both_matches_expected_forms_and_ledger(env):
    work, out, _, excel_job = env
    before = _snapshot(excel_job.input_dir)

    report = run_job(excel_job, work, gate=None)

    assert report.counts["성공"] == 2 and report.counts["확인필요"] == 0
    assert _snapshot(excel_job.input_dir) == before
    _compare_sheet(_records(out / "누적장부_PO 입력.xlsx", "PO"), _records(EXPECTED, "PO"), "원본파일", "PO")
    for expected in FORMS.glob("*.xlsx"):
        got = [r for r in openpyxl.load_workbook(out / expected.name)["SO_Upload"].iter_rows(values_only=True)
               if any(v is not None for v in r)]
        exp = [r for r in openpyxl.load_workbook(expected)["SO_Upload"].iter_rows(values_only=True)
               if any(v is not None for v in r)]
        assert got[0] == exp[0] and len(got) == len(exp)
        for g_row, e_row in zip(got[1:], exp[1:]):
            for name, g, e in zip(exp[0], g_row, e_row):
                assert _same(g, e), f"{expected.name} {name}: {g!r} != {e!r}"
        assert out / expected.name in report.output_files


def test_second_run_of_both_jobs_has_no_news(env):
    work, out, pdf_job, excel_job = env
    run_job(pdf_job, work, gate=None)
    run_job(excel_job, work, gate=None)
    ledgers = _snapshot(out)

    again = [run_job(pdf_job, work, gate=None), run_job(excel_job, work, gate=None)]

    assert [r.has_news for r in again] == [False, False]  # 서신은 하루가 안 돼 「맞지 않는 파일」로 세지 않음
    assert _snapshot(out) == ledgers


def test_excel_job_po_items_match_expected(env):
    work, out, _, excel_job = env
    run_job(excel_job, work, gate=None)
    key = lambda r: (r["문서종류"], r["문서번호"], str(r["줄"]))  # noqa: E731
    got = sorted(_records(out / "누적장부_PO 입력.xlsx", "품목"), key=key)
    expected = sorted((r for r in _records(EXPECTED, "품목") if r["문서종류"] == "PO"), key=key)
    assert len(got) == len(expected) == 7
    for g, e in zip(got, expected):
        for column, value in e.items():
            assert _same(g[column], value), f"품목 {key(e)} {column}: {g[column]!r} != {value!r}"


def test_customer_not_in_master_goes_to_review(env):
    work, out, _, excel_job = env
    wb = openpyxl.load_workbook(work / "고객마스터.xlsx")
    wb.active.delete_rows(2)  # ACME 줄을 지움
    wb.save(work / "고객마스터.xlsx")

    report = run_job(excel_job, work, gate=None)

    assert report.counts["성공"] == 1 and report.counts["확인필요"] == 1
    note = (excel_job.folder / "확인필요" / "고객오더_ACME_PO-ACME-2026-0417.xlsx.txt").read_text(encoding="utf-8")
    assert "고객마스터" in note and "ACME" in note
    assert not (out / "회사양식_SO입력_PO-ACME-2026-0417.xlsx").exists()


def test_excel_lock_file_in_input_is_ignored(env):
    work, out, _, excel_job = env
    for f in list(excel_job.input_dir.iterdir()):
        f.unlink()
    lock = excel_job.input_dir / "~$주문.xlsx"
    lock.write_bytes(b"lock")
    t = time.time() - 2 * 24 * 3600
    os.utime(lock, (t, t))

    report = run_job(excel_job, work, gate=None)

    assert not report.has_news and report.unmatched_old == 0
    assert lock.exists() and load_records(excel_job) == []
