"""파트별 마크업(설계 13-2): 마크업표.xlsx 의 파트넘버 → 마크업%. 본사 단가 = 고객 단가 × (1 + 마크업% ÷ 100)."""
import os
import shutil
import time
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest

from app.approval import REQUIRED, approval_gate, approve, form_fingerprint
from app.excel_extract import read_excel
from app.job_run import load_records, run_job
from app.jobs import create_job, load_job, save_job
from app.markup import (
    MARKUP_FILE, MarkupError, MarkupRow, apply_markup, load_markup, markup_checks, missing_parts,
    write_example_markup,
)
from app.outputs import company_rows
from app.rules import RuleError, load_rules, rules_from_text

REPO = Path(__file__).resolve().parent.parent
ORDERS = REPO / "samples" / "01_고객오더_엑셀"

# 테스트 전용 가짜 고객 PO(엑셀)와 본사 PO 시트 규칙
RULES = r'''
문서종류:
  PO:
    판별키워드: ["PURCHASE ORDER"]
    번호필드: 고객PO번호
    필드:
      고객PO번호: {패턴: "PO Number[.:]?\\s*([A-Z0-9-]+)"}
      고객명:     {패턴: "Buyer[.:]?\\s*(.+?)\\n"}
    파일명: "PO_{고객PO번호}.pdf"
엑셀양식:
  TESTCO:
    문서종류: PO
    판별키워드: ["Buyer:"]
    머리:
      고객PO번호: "오른쪽:PO Number:"
      고객명:     "오른쪽:Buyer:"
    품목:
      머리글: "Part No"
      열: {품번: A, 수량: B, 단가: C, 금액: D}
      금액열: [단가, 금액]
      수량열: [수량]
회사양식:
  파일명: "본사PO_{고객PO번호}.xlsx"
  시트: HQ_PO
  열:
    - {이름: Line No,        값: 줄번호}
    - {이름: Part No,        값: 품목.품번}
    - {이름: Qty,            값: 품목.수량}
    - {이름: Customer Price, 값: 품목.단가}
    - {이름: Markup %,       값: 마크업.율}
    - {이름: Unit Price,     값: 마크업.단가}
    - {이름: Amount,         값: 마크업.금액}
'''
NO_MARKUP_RULES = RULES.split("    - {이름: Markup %")[0]


def _po(path, number, lines, age=3600):
    """가짜 고객 PO 엑셀. lines = [(품번, 수량, 단가)]. 금액 = 수량 × 단가."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["PURCHASE ORDER"])
    ws.append(["PO Number:", number])
    ws.append(["Buyer:", "Testco Korea Co., Ltd."])
    ws.append([])
    ws.append(["Part No", "Qty", "Unit Price", "Amount"])
    for part, qty, price in lines:
        ws.append([part, qty, float(price), float(Decimal(str(price)) * qty)])
    wb.save(path)
    t = time.time() - age
    os.utime(path, (t, t))
    return path


def _markup(path, rows, header=("파트넘버", "마크업%", "반올림자리")):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(list(header))
    for row in rows:
        ws.append(list(row))
    wb.save(path)
    return path


@pytest.fixture
def env(tmp_path):
    work, inp, out = tmp_path / "work", tmp_path / "in", tmp_path / "out"
    for d in (work, inp, out):
        d.mkdir()
    job = create_job(work, "본사 PO", inp, out, "회사양식")
    job.rules_path.write_text(RULES, encoding="utf-8")
    return work, inp, out, job


def _sheet(path, sheet):
    rows = list(openpyxl.load_workbook(path, data_only=True)[sheet].iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:]]


# --- 계산 ---

def test_apply_markup_examples():
    assert apply_markup(Decimal("0.10"), MarkupRow("A", Decimal("3"), None)) == Decimal("0.103")
    assert apply_markup(Decimal("0.37"), MarkupRow("B", Decimal("3"), None)) == Decimal("0.3811")
    assert apply_markup(Decimal("0.37"), MarkupRow("B", Decimal("3"), 3)) == Decimal("0.381")
    assert apply_markup(Decimal("5.00"), MarkupRow("C", Decimal("0"), None)) == Decimal("5.00")


def test_apply_markup_rounds_half_up_only_when_places_given():
    assert apply_markup(Decimal("0.25"), MarkupRow("D", Decimal("2"), 3)) == Decimal("0.255")  # 0.255 그대로
    assert apply_markup(Decimal("0.125"), MarkupRow("E", Decimal("2"), 3)) == Decimal("0.128")  # 0.1275 → 0.128
    assert apply_markup(Decimal("0.125"), MarkupRow("E", Decimal("2"), None)) == Decimal("0.1275")
    assert apply_markup(Decimal("100"), MarkupRow("F", Decimal("2.5"), 0)) == Decimal("103")  # 102.5 → 103


# --- 마크업표 읽기 ---

def test_load_markup_reads_percent_formatted_cells(tmp_path):
    path = tmp_path / MARKUP_FILE
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["파트넘버", "마크업%"])
    ws.append(["A-100", 0.03])
    ws["B2"].number_format = "0%"
    ws.append(["B-200", 0.025])
    ws["B3"].number_format = "0.0%"
    ws.append(["  c-300 ", 2])          # 백분율 서식이 아닌 숫자 2 = 2%
    ws.append(["D-400", "1.5"])         # 글자로 적은 숫자
    wb.save(path)

    table = load_markup(path)

    assert table["a-100"].rate == Decimal("3") and table["a-100"].part == "A-100"
    assert table["b-200"].rate == Decimal("2.5")
    assert table["c-300"].rate == Decimal("2") and table["c-300"].part == "c-300"
    assert table["d-400"].rate == Decimal("1.5")
    assert all(row.places is None for row in table.values())


def test_load_markup_reads_places_and_matches_case_insensitively(tmp_path):
    path = _markup(tmp_path / MARKUP_FILE, [["VTS-A", 3, 2], ["vts-b", 0, None]])
    table = load_markup(path)
    assert table["vts-a"].places == 2 and table["vts-b"].places is None


@pytest.mark.parametrize("rows, words", [
    ([["A", 3, None], ["a ", 2, None]], ["3번째 줄", "A", "2번째 줄"]),   # 같은 파트 두 번(대소문자 · 공백 무시)
    ([["A", "삼", None]], ["2번째 줄", "마크업%"]),                         # 숫자 아님
    ([["A", None, None]], ["2번째 줄", "비어"]),                            # 비어 있음(0% 로 보지 않음)
    ([["A", -1, None]], ["2번째 줄", "0 이상 10 이하"]),
    ([["A", 10.5, None]], ["2번째 줄", "0 이상 10 이하"]),
    ([["A", 3, 1.5]], ["2번째 줄", "반올림자리"]),
    ([["A", 3, -1]], ["2번째 줄", "반올림자리"]),
    ([[None, 3, None]], ["2번째 줄", "파트넘버"]),
    ([["A", True, None]], ["2번째 줄", "마크업%"]),
])
def test_load_markup_rejects_duplicates_bad_values_out_of_range(tmp_path, rows, words):
    path = _markup(tmp_path / MARKUP_FILE, rows)
    with pytest.raises(MarkupError) as e:
        load_markup(path)
    for word in words:
        assert word in str(e.value)


def test_load_markup_accepts_bounds_and_skips_blank_lines(tmp_path):
    path = _markup(tmp_path / MARKUP_FILE, [["A", 0, None], [None, None, None], ["B", 10, None]])
    table = load_markup(path)
    assert table["a"].rate == 0 and table["b"].rate == 10


def test_load_markup_needs_header_and_file(tmp_path):
    with pytest.raises(MarkupError, match="없습니다"):
        load_markup(tmp_path / MARKUP_FILE)
    path = _markup(tmp_path / MARKUP_FILE, [["A", 3]], header=("품번", "마크업"))
    with pytest.raises(MarkupError, match="파트넘버"):
        load_markup(path)


def test_example_markup_covers_sample_parts(tmp_path):
    path = tmp_path / MARKUP_FILE
    write_example_markup(path)
    table = load_markup(path)
    rates = {row.rate for row in table.values()}
    assert {Decimal(0), Decimal(2), Decimal(3)} <= rates
    rules = load_rules(REPO / "rules" / "규칙.예시.yaml")
    for order in ORDERS.glob("*.xlsx"):
        assert missing_parts(read_excel(order, rules), rules.company_form, table) == []


# --- 규칙 ---

def test_rules_accept_markup_values_and_columns():
    rules = rules_from_text(RULES)
    assert rules.company_form.uses_markup
    assert (rules.company_form.markup.part, rules.company_form.markup.price,
            rules.company_form.markup.quantity) == ("품번", "단가", "수량")
    assert not rules_from_text(NO_MARKUP_RULES).company_form.uses_markup
    custom = RULES.replace("  시트: HQ_PO\n", "  시트: HQ_PO\n  마크업: {품번열: 품번, 단가열: 단가, 수량열: 수량}\n")
    assert rules_from_text(custom).company_form.uses_markup


def test_rules_reject_unknown_markup_value_and_option():
    with pytest.raises(RuleError, match="마크업.단가"):
        rules_from_text(RULES.replace("마크업.율", "마크업.퍼센트"))
    with pytest.raises(RuleError, match="마크업"):
        rules_from_text(RULES.replace("  시트: HQ_PO\n", "  시트: HQ_PO\n  마크업: {품명열: 품번}\n"))
    with pytest.raises(RuleError, match="단가열"):
        rules_from_text(RULES.replace("  시트: HQ_PO\n", "  시트: HQ_PO\n  마크업: {단가열: 가격}\n"))


# --- 회사 양식 행 ---

def test_company_rows_compute_markup_values():
    rules = rules_from_text(RULES)
    from app.extract import ExtractResult
    result = ExtractResult(doc_type=rules.get("PO"), fields={"고객PO번호": "PO-1"},
                           table=[{"품번": "A", "수량": 10, "단가": Decimal("0.10")},
                                  {"품번": "b", "수량": 3, "단가": Decimal("0.37")}])
    table = {"a": MarkupRow("A", Decimal("3"), None), "b": MarkupRow("B", Decimal("3"), 3)}
    header, rows = company_rows(result, None, rules.company_form, table)
    got = [dict(zip(header, r)) for r in rows]
    assert got[0]["Unit Price"] == Decimal("0.103") and got[0]["Amount"] == Decimal("1.03")
    assert got[0]["Markup %"] == Decimal("3")
    assert got[1]["Unit Price"] == Decimal("0.381") and got[1]["Amount"] == Decimal("1.143")


def test_markup_checks_catch_wrong_written_price():
    rules = rules_from_text(RULES)
    from app.extract import ExtractResult
    result = ExtractResult(doc_type=rules.get("PO"), fields={"고객PO번호": "PO-1"},
                           table=[{"품번": "A", "수량": 10, "단가": Decimal("0.10")}])
    table = {"a": MarkupRow("A", Decimal("3"), None)}
    good = [{"마크업.율": Decimal("3"), "마크업.단가": Decimal("0.103"), "마크업.금액": Decimal("1.03")}]
    assert all(c.ok for c in markup_checks(result, good, table))
    for bad in ({**good[0], "마크업.단가": Decimal("0.10")}, {**good[0], "마크업.금액": Decimal("1.0")},
                {**good[0], "마크업.율": Decimal("2")}):
        assert any(c.ok is False for c in markup_checks(result, [bad], table))
    no_price = ExtractResult(doc_type=rules.get("PO"), fields={}, table=[{"품번": "A", "수량": 10, "단가": None}])
    assert any(c.ok is False for c in markup_checks(no_price, [{}], table))


# --- 실행 ---

def test_company_form_uses_markup_columns(env):
    work, inp, out, job = env
    _markup(work / MARKUP_FILE, [["A", 3, None], ["B", 3, 3]])
    _po(inp / "po1.xlsx", "PO-1", [("A", 10, "0.10"), ("B", 3, "0.37")])

    report = run_job(job, work)

    assert report.counts["성공"] == 1, report.message
    rows = _sheet(out / "본사PO_PO-1.xlsx", "HQ_PO")
    assert Decimal(str(rows[0]["Unit Price"])) == Decimal("0.103")
    assert Decimal(str(rows[0]["Amount"])) == Decimal("1.03")
    assert Decimal(str(rows[0]["Markup %"])) == Decimal("3")
    assert Decimal(str(rows[0]["Customer Price"])) == Decimal("0.1")
    assert Decimal(str(rows[1]["Unit Price"])) == Decimal("0.381")
    assert Decimal(str(rows[1]["Amount"])) == Decimal("1.143")
    wb = openpyxl.load_workbook(out / "본사PO_PO-1.xlsx")
    ws = wb["HQ_PO"]
    shown = ws.cell(row=2, column=6).number_format  # Unit Price 0.103 이 0.10 으로 보이면 안 된다
    assert shown.endswith(".000"), shown


def test_missing_part_goes_to_review_listing_parts(env):
    work, inp, out, job = env
    _markup(work / MARKUP_FILE, [["A", 3, None]])
    _po(inp / "po1.xlsx", "PO-1", [("A", 10, "0.10"), ("B-77", 3, "0.37"), ("C-88", 1, "5.00")])

    report = run_job(job, work)

    assert report.counts["확인필요"] == 1 and report.counts["성공"] == 0
    note = (job.folder / "확인필요" / "po1.xlsx.txt").read_text(encoding="utf-8")
    assert "마크업표에 없는 파트넘버: B-77, C-88." in note  # 0% 로 넘기지 않고 빠진 파트를 모두 적는다
    assert list(out.iterdir()) == []


def test_missing_markup_file_goes_to_review(env):
    work, inp, out, job = env
    _po(inp / "po1.xlsx", "PO-1", [("A", 10, "0.10")])
    report = run_job(job, work)
    assert report.counts["확인필요"] == 1
    assert MARKUP_FILE in (job.folder / "확인필요" / "po1.xlsx.txt").read_text(encoding="utf-8")
    assert list(out.iterdir()) == []


def test_bad_markup_file_goes_to_review(env):
    work, inp, out, job = env
    _markup(work / MARKUP_FILE, [["A", 30, None]])
    _po(inp / "po1.xlsx", "PO-1", [("A", 10, "0.10")])
    report = run_job(job, work)
    assert report.counts["확인필요"] == 1
    assert "0 이상 10 이하" in (job.folder / "확인필요" / "po1.xlsx.txt").read_text(encoding="utf-8")


def test_adding_missing_part_retries_same_day(env):
    work, inp, out, job = env
    _markup(work / MARKUP_FILE, [["A", 3, None]])
    _po(inp / "po1.xlsx", "PO-1", [("A", 10, "0.10"), ("B", 3, "0.37")])
    assert run_job(job, work).counts["확인필요"] == 1
    assert run_job(job, work).counts["확인필요"] == 0  # 같은 날 · 그대로면 다시 시도하지 않는다

    _markup(work / MARKUP_FILE, [["A", 3, None], ["B", 2, None]])
    report = run_job(job, work)

    assert report.counts["성공"] == 1, report.message
    assert (out / "본사PO_PO-1.xlsx").exists()
    assert not (job.folder / "확인필요" / "po1.xlsx.txt").exists()
    assert [r["상태"] for r in load_records(job)] == ["확인필요", "성공"]


def test_markup_check_rows_in_check_file(env):
    work, inp, out, job = env
    _markup(work / MARKUP_FILE, [["A", 3, None], ["B", 3, 3]])
    _po(inp / "po1.xlsx", "PO-1", [("A", 10, "0.10"), ("B", 3, "0.37")])
    run_job(job, work)

    rows = list(openpyxl.load_workbook(job.folder / "검산" / "검산_PO-1.xlsx").active.iter_rows(values_only=True))
    lines = {r[0]: r for r in rows[2:]}
    assert lines["1줄 마크업% (A)"][1:] == ("마크업표 3%", "3", "일치")
    assert lines["1줄 본사 단가"][2:] == ("0.1030", "일치") or lines["1줄 본사 단가"][2:] == ("0.103", "일치")
    assert "0.10" in lines["1줄 본사 단가"][1] and "3" in lines["1줄 본사 단가"][1]
    assert lines["2줄 본사 단가"][2:] == ("0.381", "일치")
    assert "반올림" in lines["2줄 본사 단가"][1]
    assert lines["2줄 본사 금액"][3] == "일치"
    assert all(r[3] in ("일치", "해당 없음") for r in rows[2:])


def test_waiting_check_shows_markup_and_markup_change_remakes_it(env):
    work, inp, out, job = env
    _markup(work / MARKUP_FILE, [["A", 3, None]])
    _po(inp / "po1.xlsx", "PO-1", [("A", 10, "0.10")])
    assert run_job(job, work, gate=approval_gate).counts["승인대기"] == 1
    checks = list((job.folder / "승인대기" / "대조표").glob("*.xlsx"))
    rows = list(openpyxl.load_workbook(checks[0]).active.iter_rows(values_only=True))
    assert any(r[0] == "1줄 본사 단가" for r in rows)

    _markup(work / MARKUP_FILE, [["A", 2, None]])  # 친구가 대조표를 본 뒤 마크업이 바뀜
    report = approve(job, work)

    assert report.counts["성공"] == 0 and report.counts["승인대기"] == 1
    assert "마크업" in report.message
    assert not (out / "본사PO_PO-1.xlsx").exists()
    report = approve(job, work)  # 새 대조표를 보고 다시 승인
    assert report.counts["성공"] == 1
    rows = _sheet(out / "본사PO_PO-1.xlsx", "HQ_PO")
    assert Decimal(str(rows[0]["Unit Price"])) == Decimal("0.102")


def test_unrelated_markup_change_keeps_waiting_check(env):
    work, inp, out, job = env
    _markup(work / MARKUP_FILE, [["A", 3, None]])
    _po(inp / "po1.xlsx", "PO-1", [("A", 10, "0.10")])
    run_job(job, work, gate=approval_gate)
    _markup(work / MARKUP_FILE, [["A", 3, None], ["Z", 1, None]])  # 이 문서와 상관없는 파트를 더함
    assert approve(job, work).counts["성공"] == 1


def test_markup_change_does_not_reset_approvals(env):
    work, inp, out, job = env
    _markup(work / MARKUP_FILE, [["A", 3, None]])
    job.approvals["TESTCO"] = REQUIRED
    job.fingerprints["TESTCO"] = form_fingerprint(job.rules_path, "TESTCO")
    save_job(job)
    _po(inp / "po1.xlsx", "PO-1", [("A", 10, "0.10")])
    assert run_job(job, work, gate=approval_gate).counts["성공"] == 1

    _markup(work / MARKUP_FILE, [["A", 2, None], ["B", 1, None]])
    _po(inp / "po2.xlsx", "PO-2", [("A", 1, "1.00"), ("B", 1, "1.00")])
    report = run_job(job, work, gate=approval_gate)

    assert report.counts["성공"] == 1 and report.counts["승인대기"] == 0, report.message
    assert load_job(job.folder).approvals["TESTCO"] == REQUIRED


def test_no_markup_columns_means_no_markup_file_needed(env):
    work, inp, out, job = env
    job.rules_path.write_text(NO_MARKUP_RULES, encoding="utf-8")
    _po(inp / "po1.xlsx", "PO-1", [("NOT-IN-ANY-TABLE", 10, "0.10")])

    report = run_job(job, work)

    assert report.counts["성공"] == 1, report.message
    assert not (work / MARKUP_FILE).exists()


def test_ledger_mode_ignores_markup(env):
    work, inp, out, job = env
    job.output_mode = "누적장부"
    save_job(job)
    _po(inp / "po1.xlsx", "PO-1", [("NOT-IN-ANY-TABLE", 10, "0.10")])
    assert run_job(job, work).counts["성공"] == 1  # 누적장부는 고객 PO 값 그대로, 마크업표 없어도 됨


# --- check 명령 ---

def test_check_command_shows_markup_per_line(env, capsys):
    from app.__main__ import check
    work, inp, out, job = env
    _markup(work / MARKUP_FILE, [["A", 3, None]])
    po = _po(inp / "po1.xlsx", "PO-1", [("A", 10, "0.10"), ("B-77", 1, "1.00")])
    rules = load_rules(job.rules_path)

    check(po, rules, work, job)
    out_text = capsys.readouterr().out
    assert "B-77" in out_text and "마크업표" in out_text

    _markup(work / MARKUP_FILE, [["A", 3, None], ["B-77", 2, 2]])
    check(po, rules, work, job)
    out_text = capsys.readouterr().out
    assert "| 1줄 본사 단가 |" in out_text and "| 2줄 마크업% (B-77) | 마크업표 2% |" in out_text
