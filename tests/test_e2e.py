"""03장 '첫 실행 10분' 과 같은 시나리오: 샘플 5종 → run → 폴더·장부·로그 검증."""
import csv
import shutil
from pathlib import Path

import openpyxl
import pytest

from app.__main__ import main
from app.init import init

REPO = Path(__file__).resolve().parent.parent
SAMPLES = REPO / "samples"
CUSTOMER_DIR = Path("정리문서") / "ACME Semiconductor Korea" / "2026-10"


@pytest.fixture
def env(tmp_path):
    repo = tmp_path / "repo"
    (repo / "rules").mkdir(parents=True)
    shutil.copy(REPO / "rules" / "규칙.예시.yaml", repo / "rules" / "규칙.예시.yaml")
    work = init(work_dir=tmp_path / "work", repo_root=repo)
    return repo, work


def _drop_samples(work):
    for pdf in SAMPLES.glob("*.pdf"):
        shutil.copy(pdf, work / "받은문서" / pdf.name)


def _rows(ws):
    return [r for r in ws.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]


def _csv_rows(work):
    with open(work / "처리로그.csv", encoding="utf-8-sig", newline="") as f:
        return list(csv.reader(f))[1:]


def _pdfs(folder):
    return sorted(p.name for p in folder.glob("*.pdf"))


def test_first_run(env, capsys):
    repo, work = env
    _drop_samples(work)
    assert main([], repo_root=repo) == 0
    out = capsys.readouterr().out
    assert "성공 3 · 중복 0 · 확인필요 2" in out

    assert _pdfs(work / CUSTOMER_DIR) == ["DM_DM-5101.pdf", "Invoice_INV-2026-0001.pdf", "PO_PO-77001.pdf"]
    assert _pdfs(work / "확인필요") == ["Scan_blank.pdf", "Unknown_letter.pdf"]
    assert sorted(p.name for p in (work / "확인필요").glob("*.txt")) == ["Scan_blank.txt", "Unknown_letter.txt"]
    assert _pdfs(work / "받은문서") == []

    wb = openpyxl.load_workbook(work / "누적장부.xlsx")
    for sheet in ["Invoice", "PO", "DM"]:
        assert len(_rows(wb[sheet])) == 1, sheet
    assert len(_rows(wb["품목"])) == 3 + 2 + 2
    assert len(_rows(wb["처리로그"])) == 5
    assert len(_csv_rows(work)) == 5
    invoice = dict(zip([c.value for c in wb["Invoice"][1]], _rows(wb["Invoice"])[0]))
    assert invoice["정리된경로"] == str(CUSTOMER_DIR / "Invoice_INV-2026-0001.pdf")


def test_second_run_marks_duplicates(env, capsys):
    repo, work = env
    _drop_samples(work)
    main([], repo_root=repo)
    _drop_samples(work)
    main([], repo_root=repo)
    out = capsys.readouterr().out
    assert "성공 0 · 중복 3 · 확인필요 2" in out
    assert len(_pdfs(work / "받은문서" / "중복")) == 3
    assert _pdfs(work / "받은문서") == []
    wb = openpyxl.load_workbook(work / "누적장부.xlsx")
    assert len(_rows(wb["Invoice"])) == 1
    assert len(_rows(wb["품목"])) == 7
    assert len(_csv_rows(work)) == 10


def test_overwrite_keeps_row_count_and_updates_time(env, capsys, monkeypatch):
    repo, work = env
    _drop_samples(work)
    main([], repo_root=repo)
    first = _rows(openpyxl.load_workbook(work / "누적장부.xlsx")["Invoice"])[0][0]

    shutil.copy(SAMPLES / "ACME_Invoice_INV-2026-0001.pdf", work / "받은문서")
    import app.__main__ as cli
    monkeypatch.setattr(cli, "_now", lambda: first.replace(year=first.year + 1))
    main(["--덮어쓰기"], repo_root=repo)
    out = capsys.readouterr().out
    assert "성공 1 · 중복 0 · 확인필요 0" in out

    wb = openpyxl.load_workbook(work / "누적장부.xlsx")
    rows = _rows(wb["Invoice"])
    assert len(rows) == 1
    assert rows[0][0] > first
    assert len(_rows(wb["품목"])) == 7
    assert _pdfs(work / CUSTOMER_DIR).count("Invoice_INV-2026-0001_2.pdf") == 1


def test_locked_ledger_moves_nothing(env, capsys):
    repo, work = env
    _drop_samples(work)
    main([], repo_root=repo)
    capsys.readouterr()

    _drop_samples(work)
    with open(work / "누적장부.xlsx", "rb"):  # 엑셀이 열어 둔 상태
        code = main([], repo_root=repo)
    out = capsys.readouterr().out
    assert code != 0
    assert "엑셀을 닫고 다시 실행" in out
    assert len(_pdfs(work / "받은문서")) == 5
    assert _pdfs(work / "받은문서" / "중복") == []
    assert _pdfs(work / "확인필요") == ["Scan_blank.pdf", "Unknown_letter.pdf"]
    assert len(_csv_rows(work)) == 5 + 5  # csv 는 잠김과 무관하게 기록


def test_empty_inbox(env, capsys):
    repo, _ = env
    assert main([], repo_root=repo) == 0
    assert "받은문서에 PDF 가 없습니다" in capsys.readouterr().out
