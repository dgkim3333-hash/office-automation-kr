"""무설치판 엔진 company: 같은 결과 · 같은 회사 양식 · 같은 마크업표를 넣으면 지금 판 app\\outputs.py 와 칸마다 같은 줄을 낸다.
파일 이름(render_text)과 이름 충돌(free_name)도 app\\organize.py 와 같은 글자를 낸다.

엔진 op `company` 의 입력(action 으로 고른다):
  rows       {rules, result, markup_grid?}      -> {header, rows: [[칸 또는 null, ...] 품목마다], filename}
  template   {template, fields: {이름: 칸 또는 null}} -> {text}
  free_name  {path, reserved: [경로, ...]}      -> {path}
칸 = {"t": 값 종류, "v": 값}. 종류: decimal(자리 그대로의 글자) · int · float(파이썬 str(float)) · bool · date("YYYY-MM-DD")
  · datetime("YYYY-MM-DDTHH:MM:SS") · text. 값이 없으면 null.
result = 작업 2 · 3 의 결과 JSON(예: extract_text 의 출력). 시각이 0시인 일시는 날짜로 본다(지금 판 엑셀 읽기도 그렇게 한다).
"""
import shutil
from datetime import date, datetime
from decimal import Decimal, localcontext
from pathlib import Path

import openpyxl
import pytest
import yaml

from app import extract
from app import markup as M
from app.excel_extract import read_excel
from app.extract import ExtractResult, extract_from_text
from app.organize import free_name, render_text
from app.outputs import company_form_name, company_rows
from app.portable import rules_to_portable
from app.rules import rules_from_text

from .conftest import EXAMPLE_RULES, REPO, EngineError, example_rules, without_master
from .grid import grid_from_xlsx
from .test_engine_markup import F, markup_file

KIT = Path("C:/PO시험")
PDFS = sorted((REPO / "samples" / "03_Invoice_DM_PDF").glob("*.pdf"))
ORDERS = sorted((REPO / "samples" / "01_고객오더_엑셀").glob("*.xlsx"))
INVOICE = next(p for p in PDFS if p.name == "INV-2026-10231.pdf")
DM = next(p for p in PDFS if p.name == "DM-2026-0092.pdf")


# --- 값 · 결과를 엔진 모양으로 ---

def cell(value):
    """파이썬 값을 엔진의 칸 모양으로(엔진이 낼 모양 = 시험에서 기대하는 모양)."""
    if value is None:
        return None
    if isinstance(value, bool):  # int 의 하위 형식이라 먼저
        return {"t": "bool", "v": value}
    if isinstance(value, Decimal):
        return {"t": "decimal", "v": format(value, "f")}
    if isinstance(value, int):
        return {"t": "int", "v": value}
    if isinstance(value, float):
        return {"t": "float", "v": repr(value)}
    if isinstance(value, datetime):
        if value.time() == datetime.min.time():
            return {"t": "date", "v": value.date().isoformat()}
        return {"t": "datetime", "v": value.strftime("%Y-%m-%dT%H:%M:%S")}
    if isinstance(value, date):
        return {"t": "date", "v": value.isoformat()}
    if isinstance(value, str):
        return {"t": "text", "v": value}
    raise TypeError(type(value).__name__)


def _plain(value):
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def result_json(result) -> dict:
    """결과 객체를 엔진 결과 JSON 모양으로(작업 2 · 3 의 출력과 같은 모양)."""
    return {
        "doc_type": result.doc_type.name if result.doc_type else None,
        "form": result.source_form,
        "is_excel": result.is_excel,
        "fields": {k: _plain(v) for k, v in result.fields.items()},
        "table": None if result.table is None else [{k: _plain(v) for k, v in row.items()} for row in result.table],
        "warnings": list(result.warnings),
        "review_reason": result.review_reason,
    }


def python_view(result, rules, table=None) -> dict:
    header, rows = company_rows(result, None, rules.company_form, table)
    return {"header": header, "rows": [[cell(v) for v in row] for row in rows],
            "filename": company_form_name(result, rules.company_form)}


def engine_view(run_engine, rules, result, table_path=None, culture=None) -> dict:
    """엔진 결과(result 는 결과 JSON 사전). 마크업표는 파일 경로로 준다."""
    payload = {"rules": rules_to_portable(rules), "result": result, "action": "rows"}
    if table_path is not None:
        payload["markup_grid"] = grid_from_xlsx(table_path)
    return run_engine("company", payload, culture)


def compare(run_engine, rules, result, table_path=None, culture=None):
    """(지금 판, 엔진). result 는 파이썬 결과 객체."""
    table = M.markup_from_bytes(table_path.read_bytes()) if table_path is not None else None
    return python_view(result, rules, table), engine_view(run_engine, rules, result_json(result), table_path, culture)


def with_form(columns, filename="회사_{고객PO번호}.xlsx", markup=None):
    """예시 규칙에서 회사양식만 바꾼 규칙(고객마스터 열 없음). columns = [(이름, 값), ...]."""
    text = without_master(EXAMPLE_RULES.read_text(encoding="utf-8"))
    head = text.split("\n회사양식:")[0]
    form = {"파일명": filename, "시트": "시트1", "열": [{"이름": n, "값": v} for n, v in columns]}
    if markup is not None:
        form["마크업"] = markup
    return rules_from_text(head + "\n" + yaml.safe_dump({"회사양식": form}, allow_unicode=True))


def invoice_result(rules, **changes):
    """샘플 Invoice 를 지금 판으로 추출한 결과. changes 의 필드 값을 덮어쓴다."""
    result = extract.extract(INVOICE, rules)
    result.fields.update(changes)
    return result


def sample_markup(tmp_path, results):
    """샘플 품번 일부만 실린 마크업표(나머지는 표에 없음 → 마크업 칸이 빔)."""
    parts = []
    for result in results:
        parts += [row["품번"] for row in result.table or [] if row["품번"] not in parts]
    kinds = [lambda p: [p, 3, 2], lambda p: [p, "0.5%", None], lambda p: [p, F(0.04, "0%"), 0], lambda p: [p, 2.5, 4]]
    rows = [kinds[i % len(kinds)](p) for i, p in enumerate(parts) if i % 3 != 2]  # 3번째마다 빠뜨린다
    return markup_file(tmp_path, rows)


# --- 샘플로 비교 ---

def test_samples_present():
    assert len(PDFS) == 6 and len(ORDERS) == 2


@pytest.mark.parametrize("xlsx", ORDERS, ids=["ACME" if "ACME" in p.name else "HANIL" for p in ORDERS])
def test_example_form_on_excel_samples(xlsx, run_engine, tmp_path):
    """샘플 엑셀 2개 + 예시 회사양식(고객마스터 열 제외, 마크업 열 포함): 머리글 · 칸 값 · 파일 이름이 지금 판과 같다."""
    rules = example_rules()
    result = read_excel(xlsx, rules)
    table = sample_markup(tmp_path, [result])

    expected, got = compare(run_engine, rules, result, table)

    assert got == expected
    assert len(expected["rows"]) == len(result.table) >= 2
    assert expected["header"][3] == "Line No" and [r[3] for r in expected["rows"]][:3] == [
        {"t": "int", "v": 10}, {"t": "int", "v": 20}, {"t": "int", "v": 30}]
    kinds = {c["t"] for row in expected["rows"] for c in row if c}
    assert {"decimal", "date", "text", "int"} <= kinds
    assert any(row[-1] is None for row in expected["rows"]), "마크업표에 없는 파트가 있어야 비교가 의미 있다"
    assert any(row[-1] is not None for row in expected["rows"])
    assert any(c is None for c in expected["rows"][0]), "빈칸 · 없는 필드(담당)가 None 으로 남아야 한다"


@pytest.mark.parametrize("pdf", [INVOICE, DM], ids=["Invoice", "DM"])
def test_every_source_kind_on_pdf_samples(pdf, run_engine, tmp_path):
    """PDF 샘플: 머리 · 품목 · 줄번호 · 수량x단가 · 고정 · 빈칸 · 마크업 · 없는 필드를 모두 쓰는 양식."""
    columns = [("번호", "머리.인보이스번호"), ("DM", "머리.DM번호"), ("발행", "머리.발행일"), ("금액", "머리.합계금액"),
               ("줄", "줄번호"), ("품번", "품목.품번"), ("수량", "품목.수량"), ("단가", "품목.단가"), ("합", "수량x단가"),
               ("고정", "고정:KR-Branch / D. Lee"), ("빈칸", "빈칸"), ("없음", "머리.없는필드"), ("없는열", "품목.없음"),
               ("본사단가", "마크업.단가"), ("본사금액", "마크업.금액"), ("율", "마크업.율")]
    rules = with_form(columns, filename="{고객명}_{발행일:%Y%m%d}.xlsx")
    result = extract.extract(pdf, rules)
    table = sample_markup(tmp_path, [result])

    expected, got = compare(run_engine, rules, result, table)

    assert got == expected
    assert len(expected["rows"]) >= 1
    assert all(row[10] is None and row[11] is None and row[12] is None for row in expected["rows"])


def test_without_markup_table_cells_are_empty(run_engine):
    """마크업표를 안 주면(None) 마크업 칸은 비고 나머지는 그대로."""
    rules = example_rules()
    result = read_excel(ORDERS[0], rules)

    expected, got = compare(run_engine, rules, result)

    assert got == expected
    assert all(row[-3] is None and row[-2] is None and row[-1] is None for row in expected["rows"])


def test_zero_items_and_no_table(run_engine):
    """품목이 0줄이거나 표가 없는 문서: 머리글만 있고 줄은 0개(빈 배열 [])."""
    rules = example_rules()
    base = read_excel(ORDERS[0], rules)
    for table in ([], None):
        result = ExtractResult(base.doc_type, dict(base.fields), table, source_form=base.source_form, is_excel=True)
        expected, got = compare(run_engine, rules, result)
        assert got == expected
        assert got["rows"] == [] and len(got["header"]) == 20
        assert got["filename"].startswith("회사양식_SO입력_PO-")


def test_one_item_gives_array_of_arrays(run_engine):
    """품목 1줄 · 칸 1개짜리 양식도 [[...]] 모양 그대로(배열이 풀려 스칼라가 되지 않는다)."""
    rules = with_form([("줄", "줄번호")])
    result = invoice_result(rules)
    result.table = result.table[:1]

    expected, got = compare(run_engine, rules, result)

    assert got == expected == {"header": ["줄"], "rows": [[{"t": "int", "v": 10}]], "filename": got["filename"]}


def test_culture_de_de_same_result(run_engine, tmp_path):
    """소수점 쉼표 문화권(de-DE)에서도 같은 값 · 같은 글자."""
    rules = example_rules()
    result = read_excel(ORDERS[1], rules)
    table = sample_markup(tmp_path, [result])

    expected, got = compare(run_engine, rules, result, table, culture="de-DE")

    assert got == expected
    assert got == engine_view(run_engine, rules, result_json(result), table, culture="en-US")


# --- 값 종류와 계산 ---

def test_qty_times_price_types_and_scale(run_engine):
    """수량x단가: 정수 · 소수 수량, 글자 수량, 4자리 단가의 자리 수, 값이 없으면 빈칸. 값 종류도 지금 판과 같다."""
    rules = with_form([("합", "수량x단가"), ("수량", "품목.수량"), ("단가", "품목.단가")])
    result = invoice_result(rules)
    result.table = [
        {"수량": 3, "단가": Decimal("2.50")},               # 정수 x 소수 2자리 -> 7.50
        {"수량": 1200, "단가": Decimal("47.1234")},         # 소수 4자리
        {"수량": 2.5, "단가": Decimal("10.00")},            # 엑셀 실수 수량 -> 25.000 (자리: 2.5 는 자리 1)
        {"수량": 0.1, "단가": Decimal("3")},                # 0.1 은 유효 15자리로
        {"수량": "10", "단가": Decimal("1.5")},             # 글자 수량
        {"수량": " 7 ", "단가": Decimal("2.25")},            # 글자 수량(앞뒤 공백)
        {"수량": 0, "단가": Decimal("5.00")},            # 0 x 5.00
        {"수량": -4, "단가": Decimal("1.10")},              # 음수
        {"수량": 1e-3, "단가": 2000.0},                     # 실수 x 실수
        {"수량": None, "단가": Decimal("5")},               # 한쪽이 None -> None
        {"단가": Decimal("5")},                             # 수량 열 없음 -> None
        {"수량": 5, "단가": None},
        {},
    ]

    expected, got = compare(run_engine, rules, result)

    assert got == expected
    assert expected["rows"][0][0] == {"t": "decimal", "v": "7.50"}
    assert expected["rows"][1][0] == {"t": "decimal", "v": "56548.0800"}
    assert [r[0] for r in expected["rows"][-3:]] == [None, None, None]


@pytest.mark.parametrize("qty, price", [("1,000", Decimal("2")), ("abc", Decimal("2")), (True, Decimal("2")),
                                         (date(2026, 1, 5), Decimal("2"))], ids=["comma", "word", "bool", "date"])
def test_qty_times_price_not_a_number_both_fail(qty, price, run_engine):
    """숫자로 읽을 수 없는 수량은 지금 판도 엔진도 멈춘다(조용히 빈칸으로 넘기지 않는다)."""
    rules = with_form([("합", "수량x단가")])
    result = invoice_result(rules)
    result.table = [{"수량": qty, "단가": price}]
    with pytest.raises(Exception):
        company_rows(result, None, rules.company_form)

    with pytest.raises(EngineError):
        engine_view(run_engine, rules, result_json(result))


@pytest.mark.parametrize("quantity, price", [
    (Decimal("1234567890123456789.1234"), Decimal("3.123456789012345678")),   # 42자리: decimal 에도 안 담긴다
    (3, Decimal("20000000000000000000000000001")),                            # 29자리: decimal 에는 담기지만 파이썬은 28자리로 반올림
], ids=["too_big_for_decimal", "fits_decimal_but_python_rounds"])
def test_qty_times_price_too_many_digits_stops(quantity, price, run_engine):
    """곱이 28자리를 넘으면(파이썬은 28자리로 반올림) 엔진은 조용히 다르게 담지 않고 멈춘다."""
    rules = with_form([("합", "수량x단가")])
    result = invoice_result(rules)
    result.table = [{"수량": quantity, "단가": price}]
    rounded = company_rows(result, None, rules.company_form)[1][0][0]
    with localcontext() as exact:
        exact.prec = 100
        assert rounded != Decimal(quantity) * price, "파이썬이 반올림해야 시험이 의미 있다"

    with pytest.raises(EngineError):
        engine_view(run_engine, rules, result_json(result))


def test_plain_values_keep_their_types(run_engine):
    """머리 · 품목의 값은 종류 그대로: 글자 · 정수 · 실수 · 참거짓 · 금액 · 날짜 · 시각이 있는 일시(엑셀)."""
    rules = with_form([("a", "머리.고객PO번호"), ("b", "머리.발행일"), ("c", "머리.합계금액"), ("d", "머리.정수"),
                       ("e", "머리.실수"), ("f", "머리.참거짓"), ("g", "머리.일시"), ("h", "머리.빈글자"),
                       ("i", "품목.수량"), ("j", "품목.품명"), ("k", "고정:"), ("l", "줄번호")])
    result = invoice_result(rules, 정수=7, 실수=12.5, 참거짓=True, 빈글자="", 일시=datetime(2026, 1, 5, 13, 4, 9))
    result.is_excel, result.source_form = True, "ACME"    # 엑셀이어야 시각이 있는 일시 글자를 일시로 되살린다
    result.fields["고객PO번호"] = "PO-1"

    expected, got = compare(run_engine, rules, result)

    assert got == expected
    row = expected["rows"][0]
    assert [c and c["t"] for c in row] == ["text", "date", "decimal", "int", "float", "bool", "datetime", "text",
                                           "int", "text", "text", "int"]
    assert row[10] == {"t": "text", "v": ""}


def test_missing_fields_and_none_values(run_engine):
    """못 찾은 필드(None)는 빈칸, 파일 이름에는 미상."""
    rules = with_form([("a", "머리.인도조건"), ("b", "머리.고객명")], filename="{인도조건}_{고객명}.xlsx")
    result = invoice_result(rules, 인도조건=None, 고객명="Acme / Co.")

    expected, got = compare(run_engine, rules, result)

    assert got == expected
    assert got["filename"] == "미상_Acme _ Co.xlsx"


# --- 마크업 열 이름 · 줄번호 ---

def test_custom_markup_columns_and_many_lines(run_engine, tmp_path):
    """회사양식.마크업 으로 품번열 · 단가열 · 수량열 이름을 바꾸면 그 열로 계산한다. 줄번호는 10 씩."""
    rules = with_form([("줄", "줄번호"), ("율", "마크업.율"), ("단가", "마크업.단가"), ("금액", "마크업.금액")],
                      markup={"품번열": "품명", "단가열": "금액", "수량열": "단위"})
    result = invoice_result(rules)
    result.table = [{"품명": f"P-{i}", "금액": Decimal("10.00") + i, "단위": i + 1, "품번": "무시"} for i in range(12)]
    table = markup_file(tmp_path, [[f"p-{i}", i, i % 4] for i in range(0, 12, 2)])

    expected, got = compare(run_engine, rules, result, table)

    assert got == expected
    assert [r[0]["v"] for r in got["rows"]] == [10 * (i + 1) for i in range(12)]
    assert got["rows"][1][1] is None and got["rows"][2][1] is not None


def test_form_with_unknown_markup_table_part(run_engine, tmp_path):
    """품번이 비었거나 표에 없는 줄의 마크업 칸은 빔(확인필요는 작업 9 가 따로 정한다)."""
    rules = with_form([("품번", "품목.품번"), ("본사", "마크업.단가")])
    result = invoice_result(rules)
    result.table = [{"품번": "", "단가": Decimal("5")}, {"품번": None, "단가": Decimal("5")},
                    {"품번": "A-1", "단가": Decimal("5")}, {"품번": " a-1 ", "단가": Decimal("5")}]
    table = markup_file(tmp_path, [["A-1", 10, 2]])

    expected, got = compare(run_engine, rules, result, table)

    assert got == expected
    assert [r[1] and r[1]["v"] for r in got["rows"]] == [None, None, "5.50", "5.50"]


# --- 파일 이름: render_text ---

D = date(2026, 1, 5)
T = datetime(2026, 1, 5, 13, 4, 9)
TEMPLATE_CASES = [
    ("plain", "PO_{a}.xlsx", {"a": "X-1"}),
    ("date_ym", "{고객명}/{발주일:%Y-%m}/PO_{번호}.pdf", {"고객명": "Acme", "발주일": D, "번호": "N-1"}),
    ("date_many", "{d:%y%m%d}_{d:%Y.%m.%d}_{d:%j}_{d:%H%M%S}", {"d": D}),
    ("date_names", "{d:%B %b %A %a %h}", {"d": D}),
    ("date_ampm", "{d:%I-%p} {d:%f}", {"d": T}),
    ("date_weeks", "{d:%U/%W/%w/%u}", {"d": D}),
    ("date_iso", "{d:%F}", {"d": D}),
    ("date_percent", "{d:%%Y%Y}", {"d": D}),
    ("date_colon", "{d:%H:%M:%S}", {"d": T}),
    ("date_year_start_sunday", "{d:%U %W %w %u %j}", {"d": date(2023, 1, 1)}),
    ("date_year_end", "{d:%U %W %w %u %j}", {"d": date(2024, 12, 31)}),
    ("date_nospec", "{d}_{t}", {"d": D, "t": T}),
    ("date_korean_literal", "{d:%Y년%m월}", {"d": D}),
    ("missing", "{a}_{b}", {"a": None, "b": "x"}),
    ("missing_any_spec", "{a:%Y-%m}{a:>10}{a:,.2f}", {"a": None}),
    ("empty_value", "[{a}]", {"a": ""}),
    ("blank_value", "[{a}]", {"a": "   "}),
    ("dots_only", "[{a}]", {"a": "..."}),
    ("trailing_dot", "[{a}]", {"a": "name."}),
    ("dot_then_space", "[{a}][{b}]", {"a": "a .", "b": " a. "}),
    ("forbidden", "{a}", {"a": 'a\\b/c:d*e?f"g<h>i|j'}),
    ("forbidden_in_literal_kept", "dir/{a}:x.xlsx", {"a": "v"}),
    ("unicode_space", "[{a}]", {"a": "\u3000 x \u00a0"}),
    ("control_space", "[{a}]", {"a": "\x1c x \x1f"}),
    ("zero_width_kept", "[{a}]", {"a": "\u200b x"}),
    ("braces", "a{{b}}c{x}{{", {"x": "v"}),
    ("decimal", "{a}_{b}_{c}", {"a": Decimal("1234.50"), "b": Decimal("0.000001"), "c": Decimal("1E-8")}),
    ("int_float_bool", "{a}_{b}_{c}_{d}", {"a": 5, "b": 1.0, "c": True, "d": 1e-05}),
    ("float_big", "{a}_{b}", {"a": 1e16, "b": 123456789.123456789}),
    ("str_spec_width", "[{a:>6}][{a:<6}][{a:^6}][{a:*^7}][{a:*^6}][{a:*>6}][{a:*<6}][{a:06}]", {"a": "abc"}),
    ("str_spec_prec", "{a:.2}_{a:5.2}_{a:s}", {"a": "abc"}),
    ("num_comma", "{a:,}_{b:,.2f}_{c:_}", {"a": 1234567, "b": Decimal("1234567.895"), "c": 1234567}),
    ("num_zero", "{a:08}_{b:010.2f}_{c:05d}_{d:08.1f}", {"a": 12, "b": Decimal("-3.14159"), "c": -7, "d": 2.5}),
    ("num_float_plain", "{a:08}_{b:+}_{c:,}_{d:>9}", {"a": 2.5, "b": 0.5, "c": 1234567.25, "d": -1.5}),
    ("num_decimal_plain", "{a:08}_{b:+}_{c:,}_{d:_}", {"a": Decimal("2.50"), "b": Decimal("0.5"), "c": Decimal("1234567.25"), "d": Decimal("-9876.5")}),
    ("num_sign", "{a:+}_{b:+.1f}_{c: }_{d:-}", {"a": 5, "b": Decimal("2.25"), "c": 3, "d": -3}),
    ("num_round_half_even", "{a:.0f}_{b:.1f}_{c:.2f}_{d:.0f}_{e:.1f}", {"a": Decimal("0.5"), "b": Decimal("0.25"),
                                                                          "c": Decimal("2.345"), "d": 2.5, "e": 0.25}),
    ("num_float_exact", "{a:.2f}_{b:.1f}_{c:f}", {"a": 2.675, "b": 0.05, "c": 1.5}),
    ("num_int_f", "{a:f}_{a:.1f}_{a:d}", {"a": 7}),
    ("num_decimal_f_noprec", "{a:f}_{b:f}", {"a": Decimal("1.500"), "b": Decimal("100")}),
    ("num_align_fill", "{a:*>8}_{b:x<8}_{c:#^9}_{c:#^8}_{d:*^9}", {"a": 12, "b": Decimal("1.5"), "c": 7, "d": -2.5}),
    ("num_align_equal", "{a:*=8}_{b:0=8}", {"a": -12, "b": 5}),
    ("num_neg_zero_round", "{a:.0f}_{b:.2f}", {"a": Decimal("-0.4"), "b": -0.001}),
    ("num_big_group", "{a:,}", {"a": 123456789012345678}),
    ("repeat", "{a}{a}_{b:%Y}{b:%m}", {"a": "x", "b": D}),
    ("extra_fields_unused", "{a}", {"a": "1", "b": "2", "c": None}),
    ("empty_template", "", {"a": "1"}),
    ("literal_only", "static.xlsx", {}),
]


@pytest.mark.parametrize("template, fields", [c[1:] for c in TEMPLATE_CASES], ids=[c[0] for c in TEMPLATE_CASES])
def test_template_same_as_python(template, fields, run_engine):
    """Format-Template 은 render_text 와 같은 글자: 날짜 서식 · 금지 글자 · 빈 값 미상 · 앞뒤 공백 · 끝의 점."""
    expected = render_text(template, fields)

    got = run_engine("company", {"action": "template", "template": template,
                                 "fields": {k: cell(v) for k, v in fields.items()}})

    assert got == {"text": expected}


def test_template_de_de_same(run_engine):
    """소수점 쉼표 문화권에서도 숫자 · 날짜 서식이 같다(월 이름은 영어)."""
    fields = {"a": Decimal("1234.5"), "d": D, "f": 2.5, "g": Decimal("-9876.25"), "h": 1234567.25, "i": 7}
    template = "{a:,.2f}_{d:%B %Y}_{f}_{a}_{g:010}_{g:,}_{h:,}_{h:08.1f}_{f:08}_{i:,.2f}_{a:.1f}"
    got = run_engine("company", {"action": "template", "template": template,
                                 "fields": {k: cell(v) for k, v in fields.items()}}, "de-DE")
    assert got == {"text": render_text(template, fields)}


def test_template_forbidden_chars_all_replaced(run_engine):
    """금지 글자 9개가 모두 _ 로: 글자마다 하나씩(여러 개면 여러 개)."""
    value = '\\/:*?"<>|'
    got = run_engine("company", {"action": "template", "template": "{a}", "fields": {"a": cell(value)}})
    assert got == {"text": "_________"} == {"text": render_text("{a}", {"a": value})}


@pytest.mark.parametrize("template, fields", [
    ("{a", {"a": 1}), ("a}", {"a": 1}), ("{}", {"a": 1}), ("{0}", {"a": 1}), ("{b}", {"a": 1}),
    ("{a:%Y}", {"a": "text"}), ("{a:%Y}", {"a": Decimal("1")}), ("{a:d}", {"a": "x"}), ("{a:+}", {"a": "x"}),
    ("{a:,}", {"a": "x"}), ("{a:d}", {"a": Decimal("1")}), ("{a:%Q}", {"a": D}), ("{a:%}", {"a": D}),
    ("{a:.2}", {"a": 5}), ("{a:<08}", {"a": Decimal("1")}), ("{a:=5}", {"a": "x"}), ("{a:#}", {"a": "x"}),
], ids=["open", "close", "auto", "index", "unknown", "date_on_text", "date_on_decimal", "d_on_text", "sign_on_text",
        "comma_on_text", "d_on_decimal", "bad_directive", "lone_percent", "prec_on_int", "zero_align_decimal",
        "equal_on_text", "alt_on_text"])
def test_template_errors_stop_both(template, fields, run_engine):
    """지금 판이 오류(ValueError · KeyError · IndexError)로 멈추는 서식은 엔진도 멈춘다. 조용히 빈 글자로 넘기지 않는다."""
    with pytest.raises((ValueError, KeyError, IndexError)):
        render_text(template, fields)

    with pytest.raises(EngineError):
        run_engine("company", {"action": "template", "template": template,
                               "fields": {k: cell(v) for k, v in fields.items()}})


@pytest.mark.parametrize("template, fields", [
    ("{a!r}", {"a": "x"}), ("{a.year}", {"a": D}), ("{a[0]}", {"a": "xyz"}), ("{a:{b}}", {"a": "x", "b": 5}),
], ids=["conversion", "attribute", "index", "nested"])
def test_template_unsupported_syntax_stops(template, fields, run_engine):
    """파이썬 서식 문법 중 무설치판이 지원하지 않는 것(변환 !r · 속성 · 첨자 · 안쪽 {})은 다르게 채우지 않고 멈춘다."""
    with pytest.raises(EngineError):
        run_engine("company", {"action": "template", "template": template,
                               "fields": {k: cell(v) for k, v in fields.items()}})


def test_template_matches_every_doc_type_name_on_samples(run_engine):
    """예시 규칙의 문서종류별 파일명 템플릿을 샘플 결과에 채워도 지금 판과 같다."""
    rules = example_rules()
    results = [extract.extract(INVOICE, rules), extract.extract(DM, rules), read_excel(ORDERS[0], rules)]
    for result in results:
        template = result.doc_type.filename_tpl
        got = run_engine("company", {"action": "template", "template": template,
                                     "fields": {k: cell(v) for k, v in result.fields.items()}})
        assert got == {"text": render_text(template, result.fields)}
        assert "/20" in got["text"] and "미상" not in got["text"]


# --- 이름 충돌: free_name ---

def engine_free_name(run_engine, path, reserved=(), culture=None):
    return run_engine("company", {"action": "free_name", "path": str(path), "reserved": [str(r) for r in reserved]},
                      culture)["path"]


def touch(*paths):
    for p in paths:
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(b"x")


def test_free_name_nothing_exists(run_engine, tmp_path):
    path = tmp_path / "본사PO_A-1.xlsx"
    assert engine_free_name(run_engine, path) == str(free_name(path)) == str(path)


def test_free_name_counts_up(run_engine, tmp_path):
    path = tmp_path / "본사PO_A-1.xlsx"
    touch(path)
    assert engine_free_name(run_engine, path) == str(free_name(path)) == str(tmp_path / "본사PO_A-1_2.xlsx")
    touch(tmp_path / "본사PO_A-1_2.xlsx", tmp_path / "본사PO_A-1_3.xlsx")
    assert engine_free_name(run_engine, path) == str(free_name(path)) == str(tmp_path / "본사PO_A-1_4.xlsx")


def test_free_name_reserved(run_engine, tmp_path):
    """이미 쓰기로 한 이름(reserved)도 피한다. 대소문자 · 앞뒤 / \\ 표기가 달라도 같은 경로로 본다."""
    path = tmp_path / "a.xlsx"
    reserved = [path, tmp_path / "a_2.xlsx"]
    assert engine_free_name(run_engine, path, reserved) == str(free_name(path, reserved)) == str(tmp_path / "a_3.xlsx")
    shaped = [str(path).upper(), str(tmp_path / "A_2.XLSX").replace("\\", "/")]
    assert free_name(path, [Path(s) for s in shaped]) == tmp_path / "a_3.xlsx"
    assert engine_free_name(run_engine, path, shaped) == str(tmp_path / "a_3.xlsx")
    assert engine_free_name(run_engine, path, [tmp_path / "other.xlsx"]) == str(path)


def test_free_name_reserved_and_existing_mixed(run_engine, tmp_path):
    path = tmp_path / "a.xlsx"
    touch(path, tmp_path / "a_3.xlsx")
    reserved = [tmp_path / "a_2.xlsx"]
    assert engine_free_name(run_engine, path, reserved) == str(free_name(path, reserved)) == str(tmp_path / "a_4.xlsx")


@pytest.mark.parametrize("name, taken, expected", [
    ("noext", ["noext"], "noext_2"),
    ("a.tar.gz", ["a.tar.gz"], "a.tar_2.gz"),
    (".hidden", [".hidden"], ".hidden_2"),
    ("a b (1).xlsx", ["a b (1).xlsx"], "a b (1)_2.xlsx"),
    ("공백 이름.pdf", ["공백 이름.pdf"], "공백 이름_2.pdf"),
    ("a.xlsx", ["a.xlsx", "a_2.xlsx", "a_3.xlsx", "a_4.xlsx", "a_5.xlsx", "a_6.xlsx"], "a_7.xlsx"),
], ids=["no_ext", "multi_dot", "leading_dot", "parens", "korean", "many"])
def test_free_name_stem_and_suffix(name, taken, expected, run_engine, tmp_path):
    touch(*[tmp_path / t for t in taken])
    path = tmp_path / name
    assert str(free_name(path)) == str(tmp_path / expected)
    assert engine_free_name(run_engine, path) == str(tmp_path / expected)


def test_free_name_directory_counts_as_taken(run_engine, tmp_path):
    (tmp_path / "dir.d").mkdir()
    path = tmp_path / "dir.d"
    assert engine_free_name(run_engine, path) == str(free_name(path)) == str(tmp_path / "dir_2.d")


def test_free_name_de_de(run_engine, tmp_path):
    path = tmp_path / "a.xlsx"
    touch(path)
    assert engine_free_name(run_engine, path, culture="de-DE") == str(tmp_path / "a_2.xlsx")


# --- C:\PO시험 (대표 시험 자료, 읽기만) ---

kit_missing = not (KIT / "03_고객마스터_마크업" / "마크업표.xlsx").exists() or not (KIT / "04_정답" / "참고_검증된_규칙.yaml").exists()


@pytest.fixture
def kit(tmp_path):
    """시험 자료를 tmp_path 로 복사한다(원본은 건드리지 않는다). 규칙에서는 고객마스터 열을 뺀다."""
    rules_text = (KIT / "04_정답" / "참고_검증된_규칙.yaml").read_text(encoding="utf-8")
    assert "값: 고객마스터." in rules_text
    rules = rules_from_text(without_master(rules_text))
    original = tmp_path / "마크업표_원본복사.xlsx"
    shutil.copyfile(KIT / "03_고객마스터_마크업" / "마크업표.xlsx", original)
    # 시험 자료의 마크업표는 같은 파트가 두 번 적힌 줄이 있어 지금 판도 읽기를 거절한다. 복사본에서 뒤쪽 중복 줄만 지운다
    wb = openpyxl.load_workbook(original)
    ws = wb.worksheets[0]
    keys = [str(ws.cell(row=r, column=1).value).strip().lower() for r in range(2, ws.max_row + 1)]
    for index in range(len(keys) - 1, -1, -1):
        if keys[index] in keys[:index]:
            ws.delete_rows(index + 2)
    wb.save(tmp_path / "마크업표.xlsx")
    return rules, tmp_path / "마크업표.xlsx", original


KIT_INPUTS = ["주문서_Orion_OWK-PO-26-1187.xlsx", "PO_BPI-2026-0733.pdf", "메일_Order LDH-PO-2611-033.pdf"]


@pytest.mark.po_test_kit
@pytest.mark.skipif(kit_missing, reason="C:\\PO시험 시험 자료가 없음")
@pytest.mark.parametrize("name", KIT_INPUTS)
def test_po_test_kit_same_as_python(name, kit, run_engine):
    """시험 자료 3건: 엔진이 추출한 결과로 회사 양식 줄을 만들어도 지금 판 company_rows 와 칸마다 · 파일 이름까지 같다."""
    rules, table_path, _ = kit
    source = KIT / "01_입력예시" / name
    portable = rules_to_portable(rules)
    if source.suffix == ".xlsx":
        result = read_excel(source, rules)
        engine_result = run_engine("extract_grid", {"rules": portable, "grid": grid_from_xlsx(source)})
    else:
        result = extract.extract(source, rules)
        engine_result = run_engine("extract_text", {"rules": portable, "text": extract.read_text(source)})
    table = M.markup_from_bytes(table_path.read_bytes())
    expected = python_view(result, rules, table)

    got = run_engine("company", {"action": "rows", "rules": portable, "result": engine_result,
                                 "markup_grid": grid_from_xlsx(table_path)})

    assert got == expected
    assert len(expected["rows"]) >= 5
    assert expected["filename"].startswith("본사PO_") and expected["filename"].endswith(".xlsx")
    hq = [i for i, h in enumerate(expected["header"]) if h in ("HQ Unit Price", "HQ Amount", "Markup %")]
    assert len(hq) == 3 and all(row[i] is not None for row in expected["rows"] for i in hq), "마크업 칸이 모두 채워져야 한다"
    line = expected["header"].index("Line")
    assert [row[line] for row in expected["rows"]] == [{"t": "int", "v": 10 * (i + 1)} for i in range(len(expected["rows"]))]
    # 문서종류의 파일명 템플릿(날짜 서식)도 같다
    template = result.doc_type.filename_tpl
    named = run_engine("company", {"action": "template", "template": template,
                                   "fields": {k: cell(v) for k, v in result.fields.items()}})
    assert named == {"text": render_text(template, result.fields)}
    assert "미상" not in named["text"]


@pytest.mark.po_test_kit
@pytest.mark.skipif(kit_missing, reason="C:/PO시험 시험 자료가 없음")
def test_po_test_kit_duplicate_markup_rows_stop_both(kit, run_engine):
    """시험 자료 마크업표 원본(같은 파트 두 번)은 지금 판도 엔진도 같은 문구로 읽기를 거절한다."""
    rules, _, original = kit
    with pytest.raises(M.MarkupError) as python_error:
        M.markup_from_bytes(original.read_bytes())

    with pytest.raises(EngineError) as engine_error:
        run_engine("company", {"action": "rows", "rules": rules_to_portable(rules),
                               "result": result_json(ExtractResult(rules.doc_types[0], {}, [])),
                               "markup_grid": grid_from_xlsx(original)})

    assert str(engine_error.value) == str(python_error.value)
