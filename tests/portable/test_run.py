"""무설치판 작업 실행(작업 9): -Command 실행 이 지금 판 run_job(승인 · 고객마스터 없이)과 같은 결과를 낸다.

시험은 모두 tmp_path 의 작업 폴더 · 입력 폴더 · 출력 폴더를 쓴다(C:\\PO시험 은 읽어서 tmp_path 로 복사만).
엑셀 · 워드가 필요한 시험은 @office. 시험마다 /Automation 오피스 · %TEMP%\\무설치판_* 이 남지 않았는지 본다.
"""
import csv
import hashlib
import json
import msvcrt
import os
import re
import shutil
import subprocess
import time
from datetime import datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
import pytest

from app import job_run
from app.jobs import create_job
from app.markup import write_example_markup
from app.notify import report_text, toast_script
from app.portable import rules_to_portable
from app.rules import rules_from_text

from .conftest import DISPATCHER, EXAMPLE_RULES, REPO, decode_output, without_master
from .test_office_read import KIT, automation_pids, no_leftover, office, own_temp, temp_dir  # noqa: F401 (픽스처)
from .test_office_write import assert_same_book, book_view

PDFS = REPO / "samples" / "03_Invoice_DM_PDF"
ORDERS = REPO / "samples" / "01_고객오더_엑셀"
INVOICE = PDFS / "INV-2026-10231.pdf"
ACME = ORDERS / "고객오더_ACME_PO-ACME-2026-0417.xlsx"
TITLE = "일일 정리"
LEDGER = f"누적장부_{TITLE}.xlsx"
RUN_TIMEOUT = 900
PDF_WARNING_ALERT = "워드에서 PDF 를 한 번 열고 '이 메시지를 다시 표시 안 함'에 체크"

# 테스트 전용 PO 규칙: PDF PO(tests/fixtures 의 ACME_PO_PO-77001.pdf)와 같은 번호의 엑셀 PO(시험 안에서 만듦)
PO_RULES = r'''
문서종류:
  PO:
    판별키워드: ["PURCHASE ORDER"]
    번호필드: 고객PO번호
    필드:
      고객PO번호: {패턴: "PO Number[.:]?\\s*([A-Z0-9-]+)"}
      발주일:     {패턴: "Date[.:]?\\s*(\\d{1,2}/\\d{1,2}/\\d{4})", 형식: 날짜}
      고객명:     {패턴: "Bill To[.:]?\\s*(.+?)\\n"}
      소계:       {패턴: "Subtotal[.:]?\\s*\\$?([\\d,]+\\.\\d{2})", 형식: 금액}
      합계금액:   {패턴: "\\nTotal[.:]?\\s*\\$?([\\d,]+\\.\\d{2})", 형식: 금액}
    표:
      시작키워드: "Item Description"
      끝키워드: "Subtotal"
      행패턴: "^(\\S+)\\s+(.+?)\\s+(\\d+)\\s+([\\d,]+\\.\\d{2})\\s+([\\d,]+\\.\\d{2})$"
      열: [품번, 품명, 수량, 단가, 금액]
      금액열: [단가, 금액]
      수량열: [수량]
    검산: {소계: 소계, 합계: 합계금액}
    파일명: "PO_{고객PO번호}.pdf"
엑셀양식:
  ACMEX:
    문서종류: PO
    판별키워드: ["EXCEL PURCHASE ORDER"]
    머리:
      고객PO번호: "오른쪽:PO Number:"
      발주일:     {칸: "오른쪽:Date:", 형식: 날짜}
      고객명:     "오른쪽:Bill To:"
      소계:       {칸: "오른쪽:Subtotal:", 형식: 금액}
      합계금액:   {칸: "오른쪽:Total:", 형식: 금액}
    품목:
      머리글: "Item"
      열: {품번: A, 품명: B, 수량: C, 단가: D, 금액: E}
      금액열: [단가, 금액]
      수량열: [수량]
'''
PO_ITEMS = [("CP-200", "Chiller Pump Assembly", 1, Decimal("4500.00"), Decimal("4500.00")),
            ("HX-900", "Heat Exchanger", 1, Decimal("7845.67"), Decimal("7845.67"))]


# --- 도우미 ---

class Env:
    """무설치판 작업 폴더 하나(작업\\NN_제목\\ · 규칙.json · 작업설정.json)와 입력 · 출력 폴더."""

    def __init__(self, root: Path, mode: str, rules_text: str, number: int = 1, title: str = TITLE,
                 work: Path | None = None, inp: Path | None = None, out: Path | None = None):
        self.work = work or root / "work"
        self.inp = inp or root / f"in{number}"
        self.out = out or root / f"out{number}"
        for d in (self.work, self.inp, self.out):
            d.mkdir(parents=True, exist_ok=True)
        self.number, self.title, self.mode = number, title, mode
        self.folder = self.work / "작업" / f"{number:02d}_{title}"
        self.folder.mkdir(parents=True, exist_ok=True)
        self.set_rules(rules_text)
        settings = {"number": number, "title": title, "input_dir": str(self.inp), "output_dir": str(self.out),
                    "output_mode": mode, "time": "16:00", "method": "없음", "last_run": None}
        (self.folder / "작업설정.json").write_text(json.dumps(settings, ensure_ascii=False, indent=2), encoding="utf-8")

    def set_rules(self, rules_text: str, fingerprint: str | None = None):
        data = rules_to_portable(rules_from_text(rules_text))
        data["fingerprint"] = fingerprint or hashlib.sha256(rules_text.encode("utf-8")).hexdigest()[:16]
        (self.folder / "규칙.json").write_text(json.dumps(data, ensure_ascii=False), encoding="utf-8")

    @property
    def rules_path(self) -> Path:
        return self.folder / "규칙.json"

    def records(self) -> list[dict]:
        path = self.folder / "처리기록.csv"
        if not path.exists():
            return []
        with open(path, encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f))

    def log(self) -> list[dict]:
        path = self.folder / "처리로그.csv"
        if not path.exists():
            return []
        with open(path, encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f))

    def note(self, name: str) -> Path:
        return self.folder / "확인필요" / f"{name}.txt"


def drop(src: Path, folder: Path, name: str | None = None, age: float = 120) -> Path:
    """원본을 입력 폴더에 복사하고 수정 시각을 age 초 전으로 둔다."""
    dst = Path(folder) / (name or Path(src).name)
    shutil.copyfile(src, dst)
    t = time.time() - age
    os.utime(dst, (t, t))
    return dst


def iso(value: datetime) -> str:
    return value.strftime("%Y-%m-%dT%H:%M:%S")


def run_portable(env: Env, now: datetime | None = None, *extra: str, expect: int | None = 0) -> dict:
    """무설치판.ps1 -Command 실행 -Job N -Work … -NoNotify. 보고 JSON(-Output) 과 반환 코드를 돌려준다."""
    out = env.work.parent / f"보고_{time.monotonic_ns()}.json"
    args = ["powershell.exe", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File", str(DISPATCHER),
            "-Command", "실행", "-Job", str(env.number), "-Work", str(env.work), "-Output", str(out), "-NoNotify"]
    if now is not None:
        args += ["-Now", iso(now)]
    args += list(extra)
    done = subprocess.run(args, capture_output=True, timeout=RUN_TIMEOUT)
    assert out.exists(), decode_output(done.stderr) + decode_output(done.stdout)
    report = json.loads(out.read_text(encoding="utf-8-sig"))
    report["exit_code"] = done.returncode
    if expect is not None:
        assert done.returncode == expect, (report, decode_output(done.stderr))
    return report


def counts(report: dict) -> dict:
    return dict(report["counts"])


def files(folder: Path) -> list[str]:
    """폴더 아래 파일들(상대 경로). 장부가 저장마다 남기는 바로 전 판(.bak.xlsx)은 뺀다."""
    folder = Path(folder)
    if not folder.is_dir():
        return []
    return sorted(str(p.relative_to(folder)) for p in folder.rglob("*") if p.is_file() and not p.name.endswith(".bak.xlsx"))


def sheet_rows(path: Path, sheet: str) -> list[dict]:
    rows = list(openpyxl.load_workbook(path)[sheet].iter_rows(values_only=True))
    header = [h for h in rows[0] if h is not None]
    return [dict(zip(header, r)) for r in rows[1:] if any(v is not None for v in r)]


def same_value(a, b) -> bool:
    if hasattr(a, "year") and hasattr(b, "year"):
        return (a.year, a.month, a.day) == (b.year, b.month, b.day)
    if a is None or b is None:
        return a is None and b is None
    try:
        return Decimal(str(a)) == Decimal(str(b))
    except (InvalidOperation, ValueError):
        return str(a) == str(b)


def excel_po(path: Path, number: str = "PO-77001", total: str = "12345.67", items=PO_ITEMS, age: float = 120) -> Path:
    """PO_RULES 의 엑셀양식 ACMEX 모양의 가짜 엑셀 PO(금액 칸은 #,##0.00)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "PO"
    ws["A1"] = "EXCEL PURCHASE ORDER"
    ws["A2"], ws["B2"] = "PO Number:", number
    ws["A3"], ws["B3"] = "Date:", datetime(2026, 10, 15)
    ws["B3"].number_format = "yyyy-mm-dd"
    ws["A4"], ws["B4"] = "Bill To:", "ACME Semiconductor Korea"
    for col, label in zip("ABCDE", ("Item", "Description", "Qty", "Unit Price", "Amount")):
        ws[f"{col}6"] = label
    row = 7
    for part, desc, qty, price, amount in items:
        ws[f"A{row}"], ws[f"B{row}"], ws[f"C{row}"] = part, desc, qty
        ws[f"D{row}"], ws[f"E{row}"] = float(price), float(amount)
        ws[f"D{row}"].number_format = ws[f"E{row}"].number_format = "#,##0.00"
        row += 1
    subtotal = sum(i[4] for i in items)
    ws[f"D{row + 1}"], ws[f"E{row + 1}"] = "Subtotal:", float(subtotal)
    ws[f"D{row + 2}"], ws[f"E{row + 2}"] = "Total:", float(Decimal(total))
    ws[f"E{row + 1}"].number_format = ws[f"E{row + 2}"].number_format = "#,##0.00"
    wb.save(path)
    t = time.time() - age
    os.utime(path, (t, t))
    return path


def fixture_pdfs(tmp_path: Path) -> Path:
    import importlib.util

    spec = importlib.util.spec_from_file_location("build_fixtures", REPO / "tests" / "fixtures" / "build_fixtures.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    folder = tmp_path / "fixtures"
    module.build(folder)
    return folder


HOLDER = r"""
param([string]$Path, [string]$Ready, [string]$Stop)
$excel = New-Object -ComObject Excel.Application
try {
    $excel.Visible = $false
    $excel.DisplayAlerts = $false
    $book = $excel.Workbooks.Open($Path)
    [IO.File]::WriteAllText($Ready, 'ok')
    $deadline = [DateTime]::UtcNow.AddSeconds(600)
    while (-not (Test-Path -LiteralPath $Stop) -and [DateTime]::UtcNow -lt $deadline) { Start-Sleep -Milliseconds 200 }
    $book.Close($false)
    [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($book)
} finally {
    $excel.Quit()
    [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($excel)
}
"""


class ExcelHolder:
    """다른 powershell 이 엑셀 COM 으로 파일을 연 채 두는 동안(with 블록)."""

    def __init__(self, path: Path, scratch: Path):
        self.path, self.scratch = path, scratch

    def __enter__(self):
        script = self.scratch / "holder.ps1"
        script.write_text(HOLDER, encoding="utf-8-sig")
        self.ready, self.stop = self.scratch / "ready.txt", self.scratch / "stop.txt"
        for p in (self.ready, self.stop):
            p.unlink(missing_ok=True)
        self.before = automation_pids()
        self.proc = subprocess.Popen(["powershell.exe", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
                                      "-File", str(script), str(self.path), str(self.ready), str(self.stop)])
        deadline = time.monotonic() + 90
        while not self.ready.exists():
            assert self.proc.poll() is None, "파일을 여는 엑셀이 먼저 끝남"
            assert time.monotonic() < deadline, "엑셀이 파일을 열지 못함"
            time.sleep(0.2)
        self.mine = automation_pids() - self.before   # 이 시험용 엑셀
        return self

    def __exit__(self, *exc):
        self.stop.write_text("stop", encoding="utf-8")
        try:
            self.proc.wait(timeout=90)
        except subprocess.TimeoutExpired:
            self.proc.kill()
            for _, pid in automation_pids() - self.before:
                subprocess.run(["taskkill", "/F", "/PID", str(pid)], capture_output=True)
        wait_office_gone(self.mine)
        return False


def wait_office_gone(pids: set, wait: float = 30) -> None:
    """Quit 한 엑셀은 powershell 이 끝난 뒤에도 몇 초(이 PC 에서 약 3초) 더 남는다. 사라질 때까지 기다린다(no_leftover 가 잘못 잡지 않게)."""
    deadline = time.monotonic() + wait
    while pids & automation_pids() and time.monotonic() < deadline:
        time.sleep(0.5)


@pytest.fixture
def clean_temp(temp_dir):
    """엔진의 %TEMP% 를 시험 폴더로 바꾸고, 끝난 뒤 이 프로그램의 임시 폴더(무설치판_*)가 남았으면 실패."""
    yield temp_dir
    assert own_temp(temp_dir) == [], "임시 폴더가 남음"


def example_text() -> str:
    return without_master(EXAMPLE_RULES.read_text(encoding="utf-8"))


# --- 오피스 없이 ---

def test_lock_prevents_overlap(tmp_path):
    """실행중.lock 을 다른 프로세스가 잡고 있으면(지금 판 job_lock 과 같은 1바이트 잠금) 아무것도 하지 않고 멈춘다."""
    env = Env(tmp_path, "누적장부", example_text())
    drop(INVOICE, env.inp)
    with open(env.folder / "실행중.lock", "a+b") as f:
        f.seek(0)
        msvcrt.locking(f.fileno(), msvcrt.LK_NBLCK, 1)
        try:
            report = run_portable(env, expect=2)
        finally:
            f.seek(0)
            msvcrt.locking(f.fileno(), msvcrt.LK_UNLCK, 1)
    assert report["locked"] is True
    assert "이 작업이 이미 실행 중입니다(01_일일 정리)" in report["error"]
    assert env.records() == [] and env.log() == []
    assert files(env.out) == []
    record = (env.folder / "실행기록.txt").read_text(encoding="utf-8")
    assert "이 작업이 이미 실행 중입니다" in record


def test_missing_input_folder_stops(tmp_path):
    """입력 폴더가 없으면 아무 파일도 처리하지 않고 「중요:」 로 멈춘다(반환 코드 1, 지금 판 _preflight)."""
    env = Env(tmp_path, "누적장부", example_text())
    env.inp.rmdir()
    report = run_portable(env, expect=1)
    assert report["stopped"] is True and report["has_news"] is True
    assert report["message"].startswith(f"중요: {TITLE} 작업을 실행하지 못했습니다(아무 파일도 처리하지 않음).\n입력 폴더가 없습니다: {env.inp}")
    assert env.records() == []


def test_unknown_job(tmp_path):
    """없는 작업 번호 → 반환 코드 3, 그 글자만."""
    env = Env(tmp_path, "누적장부", example_text())
    env.number = 7
    report = run_portable(env, expect=3)
    assert report["error"].startswith("작업 7 이 없습니다") and report["locked"] is False


def test_bad_rules_stop(tmp_path):
    env = Env(tmp_path, "누적장부", example_text())
    env.rules_path.write_text("{ 깨진", encoding="utf-8")
    report = run_portable(env, expect=1)
    assert report["stopped"] is True and "규칙.json 오류" in report["message"]


@pytest.mark.parametrize("report", [
    {"counts": {"성공": 3, "중복": 0, "확인필요": 1, "보류": 2}, "stopped": False, "ledger_locked": False,
     "unmatched_old": 0, "unmatched_xls": 0, "alerts": []},
    {"counts": {"성공": 0, "중복": 1, "확인필요": 0, "보류": 0}, "stopped": True, "ledger_locked": True,
     "unmatched_old": 3, "unmatched_xls": 1, "alerts": ["같은 고객 두 줄", PDF_WARNING_ALERT]},
    {"counts": {"성공": 0, "중복": 0, "확인필요": 0, "보류": 0}, "stopped": False, "ledger_locked": False,
     "unmatched_old": 2, "unmatched_xls": 2, "alerts": []},
])
def test_report_text_same_as_python(report, run_engine):
    """알림 글자 = 지금 판 notify.report_text(승인대기 칸 없음)."""
    class Job:
        title = TITLE

    python = job_run.RunReport(counts={**report["counts"], "승인대기": 0}, unmatched_old=report["unmatched_old"],
                               unmatched_xls=report["unmatched_xls"], ledger_locked=report["ledger_locked"],
                               stopped=report["stopped"], alerts=list(report["alerts"]))
    got = run_engine("notify", {"action": "text", "title": TITLE, "report": report})
    assert got == {"text": report_text(Job, python)}


@pytest.mark.parametrize("folder", [None, "C:/사무 자동화/출력 폴더#1", "C:/out"])
def test_toast_xml_same_as_python(folder, run_engine):
    """WinRT 알림 XML = 지금 판 notify.toast_script 가 LoadXml 에 넣는 글자(누르면 그 폴더, & < > " 이스케이프)."""
    title, body = "사무자동화", f"{TITLE}: 결과 1건, <확인> & \"따옴표\""
    script = toast_script(title, body, Path(folder) if folder else None)
    quoted = re.search(r"\$xml\.LoadXml\('((?:[^']|'')*)'\)", script).group(1)
    got = run_engine("notify", {"action": "toast_xml", "title": title, "body": body,
                                "folder": str(Path(folder)) if folder else None})
    assert got == {"xml": quoted.replace("''", "'")}


# --- 오피스로 ---

@office
@pytest.mark.office
def test_same_outputs_as_python(tmp_path, clean_temp, no_leftover):
    """샘플 PDF 4개 + 엑셀 2개(2분 전) → 회사 양식 · 장부 · 검산 파일 · 확인필요 사유 · 처리로그가 지금 판 run_job(gate=None)과 같다."""
    text = example_text()
    shared_in = tmp_path / "in"
    shared_in.mkdir()
    for name in ("DM-2026-0092.pdf", "DM-2026-0097.pdf", "INV-2026-10231.pdf", "INV-2026-10278.pdf"):
        drop(PDFS / name, shared_in)
    for path in sorted(ORDERS.glob("*.xlsx")):
        drop(path, shared_in)
    before = {p.name: p.read_bytes() for p in shared_in.iterdir()}
    now = datetime.now().replace(microsecond=0)

    # 지금 판(승인 없이)
    py_work, py_out = tmp_path / "py_work", tmp_path / "py_out"
    py_work.mkdir(), py_out.mkdir()
    write_example_markup(py_work / "마크업표.xlsx")
    job = create_job(py_work, TITLE, shared_in, py_out, "둘다")
    job.rules_path.write_text(text, encoding="utf-8")
    python = job_run.run_job(job, py_work, now=now, gate=None)

    # 무설치판(같은 입력 폴더, 원본은 읽기만)
    env = Env(tmp_path, "둘다", text, inp=shared_in)
    write_example_markup(env.work / "마크업표.xlsx")
    started = time.perf_counter()
    report = run_portable(env, now)
    elapsed = time.perf_counter() - started
    print(f"\n무설치판 실행 시간: {elapsed:.1f}초 (문서 6개)")

    assert {p.name: p.read_bytes() for p in shared_in.iterdir()} == before  # 원본 그대로
    assert counts(report) == {k: python.counts[k] for k in ("성공", "중복", "확인필요", "보류")}
    assert counts(report)["성공"] >= 2 and counts(report)["확인필요"] >= 1
    assert report["ledger_locked"] is False and report["stopped"] is False
    # 출력 폴더: 회사 양식 · 장부(그리고 .bak)
    assert files(env.out) == files(py_out)
    for name in files(py_out):
        assert_same_book(book_view(env.out / name), book_view(py_out / name))
    assert (env.out / LEDGER.replace(".xlsx", ".bak.xlsx")).exists() == (py_out / LEDGER.replace(".xlsx", ".bak.xlsx")).exists()
    # 검산 파일
    assert files(env.folder / "검산") == files(job.folder / "검산")
    for name in files(job.folder / "검산"):
        assert_same_book(book_view(env.folder / "검산" / name), book_view(job.folder / "검산" / name))
    # 확인필요 사유 파일(바이트까지)
    assert files(env.folder / "확인필요") == files(job.folder / "확인필요")
    for name in files(job.folder / "확인필요"):
        assert (env.folder / "확인필요" / name).read_bytes() == (job.folder / "확인필요" / name).read_bytes()
    # 처리로그 · 처리기록
    with open(job.folder / "처리로그.csv", encoding="utf-8-sig", newline="") as f:
        assert env.log() == list(csv.DictReader(f))
    assert (env.folder / "처리로그.csv").read_bytes()[:3] == b"\xef\xbb\xbf"
    keep = ("지문", "파일", "문서종류", "번호", "상태", "날짜")
    assert [{k: r[k] for k in keep} for r in env.records()] == [{k: r[k] for k in keep} for r in job_run.load_records(job)]
    with open(env.folder / "처리기록.csv", encoding="utf-8-sig", newline="") as f:
        assert next(csv.reader(f)) == job_run.RECORD_COLUMNS
    settings = json.loads((env.folder / "작업설정.json").read_text(encoding="utf-8"))
    assert settings["last_run"] == now.date().isoformat()
    # 실행기록
    assert report["message"] in (env.folder / "실행기록.txt").read_text(encoding="utf-8")


KIT_INPUTS = ["주문서_Orion_OWK-PO-26-1187.xlsx", "PO_BPI-2026-0733.pdf", "메일_Order LDH-PO-2611-033.pdf"]
KIT_RULES = KIT / "04_정답" / "참고_검증된_규칙.yaml"
KIT_MARKUP = KIT / "03_고객마스터_마크업" / "마크업표.xlsx"
kit_missing = not KIT_RULES.exists() or not KIT_MARKUP.exists() or not all((KIT / "01_입력예시" / n).exists() for n in KIT_INPUTS)


@office
@pytest.mark.office
@pytest.mark.po_test_kit
@pytest.mark.skipif(kit_missing, reason="C:\\PO시험 시험 자료가 없음")
def test_po_test_kit_matches_answer(tmp_path, clean_temp, no_leftover):
    """C:\\PO시험 입력 3개 → 본사 PO 3개 · 누적장부가 04_정답 의 정답표와 같다(고객마스터 열은 무설치판에 없어 뺀다)."""
    env = Env(tmp_path, "둘다", without_master(KIT_RULES.read_text(encoding="utf-8")), title="PO 시험")
    # 시험 자료 마크업표는 같은 파트 두 줄이 있어(대표가 덧붙인 줄) 지금 판도 거절한다. 복사본에서 뒤쪽 중복 줄만 지운다
    wb = openpyxl.load_workbook(KIT_MARKUP)
    ws = wb.worksheets[0]
    keys = [str(ws.cell(row=r, column=1).value).strip().lower() for r in range(2, ws.max_row + 1)]
    for index in range(len(keys) - 1, -1, -1):
        if keys[index] in keys[:index]:
            ws.delete_rows(index + 2)
    wb.save(env.work / "마크업표.xlsx")
    for name in KIT_INPUTS:
        drop(KIT / "01_입력예시" / name, env.inp)

    report = run_portable(env)

    assert counts(report) == {"성공": 3, "중복": 0, "확인필요": 0, "보류": 0}, report["message"]
    answer = KIT / "04_정답" / "정답표_기대결과.xlsx"
    expected = sheet_rows(answer, "본사PO시트_기대값")
    forms = sorted(p.name for p in env.out.glob("본사PO_*.xlsx"))
    assert forms == sorted(f"본사PO_{n}.xlsx" for n in ("OWK-PO-26-1187", "BPI-2026-0733", "LDH-PO-2611-033"))
    compared = 0
    for name in forms:
        got = sheet_rows(env.out / name, "HQ_PO")
        want = [r for r in expected if r["Customer PO No"] == name[len("본사PO_"):-len(".xlsx")]]
        assert len(got) == len(want) and got
        for g, w in zip(got, want):
            for column, value in g.items():
                assert same_value(value, w[column]), (name, column, value, w[column])
                compared += 1
    assert compared > 21 * 15
    ledger_answer = KIT / "04_정답" / "누적장부_기대결과.xlsx"
    ledger = env.out / "누적장부_PO 시험.xlsx"
    for sheet in ("PO", "메일PO", "품목"):
        got, want = sheet_rows(ledger, sheet), sheet_rows(ledger_answer, sheet)
        assert len(got) == len(want), sheet
        for g, w in zip(got, want):
            for column in w:
                if column not in ("처리일시", "정리된경로"):
                    assert same_value(g.get(column), w[column]), (sheet, column, g.get(column), w[column])


@office
@pytest.mark.office
def test_rerun_skips_and_renamed_skips(tmp_path, clean_temp, no_leftover):
    """두 번째 실행 → 아무것도 처리하지 않음. 파일 이름을 바꿔도(내용 지문이 같으면) 그대로."""
    env = Env(tmp_path, "누적장부", example_text())
    drop(INVOICE, env.inp)
    first = run_portable(env)
    assert counts(first)["성공"] == 1
    ledger = env.out / LEDGER
    data = ledger.read_bytes()

    second = run_portable(env)
    assert counts(second) == {"성공": 0, "중복": 0, "확인필요": 0, "보류": 0} and second["has_news"] is False
    os.rename(env.inp / INVOICE.name, env.inp / "다시받은_인보이스.pdf")
    third = run_portable(env)
    assert counts(third) == {"성공": 0, "중복": 0, "확인필요": 0, "보류": 0}
    assert third["unmatched_old"] == 0
    assert ledger.read_bytes() == data
    assert [r["상태"] for r in env.records()] == ["성공"]
    assert env.records()[0]["지문"] == hashlib.sha256(INVOICE.read_bytes()).hexdigest()


@office
@pytest.mark.office
def test_settling_file_waits(tmp_path, clean_temp, no_leftover):
    """방금(1분 안) 바뀐 파일 → 보류(처리기록 없음), 다음 실행(시각 +2분) 때 처리."""
    env = Env(tmp_path, "누적장부", example_text())
    path = drop(ACME, env.inp, age=0)
    now = datetime.fromtimestamp(path.stat().st_mtime).replace(microsecond=0) + timedelta(seconds=10)

    first = run_portable(env, now)
    assert counts(first) == {"성공": 0, "중복": 0, "확인필요": 0, "보류": 1}
    assert env.records() == []
    assert [(r["파일"], r["결과"], r["사유"]) for r in env.log()] == [(ACME.name, "보류", job_run.SETTLING_REASON)]
    assert not (env.out / LEDGER).exists()

    second = run_portable(env, now + timedelta(minutes=2))
    assert counts(second)["성공"] == 1
    assert [r["상태"] for r in env.records()] == ["성공"]


@office
@pytest.mark.office
def test_same_po_pdf_and_excel_is_duplicate(tmp_path, clean_temp, no_leftover):
    """같은 번호의 PO 가 PDF 와 엑셀로 → 하나 성공 · 하나 중복, 장부 PO 시트 한 줄(지금 판과 같은 쪽이 성공)."""
    text = without_master(PO_RULES)
    pdf = fixture_pdfs(tmp_path) / "ACME_PO_PO-77001.pdf"
    env = Env(tmp_path, "누적장부", text)
    drop(pdf, env.inp)
    excel_po(env.inp / "ACME_PO_PO-77001_엑셀.xlsx")

    py_work, py_out = tmp_path / "py_work", tmp_path / "py_out"
    py_work.mkdir(), py_out.mkdir()
    job = create_job(py_work, TITLE, env.inp, py_out, "누적장부")
    job.rules_path.write_text(text, encoding="utf-8")
    python = job_run.run_job(job, py_work, gate=None)
    assert python.counts["성공"] == 1 and python.counts["중복"] == 1  # 시험 자료가 맞는지(지금 판)

    report = run_portable(env)

    assert counts(report) == {"성공": 1, "중복": 1, "확인필요": 0, "보류": 0}
    assert [(r["파일"], r["상태"]) for r in env.records()] == [(r["파일"], r["상태"]) for r in job_run.load_records(job)]
    assert len(sheet_rows(env.out / LEDGER, "PO")) == 1
    assert [(r["파일"], r["결과"], r["사유"]) for r in env.log()][1] == ("ACME_PO_PO-77001_엑셀.xlsx", "중복", "PO PO-77001 이미 처리됨")


TABLE_FORM = '''
회사양식:
  파일명: "SO_{고객PO번호}.xlsx"
  시트: SO
  열:
    - {이름: PO, 값: 머리.고객PO번호}
    - {이름: Part, 값: 품목.품번}
    - {이름: Amount, 값: 품목.금액}
'''


def split_row_po(path: Path, number: str) -> Path:
    """ACME_PO_PO-77001.pdf 모양인데 첫 품목 줄의 품명이 두 줄로 쪼개진 PDF(워드 변환이 표 행을 쪼갠 경우 흉내).
    워드로 읽으면 'CP-200 Chiller Pump 1 4,500.00 4,500.00' 다음에 'Assembly' 한 줄이 따로 나온다."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=letter, invariant=1)
    c.setFont("Helvetica-Bold", 20)
    c.drawString(72, 720, "PURCHASE ORDER")
    c.setFont("Helvetica", 11)
    c.drawString(72, 690, "ACME Supply Inc.")
    c.drawString(72, 660, f"PO Number: {number}")
    c.drawString(72, 645, "Date: 10/15/2026")
    c.drawString(72, 630, "Bill To: ACME Semiconductor Korea")
    y = 590
    c.setFont("Helvetica-Bold", 10)
    for x, label in [(72, "Item"), (150, "Description")]:
        c.drawString(x, y, label)
    for x, label in [(390, "Qty"), (470, "Unit Price"), (540, "Amount")]:
        c.drawRightString(x, y, label)
    c.setFont("Helvetica", 10)
    y -= 18
    c.drawString(72, y, "CP-200")
    c.drawString(150, y, "Chiller Pump")
    for x, value in [(390, "1"), (470, "4,500.00"), (540, "4,500.00")]:
        c.drawRightString(x, y, value)
    y -= 12
    c.drawString(150, y, "Assembly")
    y -= 18
    c.drawString(72, y, "HX-900")
    c.drawString(150, y, "Heat Exchanger")
    for x, value in [(390, "1"), (470, "7,845.67"), (540, "7,845.67")]:
        c.drawRightString(x, y, value)
    y -= 30
    c.drawString(390, y, "Subtotal: 12,345.67")
    c.setFont("Helvetica-Bold", 11)
    c.drawString(390, y - 18, "Total: $12,345.67")
    c.showPage()
    c.save()
    t = time.time() - 120
    os.utime(path, (t, t))
    return path


@office
@pytest.mark.office
def test_unreadable_item_table_goes_to_review(tmp_path, clean_temp, no_leftover):
    """규칙에 품목 표가 있는데 읽지 못하면(워드가 첫 품목 줄을 둘로 쪼갬 · 엑셀 금액 칸이 숫자가 아님) 「성공」 이 아니라 확인필요.
    회사 양식 · 장부 줄이 생기지 않는다. 같은 실행의 정상 PDF 는 그대로 성공(시험 자료가 맞는지)."""
    env = Env(tmp_path, "둘다", without_master(PO_RULES) + TABLE_FORM)
    drop(fixture_pdfs(tmp_path) / "ACME_PO_PO-77001.pdf", env.inp)
    split = split_row_po(env.inp / "PO 쪼갠줄 PO-77002.pdf", "PO-77002")
    bad = excel_po(env.inp / "PO 금액칸 글자 PO-77003.xlsx", number="PO-77003")
    book = openpyxl.load_workbook(bad)
    book.active["E7"] = "N/A"   # 첫 품목 금액 칸이 숫자가 아님 → 엑셀 품목 표를 통째로 버림
    book.save(bad)
    t = time.time() - 120
    os.utime(bad, (t, t))

    report = run_portable(env)

    assert counts(report) == {"성공": 1, "중복": 0, "확인필요": 2, "보류": 0}
    states = {r["파일"]: r["상태"] for r in env.records()}
    assert states == {"ACME_PO_PO-77001.pdf": "성공", split.name: "확인필요", bad.name: "확인필요"}
    note = env.note(split.name).read_text(encoding="utf-8").split("\n")
    assert note[1] == "사유: 품목 표를 읽지 못함: 표 행패턴과 맞지 않는 줄이 있어 품목 건너뜀: Assembly"
    note = env.note(bad.name).read_text(encoding="utf-8").split("\n")
    assert note[1] == "사유: 품목 표를 읽지 못함: 품목 7행 '금액' 값이 숫자가 아님, 품목 건너뜀"
    assert sorted(f for f in files(env.out) if f.endswith(".xlsx") and f.startswith("SO_")) == ["SO_PO-77001.xlsx"]
    assert [r["고객PO번호"] for r in sheet_rows(env.out / LEDGER, "PO")] == ["PO-77001"]
    assert {r["문서번호"] for r in sheet_rows(env.out / LEDGER, "품목")} == {"PO-77001"}


@office
@pytest.mark.office
def test_rules_without_table_still_succeed(tmp_path, clean_temp, no_leftover):
    """규칙에 품목 표가 없는 문서종류는 그대로 성공(품목 없이 머리만 장부에). 표 검사는 표가 정의된 문서에만 건다."""
    text = re.sub(r"\n    표:\n(?:      .*\n)+", "\n", without_master(PO_RULES))
    assert "행패턴" not in text.split("엑셀양식:")[0]
    env = Env(tmp_path, "누적장부", text)
    split = split_row_po(env.inp / "PO 쪼갠줄 PO-77002.pdf", "PO-77002")

    report = run_portable(env)

    assert counts(report)["성공"] == 1, report
    assert [r["상태"] for r in env.records()] == ["성공"]
    assert [r["고객PO번호"] for r in sheet_rows(env.out / LEDGER, "PO")] == ["PO-77002"]
    assert split.exists()


@office
@pytest.mark.office
def test_long_name_keeps_original_name(tmp_path, clean_temp, no_leftover):
    """이름이 긴(입력 폴더 경로까지 200자쯤) 엑셀 PO: 오피스는 짧은 이름의 복사본(원본.xlsx)을 열지만, 처리기록 · 처리로그 ·
    장부의 원본파일에는 원래 이름이 들어간다."""
    env = Env(tmp_path, "누적장부", without_master(PO_RULES))
    room = 200 - len(str(env.inp)) - 1 - len(".xlsx")
    name = ("메일_Order 긴 이름 PO-77001 " + "가나다 ABC " * 40)[:room].rstrip() + ".xlsx"
    path = excel_po(env.inp / name)
    assert len(str(path)) >= 195

    report = run_portable(env)

    assert counts(report)["성공"] == 1, report
    assert [(r["파일"], r["상태"]) for r in env.records()] == [(name, "성공")]
    assert [r["파일"] for r in env.log()] == [name]
    assert [r["원본파일"] for r in sheet_rows(env.out / LEDGER, "PO")] == [name]


@office
@pytest.mark.office
def test_ledger_open_holds(tmp_path, clean_temp, no_leftover):
    """장부를 엑셀로 연 채 실행 → 보류(장부 무변화, 알림 대상). 마크업표를 엑셀로 연 채 실행 → 복사본으로 읽어 정상 처리."""
    env = Env(tmp_path, "누적장부", example_text())
    drop(INVOICE, env.inp)
    assert counts(run_portable(env))["성공"] == 1
    ledger = env.out / LEDGER
    data, stamp = ledger.read_bytes(), ledger.stat().st_mtime_ns
    drop(PDFS / "INV-2026-10278.pdf", env.inp)

    with ExcelHolder(ledger, tmp_path):
        held = run_portable(env)
    assert counts(held) == {"성공": 0, "중복": 0, "확인필요": 0, "보류": 1}
    assert held["ledger_locked"] is True and held["has_news"] is True
    assert ledger.read_bytes() == data and ledger.stat().st_mtime_ns == stamp
    assert env.log()[-1]["사유"] == job_run.LEDGER_LOCKED_REASON
    assert env.records()[-1]["상태"] == "보류"
    assert "누적장부가 엑셀에서 열려 있어 파일 1개를 처리하지 못했습니다(보류)" in held["message"]
    assert counts(run_portable(env))["성공"] == 1  # 닫은 뒤 다음 실행 때 처리

    company = Env(tmp_path, "회사양식", example_text(), number=2, title="엑셀 PO", work=env.work)
    write_example_markup(env.work / "마크업표.xlsx")
    drop(ACME, company.inp)
    with ExcelHolder(env.work / "마크업표.xlsx", tmp_path):
        done = run_portable(company)
    assert counts(done) == {"성공": 1, "중복": 0, "확인필요": 0, "보류": 0}, done["message"]
    assert files(company.out) == ["회사양식_SO입력_PO-ACME-2026-0417.xlsx"]


@office
@pytest.mark.office
def test_bad_total_goes_to_review_and_retries_after_rule_change(tmp_path, clean_temp, no_leftover):
    """합계가 틀린 문서 → 확인필요 .txt(지금 판 문구), 결과 · 장부 무변화. 같은 날 다시 → 건너뜀, 규칙 지문이 바뀌거나 다음 날이면 다시 시도."""
    text = without_master(PO_RULES)
    env = Env(tmp_path, "누적장부", text)
    path = excel_po(env.inp / "PO_합계틀림.xlsx", total="12345.68")
    now = datetime.now().replace(microsecond=0)

    first = run_portable(env, now)
    assert counts(first)["확인필요"] == 1
    note = env.note(path.name).read_text(encoding="utf-8")
    lines = note.split("\n")
    assert lines[0] == f"원본: {path}" and lines[1].startswith("사유: 검산 불일치: ")
    assert note.endswith(f"\n\n{job_run.RETRY_NOTE}\n")
    assert b"\r\n" in env.note(path.name).read_bytes()
    assert files(env.out) == [] and not (env.folder / "검산").exists()

    assert counts(run_portable(env, now))["확인필요"] == 0  # 같은 날 · 같은 규칙: 건너뜀
    assert counts(run_portable(env, now + timedelta(days=1)))["확인필요"] == 1  # 다음 날: 다시 시도
    env.set_rules(text, fingerprint="0123456789abcdef")
    assert counts(run_portable(env, now + timedelta(days=1)))["확인필요"] == 1  # 규칙 지문이 바뀜: 같은 날도
    assert [r["상태"] for r in env.records()] == ["확인필요"] * 3
    assert files(env.out) == []


@office
@pytest.mark.office
def test_missing_markup_part_retries_after_markup_change(tmp_path, clean_temp, no_leftover):
    """마크업표에 없는 파트 → 확인필요. 같은 날 다시 → 건너뜀. 마크업표에 그 파트를 더하면 같은 날 다시 시도해 성공."""
    env = Env(tmp_path, "회사양식", example_text())
    markup = env.work / "마크업표.xlsx"
    write_example_markup(markup)
    wb = openpyxl.load_workbook(markup)
    ws = wb.active
    removed = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == "VTS-MFD-12")  # ACME PO 의 파트
    part = ws.cell(removed, 1).value
    ws.delete_rows(removed)
    wb.save(markup)
    drop(ACME, env.inp)
    now = datetime.now().replace(microsecond=0)

    first = run_portable(env, now)
    assert counts(first)["확인필요"] == 1, first["message"]
    assert f"마크업표에 없는 파트넘버: {part}" in env.note(ACME.name).read_text(encoding="utf-8")
    assert counts(run_portable(env, now))["확인필요"] == 0

    write_example_markup(markup)
    again = run_portable(env, now)
    assert counts(again)["성공"] == 1, again["message"]
    assert not env.note(ACME.name).exists()
    assert files(env.out) == ["회사양식_SO입력_PO-ACME-2026-0417.xlsx"]


@office
@pytest.mark.office
def test_markup_read_once_at_start(tmp_path, clean_temp, no_leftover):
    """마크업표는 실행 시작 때 한 번 읽은 바이트로 쓴다(지금 판 _read_markup). 실행 중에 친구가 마크업표를 고쳐 저장해도 이번 실행은
    시작 때의 표로 처리한다(마크업지문과 쓴 표가 어긋나지 않게)."""
    env = Env(tmp_path, "회사양식", example_text())
    markup = env.work / "마크업표.xlsx"
    write_example_markup(markup)
    changed = tmp_path / "바뀐_마크업표.xlsx"
    shutil.copyfile(markup, changed)
    wb = openpyxl.load_workbook(changed)
    ws = wb.active
    ws.delete_rows(next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value == "VTS-MFD-12"))
    wb.save(changed)
    drop(PDFS / "Letter_warranty_notice.pdf", env.inp)   # 먼저 읽히는(이름 순) 판별 안 되는 PDF: 워드가 읽는 동안 마크업표를 바꾼다
    drop(ACME, env.inp)
    out = tmp_path / "보고.json"
    before = automation_pids()
    proc = subprocess.Popen(["powershell.exe", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File",
                             str(DISPATCHER), "-Command", "실행", "-Job", "1", "-Work", str(env.work), "-Output", str(out),
                             "-NoNotify"])
    swapped = False
    try:
        deadline = time.monotonic() + 120
        while not any(name == "WINWORD.EXE" for name, _ in automation_pids() - before):   # 워드가 뜸 = 마크업표 바이트는 이미 읽음
            assert proc.poll() is None and time.monotonic() < deadline, "워드가 뜨지 않음"
            time.sleep(0.3)
        shutil.copyfile(changed, markup)
        swapped = True
    finally:
        try:
            proc.wait(timeout=RUN_TIMEOUT if swapped else 60)   # 위에서 실패했으면 오래 기다리지 않는다
        except subprocess.TimeoutExpired:
            proc.kill()
            proc.wait(timeout=60)
    assert proc.returncode == 0
    report = json.loads(out.read_text(encoding="utf-8-sig"))

    assert counts(report) == {"성공": 1, "중복": 0, "확인필요": 0, "보류": 0}, report["message"]
    assert files(env.out) == ["회사양식_SO입력_PO-ACME-2026-0417.xlsx"]
    assert markup.read_bytes() == changed.read_bytes()


@office
@pytest.mark.office
def test_company_failure_rolls_back_ledger(tmp_path, clean_temp, no_leftover):
    """둘다: 장부에 넣은 뒤 회사 양식 쓰기가 실패하면 장부를 넣기 전 판으로 되돌리고(지금 판 _rollback) 확인필요."""
    env = Env(tmp_path, "둘다", example_text())
    write_example_markup(env.work / "마크업표.xlsx")
    hanil = sorted(ORDERS.glob("*HANIL*.xlsx"))[0]
    drop(hanil, env.inp)
    assert counts(run_portable(env))["성공"] == 1
    ledger = env.out / LEDGER
    data = ledger.read_bytes()
    # 회사 양식 임시 파일 자리에 폴더를 두어 엑셀 저장이 실패하게 한다
    (env.out / ".회사양식_SO입력_PO-ACME-2026-0417.tmp.xlsx").mkdir()
    (env.out / ".회사양식_SO입력_PO-ACME-2026-0417.tmp.xlsx" / "막음.txt").write_text("x", encoding="utf-8")
    drop(ACME, env.inp)

    report = run_portable(env)

    assert counts(report) == {"성공": 0, "중복": 0, "확인필요": 1, "보류": 0}, report["message"]
    assert "결과를 쓰는 중 오류: " in env.note(ACME.name).read_text(encoding="utf-8")
    assert ledger.read_bytes() == data
    assert len(sheet_rows(ledger, "PO")) == 1
    assert not (env.out / "회사양식_SO입력_PO-ACME-2026-0417.xlsx").exists()
    assert [r["상태"] for r in env.records()] == ["성공", "확인필요"]


@office
@pytest.mark.office
def test_file_vanishing_mid_run_affects_only_that_file(tmp_path, clean_temp, no_leftover, monkeypatch):
    """지문을 구한 뒤 파일이 사라져도(옮김 · 지움) 그 파일만 영향: 지금 판처럼(os.access 거짓 → 읽기 실패) 확인필요,
    다른 문서는 처리하고 실행은 정상으로 끝난다(last_run 저장)."""
    env = Env(tmp_path, "누적장부", example_text())
    drop(INVOICE, env.inp)
    drop(ACME, env.inp)
    monkeypatch.setenv("NOINSTALL_TEST_VANISH", INVOICE.name)   # 시험용: 지문을 구한 바로 뒤 이 이름의 파일을 지운다
    now = datetime.now().replace(microsecond=0)

    report = run_portable(env, now)

    assert counts(report) == {"성공": 1, "중복": 0, "확인필요": 1, "보류": 0}, report["message"]
    assert "사유: 파일을 읽지 못함: " in env.note(INVOICE.name).read_text(encoding="utf-8")
    assert [(r["파일"], r["상태"]) for r in env.records()] == [(INVOICE.name, "확인필요"), (ACME.name, "성공")]
    settings = json.loads((env.folder / "작업설정.json").read_text(encoding="utf-8"))
    assert settings["last_run"] == now.date().isoformat()


@office
@pytest.mark.office
def test_word_timeout_only_that_doc(tmp_path, clean_temp, no_leftover):
    """워드 시간 제한 0초: PDF 는 확인필요(워드가 응답하지 않음), 엑셀은 성공. 이 실행이 띄운 오피스는 남지 않는다."""
    env = Env(tmp_path, "누적장부", example_text())
    drop(INVOICE, env.inp)
    drop(ACME, env.inp)

    report = run_portable(env, None, "-WordTimeoutSec", "0")

    assert counts(report) == {"성공": 1, "중복": 0, "확인필요": 1, "보류": 0}
    note = env.note(INVOICE.name).read_text(encoding="utf-8")
    assert "사유: 워드가 응답하지 않음\n" in note
    assert [(r["파일"], r["상태"]) for r in env.records()] == [(INVOICE.name, "확인필요"), (ACME.name, "성공")]
    assert sheet_rows(env.out / LEDGER, "PO")[0]["고객PO번호"] == "PO-ACME-2026-0417"


@office
@pytest.mark.office
def test_pdf_warning_on_leaves_pdfs_for_next_run(tmp_path, clean_temp, no_leftover):
    """워드의 PDF 변환 안내가 켜져 있으면(설계 6-4) PDF 는 열지 않고 그대로 두고(기록 없음) 엑셀만 처리, 알림으로 안내한다."""
    env = Env(tmp_path, "누적장부", example_text())
    drop(INVOICE, env.inp)
    drop(ACME, env.inp)

    report = run_portable(env, None, "-TestPdfWarningOff", "0")

    assert counts(report) == {"성공": 1, "중복": 0, "확인필요": 0, "보류": 0}
    assert [r["파일"] for r in env.records()] == [ACME.name]
    assert [r["파일"] for r in env.log()] == [ACME.name]
    assert any(PDF_WARNING_ALERT in a for a in report["alerts"]) and report["has_news"] is True
    assert "경고: " in report["message"] and PDF_WARNING_ALERT in report["message"]

    later = run_portable(env, None, "-TestPdfWarningOff", "1")
    assert counts(later)["성공"] == 1 and later["alerts"] == []
    assert [r["파일"] for r in env.records()] == [ACME.name, INVOICE.name]
