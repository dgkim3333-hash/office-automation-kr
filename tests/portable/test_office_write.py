"""무설치판 오피스 쓰기(작업 8): 엑셀 COM 으로 쓴 회사 양식 · 누적장부 · 검산 파일이 지금 판(openpyxl)이 쓴 파일과 같다.

두 파일을 openpyxl 로 읽어 시트 이름 · 순서 · 칸 값 · 숫자 형식 · 표 이름을 견준다. 날짜 형식만 글자가 다르다:
엑셀(한국 Windows)은 형식 안의 '-' 를 지역 날짜 구분자로 바꿔 저장하므로(yyyy-mm-dd → 기본 형식 14, 다른 지역에서는
10/14/2026 으로 보임) 무설치판은 '-' 를 따옴표로 감싼다(yyyy"-"mm"-"dd). 보이는 모양은 같다 → 비교 전에 따옴표를 뗀다.
엑셀이 필요한 시험은 @office(없으면 건너뜀). 시험마다 /Automation 엑셀이 남지 않았는지 · 임시 폴더가 남지 않았는지 본다.
"""
import re
import subprocess
import time
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

from app import extract
from app import markup as M
from app.excel_extract import read_excel
from app.ledger import Ledger
from app.outputs import company_rows, read_back, write_check_file, write_company_form
from app.portable import rules_to_portable
from app.verify import verify

from .conftest import REPO, EngineError, example_rules
from .grid import grid_from_xlsx
from .test_engine_company import cell, result_json, sample_markup
from .test_office_read import automation_pids, no_leftover, office, own_temp, temp_dir  # noqa: F401 (픽스처)

PDFS = REPO / "samples" / "03_Invoice_DM_PDF"
INVOICE = PDFS / "INV-2026-10231.pdf"
DM = PDFS / "DM-2026-0092.pdf"
ORDERS = sorted((REPO / "samples" / "01_고객오더_엑셀").glob("*.xlsx"))
WHEN = [datetime(2026, 10, 28, 9, 0, 5), datetime(2026, 10, 28, 9, 1, 0), datetime(2026, 10, 29, 16, 30, 59)]


# --- 도우미 ---

def fmt(number_format: str) -> str:
    """숫자 형식을 보이는 모양으로: 따옴표로 감싼 글자는 따옴표를 뗀다(yyyy"-"mm"-"dd → yyyy-mm-dd)."""
    return re.sub(r'"([^"]*)"', r"\1", number_format)


def sheet_view(ws) -> dict:
    """시트 한 장: 값이 있는 칸마다 (값, 보이는 형식), 굵은 칸, 틀 고정, 표 이름."""
    cells = {}
    bold = []
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                cells[c.coordinate] = (c.value, fmt(c.number_format))
                if c.font is not None and c.font.b:
                    bold.append(c.coordinate)
    return {"cells": cells, "bold": bold, "freeze": ws.freeze_panes, "tables": sorted(ws.tables)}


def book_view(path) -> dict:
    wb = openpyxl.load_workbook(path)
    try:
        return {"sheets": wb.sheetnames, "active": wb.active.title,
                "views": {ws.title: sheet_view(ws) for ws in wb.worksheets}}
    finally:
        wb.close()


def assert_same_book(got, want):
    """칸 값은 == 로(엑셀은 97000.00 을 정수 97000 으로 저장, openpyxl 은 실수 97000.0 → 같은 값)."""
    assert got["sheets"] == want["sheets"]
    assert got["active"] == want["active"]
    for name in want["sheets"]:
        g, w = got["views"][name], want["views"][name]
        assert g["cells"].keys() == w["cells"].keys(), name
        for key in w["cells"]:
            assert g["cells"][key] == w["cells"][key], f"{name}!{key}"
            assert type(g["cells"][key][0]) is type(w["cells"][key][0]) or (
                isinstance(g["cells"][key][0], (int, float)) and isinstance(w["cells"][key][0], (int, float))), f"{name}!{key}"
        assert g["bold"] == w["bold"], name
        assert g["freeze"] == w["freeze"], name
        assert g["tables"] == w["tables"], name


def table_refs(path) -> dict:
    wb = openpyxl.load_workbook(path)
    return {(ws.title, name): ws.tables[name].ref for ws in wb.worksheets for name in ws.tables}


def write(run_engine, action: str, culture=None, **payload) -> dict:
    """op office_write. 이 실행이 띄운 엑셀이 Quit 뒤 스스로 끝났는지(강제로 끝낸 것 0개)도 본다."""
    out = run_engine("office_write", {"action": action, **payload}, culture)
    assert out.pop("forced") == 0, "Quit 한 엑셀이 스스로 끝나지 않아 강제로 끝냄(COM 참조가 남음)"
    return out


def _iso(value: datetime) -> str:
    return value.strftime("%Y-%m-%dT%H:%M:%S")


def ledger_add(run_engine, ledger, rules, result, source, when, organized):
    """장부에 한 건. 엑셀이 연 파일 경로(opened, 시험용)는 떼고 돌려준다."""
    out = write(run_engine, "ledger_add", ledger=str(ledger), rules=rules_to_portable(rules), result=result_json(result),
                source=source, when=_iso(when), organized=str(organized))
    out.pop("opened", None)
    return out


def documents(rules):
    """샘플 PDF 2개 + 엑셀 1개의 (결과, 원본 이름, 정리된경로)."""
    return [(extract.extract(INVOICE, rules), INVOICE.name, f"정리문서/Invoice/{INVOICE.name}"),
            (extract.extract(DM, rules), DM.name, f"정리문서/DM/{DM.name}"),
            (read_excel(ORDERS[0], rules), ORDERS[0].name, f"정리문서/PO/{ORDERS[0].name}")]


def python_ledger(path, rules, docs):
    for (result, name, organized), when in zip(docs, WHEN):
        ledger = Ledger.open(path, rules)
        ledger.add(result, name, organized, processed_at=when)
        ledger.save()


def fake_checks(result):
    """검산 목록(일치 · 불일치 · 해당 없음이 모두 있게, 제어문자 섞인 글자 포함)."""
    checks = verify(result, INVOICE)
    check = type(checks[0])
    checks += [check("합계\x07금액", "1.00", "2.00", False), check("운임", "", "", None)]
    rows = [{"item": c.item, "source_value": c.source, "result_value": c.output, "ok": c.ok} for c in checks]
    return checks, rows


@pytest.fixture
def clean_temp(temp_dir):
    """엔진의 %TEMP% 를 시험 폴더로 바꾸고, 시험이 끝난 뒤 이 프로그램의 임시 폴더(무설치판_*)가 남았으면 실패."""
    yield temp_dir
    assert own_temp(temp_dir) == [], "임시 폴더가 남음"


# --- 오피스 없이 ---

def test_file_locked_free_and_missing(tmp_path, run_engine):
    """아무도 열지 않은 파일 · 없는 파일 → 잠기지 않음. 검사가 파일을 바꾸지 않는다."""
    free = tmp_path / "누적장부.xlsx"
    free.write_bytes(b"PK fake")
    stamp = free.stat().st_mtime_ns
    assert run_engine("office_write", {"action": "locked", "path": str(free)}) == {"locked": False}
    assert run_engine("office_write", {"action": "locked", "path": str(tmp_path / "없음.xlsx")}) == {"locked": False}
    assert free.read_bytes() == b"PK fake" and free.stat().st_mtime_ns == stamp
    with open(free, "rb"):   # 다른 프로그램이 연 상태(지금 판 시험과 같은 흉내)
        assert run_engine("office_write", {"action": "locked", "path": str(free)}) == {"locked": True}


# --- 엑셀로 쓰기 ---

@office
@pytest.mark.office
@pytest.mark.parametrize("culture", [None, "de-DE"])
def test_company_form_same_as_python(culture, tmp_path, clean_temp, run_engine, no_leftover):
    """샘플 엑셀 → 회사 양식: 지금 판 write_company_form 파일과 시트 · 칸 값 · 숫자 형식 · 굵은 머리글 · 틀 고정이 같다.
    소수점이 쉼표인 문화권(de-DE)에서 돌려도 같다."""
    rules = example_rules()
    result = read_excel(ORDERS[0], rules)
    table = sample_markup(tmp_path, [result])
    want = write_company_form(tmp_path / "py", result, None, rules.company_form, M.markup_from_bytes(table.read_bytes()))

    out = write(run_engine, "company_form", culture, folder=str(tmp_path / "ps"), rules=rules_to_portable(rules),
                result=result_json(result), markup_grid=grid_from_xlsx(table))

    got = Path(out["path"])
    assert got == tmp_path / "ps" / want.name
    assert_same_book(book_view(got), book_view(want))
    kinds = {type(v[0]).__name__ for v in book_view(got)["views"][rules.company_form.sheet]["cells"].values()}
    assert {"datetime", "str", "int"} <= kinds
    formats = {v[1] for v in book_view(got)["views"][rules.company_form.sheet]["cells"].values()}
    assert {"yyyy-mm-dd", "#,##0.00"} <= formats
    assert [p.name for p in (tmp_path / "ps").iterdir()] == [want.name], "임시 파일이 남음"

    # 같은 이름이 있으면 _2 (지금 판 free_name). 기존 파일은 그대로
    data = got.read_bytes()
    again = write(run_engine, "company_form", culture, folder=str(tmp_path / "ps"), rules=rules_to_portable(rules),
                  result=result_json(result), markup_grid=grid_from_xlsx(table))
    assert Path(again["path"]).name == want.stem + "_2" + want.suffix
    assert got.read_bytes() == data


@office
@pytest.mark.office
def test_ledger_same_as_python(tmp_path, clean_temp, run_engine, no_leftover):
    """PDF 2개 + 엑셀 1개를 차례로 장부에 넣기: 시트 이름 · 순서 · 머리 · 값 · 형식 · 표 이름이 지금 판 Ledger 와 같다.
    처리일시(일시) · 발행일(날짜) · 금액(자리 수 서식)이 모두 들어간다. 두 번째 저장부터 .bak.xlsx 가 생긴다."""
    rules = example_rules()
    docs = documents(rules)
    want = tmp_path / "py" / "누적장부.xlsx"
    want.parent.mkdir()
    python_ledger(want, rules, docs)

    got = tmp_path / "ps" / "누적장부.xlsx"
    got.parent.mkdir()
    backup = got.with_name("누적장부.bak.xlsx")
    for i, ((result, name, organized), when) in enumerate(zip(docs, WHEN)):
        before = got.read_bytes() if got.exists() else None
        out = ledger_add(run_engine, got, rules, result, name, when, organized)
        assert out == {"locked": False, "duplicate": False, "saved": True, "problems": [], "path": str(got),
                       "created": i == 0, "backup": None if i == 0 else str(backup)}
        if i == 0:
            assert not backup.exists()
        else:
            assert backup.read_bytes() == before   # 바로 전 판
    assert sorted(p.name for p in got.parent.iterdir()) == ["누적장부.bak.xlsx", "누적장부.xlsx"], "임시 파일이 남음"

    view, expected = book_view(got), book_view(want)
    assert_same_book(view, expected)
    assert view["sheets"] == ["Invoice", "DM", "PO", "품목", "처리로그"]
    # 표 이름은 규칙의 문서종류 순서(시트 순서가 아님) → 품목 → 처리로그
    names = [d.name for d in rules.doc_types] + ["품목", "처리로그"]
    assert [view["views"][n]["tables"] for n in names] == [[f"Ledger_{i + 1}"] for i in range(5)]
    invoice = view["views"]["Invoice"]["cells"]
    assert invoice["A2"] == (WHEN[0], "yyyy-mm-dd hh:mm:ss")
    assert {v[1] for v in invoice.values()} >= {"yyyy-mm-dd", "yyyy-mm-dd hh:mm:ss", "#,##0.00"}
    wb = openpyxl.load_workbook(got)
    assert [ws.title for ws in wb.worksheets if ws.sheet_view.tabSelected] == ["Invoice"]
    assert table_refs(got) == table_refs(want)


@office
@pytest.mark.office
def test_user_sheet_kept(tmp_path, clean_temp, run_engine, no_leftover):
    """사용자가 만든 시트 '메모'(맨 앞, 표 포함)와 문서 시트 끝의 사용자 열은 그대로 남는다.
    예전 시트순서(DM 이 앞)로 만든 장부도 새 시트순서로 문서 시트만 자리를 바꾼다(지금 판 test_ledger.py 와 같은 경우)."""
    import copy
    rules = example_rules()
    old = copy.deepcopy(rules)
    old.sheet_order = ["DM", "Invoice", "PO"]
    docs = documents(rules)
    book = tmp_path / "누적장부.xlsx"
    write(run_engine, "ledger_add", ledger=str(book), rules=rules_to_portable(old), result=result_json(docs[0][0]),
          source=docs[0][1], when=_iso(WHEN[0]), organized=docs[0][2])
    wb = openpyxl.load_workbook(book)
    assert wb.sheetnames == ["DM", "Invoice", "PO", "품목", "처리로그"]
    memo = wb.create_sheet("메모", 0)
    memo.append(["고객", "메모"])
    memo.append(["ACME", "=1+1"])
    from openpyxl.worksheet.table import Table
    memo.add_table(Table(displayName="내표", ref="A1:B2"))
    invoice = wb["Invoice"]
    width = invoice.max_column
    invoice.cell(1, width + 1, "내메모")
    invoice.cell(2, width + 1, "전화 완료")
    wb.save(book)

    out = ledger_add(run_engine, book, rules, docs[1][0], docs[1][1], WHEN[1], docs[1][2])
    assert out["saved"] is True and out["problems"] == []

    wb = openpyxl.load_workbook(book)
    assert wb.sheetnames == ["메모", "Invoice", "DM", "PO", "품목", "처리로그"]
    assert [[c.value for c in r] for r in wb["메모"].iter_rows()] == [["고객", "메모"], ["ACME", "=1+1"]]
    assert list(wb["메모"].tables) == ["내표"]
    invoice = wb["Invoice"]
    assert invoice.cell(1, width + 1).value == "내메모" and invoice.cell(2, width + 1).value == "전화 완료"
    assert invoice.tables[next(iter(invoice.tables))].ref.startswith("A1:")
    assert len([r for r in wb["DM"].iter_rows(min_row=2, values_only=True) if any(r)]) == 1
    assert wb.active.title == "Invoice"


@office
@pytest.mark.office
def test_duplicate_and_log(tmp_path, clean_temp, run_engine, no_leftover):
    """장부에 있는 문서종류 · 번호면 중복(장부는 그대로). 처리로그 줄은 일시 · 파일 · 결과 · 사유(빈 사유는 빈칸)."""
    rules = example_rules()
    portable = rules_to_portable(rules)
    result, name, organized = documents(rules)[0]
    book = tmp_path / "누적장부.xlsx"
    assert write(run_engine, "ledger_dup", ledger=str(book), rules=portable, doc_type="Invoice",
                 number="INV-2026-10231") == {"duplicate": False}   # 장부가 아직 없음
    ledger_add(run_engine, book, rules, result, name, WHEN[0], organized)
    data = book.read_bytes()

    assert write(run_engine, "ledger_dup", ledger=str(book), rules=portable, doc_type="Invoice",
                 number="INV-2026-10231") == {"duplicate": True}
    assert write(run_engine, "ledger_dup", ledger=str(book), rules=portable, doc_type="DM",
                 number="INV-2026-10231") == {"duplicate": False}
    out = ledger_add(run_engine, book, rules, result, name, WHEN[1], organized)
    assert out["duplicate"] is True and out["saved"] is False
    assert book.read_bytes() == data

    logs = [{"when": _iso(WHEN[1]), "file": "a.pdf", "outcome": "성공", "reason": ""},
            {"when": _iso(WHEN[2]), "file": "b.pdf", "outcome": "확인필요", "reason": "판별 실패"}]
    out = write(run_engine, "ledger_log", ledger=str(book), rules=portable, logs=logs)
    assert out["saved"] is True and out["locked"] is False

    py = tmp_path / "py.xlsx"
    ledger = Ledger.open(py, rules)
    ledger.add(result, name, organized, processed_at=WHEN[0])
    ledger.save()
    ledger = Ledger.open(py, rules)
    ledger.add_log(WHEN[1], "a.pdf", "성공", "")
    ledger.add_log(WHEN[2], "b.pdf", "확인필요", "판별 실패")
    ledger.save()
    assert book_view(book)["views"]["처리로그"] == book_view(py)["views"]["처리로그"]
    rows = list(openpyxl.load_workbook(book)["처리로그"].iter_rows(min_row=2, values_only=True))
    assert rows == [(WHEN[1], "a.pdf", "성공", None), (WHEN[2], "b.pdf", "확인필요", "판별 실패")]


@office
@pytest.mark.office
def test_undo_restores_previous(tmp_path, clean_temp, run_engine, no_leftover):
    """되돌리기: 있던 장부는 바로 전 판으로, 이번에 새로 만든 장부는 지운다(엑셀 없이)."""
    rules = example_rules()
    docs = documents(rules)
    book = tmp_path / "누적장부.xlsx"
    first = ledger_add(run_engine, book, rules, *docs[0][:2], WHEN[0], docs[0][2])
    before = book.read_bytes()
    second = ledger_add(run_engine, book, rules, *docs[1][:2], WHEN[1], docs[1][2])
    assert book.read_bytes() != before

    run_engine("office_write", {"action": "undo_ledger", "added": second})
    assert book.read_bytes() == before
    run_engine("office_write", {"action": "undo_ledger", "added": first})
    assert not book.exists()


@office
@pytest.mark.office
def test_read_back_failure_keeps_outputs_unchanged(tmp_path, clean_temp, run_engine, no_leftover):
    """엑셀이 쓴 값과 다르게 저장하는 값(15자리를 넘는 정수 금액) → 저장 뒤 다시 읽기가 잡는다:
    회사 양식은 파일을 남기지 않고 오류, 장부는 바꾸지 않고 problems(지금 판도 같은 값을 다르다고 본다)."""
    rules = example_rules()
    portable = rules_to_portable(rules)
    huge = "1234567890123456789.00"

    order = result_json(read_excel(ORDERS[0], rules))
    order["table"][1]["단가"] = huge
    with pytest.raises(EngineError) as error:
        write(run_engine, "company_form", folder=str(tmp_path / "out"), rules=portable, result=order)
    assert str(error.value).startswith("저장한 파일을 다시 읽어 보니 쓴 값과 다릅니다: 3행 Unit Price: 쓴 값 " + huge)
    assert not (tmp_path / "out").exists() or list((tmp_path / "out").iterdir()) == []

    result, name, organized = documents(rules)[0]
    book = tmp_path / "누적장부.xlsx"
    ledger_add(run_engine, book, rules, result, name, WHEN[0], organized)
    data = book.read_bytes()
    invoice = result_json(documents(rules)[1][0])
    invoice["fields"]["합계금액"] = huge
    out = write(run_engine, "ledger_add", ledger=str(book), rules=portable, result=invoice, source="DM.pdf",
                when=_iso(WHEN[1]), organized="x")
    assert out["saved"] is False and out["locked"] is False
    assert out["problems"] == [f"누적장부 DM 합계금액: 쓴 값 {huge}, 읽은 값 1.2345678901234568e+18"]
    assert book.read_bytes() == data and not book.with_name("누적장부.bak.xlsx").exists()
    assert sorted(p.name for p in tmp_path.iterdir() if p.suffix == ".xlsx") == ["누적장부.xlsx"]


@office
@pytest.mark.office
def test_check_file_same_as_python(tmp_path, clean_temp, run_engine, no_leftover):
    """검산 파일: 검산(원본파일 줄 · 항목 · 일치 표시) · 추출값 · 품목 · 회사양식 시트가 지금 판 write_check_file 과 같다."""
    rules = example_rules()
    result = extract.extract(INVOICE, rules)
    checks, rows = fake_checks(result)
    company = company_rows(result, None, rules.company_form)
    want = write_check_file(tmp_path / "py", result, checks, "메일_Order INV.pdf", company=company)

    out = write(run_engine, "check_file", folder=str(tmp_path / "ps"), rules=rules_to_portable(rules),
                result=result_json(result), checks=rows, source="메일_Order INV.pdf", company=True)

    got = Path(out["path"])
    assert got.name == want.name == "검산_INV-2026-10231.xlsx"
    view = book_view(got)
    assert_same_book(view, book_view(want))
    assert view["sheets"] == ["검산", "추출값", "품목", "회사양식"]
    assert view["views"]["검산"]["freeze"] == "A3"
    marks = {v[0] for k, v in view["views"]["검산"]["cells"].items() if k.startswith("D") and k != "D2"}
    assert marks == {"일치", "불일치", "해당 없음"}
    assert [p.name for p in (tmp_path / "ps").iterdir()] == [want.name]

    # 회사 양식 없이 · 품목 없는 결과: 검산 · 추출값 두 시트
    out = write(run_engine, "check_file", folder=str(tmp_path / "ps"), rules=rules_to_portable(rules),
                result={**result_json(result), "table": []}, checks=[], source="x.pdf", company=False)
    assert openpyxl.load_workbook(out["path"]).sheetnames == ["검산", "추출값"]
    assert Path(out["path"]).name == "검산_INV-2026-10231_2.xlsx"


@office
@pytest.mark.office
def test_read_back_detects_change(tmp_path, clean_temp, run_engine, no_leftover):
    """저장 뒤 칸 하나를 바꾸면 Test-ReadBack 이 그 칸을 알려 준다. 문구는 지금 판 read_back 과 같다."""
    rules = example_rules()
    result = read_excel(ORDERS[0], rules)
    out = write(run_engine, "company_form", folder=str(tmp_path), rules=rules_to_portable(rules),
                result=result_json(result))
    path = Path(out["path"])
    header, rows = company_rows(result, None, rules.company_form)
    sheet = rules.company_form.sheet
    payload = {"path": str(path), "sheet": sheet, "header": header, "rows": [[cell(v) for v in row] for row in rows]}

    assert write(run_engine, "read_back", **payload) == {"problems": []}
    assert read_back(path, sheet, header, rows) == []

    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]
    qty = header.index("Qty") + 1
    ws.cell(3, qty).value = 999
    ws.cell(2, 1).value = "바꾼 날짜"
    ws.cell(1, 2).value = "바꾼 머리"
    wb.save(path)

    problems = write(run_engine, "read_back", **payload)["problems"]
    assert problems == read_back(path, sheet, header, rows)
    assert len(problems) == 3
    assert problems[0] == f"머리글 2열: 쓴 값 {header[1]}, 읽은 값 바꾼 머리"
    assert problems[1].startswith("2행 Order Date: 쓴 값 2026-") and problems[1].endswith("읽은 값 바꾼 날짜")
    assert problems[2].startswith("3행 Qty: 쓴 값 ") and problems[2].endswith("읽은 값 999")

    assert write(run_engine, "read_back", **{**payload, "sheet": "없는시트"}) == {"problems": ["시트 '없는시트' 가 없습니다"]}


@office
@pytest.mark.office
def test_timeout_leaves_nothing(tmp_path, clean_temp, run_engine, no_leftover):
    """엑셀이 시간 안에 끝내지 못하면(시간 제한 0초) 「엑셀이 응답하지 않음」: 이 실행이 띄운 엑셀은 끝내고,
    회사 양식 폴더에는 아무 파일도(임시 파일도) 없고, 장부 · .bak 은 그대로."""
    rules = example_rules()
    portable = rules_to_portable(rules)
    order = result_json(read_excel(ORDERS[0], rules))
    with pytest.raises(EngineError) as error:
        run_engine("office_write", {"action": "company_form", "folder": str(tmp_path / "out"), "rules": portable,
                                    "result": order, "timeout": 0})
    assert str(error.value) == "엑셀이 응답하지 않음"
    assert not (tmp_path / "out").exists() or list((tmp_path / "out").iterdir()) == []

    result, name, organized = documents(rules)[0]
    book = tmp_path / "누적장부.xlsx"
    ledger_add(run_engine, book, rules, result, name, WHEN[0], organized)
    data = book.read_bytes()
    with pytest.raises(EngineError) as error:
        run_engine("office_write", {"action": "ledger_add", "ledger": str(book), "rules": portable,
                                    "result": result_json(documents(rules)[1][0]), "source": "DM.pdf",
                                    "when": _iso(WHEN[1]), "timeout": 0})
    assert str(error.value) == "엑셀이 응답하지 않음"
    assert book.read_bytes() == data
    assert sorted(p.name for p in tmp_path.iterdir() if p.suffix == ".xlsx") == ["누적장부.xlsx"]



# --- 검토 수정 1회차 ---

REPLACE_FAIL = "NOINSTALL_TEST_REPLACE_FAIL"   # 시험용 주입: 장부 바꾸기(ReplaceFile)가 중간에 실패한 상태를 흉내


def data_rows(path, sheet):
    return [r for r in openpyxl.load_workbook(path)[sheet].iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]


@office
@pytest.mark.office
def test_replace_failure_keeps_ledger(tmp_path, clean_temp, run_engine, no_leftover, monkeypatch):
    """장부 바꾸기가 중간에 실패해도(교체될 파일이 제 이름에서 옮겨진 상태) 장부는 원래 이름으로 넣기 전 판 그대로 있고
    이력이 사라지지 않는다. 다시 실행하면 정상으로 넣는다."""
    rules = example_rules()
    docs = documents(rules)
    book = tmp_path / "누적장부.xlsx"
    ledger_add(run_engine, book, rules, *docs[0][:2], WHEN[0], docs[0][2])
    data = book.read_bytes()

    monkeypatch.setenv(REPLACE_FAIL, "1")
    with pytest.raises(EngineError) as error:
        ledger_add(run_engine, book, rules, *docs[1][:2], WHEN[1], docs[1][2])
    assert "누적장부를 저장하지 못함" in str(error.value)
    assert book.exists() and book.read_bytes() == data
    assert sorted(p.name for p in tmp_path.iterdir() if p.name != "temp" and not p.name.endswith(".json")) == [
        "누적장부.bak.xlsx", "누적장부.xlsx"]

    monkeypatch.delenv(REPLACE_FAIL)
    out = ledger_add(run_engine, book, rules, *docs[1][:2], WHEN[1], docs[1][2])
    assert out["saved"] is True
    assert len(data_rows(book, "Invoice")) == 1 and len(data_rows(book, "DM")) == 1


@office
@pytest.mark.office
def test_missing_ledger_with_backup_stops(tmp_path, clean_temp, run_engine, no_leftover):
    """장부가 없는데 .bak.xlsx 가 있으면 새 장부를 만들지 않고 멈춘다(.bak 의 이력을 덮지 않게)."""
    rules = example_rules()
    docs = documents(rules)
    book = tmp_path / "누적장부.xlsx"
    backup = book.with_name("누적장부.bak.xlsx")
    ledger_add(run_engine, book, rules, *docs[0][:2], WHEN[0], docs[0][2])
    ledger_add(run_engine, book, rules, *docs[1][:2], WHEN[1], docs[1][2])
    book.unlink()
    data = backup.read_bytes()

    with pytest.raises(EngineError) as error:
        ledger_add(run_engine, book, rules, *docs[2][:2], WHEN[2], docs[2][2])
    assert "누적장부.bak.xlsx" in str(error.value) and "없" in str(error.value)
    with pytest.raises(EngineError):
        write(run_engine, "ledger_log", ledger=str(book), rules=rules_to_portable(rules),
              logs=[{"when": _iso(WHEN[2]), "file": "a.pdf", "outcome": "성공", "reason": ""}])
    assert not book.exists() and backup.read_bytes() == data


@office
@pytest.mark.office
def test_midnight_processed_at_is_datetime(tmp_path, clean_temp, run_engine, no_leftover):
    """처리일시 · 처리로그 일시가 꼭 0시여도 일시 형식(지금 판과 같음). 날짜 필드(발행일)는 날짜 형식 그대로."""
    rules = example_rules()
    result, name, organized = documents(rules)[0]
    midnight = datetime(2026, 10, 28, 0, 0, 0)
    book = tmp_path / "누적장부.xlsx"
    ledger_add(run_engine, book, rules, result, name, midnight, organized)
    write(run_engine, "ledger_log", ledger=str(book), rules=rules_to_portable(rules),
          logs=[{"when": _iso(midnight), "file": "a.pdf", "outcome": "성공", "reason": ""}])
    py = tmp_path / "py.xlsx"
    ledger = Ledger.open(py, rules)
    ledger.add(result, name, organized, processed_at=midnight)
    ledger.add_log(midnight, "a.pdf", "성공", "")
    ledger.save()

    got, want = book_view(book)["views"], book_view(py)["views"]
    assert got["Invoice"]["cells"]["A2"] == want["Invoice"]["cells"]["A2"] == (midnight, "yyyy-mm-dd hh:mm:ss")
    assert got["처리로그"]["cells"]["A2"] == want["처리로그"]["cells"]["A2"] == (midnight, "yyyy-mm-dd hh:mm:ss")
    date_col = [k for k, v in got["Invoice"]["cells"].items() if k.endswith("2") and v[1] == "yyyy-mm-dd"]
    assert date_col, "발행일 같은 날짜 필드는 날짜 형식"


@office
@pytest.mark.office
def test_company_form_exact_name(tmp_path, clean_temp, run_engine, no_leftover):
    """exact: 같은 이름의 회사 양식이 이미 있으면 _2 로 쓰지 않고 멈춘다(지금 판 job_run 은 확인필요). 있던 파일은 그대로."""
    rules = example_rules()
    portable = rules_to_portable(rules)
    order = result_json(read_excel(ORDERS[0], rules))
    first = Path(write(run_engine, "company_form", folder=str(tmp_path / "out"), rules=portable, result=order, exact=True)["path"])
    data = first.read_bytes()

    with pytest.raises(EngineError) as error:
        write(run_engine, "company_form", folder=str(tmp_path / "out"), rules=portable, result=order, exact=True)
    assert str(error.value) == f"출력 폴더에 {first.name} 이 이미 있습니다"
    assert [p.name for p in (tmp_path / "out").iterdir()] == [first.name] and first.read_bytes() == data


@office
@pytest.mark.office
def test_ledger_opened_from_copy(tmp_path, clean_temp, run_engine, no_leftover):
    """있던 장부는 엑셀이 원본이 아니라 임시 복사본(%TEMP%\무설치판_*\같은 이름)을 연다. 복사본 폴더는 끝나면 지운다."""
    rules = example_rules()
    docs = documents(rules)
    book = tmp_path / "누적장부.xlsx"
    first = write(run_engine, "ledger_add", ledger=str(book), rules=rules_to_portable(rules), result=result_json(docs[0][0]),
                  source=docs[0][1], when=_iso(WHEN[0]), organized=docs[0][2])
    assert first["opened"] is None   # 새 장부는 열 파일이 없음
    second = write(run_engine, "ledger_add", ledger=str(book), rules=rules_to_portable(rules), result=result_json(docs[1][0]),
                   source=docs[1][1], when=_iso(WHEN[1]), organized=docs[1][2])
    opened = Path(second["opened"])
    assert second["saved"] is True
    assert opened.name == book.name and opened.parent.parent == clean_temp and opened.parent.name.startswith("무설치판_")


HOLDER = r"""
param([string]$Path, [string]$Ready, [string]$Stop)
$excel = New-Object -ComObject Excel.Application
try {
    $excel.Visible = $false
    $excel.DisplayAlerts = $false
    $book = $excel.Workbooks.Open($Path)
    [IO.File]::WriteAllText($Ready, 'ok')
    $deadline = [DateTime]::UtcNow.AddSeconds(100)
    while (-not (Test-Path -LiteralPath $Stop) -and [DateTime]::UtcNow -lt $deadline) { Start-Sleep -Milliseconds 200 }
    $book.Close($false)
    [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($book)
} finally {
    $excel.Quit()
    [void][Runtime.InteropServices.Marshal]::FinalReleaseComObject($excel)
}
"""


@office
@pytest.mark.office
def test_locked_ledger(tmp_path, clean_temp, run_engine, no_leftover):
    """엑셀 COM 으로 장부를 연 채: Test-FileLocked → 참, 장부에 넣기는 locked(장부 · .bak 무변화). 닫으면 다시 거짓."""
    rules = example_rules()
    docs = documents(rules)
    book = tmp_path / "누적장부.xlsx"
    # 이 시험은 잠김을 본다. Quit 뒤 스스로 끝났는지(forced == 0)는 전체 시험 부하에서 흔들리므로 보지 않고,
    # 띄운 엑셀이 남지 않았는지만 본다(남은 것은 엔진이 이미 끝냈어야 한다)
    start = automation_pids()
    first = run_engine("office_write", {"action": "ledger_add", "ledger": str(book), "rules": rules_to_portable(rules),
                                        "result": result_json(docs[0][0]), "source": docs[0][1], "when": _iso(WHEN[0]),
                                        "organized": str(docs[0][2])})
    assert first["saved"] is True and first["locked"] is False
    assert automation_pids() - start == set(), "장부를 쓴 엑셀이 남음"
    data, stamp = book.read_bytes(), book.stat().st_mtime_ns

    script = tmp_path / "holder.ps1"
    script.write_text(HOLDER, encoding="utf-8-sig")
    ready, stop = tmp_path / "ready.txt", tmp_path / "stop.txt"
    before = automation_pids()
    holder = subprocess.Popen(["powershell.exe", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass",
                               "-File", str(script), str(book), str(ready), str(stop)])
    try:
        deadline = time.monotonic() + 60
        while not ready.exists():
            assert holder.poll() is None, "장부를 여는 엑셀이 먼저 끝남"
            assert time.monotonic() < deadline, "엑셀이 장부를 열지 못함"
            time.sleep(0.2)

        assert run_engine("office_write", {"action": "locked", "path": str(book)}) == {"locked": True}
        out = ledger_add(run_engine, book, rules, *docs[1][:2], WHEN[1], docs[1][2])
        assert out["locked"] is True and out["saved"] is False
        out = write(run_engine, "ledger_log", ledger=str(book), rules=rules_to_portable(rules),
                    logs=[{"when": _iso(WHEN[1]), "file": "a.pdf", "outcome": "성공", "reason": ""}])
        assert out["locked"] is True and out["saved"] is False
    finally:
        stop.write_text("stop", encoding="utf-8")
        try:
            holder.wait(timeout=60)
        except subprocess.TimeoutExpired:
            holder.kill()
            for name, pid in automation_pids() - before:
                subprocess.run(["taskkill", "/F", "/PID", str(pid)], capture_output=True)
    assert book.read_bytes() == data and book.stat().st_mtime_ns == stamp
    assert not book.with_name("누적장부.bak.xlsx").exists()
    assert sorted(p.name for p in tmp_path.iterdir() if p.suffix == ".xlsx") == ["누적장부.xlsx"]
    assert run_engine("office_write", {"action": "locked", "path": str(book)}) == {"locked": False}
