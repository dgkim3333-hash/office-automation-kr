"""무설치판 엔진 verify: 같은 결과 · 같은 원본을 넣으면 지금 판 verify() 와 같은 검산 목록(항목 이름 · 값 글자 · 순서 · ok)을 낸다.

결과 객체는 JSON 을 한 번 거친 모양(금액은 글자, 날짜는 ISO)으로 엔진에 넣는다.
PDF 는 지금 판이 pdfplumber 로 읽은 글자를 엔진에도 준다(시험에서는 read_text 를 가짜로 바꿔 같은 글자를 쓴다).
엑셀은 openpyxl 로 격자 JSON(tests\\portable\\grid.py)으로 옮겨 넣는다.
"""
import json
import random
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
import pytest
import yaml

from app import extract
from app import verify as verify_module
from app.excel_extract import read_excel
from app.extract import ExtractResult, extract_from_text
from app.portable import rules_to_portable
from app.rules import rules_from_text
from app.verify import DATE, NUMBER, verify

from .conftest import PORTABLE, REPO, decode_output, example_rules, run_powershell
from .grid import grid_from_xlsx

PDFS = sorted((REPO / "samples" / "03_Invoice_DM_PDF").glob("*.pdf"))
ORDERS = sorted((REPO / "samples" / "01_고객오더_엑셀").glob("*.xlsx"))
ACME = next(p for p in ORDERS if "ACME" in p.name)
INVOICE = next(p for p in PDFS if p.name == "INV-2026-10231.pdf")
DM = next(p for p in PDFS if p.name == "DM-2026-0092.pdf")

AMOUNT = "\\$?([\\d,]+\\.\\d+)"
TEST_RULES = {
    "문서종류": {
        # 표 · 검산 없는 짧은 문서
        "Memo": {
            "판별키워드": ["Memo No"],
            "번호필드": "메모번호",
            "필드": {
                "메모번호": {"패턴": "Memo No[.:]?\\s*(\\S+)"},
                "발행일": {"패턴": "Date[.:]?\\s*(\\d{1,2}/\\d{1,2}/\\d{4})", "형식": "날짜"},
                "합계금액": {"패턴": "Total[.:]?\\s*" + AMOUNT, "형식": "금액"},
            },
            "파일명": "Memo_{메모번호}.pdf",
        },
        # 품목 표 + 소계 · 운임 · 합계
        "Bill": {
            "판별키워드": ["Bill No"],
            "번호필드": "청구번호",
            "필드": {
                "청구번호": {"패턴": "Bill No[.:]?\\s*([A-Z0-9-]+)"},
                "소계": {"패턴": "Subtotal[.:]?\\s*" + AMOUNT, "형식": "금액"},
                "운임": {"패턴": "Freight[.:]?\\s*" + AMOUNT, "형식": "금액"},
                "합계금액": {"패턴": "\\nTotal[.:]?\\s*" + AMOUNT, "형식": "금액"},
            },
            "표": {
                "시작키워드": "Line Item",
                "끝키워드": "Subtotal",
                "행패턴": "^(\\d+)\\s+(\\S+)\\s+(\\d+)\\s+([\\d,]+\\.\\d+)\\s+([\\d,]+\\.\\d+)$",
                "열": ["줄", "품번", "수량", "단가", "금액"],
                "금액열": ["단가", "금액"],
                "수량열": ["수량"],
            },
            "검산": {"소계": "소계", "부대비용": ["운임"], "합계": "합계금액"},
            "파일명": "Bill_{청구번호}.pdf",
        },
        # 표는 있지만 검산 항목이 없는 문서(기본 열 이름 수량 · 단가 · 금액으로 줄 검산만 한다)
        "Bare": {
            "판별키워드": ["Bare No"],
            "번호필드": "번호",
            "필드": {"번호": {"패턴": "Bare No[.:]?\\s*([A-Z0-9-]+)"}},
            "표": {
                "시작키워드": "Line Item",
                "끝키워드": "End",
                "행패턴": "^(\\d+)\\s+([\\d,]+\\.\\d+)\\s+([\\d,]+\\.\\d+)$",
                "열": ["수량", "단가", "금액"],
                "금액열": ["단가", "금액"],
                "수량열": ["수량"],
            },
            "파일명": "Bare_{번호}.pdf",
        },
        # 합계만 검산하는 문서(소계 · 부대비용 없음) + 열 이름을 따로 지정
        "Plain": {
            "판별키워드": ["Plain No"],
            "번호필드": "번호",
            "필드": {
                "번호": {"패턴": "Plain No[.:]?\\s*([A-Z0-9-]+)"},
                "합계금액": {"패턴": "\\nTotal[.:]?\\s*" + AMOUNT, "형식": "금액"},
            },
            "표": {
                "시작키워드": "Line Item",
                "끝키워드": "Total",
                "행패턴": "^(\\S+)\\s+(\\d+)\\s+([\\d,]+\\.\\d+)\\s+([\\d,]+\\.\\d+)$",
                "열": ["품번", "Qty", "Price", "Amount"],
                "금액열": ["Price", "Amount"],
                "수량열": ["Qty"],
            },
            "검산": {"합계": "합계금액", "수량열": "Qty", "단가열": "Price", "금액열": "Amount"},
            "파일명": "Plain_{번호}.pdf",
        },
    }
}


@pytest.fixture(scope="module")
def rules():
    return rules_from_text(yaml.safe_dump(TEST_RULES, allow_unicode=True))


def memo_text(total="$1,234.50", number="M-2026-01", extra=()):
    return "\n".join(["MEMO", f"Memo No: {number}", "Date: 10/15/2026", f"Total: {total}", *extra])


def bill_text(items=(("2", "50.00", "100.00"),), subtotal="100.00", freight="15.00", total="115.00"):
    lines = ["BILL", "Bill No: B-2026-07", "Line Item Part Qty Price Amount"]
    lines += [f"{n} P-{n} {q} {p} {a}" for n, (q, p, a) in enumerate(items, 1)]
    lines += [f"Subtotal: ${subtotal}", f"Freight: ${freight}", f"Total: ${total}"]
    return "\n".join(lines)


def plain_text(items=(("P-1", "2", "50.00", "100.00"),), total="100.00"):
    lines = ["PLAIN", "Plain No: PL-1", "Line Item Part Qty Price Amount"]
    lines += [" ".join(item) for item in items]
    lines += [f"Total: ${total}"]
    return "\n".join(lines)


# --- 지금 판과 엔진을 같은 입력으로 돌려 목록을 비교하는 도우미 ---

def _plain(value):
    """파이썬 값을 엔진 JSON 표기로: 금액은 자리 그대로의 글자, 날짜 · 일시는 ISO, 나머지는 그대로."""
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def result_json(result) -> dict:
    """결과 객체를 엔진 결과 JSON 모양으로(작업 2 · 3 의 출력과 같은 모양)."""
    return {
        "doc_type": result.doc_type.name if result.doc_type else None,
        "form": result.source_form,
        "is_excel": result.is_excel,
        "fields": {k: _plain(v) for k, v in result.fields.items()},
        "table": None if result.table is None else [{k: _plain(v) for k, v in row.items()} for row in result.table],
        "warnings": list(result.warnings),
        "review_reason": result.review_reason,
    }


def view(checks) -> list[dict]:
    return [{"item": c.item, "source_value": c.source, "result_value": c.output, "ok": c.ok} for c in checks]


def line_items(checks) -> list[str]:
    """줄 검산 항목(「1줄 금액」 …)만. 「품목 1줄 금액」은 원문 대조 항목이다."""
    return [c["item"] for c in checks if re.fullmatch(r"\d+줄 금액", c["item"])]


def failed_items(checks) -> list[str]:
    return [c.item for c in checks if c.ok is False]


def compare_text(run_engine, monkeypatch, rules, result, text, culture=None):
    """(지금 판 목록, 엔진 목록). 지금 판은 read_text 를 가짜로 바꿔 같은 글자를 읽게 한다."""
    monkeypatch.setattr(verify_module, "read_text", lambda path: text)
    expected = view(verify(result, Path("원본.pdf")))
    got = run_engine("verify", {"rules": rules_to_portable(rules), "result": result_json(result),
                                "source": {"kind": "text", "text": text}}, culture)
    return expected, got


def compare_grid(run_engine, rules, result, path, culture=None):
    expected = view(verify(result, path))
    got = run_engine("verify", {"rules": rules_to_portable(rules), "result": result_json(result),
                                "source": {"kind": "grid", "grid": grid_from_xlsx(path)}}, culture)
    return expected, got


def text_case(run_engine, monkeypatch, rules, text, mutate=None, culture=None):
    """text 를 지금 판으로 추출 → (선택) 값을 바꿈 → 지금 판 verify 와 엔진 verify 를 비교용으로 돌려준다."""
    result = extract_from_text(text, rules)
    if mutate is not None:
        mutate(result)
    return compare_text(run_engine, monkeypatch, rules, result, text, culture)


# --- 샘플 ---

def test_samples_present():
    assert len(PDFS) == 6 and len(ORDERS) == 2


@pytest.mark.parametrize("pdf", PDFS, ids=[p.name for p in PDFS])
def test_same_as_python_on_pdf_samples(pdf, run_engine):
    """샘플 PDF 6개: 엔진 extract_text 가 낸 결과를 그대로 넣어도 지금 판 verify(result, path) 와 같은 목록(순서 · ok 포함)."""
    rules = example_rules()
    portable = rules_to_portable(rules)
    result = extract.extract(pdf, rules)
    text = extract.read_text(pdf)
    expected = view(verify(result, pdf))

    engine_result = run_engine("extract_text", {"rules": portable, "text": text})
    got = run_engine("verify", {"rules": portable, "result": engine_result, "source": {"kind": "text", "text": text}})

    assert got == expected
    if pdf.name.startswith(("INV-", "DM-")):
        assert expected and all(c["ok"] for c in expected), "샘플은 모든 검산이 통과해야 비교가 의미 있다"
        assert any(c["item"] == "합계금액" for c in expected)
    else:
        assert got == []   # 판별 실패 · 글자 없음 → 검산할 것이 없다


@pytest.mark.parametrize("xlsx", ORDERS, ids=["ACME" if "ACME" in p.name else "HANIL" for p in ORDERS])
def test_same_as_python_on_excel_samples(xlsx, run_engine):
    """샘플 엑셀 2개: 엔진 extract_grid 결과 + 격자 원본으로 검산해도 지금 판과 같은 목록."""
    rules = example_rules()
    portable = rules_to_portable(rules)
    grid = grid_from_xlsx(xlsx)
    expected = view(verify(read_excel(xlsx, rules), xlsx))

    engine_result = run_engine("extract_grid", {"rules": portable, "grid": grid})
    got = run_engine("verify", {"rules": portable, "result": engine_result, "source": {"kind": "grid", "grid": grid}})

    assert got == expected
    assert expected and all(c["ok"] for c in expected)
    assert any(c["item"] == "1줄 금액" for c in expected)


def test_check_list_shape(run_engine):
    """검산 항목은 item · source_value · result_value · ok 네 키 뿐이고 ok 는 참 · 거짓 · null."""
    rules = example_rules()
    text = extract.read_text(INVOICE)
    out = run_engine("verify", {"rules": rules_to_portable(rules), "result": result_json(extract.extract(INVOICE, rules)),
                                "source": {"kind": "text", "text": text}})
    assert isinstance(out, list) and out
    assert all(list(c) == ["item", "source_value", "result_value", "ok"] for c in out)
    by_item = {c["item"]: c for c in out}
    assert by_item["합계금액"] == {"item": "합계금액", "source_value": "115,440.00", "result_value": "115440.00", "ok": True}
    assert by_item["품목 합 = 소계"]["ok"] is True


def test_culture_de_de_same_result(run_engine, monkeypatch):
    """독일어 문화권(소수점이 쉼표)으로 돌려도 PDF · 엑셀 검산 목록이 기본 문화권 · 지금 판과 같다."""
    rules = example_rules()
    portable = rules_to_portable(rules)
    text = extract.read_text(INVOICE)
    pdf_payload = {"rules": portable, "result": result_json(extract.extract(INVOICE, rules)),
                   "source": {"kind": "text", "text": text}}
    xlsx_payload = {"rules": portable, "result": result_json(read_excel(ACME, rules)),
                    "source": {"kind": "grid", "grid": grid_from_xlsx(ACME)}}

    for payload, path in ((pdf_payload, INVOICE), (xlsx_payload, ACME)):
        expected = view(verify(extract.extract(path, rules) if path.suffix == ".pdf" else read_excel(path, rules), path))
        default = run_engine("verify", payload)
        german = run_engine("verify", payload, culture="de-DE")
        assert german == default == expected
        if path == INVOICE:
            assert any("." in c["result_value"] for c in default)   # 소수점은 마침표 그대로
            assert any("." in c["source_value"] for c in default)


# --- 원문 대조 · 소수 자리 (PDF 글자) ---

# (이름, 글자, 값 바꾸기, 지금 판이 거짓으로 보는 항목)
def _set(**fields):
    def mutate(result):
        result.fields.update(fields)
    return mutate


MEMO_CASES = [
    ("pass", memo_text(), None, []),
    ("digits_differ", memo_text(), _set(합계금액=Decimal("12345.00")), ["합계금액"]),
    ("digits_and_places_differ", memo_text(), _set(합계금액=Decimal("12345.0")), ["합계금액"]),
    ("decimal_point_shifted", memo_text(), _set(합계금액=Decimal("123.450")), ["합계금액"]),
    ("notation_only", memo_text("$1,234.50"), None, []),
    ("notation_without_comma_value", memo_text("$1234.50"), _set(합계금액=Decimal("1234.50")), []),
    ("places_differ", memo_text("$12.50"), _set(합계금액=Decimal("12.5")), ["합계금액 소수 자리"]),
    ("places_differ_with_comma", memo_text(), _set(합계금액=Decimal("1234.5")), ["합계금액 소수 자리"]),
    ("places_any_token_same", memo_text("$12.50", extra=["Note 12.5 and 12.50"]), _set(합계금액=Decimal("12.5")), []),
    ("wrong_date", memo_text(), _set(발행일=date(2026, 10, 16)), ["발행일"]),
    ("date_other_format_in_text", memo_text(extra=["Paid 2026-10-16"]), _set(발행일=date(2026, 10, 16)), []),
    ("number_missing", memo_text(), _set(메모번호="M-2026-07"), ["메모번호"]),
    ("number_part_of_longer", memo_text(), _set(메모번호="M-2026"), ["메모번호"]),
    ("number_part_digits_only", memo_text(), _set(메모번호="2026"), ["메모번호"]),
    ("number_part_of_longer_after", memo_text(number="M-2026-017"), _set(메모번호="M-2026-01"), ["메모번호"]),
    ("number_before_sentence_period", memo_text(number="M-2026-01."), None, []),
    ("number_after_label_period", "MEMO\nMemo No.M-2026-01\nDate: 10/15/2026\nTotal: $1,234.50", None, []),
    ("number_after_digit_period", memo_text(number="M-1", extra=["Ref 1.2026"]), _set(메모번호="2026"), ["메모번호"]),
    ("number_followed_by_dot_letter", memo_text(number="M-2026-01.x"), _set(메모번호="M-2026-01"), ["메모번호"]),
    ("number_blank", memo_text(), _set(메모번호="  "), ["메모번호"]),
    ("number_none", memo_text(), _set(메모번호=None), []),
    ("number_with_regex_specials", memo_text(number="A.B+C(1)[x]{2}|^$"), _set(메모번호="A.B+C(1)[x]{2}|^$"), []),
    ("number_regex_specials_not_pattern", memo_text(number="AxB+C(1)[x]{2}|^$"), _set(메모번호="A.B+C(1)[x]{2}|^$"), ["메모번호"]),
    ("number_with_space", memo_text(extra=["Ref M 2026 01"]), _set(메모번호="M 2026 01"), []),
    ("negative_signed", memo_text("-12.50"), _set(합계금액=Decimal("-12.50")), []),
    ("negative_not_in_text", memo_text("12.50"), _set(합계금액=Decimal("-12.50")), ["합계금액"]),
    ("positive_from_negative_token", memo_text("-12.50"), _set(합계금액=Decimal("12.50")), []),
    ("int_and_float_fields", memo_text(extra=["Qty 7 Ratio 0.5"]), _set(수량합=7, 비율=0.5), []),
    ("int_and_float_not_found", memo_text(), _set(수량합=7, 비율=0.5), ["수량합", "비율"]),
    ("float_rounds_to_15_digits", memo_text(extra=["Ratio 0.3"]), _set(비율=0.1 + 0.2), []),
    ("bool_and_plain_text_skipped", memo_text(), _set(체크=True, 비고="Note"), []),
    ("missing_total", memo_text(), _set(합계금액=None), []),
    ("zero_amount", memo_text("0.00"), _set(합계금액=Decimal("0.00")), []),
    ("many_tokens_signed", memo_text("$1,234.50", extra=["A -5,000.25 B 7,0 C 3."]), _set(합계금액=Decimal("5000.25")), []),
]


@pytest.mark.parametrize("text, mutate, bad", [c[1:] for c in MEMO_CASES], ids=[c[0] for c in MEMO_CASES])
def test_memo_cases_same_as_python(text, mutate, bad, run_engine, monkeypatch, rules):
    """원문 대조 · 소수 자리 · 번호필드: 값을 하나씩 어긋나게 해도 항목 · 값 글자 · ok 가 지금 판과 같다."""
    result = extract_from_text(text, rules)
    if mutate is not None:
        mutate(result)
    expected, got = compare_text(run_engine, monkeypatch, rules, result, text)

    assert failed_items(verify_module.verify(result, Path("원본.pdf"))) == bad, "지금 판 결과가 의도한 상황이어야 시험이 의미 있다"
    assert got == expected


# --- 줄 검산 · 합계 검산 (PDF 글자) ---

def _table(index, **values):
    def mutate(result):
        result.table[index].update(values)
    return mutate


def _drop_column(column):
    def mutate(result):
        for row in result.table:
            row.pop(column, None)
    return mutate


def _table_is(table):
    def mutate(result):
        result.table = table
    return mutate


TWO = (("2", "50.00", "100.00"), ("1", "5.00", "5.00"))

BILL_CASES = [
    ("pass", bill_text(), None, []),
    ("two_lines_pass", bill_text(TWO, "105.00", "15.00", "120.00"), None, []),
    ("line_amount_wrong", bill_text((("2", "50.00", "100.01"),), "100.01", "15.00", "115.01"), None, ["1줄 금액"]),
    ("line_wrong_in_second", bill_text(TWO, "105.00", "15.00", "120.00"), _table(1, 금액=Decimal("6.00")),
     ["품목 2줄 금액", "2줄 금액", "품목 합 = 소계"]),
    ("rounding_allowed", bill_text((("3", "0.333", "1.00"),), "1.00", "0.00", "1.00"), None, []),
    ("rounding_too_far", bill_text((("3", "0.333", "1.10"),), "1.10", "0.00", "1.10"), None, ["1줄 금액"]),
    ("half_even_also_passes", bill_text((("1", "0.125", "0.12"),), "0.12", "0.00", "0.12"), None, []),
    ("half_up_passes", bill_text((("1", "0.125", "0.13"),), "0.13", "0.00", "0.13"), None, []),
    ("more_places_than_product", bill_text((("2", "0.50", "1.000"),), "1.000", "0.000", "1.000"), None, []),
    ("subtotal_wrong", bill_text(TWO, "100.00", "15.00", "115.00"), None, ["품목 합 = 소계"]),
    ("total_wrong", bill_text(total="116.00"), None, ["소계 + 부대비용 = 합계"]),
    ("total_without_extras_text", bill_text(freight="0.00", total="100.00"), None, []),
    ("freight_missing_counts_zero", bill_text(), _set(운임=None), ["소계 + 부대비용 = 합계"]),
    ("total_missing", bill_text(), _set(합계금액=None), []),
    ("subtotal_missing_falls_back", bill_text(), _set(소계=None), []),
    ("subtotal_missing_and_item_wrong", bill_text(), lambda r: (r.fields.update(소계=None), r.table[0].update(금액=Decimal("99.00"))),
     ["품목 1줄 금액", "1줄 금액", "소계 + 부대비용 = 합계"]),
    ("unit_price_missing_not_applicable", bill_text(), _drop_column("단가"), []),
    ("quantity_missing_not_applicable", bill_text(), _drop_column("수량"), []),
    ("amount_missing_not_applicable", bill_text(), _drop_column("금액"), []),
    ("empty_table", bill_text(), _table_is([]), []),
    ("table_none", bill_text(), _table_is(None), []),
    ("single_row_table", bill_text(), None, []),
    ("wrong_table_value", bill_text(), _table(0, 단가=Decimal("5.00")), ["품목 1줄 단가", "1줄 금액"]),
    ("table_value_notation", bill_text((("2", "1,000.00", "2,000.00"),), "2,000.00", "0.00", "2,000.00"), None, []),
    ("big_numbers", bill_text((("3", "1234567890123.45", "3703703670370.35"),), "3703703670370.35", "0.00", "3703703670370.35"), None, []),
]


@pytest.mark.parametrize("text, mutate, bad", [c[1:] for c in BILL_CASES], ids=[c[0] for c in BILL_CASES])
def test_bill_cases_same_as_python(text, mutate, bad, run_engine, monkeypatch, rules):
    """줄 검산(수량 x 단가 = 금액) · 합계 검산(품목 합 = 소계, 소계 + 부대비용 = 합계)을 어긋나게 해도 지금 판과 같다."""
    result = extract_from_text(text, rules)
    if mutate is not None:
        mutate(result)
    expected, got = compare_text(run_engine, monkeypatch, rules, result, text)

    assert failed_items(verify_module.verify(result, Path("원본.pdf"))) == bad, "지금 판 결과가 의도한 상황이어야 시험이 의미 있다"
    assert got == expected


def test_empty_and_single_row_tables(run_engine, monkeypatch, rules):
    """빈 표 [] 는 품목 합이 없어 「품목 합 = 소계」가 해당 없음(null), 한 줄짜리 표는 한 줄 검산."""
    text = bill_text()
    empty, got_empty = compare_text(run_engine, monkeypatch, rules, _with(extract_from_text(text, rules), table=[]), text)
    assert got_empty == empty
    by_item = {c["item"]: c for c in got_empty}
    assert by_item["품목 합 = 소계"] == {"item": "품목 합 = 소계", "source_value": "", "result_value": "100.00", "ok": None}
    assert line_items(got_empty) == []

    one, got_one = compare_text(run_engine, monkeypatch, rules, extract_from_text(text, rules), text)
    assert got_one == one
    assert line_items(got_one) == ["1줄 금액"]
    assert next(c for c in got_one if c["item"] == "1줄 금액") == {
        "item": "1줄 금액", "source_value": "2 x 50.00 = 100.00", "result_value": "100.00", "ok": True}


def _with(result, **changes):
    for key, value in changes.items():
        setattr(result, key, value)
    return result


def test_named_columns_and_total_only(run_engine, monkeypatch, rules):
    """검산에 열 이름을 따로 적은 문서(Qty · Price · Amount), 합계만 있는 검산."""
    text = plain_text()
    expected, got = text_case(run_engine, monkeypatch, rules, text)
    assert got == expected and all(c["ok"] for c in got)
    assert [c["item"] for c in got if "=" in c["item"]] == ["품목 합 = 합계", ]

    wrong = plain_text((("P-1", "2", "50.00", "100.01"),), "100.01")
    expected, got = text_case(run_engine, monkeypatch, rules, wrong)
    assert got == expected
    assert [c["item"] for c in got if c["ok"] is False] == ["1줄 금액"]


def test_no_check_spec_uses_default_columns(run_engine, monkeypatch, rules):
    """검산 항목이 없는 문서종류(표만 있음)는 기본 열 이름(수량 · 단가 · 금액)으로 줄 검산만 하고 합계 검산은 없다."""
    text = "BARE\nBare No: B-1\nLine Item\n2 50.00 100.00\n3 1.10 3.00\nEnd"
    expected, got = text_case(run_engine, monkeypatch, rules, text)
    assert got == expected
    assert line_items(got) == ["1줄 금액", "2줄 금액"]
    assert [c["ok"] for c in got if c["item"] in line_items(got)] == [True, False]
    assert not any("합" in c["item"] for c in got)   # 합계 검산은 없다


def test_doc_type_none_gives_empty_list(run_engine, monkeypatch, rules):
    """판별 실패(doc_type null)는 빈 목록 [] (null 이 아니다)."""
    result = ExtractResult(doc_type=None)
    expected, got = compare_text(run_engine, monkeypatch, rules, result, "아무 글자")
    assert expected == [] and got == []


def test_single_check_is_array(run_engine, monkeypatch, rules):
    """항목이 하나뿐이어도 JSON 배열로 나온다."""
    text = "MEMO\nMemo No: M-1"
    result = ExtractResult(doc_type=rules.get("Memo"), fields={"메모번호": "M-1", "합계금액": None, "발행일": None})
    expected, got = compare_text(run_engine, monkeypatch, rules, result, text)
    assert len(expected) == 1
    assert got == expected


# --- 숫자 글자 · 날짜 글자 찾기 (.NET 정규식이 파이썬과 같은 뜻인지) ---

REGEX_SAMPLES = [
    "Total: $1,234.50 and -5,000.25, 7,0 3. 1.2026 .5 5. 1,,2 -0 -",
    "Date 12/31/2026 and 2026-01-05, 1/2/3/4 5-6-7-8 12/31/2026-12/31/2027",
    "A1/2/2026B x2026-01-05y 20260105 1234/1/1 12345/1/1 1/2/12345 /1/2/2026",
    "-12.50 +12.50 12.5.6 1-2-3 9-9-99999",
    "٣١٢ １２３ 12٣ 2026-１-05",
    "",
    "no numbers here",
    "1,234,567,890.123456789 0.0000000001 00012 -00.10",
    "5/6/2026/7/8/2027",
]

REGEX_SCRIPT = """\
param([string]$Lib, [string]$In, [string]$Out)
. (Join-Path $Lib 'Values.ps1')
. (Join-Path $Lib 'Rules.ps1')
. (Join-Path $Lib 'ExtractText.ps1')
. (Join-Path $Lib 'ExtractGrid.ps1')
. (Join-Path $Lib 'Verify.ps1')
$texts = [IO.File]::ReadAllText($In, [Text.Encoding]::UTF8) | ConvertFrom-Json
$found = New-Object 'System.Collections.Generic.List[object]'
foreach ($text in @($texts)) {
    $numbers = @($script:NumberToken.Matches([string]$text) | ForEach-Object { $_.Value })
    $dates = @($script:DateToken.Matches([string]$text) | ForEach-Object { $_.Groups[1].Value })
    $found.Add([ordered]@{ numbers = [object[]]$numbers; dates = [object[]]$dates })
}
[IO.File]::WriteAllText($Out, (ConvertTo-Json -InputObject ([object[]]$found.ToArray()) -Depth 5), (New-Object Text.UTF8Encoding $false))
"""


def test_number_and_date_regex_same_as_python(powershell_exe, tmp_path):
    """NUMBER · DATE 정규식(겹쳐 찾는 DATE 포함)이 .NET 에서도 파이썬과 같은 조각을 같은 순서로 찾는다."""
    script = tmp_path / "regex.ps1"
    script.write_bytes(b"\xef\xbb\xbf" + REGEX_SCRIPT.encode("utf-8"))
    source = tmp_path / "texts.json"
    source.write_text(json.dumps(REGEX_SAMPLES, ensure_ascii=False), encoding="utf-8")
    out = tmp_path / "found.json"

    done = run_powershell(powershell_exe, script, "-Lib", str(PORTABLE / "lib"), "-In", str(source), "-Out", str(out))

    assert done.returncode == 0, decode_output(done.stderr)
    found = json.loads(out.read_text(encoding="utf-8-sig"))
    expected = [{"numbers": [m.group() for m in NUMBER.finditer(t)], "dates": [m.group(1) for m in DATE.finditer(t)]}
                for t in REGEX_SAMPLES]
    assert found == expected
    assert any(len(e["dates"]) >= 4 for e in expected), "겹치는 날짜 시료가 있어야 시험이 의미 있다"


# --- 날짜 겹쳐 찾기가 검산 결과에 드러나는 경우 ---

def test_overlapping_dates_found_in_text(run_engine, monkeypatch, rules):
    """12/31/2026-12/31/2027 처럼 붙어 있는 날짜도 각각 찾는다(겹쳐 찾기)."""
    text = memo_text(extra=["Span 12/31/2026-12/31/2027 and 5/6/2026/7/8/2027"])
    for found in (date(2026, 12, 31), date(2027, 12, 31), date(2026, 5, 6), date(2026, 6, 7), date(2027, 8, 7)):
        expected, got = text_case(run_engine, monkeypatch, rules, text, _set(발행일=found))
        assert got == expected


# --- 엑셀 ---

def acme_file(tmp_path, edit=None, name="po.xlsx"):
    wb = openpyxl.load_workbook(ACME)
    if edit is not None:
        edit(wb.active)
    path = tmp_path / name
    wb.save(path)
    return path


def raw_float(path, placeholder, raw):
    """openpyxl 은 float 을 짧게 쓰므로, 엑셀이 저장하는 긴 값(=3*0.1 → 0.30000000000000004)을 XML 에 직접 넣는다."""
    with zipfile.ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
    sheet = "xl/worksheets/sheet1.xml"
    assert files[sheet].count(f"<v>{placeholder}</v>".encode()) == 1
    files[sheet] = files[sheet].replace(f"<v>{placeholder}</v>".encode(), f"<v>{raw}</v>".encode())
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for n, data in files.items():
            z.writestr(n, data)


def excel_case(run_engine, path, mutate=None, culture=None):
    rules = example_rules()
    result = read_excel(path, rules)
    if mutate is not None:
        mutate(result)
    expected, got = compare_grid(run_engine, rules, result, path, culture)
    return expected, got, failed_items(verify(result, path))


EXCEL_EDITS = [
    ("pass", None, []),
    ("int_amount_two_places", lambda ws: (ws.__setitem__("D14", 2), ws.__setitem__("F14", 50.2), ws.__setitem__("G14", 100)),
     ["1줄 금액", "품목 합 = 합계"]),
    ("amount_text_cells", lambda ws: (ws.__setitem__("D14", "2"), ws.__setitem__("F14", "48,500.00"), ws.__setitem__("G14", "$97,000.50")),
     ["1줄 금액", "품목 합 = 합계"]),
    ("float_cell_places", lambda ws: (ws.__setitem__("F14", 212.3375), ws.__setitem__("G14", 0.1 + 0.2)),
     ["1줄 금액", "품목 합 = 합계"]),
    ("date_in_plain_cell", lambda ws: ws.__setitem__("B11", datetime(2026, 9, 18)), []),
    ("datetime_in_plain_cell", lambda ws: ws.__setitem__("B11", datetime(2026, 9, 18, 13, 45, 30)), []),
    ("plain_number_and_bool", lambda ws: (ws.__setitem__("B9", True), ws.__setitem__("B10", 12.5), ws.__setitem__("B11", 7)), []),
    ("whitespace_values", lambda ws: (ws.__setitem__("B10", "  USD  "), ws.__setitem__("B9", "   ")), []),
    ("date_as_text", lambda ws: (ws.__setitem__("B5", "2026-09-18"), ws.__setitem__("B8", " 09/20/2026 ")), []),
    ("total_blank", lambda ws: ws.__setitem__("G18", None), []),
    ("zero_part_continues", lambda ws: ws.__setitem__("B16", 0), []),
]


@pytest.mark.parametrize("edit, bad", [c[1:] for c in EXCEL_EDITS], ids=[c[0] for c in EXCEL_EDITS])
def test_excel_variants_same_as_python(edit, bad, run_engine, tmp_path):
    """샘플 ACME 를 한 군데씩 고쳐도(실수 · 글자 금액 · 날짜 · 참거짓) 검산 목록이 지금 판과 같다."""
    path = acme_file(tmp_path, edit)
    expected, got, failed = excel_case(run_engine, path)

    assert failed == bad, "지금 판 결과가 의도한 상황이어야 시험이 의미 있다"
    assert got == expected
    if bad == [] or bad == ["품목 합 = 합계"]:
        assert expected


def test_excel_formula_float_line_check(run_engine, tmp_path):
    """=3*0.1 의 저장값 0.30000000000000004 는 엑셀처럼 0.3 으로 보고 줄 검산에 통과한다(글자도 지금 판과 같다)."""
    path = acme_file(tmp_path, lambda ws: (ws.__setitem__("D14", 3), ws.__setitem__("F14", 0.1), ws.__setitem__("G14", 0.375)))
    raw_float(path, "0.375", "0.30000000000000004")
    expected, got, failed = excel_case(run_engine, path)

    assert got == expected
    line = next(c for c in got if c["item"] == "1줄 금액")
    assert line["ok"] is True and line["result_value"] == "0.3"
    assert failed == ["품목 합 = 합계"]


def test_excel_wrong_scale_fails(run_engine, tmp_path):
    """엑셀 12.3(0.00 서식)은 Decimal 12.30 과 같은 값이라 통과, 123.0 은 다른 값."""
    def edit(ws):
        ws["G18"] = 12.3
        ws["G18"].number_format = "0.00"
    path = acme_file(tmp_path, edit)

    def same(result):
        result.fields["합계금액"] = Decimal("12.30")

    def wrong(result):
        result.fields["합계금액"] = Decimal("123.0")

    expected, got, failed = excel_case(run_engine, path, same)
    assert got == expected and "합계금액" not in failed
    expected, got, failed = excel_case(run_engine, path, wrong)
    assert got == expected and "합계금액" in failed


def test_excel_values_on_other_sheets_and_string_numbers(run_engine, tmp_path):
    """다른 시트의 값 · 글자로 적힌 숫자("1,500.25") · 날짜 칸도 원본으로 센다."""
    rules = example_rules()
    wb = openpyxl.Workbook()
    wb.active["A1"] = "PO-T-2"
    other = wb.create_sheet("other")
    other["B2"] = "1,500.25"
    other["C3"] = date(2026, 9, 18)
    other["D4"] = datetime(2026, 9, 19, 8, 30, 15)
    path = tmp_path / "po.xlsx"
    wb.save(path)
    result = ExtractResult(doc_type=rules.get("PO"), is_excel=True, source_form="ACME", fields={
        "고객PO번호": "PO-T-2", "합계금액": Decimal("1500.25"), "발주일": date(2026, 9, 18),
        "납기요청일": date(2026, 9, 19), "담당": datetime(2026, 9, 19, 8, 30, 15), "지연": datetime(2026, 9, 19, 8, 30, 16)})

    expected, got = compare_grid(run_engine, rules, result, path)
    assert got == expected
    assert [c["item"] for c in got if c["ok"] is False] == ["납기요청일", "지연"]   # 날짜 칸의 시각이 있으면 그 날 0시와 다르다

    result.fields["발주일"] = date(2026, 9, 19)
    expected, got = compare_grid(run_engine, rules, result, path)
    assert got == expected
    assert "발주일" in [c["item"] for c in got if c["ok"] is False]


def test_excel_number_field_must_equal_a_whole_cell(run_engine, tmp_path):
    """엑셀 번호필드는 칸 하나의 값(앞뒤 공백 뺌)과 같아야 한다. 일부만 같으면 거짓."""
    rules = example_rules()
    wb = openpyxl.Workbook()
    wb.active["A1"], wb.active["B1"] = "PO Number:", "  PO-T-12 "
    path = tmp_path / "po.xlsx"
    wb.save(path)
    result = ExtractResult(doc_type=rules.get("PO"), is_excel=True, source_form="ACME", fields={"고객PO번호": "PO-T-12"})
    expected, got = compare_grid(run_engine, rules, result, path)
    assert got == expected and [c["ok"] for c in got] == [True, None]   # 번호 일치, 품목 합 = 합계는 값이 없어 해당 없음
    result.fields["고객PO번호"] = "PO-T-1"
    expected, got = compare_grid(run_engine, rules, result, path)
    assert got == expected and [c["item"] for c in got if c["ok"] is False] == ["고객PO번호"]


def test_excel_number_field_numeric_cell(run_engine, tmp_path):
    """번호필드가 숫자 칸(글자가 아님)이어도 지금 판과 같다."""
    rules = example_rules()
    wb = openpyxl.Workbook()
    wb.active["A1"], wb.active["B1"] = "PO Number:", 20260417
    path = tmp_path / "po.xlsx"
    wb.save(path)
    result = ExtractResult(doc_type=rules.get("PO"), is_excel=True, source_form="ACME", fields={"고객PO번호": 20260417})
    expected, got = compare_grid(run_engine, rules, result, path)
    assert got == expected and [c["ok"] for c in got] == [True, None]


def test_excel_empty_and_single_row_table(run_engine, tmp_path):
    """엑셀 PO 도 빈 표 [] / 한 줄짜리 표를 지금 판과 같게 검산한다."""
    rules = example_rules()
    one = acme_file(tmp_path, lambda ws: [ws.__setitem__(f"{c}{r}", None) for r in range(15, 18) for c in "ABCDEFGH"], "one.xlsx")
    empty = acme_file(tmp_path, lambda ws: [ws.__setitem__(f"{c}{r}", None) for r in range(14, 18) for c in "ABCDEFGH"], "empty.xlsx")
    for path, rows in ((one, 1), (empty, 0)):
        result = read_excel(path, rules)
        assert len(result.table) == rows
        expected, got = compare_grid(run_engine, rules, result, path)
        assert got == expected
        assert len(line_items(got)) == rows


# --- 실수(float) 글자가 파이썬 str(float) 과 같은가 ---

def _floats():
    rng = random.Random(20261029)
    values = [0.5, 12.5, 0.1 + 0.2, 100.0, 5.0, 1e15, 1e16, 1.5e22, 1e-4, 1e-5, 1e-7, 2.5e-5, 0.000123, 123456789.125,
              1 / 3, 2 / 3, 1234567890123456.0, 12345678901234567.0, 9007199254740993.0, 0.1, 100000000000001.5, 1e21, 1e22,
              -0.5, -1e-7, -123456.789, 5e-324, 1.7976931348623157e308]
    for _ in range(150):
        kind = rng.randrange(4)
        if kind == 0:
            values.append(round(rng.uniform(-1e6, 1e6), rng.randint(0, 8)))
        elif kind == 1:
            values.append(rng.random() * 10 ** rng.randint(-12, 22))
        elif kind == 2:
            values.append(sum(rng.choice((0.1, 0.01, 0.07, 0.3)) for _ in range(rng.randint(2, 9))))
        else:
            values.append(1 / rng.randint(1, 100000))
    return values


def test_float_fields_text_same_as_python(run_engine, monkeypatch, rules):
    """서식 없는 칸의 실수 값은 result_value 글자가 파이썬 str(float) 과 같다(5.0 · 1e-07 · 1e+16 · 0.30000000000000004 …)."""
    values = _floats()
    result = ExtractResult(doc_type=rules.get("Memo"), fields={"메모번호": "M-1"})
    result.fields.update({f"f{i}": v for i, v in enumerate(values)})
    text = "MEMO\nMemo No: M-1\n0.3 12.5 0.5 100 5 1000000000000000"
    expected, got = compare_text(run_engine, monkeypatch, rules, result, text)

    wrong = [(e, g) for e, g in zip(expected, got) if e != g]
    assert not wrong, f"{len(wrong)}개 다름 (지금 판, 엔진): {wrong[:3]}"
    assert got == expected
    assert any(c["ok"] for c in got[1:]) and any(c["ok"] is False for c in got[1:])


# --- 엑셀 원본 숫자: 유효 15자리로 맞춘 값으로 찾는가 ---

def test_excel_source_numbers_use_15_digits(run_engine, tmp_path):
    """원본 엑셀의 실수 칸은 엑셀처럼 유효 15자리로 본다: 결과 값이 그 값이면 찾고, 17자리 전체 값이면 못 찾는다(지금 판과 같음)."""
    from app.excel_extract import excel_number
    rules = example_rules()
    numbers = [n for n in _floats() if 1e-20 < abs(n) < 1e25] + [123456789012345678, 9007199254740993, 1000000000000005.0,
                                                               1000000000000015.0, 100000000000001.5, 100000000000002.5]
    wb = openpyxl.Workbook()
    ws = wb.active
    for i, number in enumerate(numbers, 1):
        ws.cell(row=i, column=1, value=number)
    path = tmp_path / "numbers.xlsx"
    wb.save(path)
    # 엑셀양식 ACME 의 금액열(단가)에 한 줄씩. 금액 칸은 비워 품목 합(28자리를 넘는 합의 반올림)이 끼어들지 않게 한다.
    # 앞쪽 줄 = 엑셀이 보는 값(찾아야 한다), 뒤쪽 줄 = 17자리 전체 값(15자리와 다르면 못 찾는다)
    table = [{"단가": excel_number(number)} for number in numbers] + [{"단가": Decimal(repr(float(number)))} for number in numbers]
    result = ExtractResult(doc_type=rules.get("PO"), is_excel=True, source_form="ACME", fields={"고객PO번호": "PO-T-1"}, table=table)

    expected, got = compare_grid(run_engine, rules, result, path)

    wrong = [(e, g) for e, g in zip(expected, got) if e != g]
    assert not wrong, f"{len(wrong)}개 다름 (지금 판, 엔진): {wrong[:3]}"
    assert got == expected
    oks = [c["ok"] for c in got]
    assert True in oks and False in oks
