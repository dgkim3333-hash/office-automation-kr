"""무설치판 엔진 extract_grid: 같은 엑셀을 넣으면 지금 판 read_excel 과 같은 결과를 낸다.

엑셀은 openpyxl 로 격자 JSON 으로 옮겨(tests\\portable\\grid.py) 엔진에 넣는다. 시험용 엑셀은 샘플 ACME 를 tmp_path 에 고쳐 저장해 만든다.
"""
import random
from datetime import date, datetime
from decimal import Decimal

import openpyxl
import pytest

from app.excel_extract import excel_number, read_excel
from app.extract import NO_TEXT
from app.portable import rules_to_portable

from .conftest import REPO, example_rules
from .grid import grid_from_xlsx

SAMPLES = sorted((REPO / "samples" / "01_고객오더_엑셀").glob("*.xlsx"))
ACME = next(p for p in SAMPLES if "ACME" in p.name)
NO_FORM = "판별 실패 (규칙.yaml 의 엑셀양식 중 맞는 고객이 없음. 05장대로 이 고객의 엑셀 양식을 가르치세요)"


def _plain(value):
    """파이썬 값을 엔진 JSON 표기로: 금액은 자리 그대로의 글자(지수 표기 없이), 날짜 · 일시는 ISO, 나머지는 그대로."""
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return value


def python_view(result) -> dict:
    return {
        "doc_type": result.doc_type.name if result.doc_type else None,
        "form": result.source_form,
        "is_excel": result.is_excel,
        "fields": {k: _plain(v) for k, v in result.fields.items()},
        "table": None if result.table is None else [{k: _plain(v) for k, v in row.items()} for row in result.table],
        "warnings": list(result.warnings),
        "review_reason": result.review_reason,
    }


def engine_view(out: dict) -> dict:
    return {k: out[k] for k in ("doc_type", "form", "is_excel", "fields", "table", "warnings", "review_reason")}


def acme_file(tmp_path, edit=None, name="po.xlsx"):
    """샘플 ACME 를 고쳐 tmp_path 에 저장한 경로."""
    wb = openpyxl.load_workbook(ACME)
    if edit is not None:
        edit(wb.active)
    path = tmp_path / name
    wb.save(path)
    return path


def compare(run_engine, path):
    """지금 판 read_excel 과 엔진 결과를 같은 모양으로 돌려준다: (지금 판, 엔진)."""
    rules = example_rules()
    expected = python_view(read_excel(path, rules))
    out = run_engine("extract_grid", {"rules": rules_to_portable(rules), "grid": grid_from_xlsx(path)})
    return expected, engine_view(out)


def test_samples_present():
    assert len(SAMPLES) == 2


@pytest.mark.parametrize("xlsx", SAMPLES, ids=["ACME" if "ACME" in p.name else "HANIL" for p in SAMPLES])
def test_same_as_python_on_samples(xlsx, run_engine):
    """샘플 엑셀 2개 각각: form · fields · table · warnings · review_reason 이 지금 판과 같다."""
    expected, got = compare(run_engine, xlsx)

    assert got == expected
    assert got["form"] == ("ACME" if "ACME" in xlsx.name else "HANIL")
    assert got["is_excel"] is True and got["doc_type"] == "PO"
    assert got["review_reason"] is None
    assert got["table"], "품목 표가 비어 있으면 비교가 약하다"
    assert all(type(row["수량"]) is int for row in got["table"])  # 글자 "2" 가 아니라 숫자
    assert got["fields"]["발주일"].count("-") == 2  # "YYYY-MM-DD"
    assert isinstance(got["fields"]["합계금액"], str)  # 금액은 글자


def test_no_form_matches(run_engine, tmp_path):
    """판별키워드가 없는 엑셀(칸 하나뿐인 시트)은 사유가 지금 판과 같은 글자다."""
    wb = openpyxl.Workbook()
    wb.active["A1"] = "just a memo"
    path = tmp_path / "memo.xlsx"
    wb.save(path)

    expected, got = compare(run_engine, path)

    assert got == expected
    assert got["review_reason"] == NO_FORM
    assert got["doc_type"] is None and got["form"] is None and got["is_excel"] is True
    assert got["fields"] == {} and got["table"] is None


def test_empty_sheet_has_no_text(run_engine, tmp_path):
    """칸이 하나도 없는 시트는 「텍스트 없음」 사유(cells 가 빈 배열이어도 오류 없이)."""
    path = tmp_path / "empty.xlsx"
    openpyxl.Workbook().save(path)
    assert grid_from_xlsx(path)["sheets"][0]["cells"] == []

    expected, got = compare(run_engine, path)

    assert got == expected
    assert got["review_reason"] == NO_TEXT and got["is_excel"] is True


def test_float_cell_places(run_engine, tmp_path):
    """단가 칸의 double 212.3375 는 결과 글자 "212.3375" (엑셀처럼 유효 15자리, 뒤의 0 없이). 수식 잡음 0.1+0.2 는 0.3."""
    def edit(ws):
        ws["F14"] = 212.3375
        ws["G14"] = 0.1 + 0.2
        ws["F15"] = 100.0
    expected, got = compare(run_engine, acme_file(tmp_path, edit))

    assert got == expected
    first = got["table"][0]
    assert first["단가"] == "212.3375" and first["금액"] == "0.3"
    assert got["table"][1]["단가"] == "100"


def test_integral_float_in_grid_has_no_places(run_engine, tmp_path):
    """격자 JSON 에 48500.0 처럼 적힌 숫자(openpyxl 은 정수로 읽지만 다른 곳에서 올 수 있다)도 「48500」: 지금 판 excel_number(48500.0)."""
    rules = example_rules()
    grid = grid_from_xlsx(acme_file(tmp_path))
    cell = next(c for c in grid["sheets"][0]["cells"] if c[:2] == [14, 6])
    assert cell[2]["v"] == 48500
    cell[2]["v"] = 48500.0   # JSON 글자는 48500.0

    out = run_engine("extract_grid", {"rules": rules_to_portable(rules), "grid": grid})

    assert excel_number(48500.0) == Decimal(48500)
    assert out["table"][0]["단가"] == "48500"


# 격자에 넣을 숫자: 지금 판 excel_number 와 같은 글자가 나와야 하는 것들
SPECIAL_NUMBERS = [
    212.3375, 0.1 + 0.2, 3 * 0.1, 1 / 3, 2 / 3, -212.3375, -(0.1 + 0.2),
    1e-7, 1e-5, 2.5e-5, 0.000123,                                   # 아주 작은 수 (지수 표기로 저장되는 것)
    1e15, 1e16, 1.5e22, 123456789012345678, 9007199254740993,      # 큰 수 (실수 · 정수)
    5.0, 100.0, 999999999999999.9,                                  # 정수로 보이는 실수 · 올림으로 자리가 늘어나는 수
    1000000000000005.0, 1000000000000015.0,                         # 15자리 다음이 꼭 5 인 수(짝수로 맞춤)
    100000000000001.5, 100000000000002.5, -100000000000001.5,
]


def _random_numbers(count):
    rng = random.Random(20261029)
    numbers = []
    for _ in range(count):
        kind = rng.randrange(6)
        if kind == 0:
            numbers.append(round(rng.uniform(0, 1e6), rng.randint(0, 6)))
        elif kind == 1:
            numbers.append(round(rng.uniform(0, 1000), 2) * rng.randint(1, 500))      # 단가 x 수량의 잡음
        elif kind == 2:
            numbers.append(sum(rng.choice((0.1, 0.01, 0.07, 0.3)) for _ in range(rng.randint(2, 9))))
        elif kind == 3:
            numbers.append(rng.random() * 10 ** rng.randint(-8, 16))                  # 17자리 가까이 다 찬 수
        elif kind == 4:
            numbers.append(1 / rng.randint(1, 10000))
        else:                                                                         # 15자리 다음이 5 인 16자리 수
            numbers.append(float(rng.randrange(10 ** 15, 2 ** 53) // 10 * 10 + 5) / 10 ** rng.randint(0, 3))
    return numbers


def _sweep_file(tmp_path, numbers):
    """ACME 모양으로 단가 · 금액 칸에 숫자를 한 줄에 하나씩 넣은 엑셀."""
    def edit(ws):
        for r in range(14, 19):
            for c in "ABCDEFGH":
                ws[f"{c}{r}"] = None
        for i, number in enumerate(numbers):
            r = 14 + i
            ws[f"A{r}"], ws[f"B{r}"], ws[f"D{r}"], ws[f"F{r}"], ws[f"G{r}"] = i + 1, f"P-{i + 1}", 1, number, number
    return acme_file(tmp_path, edit, "sweep.xlsx")


@pytest.mark.parametrize("label, numbers", [
    ("special", SPECIAL_NUMBERS),
    ("random", _random_numbers(1200)),
])
def test_number_text_same_as_python(label, numbers, run_engine, tmp_path):
    """숫자 칸이 어떤 값이어도 결과 글자가 지금 판 excel_number(유효 15자리)와 같다."""
    path = _sweep_file(tmp_path, numbers)
    expected, got = compare(run_engine, path)

    assert expected["table"] is not None and len(expected["table"]) == len(numbers)
    assert got["table"] is not None and len(got["table"]) == len(numbers)
    wrong = [(n, e["단가"], g["단가"]) for n, e, g in zip(numbers, expected["table"], got["table"]) if e["단가"] != g["단가"]]
    assert not wrong, f"{len(wrong)}개 다름 (값, 지금 판, 엔진): {wrong[:5]}"
    assert got == expected


# (이름, 샘플 ACME 를 고치는 함수, 지금 판 결과에 그 상황이 실제로 나왔는지 확인하는 함수)
def _blank(ws, rows, cols="ABCDEFGH"):
    for r in rows:
        for c in cols:
            ws[f"{c}{r}"] = None


def _general(ws, coord, value):
    ws[coord] = value
    ws[coord].number_format = "General"


VARIANTS = [
    ("label_missing", lambda ws: ws.__setitem__("A5", None),
     lambda e: "머리 '발주일' 칸(오른쪽:PO Date:)을 못 찾음" in e["warnings"]),
    ("number_label_missing", lambda ws: ws.__setitem__("A4", None),
     lambda e: e["review_reason"] == "번호필드 미추출 ('고객PO번호' 패턴이 맞지 않음)"),
    ("no_item_rows", lambda ws: _blank(ws, range(14, 18)),
     lambda e: e["table"] == []),
    ("items_header_missing", lambda ws: ws.__setitem__("B13", "Part"),
     lambda e: e["table"] is None and "품목 머리글 'Supplier Part No' 를 못 찾음, 품목 건너뜀" in e["warnings"]),
    ("text_quantity", lambda ws: ws.__setitem__("D15", "abc"),
     lambda e: e["table"] is None and "품목 15행 '수량' 값이 숫자가 아님, 품목 건너뜀" in e["warnings"]),
    ("fraction_quantity", lambda ws: ws.__setitem__("D15", 2.5),
     lambda e: e["table"] is None),
    ("integral_float_quantity", lambda ws: ws.__setitem__("D15", 4.0),
     lambda e: e["table"][1]["수량"] == 4),
    ("amounts_as_text", lambda ws: (ws.__setitem__("D14", "2"), ws.__setitem__("F14", "48,500.00"), ws.__setitem__("G14", "$97,000.50")),
     lambda e: e["table"][0]["금액"] == "97000.50"),
    ("blank_amount", lambda ws: ws.__setitem__("F15", None),
     lambda e: e["table"] is None and "품목 15행 '단가' 값이 숫자가 아님, 품목 건너뜀" in e["warnings"]),
    ("bool_amount", lambda ws: ws.__setitem__("G18", True),
     lambda e: "'합계금액' 값 'True' 를 금액로 바꾸지 못함" in e["warnings"]),
    ("total_blank", lambda ws: ws.__setitem__("G18", None),
     lambda e: e["fields"]["합계금액"] is None and not any("합계금액" in w for w in e["warnings"])),
    ("date_as_text", lambda ws: (ws.__setitem__("B5", "2026-09-18"), ws.__setitem__("B8", " 09/20/2026 ")),
     lambda e: e["fields"]["발주일"] == "2026-09-18" and e["fields"]["납기요청일"] == "2026-09-20"),
    ("date_as_bad_text", lambda ws: ws.__setitem__("B5", "18.09.2026"),
     lambda e: "'발주일' 값 '18.09.2026' 를 날짜로 바꾸지 못함" in e["warnings"]),
    ("date_as_number", lambda ws: _general(ws, "B5", 46283),
     lambda e: "'발주일' 값 '46283' 를 날짜로 바꾸지 못함" in e["warnings"]),
    ("date_with_time", lambda ws: ws.__setitem__("B5", datetime(2026, 9, 18, 13, 45)),
     lambda e: e["fields"]["발주일"] == "2026-09-18"),
    ("datetime_in_plain_cell", lambda ws: ws.__setitem__("B11", datetime(2026, 9, 18, 13, 45, 30)),
     lambda e: e["fields"]["담당"] == "2026-09-18T13:45:30"),
    ("date_in_plain_cell", lambda ws: ws.__setitem__("B11", datetime(2026, 9, 18)),
     lambda e: e["fields"]["담당"] == "2026-09-18"),
    ("plain_number_and_bool", lambda ws: (ws.__setitem__("B9", True), ws.__setitem__("B10", 12.5), ws.__setitem__("B11", 7)),
     lambda e: e["fields"]["인도조건"] is True and e["fields"]["통화"] == 12.5 and e["fields"]["담당"] == 7),
    ("whitespace_values", lambda ws: (ws.__setitem__("B10", "  USD  "), ws.__setitem__("B9", "   ")),
     lambda e: e["fields"]["통화"] == "USD" and e["fields"]["인도조건"] is None),
    ("first_label_wins", lambda ws: (ws.__setitem__("A20", "PO Number:"), ws.__setitem__("B20", "PO-OTHER")),
     lambda e: e["fields"]["고객PO번호"] == "PO-ACME-2026-0417"),
    ("blank_text_ends_items", lambda ws: ws.__setitem__("B16", "   "),
     lambda e: len(e["table"]) == 2),
    ("zero_part_continues", lambda ws: ws.__setitem__("B16", 0),
     lambda e: len(e["table"]) == 4 and e["table"][2]["품번"] == 0),
    ("keyword_case_ignored", lambda ws: ws.__setitem__("A3", "BUYER:"),
     lambda e: e["form"] == "ACME" and any("고객명" in w for w in e["warnings"])),
    ("header_text_first_match", lambda ws: ws.__setitem__("A2", "Supplier Part No"),
     lambda e: e["table"] is None and "품목 3행 '수량' 값이 숫자가 아님, 품목 건너뜀" in e["warnings"]),
]


@pytest.mark.parametrize("edit, check", [(v[1], v[2]) for v in VARIANTS], ids=[v[0] for v in VARIANTS])
def test_variants_same_as_python(edit, check, run_engine, tmp_path):
    """샘플 ACME 를 한 군데씩 어긋나게 고쳐도 결과 · 경고 · 사유가 지금 판과 같다."""
    expected, got = compare(run_engine, acme_file(tmp_path, edit))

    assert check(expected), "이 상황이 지금 판 결과에 나타나야 시험이 의미가 있다"
    assert got == expected


def test_active_sheet_is_read(run_engine, tmp_path):
    """활성 시트(wb.active)만 읽는다: 앞 시트가 달라도, 활성이 뒤 시트면 뒤 시트."""
    wb = openpyxl.load_workbook(ACME)
    real = wb.active
    front = wb.create_sheet("Memo", 0)
    front["A1"] = "just a memo"
    wb.active = wb.worksheets.index(real)
    path = tmp_path / "two.xlsx"
    wb.save(path)
    grid = grid_from_xlsx(path)
    assert grid["active"] == 1 and [s["name"] for s in grid["sheets"]] == ["Memo", real.title]

    expected, got = compare(run_engine, path)

    assert got == expected
    assert got["form"] == "ACME"


@pytest.mark.parametrize("row, filler, matched", [
    (20, range(1, 20), True),     # 칸이 있는 줄의 20번째
    (21, range(1, 21), False),    # 21번째 줄은 판별 범위 밖
    (28, range(1, 5), True),      # 빈 줄은 세지 않는다: 28행이지만 5번째 줄
    (31, range(1, 5), False),     # 31행부터는 읽지 않는다
], ids=["20th_line", "21st_line", "sparse_row_28", "row_31"])
def test_detect_window(row, filler, matched, run_engine, tmp_path):
    """판별은 1~30행 중 칸이 있는 첫 20줄만 본다(지금 판 DETECT_ROWS = 20)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    for r in filler:
        ws[f"A{r}"] = f"memo {r}"
    ws[f"A{row}"] = "Buyer:"
    path = tmp_path / "detect.xlsx"
    wb.save(path)

    expected, got = compare(run_engine, path)

    assert (expected["form"] == "ACME") is matched
    assert got == expected
    if not matched:
        assert got["review_reason"] == NO_FORM
