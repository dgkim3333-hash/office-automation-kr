"""무설치판 오피스 읽기(작업 7): 워드로 PDF 글자 · 엑셀로 격자를 읽어도 지금 판과 같은 엔진 결과가 나온다.

오피스가 없는 PC 에서는 @office 시험을 건너뛴다. 원본은 tmp_path 의 복사본만 쓰고(C:\\PO시험 은 읽기만), 시험마다
이 시험 전에 없던 /Automation 워드 · 엑셀이 남지 않았는지 확인한다.
"""
import re
import shutil
import subprocess
import time
from datetime import datetime
from pathlib import Path

import openpyxl
import pytest

try:
    import msvcrt
    import winreg
except ImportError:   # Windows 가 아닌 곳(클라우드 등)에서는 오피스 시험을 건너뜀
    msvcrt = winreg = None

from app import extract
from app.extract import extract_from_text
from app.portable import rules_to_portable
from app.rules import rules_from_text

from .conftest import REPO, EngineError, example_rules, without_master
from .grid import grid_from_xlsx
from .test_engine_text import engine_view, python_view

SAMPLE_PDFS = sorted(p for p in (REPO / "samples" / "03_Invoice_DM_PDF").glob("*.pdf") if not p.name.startswith("Scan_"))
SAMPLE_XLSX = sorted((REPO / "samples" / "01_고객오더_엑셀").glob("*.xlsx"))
KIT = Path("C:/PO시험")
KIT_RULES = KIT / "04_정답" / "참고_검증된_규칙.yaml"
KIT_PDFS = ["PO_BPI-2026-0733.pdf", "메일_Order LDH-PO-2611-033.pdf"]
KIT_XLSX = "주문서_Orion_OWK-PO-26-1187.xlsx"
kit_missing = not KIT_RULES.exists() or not all((KIT / "01_입력예시" / n).exists() for n in KIT_PDFS + [KIT_XLSX])

WORD_TIMEOUT = "워드가 응답하지 않음"
EXCEL_TIMEOUT = "엑셀이 응답하지 않음"
READ_TIMEOUT = 90   # 시험에서 파일 하나에 주는 시간(엔진 부르기 도우미의 120초 안)


def _installed(progid: str) -> bool:
    if winreg is None:
        return False
    try:
        winreg.CloseKey(winreg.OpenKey(winreg.HKEY_CLASSES_ROOT, progid + r"\CurVer"))
        return True
    except OSError:
        return False


office = pytest.mark.skipif(not (_installed("Word.Application") and _installed("Excel.Application")),
                            reason="워드 · 엑셀이 없어 오피스 시험을 건너뜀")


def automation_pids() -> set[tuple[str, int]]:
    """지금 떠 있는 /Automation 워드 · 엑셀 (이름, PID)."""
    script = ("Get-CimInstance Win32_Process -Filter \"Name='WINWORD.EXE' or Name='EXCEL.EXE'\" | "
              "Where-Object { $_.CommandLine -like '*/Automation*' } | "
              "ForEach-Object { $_.Name + ' ' + $_.ProcessId }")
    done = subprocess.run(["powershell.exe", "-NoProfile", "-NonInteractive", "-Command", script],
                          capture_output=True, text=True, timeout=60)
    found = set()
    for line in done.stdout.splitlines():
        if line.strip():
            name, pid = line.split()
            found.add((name.upper(), int(pid)))
    return found


@pytest.fixture
def no_leftover():
    """시험이 끝난 뒤 시험 전에 없던 /Automation 워드 · 엑셀이 남아 있으면 실패."""
    before = automation_pids()
    yield
    left = automation_pids() - before
    assert not left, f"시험 뒤 남은 오피스 프로세스: {sorted(left)}"


@pytest.fixture
def temp_dir(tmp_path, monkeypatch):
    """엔진의 %TEMP% 를 시험 폴더로 바꾼다(임시 폴더가 남는지 보려고)."""
    temp = tmp_path / "temp"
    temp.mkdir()
    monkeypatch.setenv("TEMP", str(temp))
    monkeypatch.setenv("TMP", str(temp))
    return temp


def own_temp(temp) -> list:
    """이 프로그램이 만든 임시 폴더(무설치판_*). 오피스가 스스로 남기는 파일(Diagnostics 등)은 세지 않는다."""
    return sorted(p.name for p in temp.iterdir() if p.name.startswith("무설치판_"))


def kit_rules():
    return rules_from_text(without_master(KIT_RULES.read_text(encoding="utf-8")))


def read_word(run_engine, path, timeout=READ_TIMEOUT) -> dict:
    """{text, pages}. 문서를 오류 없이 닫았는지, 워드가 Quit 뒤 스스로 끝났는지(강제로 끝낸 것 0개)도 본다."""
    out = run_engine("office", {"action": "read_word", "path": str(path), "timeout": timeout})
    assert out.pop("close_error") is None, "워드 문서 닫기가 실패함"
    assert out.pop("forced") == 0, "Quit 한 워드가 스스로 끝나지 않아 강제로 끝냄"
    assert out.pop("left") is None, "남의 문서가 없는데 워드를 넘겨 줌"
    return out


def read_excel(run_engine, path, timeout=READ_TIMEOUT) -> dict:
    """격자. 엑셀이 Quit 뒤 스스로 끝났는지(COM 참조를 다 놓았는지)도 본다."""
    out = run_engine("office", {"action": "read_excel", "path": str(path), "timeout": timeout})
    assert out["forced"] == 0, "Quit 한 엑셀이 스스로 끝나지 않아 강제로 끝냄(COM 참조가 남음)"
    assert out["left"] is None, "남의 문서가 없는데 엑셀을 넘겨 줌"
    return out["grid"]


# --- 오피스 없이 ---------------------------------------------------------------------------------

def test_select_own_pids(run_engine):
    """실행 뒤에 새로 생겼고 명령줄에 /Automation 이 있는 것만 고른다."""
    before = [{"pid": 1, "commandline": r'"C:\Office\WINWORD.EXE"'},
              {"pid": 2, "commandline": r'"C:\Office\WINWORD.EXE" /Automation -Embedding'}]
    after = before + [{"pid": 3, "commandline": r'"C:\Office\WINWORD.EXE" /Automation -Embedding'},
                      {"pid": 4, "commandline": r'"C:\Office\WINWORD.EXE" /n "문서.docx"'}]

    assert run_engine("office", {"action": "select_pids", "before": before, "after": after}) == [3]


def test_select_own_pids_edge_cases(run_engine):
    """명령줄을 못 읽은 프로세스(null)는 고르지 않는다. 고를 것이 없거나 여럿이어도 배열."""
    auto = r'"C:\Office\EXCEL.EXE" /Automation -Embedding'
    run = lambda before, after: run_engine("office", {"action": "select_pids", "before": before, "after": after})
    assert run([], [{"pid": 7, "commandline": None}]) == []
    assert run([], []) == []
    assert run([{"pid": 5, "commandline": auto}], [{"pid": 5, "commandline": auto}]) == []
    assert run([], [{"pid": 8, "commandline": auto}, {"pid": 9, "commandline": auto}]) == [8, 9]


def _para(text, list_string="", page=1, y=0.0, x=0.0):
    return {"text": text, "list": list_string, "page": page, "y": y, "x": x}


def test_normalize_word_text(run_engine):
    """정규화: 목록 번호 붙이기 · 탭과 연속 공백 → 한 칸 · 앞뒤 공백 · '/' 줄 · 빈 줄 지우기 · 표 행은 높이가 같은 칸끼리 한 줄."""
    blocks = [
        {"kind": "p", "text": "Vertex  Thermal\tINVOICE\r", "list": ""},
        {"kind": "p", "text": "/\r", "list": ""},
        {"kind": "p", "text": "   \r", "list": ""},
        {"kind": "p", "text": "VTS-RDHX-42U  Rear Door  18  9,875.00\r", "list": "1"},
        {"kind": "p", "text": "Delivery within 35 days.\r", "list": "2."},
        {"kind": "p", "text": "Authorized by: K. Lim    Page 1 of 2\f\r", "list": ""},
        {"kind": "p", "text": "first\vsecond\r", "list": ""},
        # 표 한 행: 칸 안의 문단이 여러 줄이면 높이(y)가 같은 문단끼리 왼쪽(x)부터 한 줄로
        {"kind": "row", "cells": [
            _para("Reason:\tFreight surcharge \r", y=228.0, x=57.6),
            _para("Line\tItem\r\a", y=265.2, x=57.6),
            _para("- expedited air\r", y=229.2, x=174.6),
            _para("Description\r\a", y=265.2, x=174.6),
            _para("Qty\tUnit Price\tAmount\r\a", y=265.2, x=397.8),
        ]},
        {"kind": "row", "cells": [_para("From:\r\a", y=87, x=49.8), _para("Daniel Kim\r\a", y=87, x=120.0),
                                  _para("\r\a", y=80, x=300.0)]},
        {"kind": "p", "text": "Total:\t$1.00\r", "list": ""},
    ]

    out = run_engine("office", {"action": "normalize_word", "blocks": blocks})

    assert out == {"text": "\n".join([
        "Vertex Thermal INVOICE",
        "1 VTS-RDHX-42U Rear Door 18 9,875.00",
        "2. Delivery within 35 days.",
        "Authorized by: K. Lim Page 1 of 2",
        "first",
        "second",
        "Reason: Freight surcharge - expedited air",
        "Line Item Description Qty Unit Price Amount",
        "From: Daniel Kim",
        "Total: $1.00",
    ])}


def test_normalize_word_row_without_positions(run_engine):
    """위치를 못 읽은 표 행(y 가 음수)은 칸을 차례대로 한 칸 띄어 한 줄로."""
    blocks = [{"kind": "row", "cells": [_para("A:\r\a", y=-1), _para("1\r\a", y=-1), _para("B\r\a", y=-1)]}]
    assert run_engine("office", {"action": "normalize_word", "blocks": blocks}) == {"text": "A: 1 B"}


def test_normalize_word_rows_on_two_pages(run_engine):
    """다른 쪽의 문단은 높이가 같아도 다른 줄이다(쪽 번호 먼저)."""
    blocks = [{"kind": "row", "cells": [_para("second page\r", page=2, y=50), _para("first page\r", page=1, y=50)]}]
    assert run_engine("office", {"action": "normalize_word", "blocks": blocks}) == {"text": "first page\nsecond page"}


def test_new_copy_fresh_folder_without_internet_mark(tmp_path, temp_dir, run_engine):
    """복사본: 실행마다 새 이름의 짧은 폴더(무설치판_<임의 8자>), 파일 이름은 원본<확장자>(긴 이름이 경로 한도를 넘지 않게),
    바이트는 그대로, 인터넷 표시는 복사본에서만 지운다."""
    original = tmp_path / "메일_Order 시험-PO 1.pdf"
    original.write_bytes(b"%PDF-1.4 fake")
    zone = b"[ZoneTransfer]\r\nZoneId=3\r\n"
    with open(f"{original}:Zone.Identifier", "wb") as f:
        f.write(zone)

    first = run_engine("office", {"action": "new_copy", "path": str(original)})
    second = run_engine("office", {"action": "new_copy", "path": str(original)})

    for made in (first, second):
        folder, copy = Path(made["folder"]), Path(made["path"])
        assert folder.parent == temp_dir
        assert re.fullmatch(r"무설치판_[0-9a-f]{8}", folder.name), folder.name
        assert copy == folder / "원본.pdf" and copy.read_bytes() == original.read_bytes()
        with pytest.raises(FileNotFoundError):
            open(f"{copy}:Zone.Identifier", "rb")
    assert first["folder"] != second["folder"]
    with open(f"{original}:Zone.Identifier", "rb") as f:
        assert f.read() == zone


@pytest.mark.parametrize("app", ["word", "excel"])
def test_office_opens_the_copy(app, tmp_path, temp_dir, run_engine):
    """오피스 일에 넘기는 경로는 원본이 아니라 새 임시 폴더의 복사본(원본<확장자>)이다. 일이 끝나면 폴더는 지운다(오피스 없이)."""
    original = tmp_path / ("메일_Order 시험-PO 1.pdf" if app == "word" else "주문 1.xlsx")
    original.write_bytes(b"fake")

    out = run_engine("office", {"action": "work_path", "app": app, "path": str(original)})

    given = Path(out["path"])
    assert given != original and given.name == "원본" + original.suffix
    assert given.parent.parent == temp_dir and given.parent.name.startswith("무설치판_")
    assert own_temp(temp_dir) == []


def test_new_copy_failure_leaves_no_folder(tmp_path, temp_dir, run_engine):
    """복사가 실패하면(없는 파일 · 다른 프로그램이 잠근 파일) 오류를 내고 만든 임시 폴더는 남기지 않는다."""
    with pytest.raises(EngineError):
        run_engine("office", {"action": "new_copy", "path": str(tmp_path / "없음.pdf")})
    assert own_temp(temp_dir) == []

    locked = tmp_path / "잠김.pdf"
    locked.write_bytes(b"%PDF-1.4 " + b"x" * 1000)
    with open(locked, "rb") as f:
        msvcrt.locking(f.fileno(), msvcrt.LK_NBLCK, 1000)   # 읽기를 막는 잠금(복사가 도중에 실패)
        try:
            with pytest.raises(EngineError):
                run_engine("office", {"action": "new_copy", "path": str(locked)})
        finally:
            f.seek(0)
            msvcrt.locking(f.fileno(), msvcrt.LK_UNLCK, 1000)
    assert own_temp(temp_dir) == []


def test_stop_own_refuses_other_names(run_engine):
    """시험용 stop_own 은 워드 · 엑셀(과 시험용 powershell.exe) 이름만 받는다."""
    for name in ("notepad.exe", "explorer.exe", "WINWORD"):
        with pytest.raises(EngineError) as error:
            run_engine("office", {"action": "stop_own", "name": name, "before": [], "grace": 0})
        assert "끝낼 수 없는 프로세스 이름" in str(error.value)


def _powershell_pids() -> list[dict]:
    out = subprocess.run(["powershell.exe", "-NoProfile", "-Command", "Get-Process powershell | ForEach-Object { $_.Id }"],
                         capture_output=True, text=True, timeout=60).stdout
    return [{"pid": int(x)} for x in out.split()]


def test_stop_own_waits_then_kills_only_own(run_engine):
    """Stop-OwnOffice: 실행 뒤 생긴 /Automation 프로세스만, 스스로 끝나기를 기다렸다가(grace) 남은 것만 강제로 끝낸다."""
    before = _powershell_pids()
    sleep = lambda sec, mark: subprocess.Popen(["powershell.exe", "-NoProfile", "-Command", f"Start-Sleep -Seconds {sec} {mark}"])
    quick = sleep(2, "# /Automation")    # grace 안에 스스로 끝남
    stuck = sleep(120, "# /Automation")  # 끝나지 않음 → 강제로
    other = sleep(120, "# user")         # 사용자 것(명령줄에 /Automation 없음) → 건드리지 않음
    try:
        time.sleep(1)
        out = run_engine("office", {"action": "stop_own", "name": "powershell.exe", "before": before, "grace": 8})

        assert out == {"forced": [stuck.pid]}
        assert quick.wait(timeout=5) == 0
        assert stuck.poll() is not None and stuck.returncode != 0
        assert other.poll() is None
    finally:
        for p in (quick, stuck, other):
            if p.poll() is None:
                p.kill()


# --- 워드 ---------------------------------------------------------------------------------------

def _word_cases():
    cases = [pytest.param(p, "samples", id=p.name) for p in SAMPLE_PDFS]
    kit = [pytest.mark.po_test_kit, pytest.mark.skipif(kit_missing, reason="C:\\PO시험 시험 자료가 없음")]
    cases += [pytest.param(KIT / "01_입력예시" / n, "kit", id=n, marks=kit) for n in KIT_PDFS]
    return cases


def test_sample_pdfs_present():
    assert len(SAMPLE_PDFS) == 5


@office
@pytest.mark.office
@pytest.mark.parametrize("source,kind", _word_cases())
def test_word_text_same_engine_result(source, kind, tmp_path, run_engine, no_leftover):
    """워드로 읽은 글자(정규화)로 낸 엔진 결과가 pdfplumber 글자로 낸 엔진 결과(= 지금 판)와 같다."""
    pdf = tmp_path / source.name
    shutil.copyfile(source, pdf)   # C:\PO시험 원본은 읽기만
    rules = example_rules() if kind == "samples" else kit_rules()
    portable = rules_to_portable(rules)
    plumber_text = extract.read_text(pdf)

    word = read_word(run_engine, pdf)

    assert isinstance(word["pages"], int) and word["pages"] >= 1
    expected = python_view(extract_from_text(plumber_text, rules))
    from_plumber = run_engine("extract_text", {"rules": portable, "text": plumber_text})
    from_word = run_engine("extract_text", {"rules": portable, "text": word["text"]})
    assert engine_view(from_plumber) == expected
    assert engine_view(from_word) == engine_view(from_plumber)
    if kind == "kit":
        # 줄번호가 워드의 글머리 번호 목록이 되는 PDF: 번호를 되살려 줄 열이 1, 2, 3 …
        assert [row["줄"] for row in from_word["table"]] == [str(i + 1) for i in range(len(from_word["table"]))]
        assert len(from_word["table"]) >= 5
    if source.name.startswith(("INV-", "DM-")):
        assert from_word["table"], "품목 표가 비어 있으면 비교가 약하다"


@office
@pytest.mark.office
@pytest.mark.parametrize("suffix", [".pdf", ".xlsx"])
def test_korean_name_and_original_untouched(suffix, tmp_path, temp_dir, run_engine, no_leftover):
    """한글 · 공백 이름의 파일도 읽고, 원본(바이트 · 수정 시각 · 인터넷 표시)은 그대로, 임시 폴더는 남지 않는다."""
    source = SAMPLE_PDFS[2] if suffix == ".pdf" else SAMPLE_XLSX[0]
    assert source.name.startswith(("INV-", "고객오더_"))
    original = tmp_path / f"메일_Order 시험-PO 1{suffix}"
    shutil.copyfile(source, original)
    zone = b"[ZoneTransfer]\r\nZoneId=3\r\n"
    with open(f"{original}:Zone.Identifier", "wb") as f:   # 메일 첨부처럼 인터넷 표시
        f.write(zone)
    data = original.read_bytes()
    stamp = original.stat().st_mtime_ns

    if suffix == ".pdf":
        out = read_word(run_engine, original)
        assert "Invoice No: INV-2026-10231" in out["text"]
    else:
        out = read_excel(run_engine, original)
        assert out["sheets"][0]["name"] == "Purchase Order"

    assert original.read_bytes() == data
    assert original.stat().st_mtime_ns == stamp
    with open(f"{original}:Zone.Identifier", "rb") as f:
        assert f.read() == zone
    assert own_temp(temp_dir) == [], "임시 폴더가 남음"


def long_named(tmp_path: Path, source: Path, total: int = 240) -> Path:
    """전체 경로가 total 자(한글 · 띄어쓰기 섞임)인 원본 복사. 메일 첨부의 긴 이름 흉내."""
    folder = tmp_path / "메일 첨부 받은 폴더"
    folder.mkdir(exist_ok=True)
    stem = "메일_Order 긴 이름 시험 "
    room = total - len(str(folder)) - 1 - len(source.suffix)
    assert room > len(stem) + 20, "tmp_path 가 너무 길어 시험할 수 없음"
    name = (stem + "가나다 ABC-123 " * 40)[:room].rstrip() + source.suffix
    path = folder / name
    shutil.copyfile(source, path)
    assert len(str(path)) >= total - 2
    return path


@office
@pytest.mark.office
@pytest.mark.parametrize("suffix", [".pdf", ".xlsx"])
def test_long_name_reads(suffix, tmp_path, temp_dir, run_engine, no_leftover):
    """전체 경로가 240 자쯤인 원본도 읽는다: 오피스가 여는 복사본은 짧은 경로(%TEMP%\\무설치판_<임의>\\원본<확장자>)라
    엑셀 218 자 · 워드 MAX_PATH 한도에 걸리지 않는다. 읽은 값은 짧은 이름의 같은 파일과 같다."""
    source = SAMPLE_PDFS[2] if suffix == ".pdf" else SAMPLE_XLSX[0]
    original = long_named(tmp_path, source)
    short = tmp_path / f"짧음{suffix}"
    shutil.copyfile(source, short)

    if suffix == ".pdf":
        assert read_word(run_engine, original) == read_word(run_engine, short)
    else:
        assert read_excel(run_engine, original) == read_excel(run_engine, short)
    assert own_temp(temp_dir) == [], "임시 폴더가 남음"


@office
@pytest.mark.office
@pytest.mark.parametrize("action,message", [("read_word", WORD_TIMEOUT), ("read_excel", EXCEL_TIMEOUT)])
def test_timeout_cleans_up(action, message, tmp_path, temp_dir, run_engine, no_leftover):
    """시간 제한 0초 → OfficeTimeout(문구), 이 실행이 띄운 /Automation 오피스는 끝내고 임시 폴더도 지운다."""
    source = SAMPLE_PDFS[2] if action == "read_word" else SAMPLE_XLSX[0]
    path = tmp_path / source.name
    shutil.copyfile(source, path)

    with pytest.raises(EngineError) as error:
        run_engine("office", {"action": action, "path": str(path), "timeout": 0})

    assert str(error.value) == message
    assert own_temp(temp_dir) == []


@office
@pytest.mark.office
@pytest.mark.parametrize("app", ["word", "excel"])
def test_stuck_office_is_killed(app, tmp_path, temp_dir, run_engine, no_leftover):
    """파일을 연 채 멈춘 오피스(스스로는 끝나지 않음) → 시간 제한 오류, 이 실행이 띄운 오피스는 끝내고 임시 폴더도 지운다."""
    source = SAMPLE_PDFS[2] if app == "word" else SAMPLE_XLSX[0]
    path = tmp_path / source.name
    shutil.copyfile(source, path)

    with pytest.raises(EngineError) as error:
        run_engine("office", {"action": "stuck", "app": app, "path": str(path), "timeout": 15})

    assert str(error.value) == (WORD_TIMEOUT if app == "word" else EXCEL_TIMEOUT)
    assert own_temp(temp_dir) == []


@office
@pytest.mark.office
def test_reopen_after_crash(tmp_path, run_engine, no_leftover):
    """파일을 연 채 멈춰 강제로 끝낸 워드 다음에도, 같은 파일을 바로 다시 읽으면 「지난번 오류」 창에 막히지 않는다."""
    pdf = tmp_path / SAMPLE_PDFS[2].name
    shutil.copyfile(SAMPLE_PDFS[2], pdf)

    with pytest.raises(EngineError) as error:
        run_engine("office", {"action": "stuck", "app": "word", "path": str(pdf), "timeout": 15})
    assert str(error.value) == WORD_TIMEOUT

    out = read_word(run_engine, pdf, timeout=60)
    assert "Invoice No: INV-2026-10231" in out["text"]


# --- 남의 문서 보호(최종 검토 I3): 숨은 오피스에 친구가 연 문서가 들어온 경우 ---

def friend_document(tmp_path: Path, app: str) -> Path:
    """친구가 연 문서 흉내(엑셀 통합 문서 · 워드가 여는 RTF). 우리 인스턴스에 이 파일을 연다(두 번 클릭 흉내)."""
    if app == "excel":
        path = tmp_path / "친구 문서.xlsx"
        book = openpyxl.Workbook()
        book.active["A1"] = "친구가 쓰던 표"
        book.save(path)
    else:
        path = tmp_path / "친구 문서.rtf"
        path.write_bytes(b"{\\rtf1\\ansi friend's notes\\par}")
    return path


def is_locked(path: Path) -> bool:
    """오피스가 그 파일을 열고 있으면(쓰기로 열 수 없음) True."""
    try:
        with open(path, "r+b"):
            return False
    except PermissionError:
        return True


def window_shown(pid: int, wait: float = 15) -> bool:
    """그 프로세스에 화면 창(MainWindowHandle)이 생겼는지(최대 wait 초 기다림)."""
    deadline = time.monotonic() + wait
    while True:
        done = subprocess.run(["powershell.exe", "-NoProfile", "-NonInteractive", "-Command",
                               f"(Get-Process -Id {pid} -ErrorAction SilentlyContinue).MainWindowHandle"],
                              capture_output=True, text=True, timeout=60)
        if done.stdout.strip() not in ("", "0"):
            return True
        if time.monotonic() > deadline:
            return False
        time.sleep(0.5)


def kill_new_office(before: set) -> None:
    """시험이 띄운(시험 전에 없던) /Automation 오피스를 끝내고 사라질 때까지 기다린다."""
    for _, pid in automation_pids() - before:
        subprocess.run(["taskkill", "/F", "/PID", str(pid)], capture_output=True)
    deadline = time.monotonic() + 30
    while automation_pids() - before and time.monotonic() < deadline:
        time.sleep(0.5)


HANDOVER = {
    "excel": "참고: 숨은 엑셀에 이 프로그램이 열지 않은 통합 문서(친구 문서.xlsx)가 열려 있어 엑셀을 끝내지 않고 화면에 보이게 넘겼습니다(저장은 직접 하세요).",
    "word": "참고: 숨은 워드에 이 프로그램이 열지 않은 문서(친구 문서.rtf)가 열려 있어 워드를 끝내지 않고 화면에 보이게 넘겼습니다(저장은 직접 하세요).",
}
FOREIGN = {
    "excel": "새로 띄운 엑셀에 이미 다른 통합 문서가 열려 있어(사용자가 연 엑셀일 수 있음) 손대지 않았습니다.",
    "word": "새로 띄운 워드에 이미 다른 문서가 열려 있어(사용자가 연 워드일 수 있음) 손대지 않았습니다.",
}


@office
@pytest.mark.office
@pytest.mark.parametrize("app", ["excel", "word"])
def test_hidden_office_hands_over_friends_document(app, tmp_path, temp_dir, run_engine, no_leftover):
    """읽는 동안 친구가 연 문서가 우리 숨은 인스턴스에 들어오면: 우리 문서만 닫고 Quit 하지 않는다. 인스턴스는 살아서 화면에 보이고
    친구 문서는 열린 채, 넘겨 준 기록(left)이 남는다. 읽기 결과는 그대로."""
    friend = friend_document(tmp_path, app)
    before = automation_pids()
    try:
        if app == "excel":
            out = run_engine("office", {"action": "read_excel", "path": str(SAMPLE_XLSX[0]), "test_extra": str(friend)})
            assert out["grid"]["sheets"][0]["name"] == "Purchase Order"
        else:
            out = run_engine("office", {"action": "read_word", "path": str(SAMPLE_PDFS[2]), "test_extra": str(friend)})
            assert "Invoice No: INV-2026-10231" in out["text"]
        mine = automation_pids() - before
        assert len(mine) == 1, f"우리 인스턴스가 끝나 버림(친구 문서도 함께 닫힘): {sorted(mine)}"
        assert is_locked(friend), "친구 문서가 닫힘"
        assert window_shown(next(iter(mine))[1]), "넘겨 준 인스턴스가 화면에 보이지 않음"
        assert out["left"] == HANDOVER[app]
        assert out["forced"] == 0
        assert own_temp(temp_dir) == [], "임시 폴더가 남음"
    finally:
        kill_new_office(before)


@office
@pytest.mark.office
@pytest.mark.parametrize("app", ["excel", "word"])
def test_office_with_open_document_is_not_ours(app, tmp_path, temp_dir, run_engine, no_leftover):
    """만든 직후 이미 문서가 열린 인스턴스(사용자 오피스에 붙은 경우 흉내)는 우리 것이 아니다: 보이기 · Quit · 강제 종료 없이
    그 읽기만 실패한다(분명한 한국어 사유)."""
    friend = friend_document(tmp_path, app)
    before = automation_pids()
    try:
        action, source = ("read_excel", SAMPLE_XLSX[0]) if app == "excel" else ("read_word", SAMPLE_PDFS[2])
        with pytest.raises(EngineError) as error:
            run_engine("office", {"action": action, "path": str(source), "test_preopen": str(friend)})
        assert FOREIGN[app] in str(error.value)
        mine = automation_pids() - before
        assert len(mine) == 1, f"그 인스턴스를 끝내 버림: {sorted(mine)}"
        assert not window_shown(next(iter(mine))[1], wait=2), "건드리지 말아야 할 인스턴스를 보이게 함"
        assert is_locked(friend)
        assert own_temp(temp_dir) == [], "임시 폴더가 남음"
    finally:
        kill_new_office(before)


@office
@pytest.mark.office
def test_pdf_warning_setting_read_only(run_engine):
    """DisableConvertPdfWarning 을 읽기만 한다: 엔진 값 = 레지스트리 값, 시험 뒤에도 값이 그대로."""
    def registry():
        version = winreg.QueryValue(winreg.HKEY_CLASSES_ROOT, r"Word.Application\CurVer").rsplit(".", 1)[1] + ".0"
        try:
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, rf"Software\Microsoft\Office\{version}\Word\Options") as key:
                return winreg.QueryValueEx(key, "DisableConvertPdfWarning")[0]
        except OSError:
            return None

    before = registry()
    out = run_engine("office", {"action": "pdf_warning_off"})
    assert out == {"off": before == 1}
    assert registry() == before


# --- 엑셀 ---------------------------------------------------------------------------------------

def _cells(grid):
    """시트마다 {(행, 열): (t, v)} (형식 f 는 빼고)."""
    return [{(r, c): (cell["t"], cell["v"]) for r, c, cell in sheet["cells"]} for sheet in grid["sheets"]]


def assert_same_grid(got, expected):
    assert got["active"] == expected["active"]
    assert [s["name"] for s in got["sheets"]] == [s["name"] for s in expected["sheets"]]
    for mine, theirs in zip(_cells(got), _cells(expected)):
        assert mine.keys() == theirs.keys()
        for key, (t, v) in theirs.items():
            assert mine[key][0] == t, key
            assert mine[key][1] == v, key
            if t == "n":   # openpyxl 이 정수로 읽는 칸은 정수로(엔진 글자가 '11.0' 이 아니라 '11')
                assert type(mine[key][1]) is type(v), (key, mine[key][1], v)


def _excel_cases():
    cases = [pytest.param(p, "samples", id=p.name) for p in SAMPLE_XLSX]
    kit = [pytest.mark.po_test_kit, pytest.mark.skipif(kit_missing, reason="C:\\PO시험 시험 자료가 없음")]
    cases.append(pytest.param(KIT / "01_입력예시" / KIT_XLSX, "kit", id=KIT_XLSX, marks=kit))
    return cases


def test_sample_xlsx_present():
    assert len(SAMPLE_XLSX) == 2


@office
@pytest.mark.office
@pytest.mark.parametrize("source,kind", _excel_cases())
def test_excel_grid_same_as_openpyxl(source, kind, tmp_path, run_engine, no_leftover):
    """엑셀로 읽은 격자가 grid_from_xlsx 와 칸마다 같고(숫자 · 글자 · 날짜), 그 격자로 낸 엔진 결과도 같다."""
    path = tmp_path / source.name
    shutil.copyfile(source, path)
    rules = example_rules() if kind == "samples" else kit_rules()
    portable = rules_to_portable(rules)
    expected_grid = grid_from_xlsx(path)

    grid = read_excel(run_engine, path)

    assert_same_grid(grid, expected_grid)
    from_openpyxl = run_engine("extract_grid", {"rules": portable, "grid": expected_grid})
    from_excel = run_engine("extract_grid", {"rules": portable, "grid": grid})
    assert from_excel == from_openpyxl
    assert from_excel["table"], "품목 표가 비어 있으면 비교가 약하다"
    # 작업 9 처럼 JSON 을 거치지 않고 메모리의 격자로 바로 추출해도 같다
    in_memory = run_engine("office", {"action": "excel_extract", "path": str(path), "rules": portable, "timeout": READ_TIMEOUT})
    assert in_memory == from_openpyxl


@office
@pytest.mark.office
def test_excel_grid_value_kinds(tmp_path, run_engine, no_leftover):
    """여러 종류의 칸(정수 · 소수 · 참거짓 · 날짜 · 일시 · 글자 · 백분율)과 활성 시트(두 번째)."""
    wb = openpyxl.Workbook()
    first = wb.active
    first.title = "Cover"
    first["A1"] = "cover"
    ws = wb.create_sheet("Data")
    values = [11, -3, 0.1, 0.30000000000000004, 1e-07, 123456789012, 2.5, True, False,
              datetime(2026, 10, 29), datetime(2026, 10, 29, 13, 45, 30), "  spaced text ", "12345", 0.05]
    for row, value in enumerate(values, start=1):
        ws.cell(row=row, column=2, value=value)
    ws.cell(row=14, column=2).number_format = "0%"
    ws.cell(row=10, column=2).number_format = "yyyy-mm-dd"
    ws["D20"] = "far cell"
    wb.active = 1
    path = tmp_path / "칸 종류.xlsx"
    wb.save(path)

    grid = read_excel(run_engine, path)

    assert_same_grid(grid, grid_from_xlsx(path))
    assert grid["active"] == 1


@office
@pytest.mark.office
def test_excel_integer_cell_in_markup_error(tmp_path, run_engine, no_leftover):
    """마크업표의 정수 칸(반올림자리 11)이 오류 문구에 '11' 로 찍힌다(작업 5 이월: COM 의 11.0 을 정수로)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["파트넘버", "마크업%", "반올림자리"])
    ws.append(["P-1", 3, 11])
    path = tmp_path / "마크업표.xlsx"
    wb.save(path)

    errors = []
    with pytest.raises(EngineError) as error:
        run_engine("markup", {"action": "read", "grid": grid_from_xlsx(path)})
    errors.append(str(error.value))
    with pytest.raises(EngineError) as error:   # 메모리의 엑셀 격자 그대로(JSON 을 거치면 11.0 이 11 로 바뀌어 차이가 안 보임)
        run_engine("office", {"action": "excel_markup", "path": str(path), "timeout": READ_TIMEOUT})
    errors.append(str(error.value))

    assert errors[0] == errors[1]
    assert "반올림자리 '11'" in errors[0]
