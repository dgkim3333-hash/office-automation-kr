"""작업 실행(run_job): 입력 폴더의 원본은 그대로 두고, 처리기록 · 결과 · 확인필요 · 승인대기를 만든다."""
import csv
import hashlib
import importlib.util
import os
import shutil
import time
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
import pytest
import yaml

import app.job_run as job_run
from app.job_run import (
    RECORD_FILE, JobLocked, RunReport, file_fingerprint, job_lock, load_records, rules_fingerprint, run_job,
)
from app.jobs import create_job, load_job
from app.master import write_example
from app.verify import Check

REPO = Path(__file__).resolve().parent.parent
PDFS = REPO / "samples" / "03_Invoice_DM_PDF"
ORDERS = REPO / "samples" / "01_고객오더_엑셀"
EXAMPLE_RULES = REPO / "rules" / "규칙.예시.yaml"
INV = PDFS / "INV-2026-10231.pdf"
ACME_PO = ORDERS / "고객오더_ACME_PO-ACME-2026-0417.xlsx"
TITLE = "일일 정리"
LEDGER = f"누적장부_{TITLE}.xlsx"

# 테스트 전용 PDF PO(tests/fixtures 의 ACME_PO_PO-77001.pdf)에 맞춘 규칙
PDF_PO_RULES = r'''
문서종류:
  PO:
    판별키워드: ["PURCHASE ORDER"]
    번호필드: 고객PO번호
    필드:
      고객PO번호: {패턴: "PO Number[.:]?\\s*([A-Z0-9-]+)"}
      발주일:     {패턴: "Date[.:]?\\s*(\\d{1,2}/\\d{1,2}/\\d{4})", 형식: 날짜}
      고객명:     {패턴: "Bill To[.:]?\\s*(.+?)\\n"}
      소계:       {패턴: "Subtotal[.:]?\\s*\\$?([\\d,]+\\.\\d{2})", 형식: 금액}
      합계금액:   {패턴: "\\nTotal[.:]?\\s*\\$?([\\d,]+\\.\\d{2})", 형식: 금액}
    표:
      시작키워드: "Item Description"
      끝키워드: "Subtotal"
      행패턴: "^(\\S+)\\s+(.+?)\\s+(\\d+)\\s+([\\d,]+\\.\\d{2})\\s+([\\d,]+\\.\\d{2})$"
      열: [품번, 품명, 수량, 단가, 금액]
      금액열: [단가, 금액]
      수량열: [수량]
    검산: {소계: 소계, 합계: 합계금액}
    파일명: "PO_{고객PO번호}.pdf"
회사양식:
  파일명: "회사양식_SO입력_{고객PO번호}.xlsx"
  시트: SO_Upload
  열:
    - {이름: Order Date,     값: 머리.발주일}
    - {이름: Customer Code,  값: 고객마스터.고객코드}
    - {이름: Customer PO No, 값: 머리.고객PO번호}
    - {이름: Line No,        값: 줄번호}
    - {이름: Part No,        값: 품목.품번}
    - {이름: Qty,            값: 품목.수량}
    - {이름: Unit Price,     값: 품목.단가}
    - {이름: Ext Price,      값: 수량x단가}
'''


@pytest.fixture
def env(tmp_path):
    work, inp, out = tmp_path / "work", tmp_path / "in", tmp_path / "out"
    for d in (work, inp, out):
        d.mkdir()
    write_example(work / "고객마스터.xlsx")
    return work, inp, out


def _job(work, inp, out, mode="누적장부", rules_text=None, title=TITLE):
    job = create_job(work, title, inp, out, mode)
    if rules_text is None:
        shutil.copy(EXAMPLE_RULES, job.rules_path)
    else:
        job.rules_path.write_text(rules_text, encoding="utf-8")
    return job


def _only(doc_type):
    """예시 규칙에서 문서종류 하나만 남긴 규칙 글자."""
    data = yaml.safe_load(EXAMPLE_RULES.read_text(encoding="utf-8"))
    data = {"문서종류": {doc_type: data["문서종류"][doc_type]}}
    return yaml.safe_dump(data, allow_unicode=True, sort_keys=False)


def _drop(src, folder, name=None, age=3600):
    """원본을 입력 폴더에 복사하고 수정 시각을 age 초 전으로 둔다."""
    dst = Path(folder) / (name or Path(src).name)
    shutil.copy(src, dst)
    t = time.time() - age
    os.utime(dst, (t, t))
    return dst


def _snapshot(folder):
    return {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in Path(folder).iterdir() if p.is_file()}


def _files(folder):
    """폴더의 파일 이름들. 장부가 저장할 때마다 남기는 바로 전 판(.bak.xlsx)은 뺀다."""
    folder = Path(folder)
    if not folder.is_dir():
        return []
    return sorted(p.name for p in folder.iterdir() if p.is_file() and not p.name.endswith(".bak.xlsx"))


def _rows(path, sheet):
    rows = list(openpyxl.load_workbook(path)[sheet].iter_rows(values_only=True))
    header = [h for h in rows[0] if h is not None]
    return [dict(zip(header, r)) for r in rows[1:] if any(v is not None for v in r)]


def _fixtures(tmp_path):
    spec = importlib.util.spec_from_file_location("build_fixtures", REPO / "tests" / "fixtures" / "build_fixtures.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    folder = tmp_path / "fixtures"
    module.build(folder)
    return folder


def test_runs_only_detected_files_and_leaves_others(env):
    work, inp, out = env
    job = _job(work, inp, out)
    _drop(PDFS / "Letter_warranty_notice.pdf", inp)
    _drop(INV, inp)
    before = _snapshot(inp)

    report = run_job(job, work)

    assert report.counts["성공"] == 1 and report.counts["확인필요"] == 0
    assert report.has_news
    assert _snapshot(inp) == before  # 원본은 이름 · 내용 모두 그대로
    records = load_records(job)
    assert [(r["파일"], r["상태"], r["문서종류"], r["번호"]) for r in records] == [
        ("INV-2026-10231.pdf", "성공", "Invoice", "INV-2026-10231")]
    assert records[0]["지문"] == file_fingerprint(inp / "INV-2026-10231.pdf")
    assert records[0]["규칙지문"] == rules_fingerprint(job) and len(rules_fingerprint(job)) == 16
    raw = (job.folder / RECORD_FILE).read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # utf-8-sig
    with open(job.folder / RECORD_FILE, encoding="utf-8-sig", newline="") as f:
        assert next(csv.reader(f)) == ["지문", "파일", "문서종류", "번호", "상태", "규칙지문", "날짜"]
    assert load_job(job.folder).last_run == datetime.now().date().isoformat()


def test_second_run_skips_processed_even_if_renamed(env):
    work, inp, out = env
    job = _job(work, inp, out)
    _drop(INV, inp)
    run_job(job, work)
    os.rename(inp / "INV-2026-10231.pdf", inp / "다시받은_인보이스.pdf")

    report = run_job(job, work)

    assert all(v == 0 for v in report.counts.values())
    assert report.unmatched_old == 0 and not report.has_news
    assert len(_rows(out / LEDGER, "Invoice")) == 1


def test_review_file_retried_next_day_or_after_rule_change(env):
    work, inp, out = env
    job = _job(work, inp, out)
    _drop(PDFS / "Scan_signed_packing_list.pdf", inp)
    now = datetime.now().replace(microsecond=0)

    first = run_job(job, work, now=now)
    assert first.counts["확인필요"] == 1
    note = job.folder / "확인필요" / "Scan_signed_packing_list.txt"
    assert note.is_file()
    text = note.read_text(encoding="utf-8")
    assert str(inp / "Scan_signed_packing_list.pdf") in text and "텍스트 없음" in text
    assert _files(job.folder / "확인필요") == ["Scan_signed_packing_list.txt"]  # 원본 복사 없음

    assert run_job(job, work, now=now).counts["확인필요"] == 0  # 같은 날: 건너뜀
    assert run_job(job, work, now=now + timedelta(days=1)).counts["확인필요"] == 1  # 다음 날: 다시 시도

    with open(job.rules_path, "a", encoding="utf-8") as f:
        f.write("\n# 규칙을 고침\n")
    assert run_job(job, work, now=now + timedelta(days=1)).counts["확인필요"] == 1  # 규칙이 바뀌면 같은 날도
    assert [r["상태"] for r in load_records(job)] == ["확인필요"] * 3
    assert (inp / "Scan_signed_packing_list.pdf").is_file()


def test_new_file_with_same_number_is_duplicate(env):
    work, inp, out = env
    job = _job(work, inp, out)
    _drop(INV, inp)
    run_job(job, work)
    changed = inp / "INV-2026-10231_재발송.pdf"
    changed.write_bytes(INV.read_bytes() + b"\n% resent\n")  # 내용이 다른 파일, 같은 인보이스번호
    os.utime(changed, (time.time() - 3600,) * 2)
    out_before = _snapshot(out)

    report = run_job(job, work)

    assert report.counts["중복"] == 1 and report.counts["성공"] == 0
    assert report.has_news and report.output_files == []
    assert _snapshot(out) == out_before
    assert load_records(job)[-1]["상태"] == "중복"
    assert _files(job.folder / "검산") == ["검산_INV-2026-10231.xlsx"]


def test_duplicate_found_without_ledger_in_company_mode(env):
    work, inp, out = env
    job = _job(work, inp, out, mode="회사양식")
    _drop(ACME_PO, inp)
    run_job(job, work)
    resent = inp / "재발송_PO.xlsx"
    wb = openpyxl.load_workbook(ACME_PO)
    wb.active["Z1"] = "재발송"  # 내용은 다르지만 같은 PO 번호
    wb.save(resent)
    os.utime(resent, (time.time() - 3600,) * 2)

    report = run_job(job, work)

    assert report.counts["중복"] == 1 and report.counts["성공"] == 0
    assert _files(out) == ["회사양식_SO입력_PO-ACME-2026-0417.xlsx"]


def test_mode_ledger_only_writes_ledger_and_no_company_file(env):
    work, inp, out = env
    job = _job(work, inp, out, mode="누적장부")
    _drop(INV, inp)

    report = run_job(job, work)

    assert _files(out) == [LEDGER]
    assert report.output_files == [out / LEDGER]
    rows = _rows(out / LEDGER, "Invoice")
    assert [r["인보이스번호"] for r in rows] == ["INV-2026-10231"]
    assert {r["문서번호"] for r in _rows(out / LEDGER, "품목")} == {"INV-2026-10231"}
    assert _files(job.folder / "검산") == ["검산_INV-2026-10231.xlsx"]


def test_mode_company_only_writes_company_file_and_no_ledger(env):
    work, inp, out = env
    job = _job(work, inp, out, mode="회사양식")
    _drop(ACME_PO, inp)

    report = run_job(job, work)

    assert report.counts["성공"] == 1
    assert _files(out) == ["회사양식_SO입력_PO-ACME-2026-0417.xlsx"]
    assert report.output_files == [out / "회사양식_SO입력_PO-ACME-2026-0417.xlsx"]
    wb = openpyxl.load_workbook(out / "회사양식_SO입력_PO-ACME-2026-0417.xlsx")
    assert wb.sheetnames == ["SO_Upload"]  # 회사 양식에 시트를 더하지 않는다
    assert _files(job.folder / "검산") == ["검산_PO-ACME-2026-0417.xlsx"]


def test_mode_both_from_excel_po(env):
    work, inp, out = env
    job = _job(work, inp, out, mode="둘다")
    _drop(ACME_PO, inp)
    before = _snapshot(inp)

    report = run_job(job, work)

    assert report.counts["성공"] == 1
    assert _files(out) == sorted(["회사양식_SO입력_PO-ACME-2026-0417.xlsx", LEDGER])
    assert [r["고객PO번호"] for r in _rows(out / LEDGER, "PO")] == ["PO-ACME-2026-0417"]
    assert _snapshot(inp) == before


def test_company_form_from_pdf_po_when_mode_company(env, tmp_path):
    work, inp, out = env
    master = openpyxl.Workbook()
    master.active.append(["고객명", "고객코드", "결제조건"])
    master.active.append(["ACME Semiconductor Korea", "C-KR-0001", "Net 30"])
    master.save(work / "고객마스터.xlsx")
    job = _job(work, inp, out, mode="회사양식", rules_text=PDF_PO_RULES)
    _drop(_fixtures(tmp_path) / "ACME_PO_PO-77001.pdf", inp)

    report = run_job(job, work)

    assert report.counts["성공"] == 1, report.message
    path = out / "회사양식_SO입력_PO-77001.xlsx"
    rows = _rows(path, "SO_Upload")
    assert [(r["Customer Code"], r["Line No"], r["Part No"], r["Qty"]) for r in rows] == [
        ("C-KR-0001", 10, "CP-200", 1), ("C-KR-0001", 20, "HX-900", 1)]
    assert [str(r["Ext Price"]) for r in rows] == ["4500", "7845.67"]


def test_verify_failure_goes_to_review_and_nothing_written(env, monkeypatch):
    work, inp, out = env
    job = _job(work, inp, out, mode="둘다")
    _drop(INV, inp)
    monkeypatch.setattr(job_run, "verify", lambda result, path: [Check("합계금액", "100.00", "1000.00", False)])

    report = run_job(job, work)

    assert report.counts["확인필요"] == 1 and report.counts["성공"] == 0
    assert _files(out) == []
    assert _files(job.folder / "검산") == []
    text = (job.folder / "확인필요" / "INV-2026-10231.txt").read_text(encoding="utf-8")
    assert "합계금액" in text
    assert load_records(job)[-1]["상태"] == "확인필요"
    assert (inp / "INV-2026-10231.pdf").is_file()


def test_read_back_mismatch_rolls_back(env, monkeypatch):
    work, inp, out = env
    job = _job(work, inp, out, mode="둘다")
    _drop(ACME_PO, inp)
    monkeypatch.setattr(job_run, "read_back", lambda *a, **k: ["2행 Qty: 쓴 값 1, 읽은 값 2"])
    seen = []

    report = run_job(job, work, on_success=lambda job, result: seen.append(result))

    assert report.counts["확인필요"] == 1 and report.counts["성공"] == 0
    assert not (out / "회사양식_SO입력_PO-ACME-2026-0417.xlsx").exists()
    assert _rows(out / LEDGER, "PO") == [] and _rows(out / LEDGER, "품목") == []
    assert report.output_files == [] and seen == []
    assert "Qty" in (job.folder / "확인필요" / "고객오더_ACME_PO-ACME-2026-0417.txt").read_text(encoding="utf-8")
    assert load_records(job)[-1]["상태"] == "확인필요"


def test_ledger_read_back_mismatch_rolls_back(env, monkeypatch):
    work, inp, out = env
    job = _job(work, inp, out, mode="누적장부")
    _drop(INV, inp)
    monkeypatch.setattr(job_run, "_ledger_problems", lambda *a, **k: ["Invoice 인보이스번호: 다름"])

    report = run_job(job, work)

    assert report.counts["확인필요"] == 1
    assert _rows(out / LEDGER, "Invoice") == [] and _rows(out / LEDGER, "품목") == []


def test_ledger_read_back_finds_written_rows(env):
    work, inp, out = env
    job = _job(work, inp, out, mode="누적장부")
    _drop(INV, inp)
    _drop(PDFS / "INV-2026-10278.pdf", inp)
    assert run_job(job, work).counts["성공"] == 2  # 장부에 두 번째 문서를 더해도 다시 읽기가 맞다


def test_gate_true_copies_to_waiting_with_check_file(env):
    work, inp, out = env
    job = _job(work, inp, out, mode="둘다")
    _drop(INV, inp)
    before = _snapshot(inp)
    calls = []

    def gate(job_, rules, result):
        calls.append(result.number)
        return True

    report = run_job(job, work, gate=gate)

    assert calls == ["INV-2026-10231"]
    assert report.counts["승인대기"] == 1 and report.has_news
    assert _snapshot(inp) == before
    waiting = job.folder / "승인대기"
    assert (waiting / "INV-2026-10231.pdf").read_bytes() == INV.read_bytes()
    assert _files(waiting / "대조표") == ["검산_INV-2026-10231.xlsx"]
    assert _files(out) == [] and _files(job.folder / "검산") == []
    assert load_records(job)[-1]["상태"] == "승인대기"
    assert run_job(job, work, gate=gate).counts["승인대기"] == 0  # 다시 돌려도 또 보내지 않는다


def test_approval_run_reads_waiting_copy_and_keeps_it(env):
    work, inp, out = env
    job = _job(work, inp, out, mode="누적장부")
    _drop(INV, inp)
    run_job(job, work, gate=lambda *a: True)
    seen = []

    report = run_job(job, work, source_dir=job.folder / "승인대기", on_success=lambda j, r: seen.append(r.number))

    assert report.counts["성공"] == 1 and seen == ["INV-2026-10231"]
    assert (job.folder / "승인대기" / "INV-2026-10231.pdf").is_file()  # 복사본은 승인 쪽(W5)이 지운다
    assert [r["인보이스번호"] for r in _rows(out / LEDGER, "Invoice")] == ["INV-2026-10231"]
    assert run_job(job, work).counts == dict.fromkeys(report.counts, 0)  # 입력 폴더의 원본은 이미 처리됨


def test_file_modified_within_a_minute_is_left_for_next_run(env):
    work, inp, out = env
    job = _job(work, inp, out)
    _drop(INV, inp, age=0)

    report = run_job(job, work)

    assert report.counts["성공"] == 0 and report.counts["보류"] == 1
    assert not report.has_news
    assert load_records(job) == []
    later = run_job(job, work, now=datetime.now() + timedelta(minutes=2))
    assert later.counts["성공"] == 1


def test_open_file_is_held(env, monkeypatch):
    work, inp, out = env
    job = _job(work, inp, out)
    held = _drop(INV, inp)
    monkeypatch.setattr(job_run, "is_locked", lambda path: Path(path) == held)

    report = run_job(job, work)

    assert report.counts["보류"] == 1 and report.counts["성공"] == 0
    assert load_records(job) == []


def test_missing_output_folder_stops_without_moving(env):
    work, inp, out = env
    job = _job(work, inp, out)
    _drop(INV, inp)
    before = _snapshot(inp)
    out.rmdir()

    report = run_job(job, work)

    assert "출력 폴더가 없습니다" in report.message
    assert all(v == 0 for v in report.counts.values()) and not report.has_news
    assert _snapshot(inp) == before
    assert not (job.folder / RECORD_FILE).exists()
    assert load_job(job.folder).last_run == ""


def test_missing_input_folder_stops(env):
    work, inp, out = env
    job = _job(work, inp, out)
    inp.rmdir()

    report = run_job(job, work)

    assert "입력 폴더가 없습니다" in report.message
    assert _files(out) == []


def test_second_run_while_locked_exits_with_message(env):
    work, inp, out = env
    job = _job(work, inp, out)
    _drop(INV, inp)
    with job_lock(job):
        with pytest.raises(JobLocked, match="이미 실행 중"):
            run_job(job, work)
    assert load_records(job) == []
    assert run_job(job, work).counts["성공"] == 1  # 잠금이 풀리면 다시 돈다


def test_two_jobs_sharing_input_folder_take_their_own_files(env):
    work, inp, out = env
    invoices = _job(work, inp, out, rules_text=_only("Invoice"), title="인보이스")
    memos = _job(work, inp, out, rules_text=_only("DM"), title="DM")
    for name in ("INV-2026-10231.pdf", "INV-2026-10278.pdf", "DM-2026-0092.pdf", "DM-2026-0097.pdf"):
        _drop(PDFS / name, inp)
    before = _snapshot(inp)

    first = run_job(invoices, work)
    second = run_job(memos, work)

    assert first.counts["성공"] == 2 and first.counts["확인필요"] == 0
    assert second.counts["성공"] == 2 and second.counts["확인필요"] == 0
    assert sorted(r["번호"] for r in load_records(invoices)) == ["INV-2026-10231", "INV-2026-10278"]
    assert sorted(r["번호"] for r in load_records(memos)) == ["DM-2026-0092", "DM-2026-0097"]
    assert _files(out) == ["누적장부_DM.xlsx", "누적장부_인보이스.xlsx"]
    assert "DM" not in openpyxl.load_workbook(out / "누적장부_인보이스.xlsx").sheetnames
    assert _snapshot(inp) == before


def test_unmatched_old_file_is_counted(env):
    work, inp, out = env
    invoices = _job(work, inp, out, rules_text=_only("Invoice"), title="인보이스")
    memos = _job(work, inp, out, rules_text=_only("DM"), title="DM")
    day = 25 * 3600
    _drop(INV, inp, age=day)
    _drop(PDFS / "DM-2026-0092.pdf", inp, age=day)
    _drop(PDFS / "Letter_warranty_notice.pdf", inp, age=day)                       # 어느 작업에도 안 맞음
    _drop(PDFS / "Letter_warranty_notice.pdf", inp, name="새편지.pdf", age=3600)     # 하루가 안 됨
    _drop(PDFS / "Letter_warranty_notice.pdf", inp, name="~$임시.xlsx", age=day)     # 엑셀 임시 파일

    run_job(memos, work)
    report = run_job(invoices, work)

    assert report.counts["성공"] == 1
    assert report.unmatched_old == 1  # DM 은 다른 작업이 처리했으므로 세지 않는다
    assert report.has_news
    assert "어느 작업에도 맞지 않는 파일 1개" in report.message


def test_run_report_has_news():
    empty = dict.fromkeys(("성공", "중복", "확인필요", "승인대기", "보류"), 0)
    assert not RunReport(counts=dict(empty), unmatched_old=0, message="", output_files=[]).has_news
    assert not RunReport(counts={**empty, "보류": 2}, unmatched_old=0, message="", output_files=[]).has_news
    assert RunReport(counts={**empty, "중복": 1}, unmatched_old=0, message="", output_files=[]).has_news
    assert RunReport(counts=dict(empty), unmatched_old=1, message="", output_files=[]).has_news
